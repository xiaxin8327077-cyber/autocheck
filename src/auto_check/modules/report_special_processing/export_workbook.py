from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from math import ceil
from typing import Any, Mapping, Sequence

from .contracts import STATUS_LABELS, RecordStatus
from .bilingual_names import parse_bilingual_groups, parse_bilingual_items, serialize_bilingual_items
from .ledger_display import display_width, ledger_display


EXPORT_HEADERS = (
    "报送期", "处理表名", "修改字段", "修改前", "修改后", "所属业务系统", "关联报送", "处理摘要", "状态", "处理人", "处理时间",
)
MAX_EXPORT_ROWS = 5_000
_COLUMN_WIDTHS = (16, 28, 42, 32, 32, 20, 30, 46, 12, 16, 23)


def _display(value: Any) -> str:
    text = "" if value is None else str(value)
    return text if text.strip() else "—"


def _period_text(value: Any) -> str:
    if isinstance(value, datetime):
        value = value.date()
    return value.isoformat() if isinstance(value, date) else _display(value)


def _table_sections(record: Mapping[str, Any]) -> list[tuple[str, Mapping[str, Any]]]:
    """导出按表保留字段归属；每个表内部复用列表的同值分组规则。"""
    structured = record.get("structured_content")
    tables = structured.get("tables") if isinstance(structured, Mapping) else None
    if isinstance(tables, list) and tables:
        return [
            (
                _display(str(table.get("chinese_table_name") or "").strip() or str(table.get("table_name") or "").strip()),
                {**record, "structured_content": {"tables": [table]}},
            )
            for table in tables if isinstance(table, Mapping)
        ] or [("—", record)]
    table_text = str(record.get("table_name") or "").strip()
    try:
        names = parse_bilingual_items(table_text)
    except ValueError:
        return [(_display(table_text), record)]
    table_label = "\n".join(zh or en for zh, en in names)
    try:
        field_groups = parse_bilingual_groups(str(record.get("field_name") or ""))
    except ValueError:
        return [(_display(table_label), record)]
    if not names or len(names) != len(field_groups):
        return [(_display(table_label), record)]
    before = "" if record.get("value_before") is None else str(record["value_before"])
    after = "" if record.get("value_after") is None else str(record["value_after"])
    before_lines, after_lines = before.split("\n"), after.split("\n")
    global_values = len(before_lines) == len(after_lines) == 1
    field_count = sum(len(fields) for fields in field_groups)
    if not global_values and not (field_count == len(before_lines) == len(after_lines)):
        return [(_display(table_label), record)]
    sections = []
    offset = 0
    for (zh, en), fields in zip(names, field_groups):
        length = len(fields)
        sections.append((zh or en, {
            **record,
            "structured_content": None,
            "field_name": serialize_bilingual_items(fields),
            "value_before": before if global_values else "\n".join(before_lines[offset:offset + length]),
            "value_after": after if global_values else "\n".join(after_lines[offset:offset + length]),
        }))
        offset += length
    return sections


def _record_rows(record: Mapping[str, Any]) -> list[list[str]]:
    display = ledger_display(record)
    code = str(record.get("status") or "")
    try:
        status = STATUS_LABELS[RecordStatus(code)]
    except ValueError:
        status = code
    metadata = [
        _display(record.get("business_system_name_snapshot")),
        _display("\n".join(display["process_names"])),
        _display(record.get("summary")),
        _display(status),
        _display(record.get("handler_display_name_snapshot") or record.get("handler_username_snapshot")),
        _display(display["handled_at"]),
    ]
    rows = []
    for table_name, section in _table_sections(record):
        groups = ledger_display(section)["change_groups"]
        rows.extend(
            [_period_text(record.get("report_period")), table_name, "\n".join(group["fields"]), _display(group["before"]), _display(group["after"]), *metadata]
            for group in groups
        )
    return rows


def export_rows(records: Sequence[Mapping[str, Any]]) -> list[list[Any]]:
    return [row for record in records for row in _record_rows(record)]


def _line_count(value: str, width: float) -> int:
    return sum(max(1, ceil(display_width(line) / (width - 2))) for line in value.split("\n"))


def build_export_xlsx(records: Sequence[Mapping[str, Any]], *, title: str = "报表特殊处理") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.properties.title = title
    sheet = workbook.active
    sheet.title = "报表特殊处理"
    sheet.append(list(EXPORT_HEADERS))
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    edge = Side(style="thin", color="AEB8C8")
    border = Border(left=edge, right=edge, top=edge, bottom=edge)
    for index, width in enumerate(_COLUMN_WIDTHS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.font = Font(name="微软雅黑", size=11, bold=True, color="244578")
        cell.fill = PatternFill("solid", fgColor="F0F5FA")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    sheet.row_dimensions[1].height = 30

    for record in records:
        rows = _record_rows(record)
        start = sheet.max_row + 1
        metadata_columns = (0, *range(5, len(EXPORT_HEADERS)))
        metadata_lines = max(_line_count(rows[0][index], _COLUMN_WIDTHS[index]) for index in metadata_columns)
        for values in rows:
            sheet.append(values)
            row_number = sheet.max_row
            lines = max(_line_count(value, _COLUMN_WIDTHS[index]) for index, value in enumerate(values[1:5], 1))
            lines = max(lines, ceil(metadata_lines / len(rows)))
            sheet.row_dimensions[row_number].height = min(409, max(36, lines * 16 + 12))
            for cell in sheet[row_number]:
                # 所有业务值均为文本，保留零、前导零和以等号开头的原始内容。
                cell.data_type = "s"
                cell.number_format = "@"
                cell.font = Font(name="微软雅黑", size=11, color="244578")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = border
        if len(rows) > 1:
            for column in (index + 1 for index in metadata_columns):
                sheet.merge_cells(start_row=start, end_row=sheet.max_row, start_column=column, end_column=column)
        status_cell = sheet.cell(start, EXPORT_HEADERS.index("状态") + 1)
        status_color = {"completed": "008C69", "pending": "C67500", "voided": "C53B46"}.get(str(record.get("status")), "244578")
        status_cell.font = Font(name="微软雅黑", size=11, color=status_color)

    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
