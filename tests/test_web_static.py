from __future__ import annotations

import re
import subprocess
import textwrap
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
INDEX_HTML = ROOT / "src" / "auto_check" / "web" / "index.html"
APP_JS = ROOT / "src" / "auto_check" / "web" / "app.js"
EXPORT_DETAIL_JS = ROOT / "src" / "auto_check" / "web" / "export_detail.js"
SERVER_PY = ROOT / "src" / "auto_check" / "app" / "server.py"
STYLES_CSS = ROOT / "src" / "auto_check" / "web" / "styles.css"
README_MD = ROOT / "README.md"
PYINSTALLER_SPEC = ROOT / "auto-check.spec"
MODULE_HOST_JS = ROOT / "src" / "auto_check" / "web" / "module_host.js"
MODULE_HOST_CSS = ROOT / "src" / "auto_check" / "web" / "module_host.css"


def _read(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def test_module_host_assets_are_loaded_before_legacy_app():
    html = _read(INDEX_HTML)

    assert html.count('href="/module_host.css"') == 1
    assert html.count('src="/module_host.js') == 1
    assert html.index('src="/module_host.js') < html.index('src="/app.js')
    assert 'id="moduleSideNavigation"' in html
    assert 'id="moduleTopNavigation"' in html
    assert 'id="modulePageHost"' in html


def test_module_host_assets_keep_module_styles_scoped():
    css = _read(MODULE_HOST_CSS)
    script = _read(MODULE_HOST_JS)

    assert ".auto-check-module" in css
    assert "\nbutton {" not in css
    assert "\ninput {" not in css
    assert "\ntable {" not in css
    assert "Object.freeze(context)" in script
    assert "hashchange" in script
    assert "await window.AutoCheckModuleHost?.deactivate();" in _read(APP_JS)


def test_module_host_grouped_top_navigation_keeps_accessibility_and_theme_boundaries():
    html = _read(INDEX_HTML)
    css = _read(MODULE_HOST_CSS)
    script = _read(MODULE_HOST_JS)

    for fragment in [
        "module-top-nav-group-toggle",
        "aria-expanded",
        "aria-controls",
        "aria-current",
        "group_order",
        "group_id",
    ]:
        assert fragment in script
    assert "border-radius: var(--ui-radius)" in css
    assert "z-index: 50" in css
    assert "box-shadow: none" in css
    assert "width: max-content" in css
    assert "min-width: 100%" in css
    assert "white-space: nowrap" in css
    assert "font-family: inherit" in css
    assert "font: inherit" not in css
    assert "width: min(280px" not in css
    assert html.index('data-nav-group="smart-reconcile"', html.index('<nav class="top-nav-tabs">')) < html.index('id="moduleTopNavigation"')
    assert html.index('id="moduleTopNavigation"') < html.index('data-page="tools"', html.index('<nav class="top-nav-tabs">'))
    assert 'wrapper.addEventListener("mouseenter"' in script
    assert 'wrapper.addEventListener("mouseleave"' in script
    assert "box-shadow: 0 0" not in css


def test_semantic_action_tokens_and_disabled_priority_are_centralized():
    css = _read(STYLES_CSS)

    for declaration in [
        "--action-danger: #BA1A1A;",
        "--action-warning: #B45309;",
        "--action-success: #137333;",
        "--action-danger: #FFB4AB;",
        "--action-warning: #FBBF24;",
        "--action-success: #6DDB9C;",
    ]:
        assert declaration in css

    for selector in [
        '[data-action-tone="primary"][data-action-variant="solid"]',
        '[data-action-tone="danger"][data-action-variant="solid"]',
        '[data-action-tone="warning"][data-action-variant="solid"]',
        '[data-action-tone="success"][data-action-variant="solid"]',
        '[data-action-tone="danger"][data-action-variant="weak"]',
        '[data-action-tone="warning"][data-action-variant="weak"]',
        '[data-action-tone="success"][data-action-variant="weak"]',
        '[data-action-tone][data-action-variant]:disabled',
    ]:
        assert selector in css

    disabled_rule = re.search(
        r'(?m)^:is\([^\n]*\[data-action-tone\][^\n]*\):disabled[^\{]*\{(?P<body>.*?)\}',
        css,
        re.S,
    )
    assert disabled_rule is not None
    disabled_body = disabled_rule.group("body")
    assert "color: var(--on-surface-variant)" in disabled_body
    assert "background: var(--surface-variant)" in disabled_body
    assert "box-shadow: none" in disabled_body
    assert "transform: none" in disabled_body


def test_semantic_button_inventory_classifies_key_static_and_dynamic_actions():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)

    static_actions = {
        "stopRunBtn": ("danger", "solid"),
        "pbcClearFilesBtn": ("danger", "weak"),
        "pbcFinishBtn": ("success", "solid"),
        "dbValidationDownloadBtn": ("neutral", "weak"),
        "flowCancelBtn": ("danger", "solid"),
        "resetInterfaceSettingsBtn": ("warning", "weak"),
        "resetSettingsBtn": ("warning", "weak"),
        "userModalCancel": ("neutral", "weak"),
        "userModalSave": ("primary", "solid"),
    }
    for element_id, (tone, variant) in static_actions.items():
        button = re.search(rf'<button[^>]*id="{element_id}"[^>]*>', html)
        assert button is not None, element_id
        assert f'data-action-tone="{tone}"' in button.group(0), element_id
        assert f'data-action-variant="{variant}"' in button.group(0), element_id

    for required_fragment in [
        'class="user-icon-action edit-user" data-action-tone="primary" data-action-variant="weak"',
        'class="user-icon-action toggle-user" data-action-tone="${enabled ? "warning" : "success"}" data-action-variant="weak"',
        'class="user-icon-action delete-user" data-action-tone="danger" data-action-variant="weak"',
        'class="btn-outline btn-xs btn-danger delete-history" data-action-tone="danger" data-action-variant="weak"',
        'class="btn-outline btn-xs btn-danger del-cfg" data-action-tone="danger" data-action-variant="weak"',
        'class="btn-outline btn-sm flow-chain-remove" data-action="remove-chain" data-action-tone="danger" data-action-variant="weak"',
        'class="pbc-file-remove-btn" data-action-tone="danger" data-action-variant="weak"',
        'data-action-tone="${canRestore ? "neutral" : "danger"}" data-action-variant="weak"',
    ]:
        assert required_fragment in app_js


def test_show_confirm_normalizes_and_resets_explicit_action_tone():
    app_js = _read(APP_JS)

    assert "function showConfirm(title, message, options = {})" in app_js
    confirm_body = app_js[
        app_js.index("function showConfirm(title, message, options = {})") :
        app_js.index("function showPrompt", app_js.index("function showConfirm(title, message, options = {})"))
    ]
    assert 'const allowedTones = new Set(["primary", "danger", "warning", "success"]);' in confirm_body
    assert 'const requestedTone = allowedTones.has(resolvedOptions.tone) ? resolvedOptions.tone : "primary";' in confirm_body
    assert 'const tone = requestedTone === "danger" ? "danger" : "primary";' in confirm_body
    assert "okBtn.dataset.actionTone = tone;" in confirm_body
    assert 'okBtn.dataset.actionVariant = "solid";' in confirm_body
    assert "delete okBtn.dataset.actionTone;" in confirm_body
    assert "delete okBtn.dataset.actionVariant;" in confirm_body

    for required_fragment in [
        'showConfirm("删除用户", `确定删除用户 ${targetUser.username} 吗？`, { tone: "danger" })',
        'showConfirm("删除历史记录", "确定删除这条历史记录吗？", { tone: "danger" })',
        'showConfirm("删除数据源", `确定删除“${cfg?.name || b.dataset.id}”吗？`, { tone: "danger" })',
        'showConfirm("初始化表字段配置", "将使用服务端 reconcile-schema.yaml 覆盖当前页面配置。是否继续？", { tone: "warning" })',
        '{ tone: nextEnabled ? "success" : "warning" }',
        '{ tone: "danger" }',
    ]:
        assert required_fragment in app_js


def _run_interface_radius_node_scenario(tmp_path: Path, scenario_source: str) -> None:
    app_js = _read(APP_JS)
    block = re.search(
        r"// Interface radius start.*?// Interface radius end",
        app_js,
        re.S,
    )
    assert block is not None
    ensure_authenticated = re.search(
        r"async function ensureAuthenticated\(\) \{.*?\n\}",
        app_js,
        re.S,
    )
    logout = re.search(
        r"async function logout\(\) \{.*?\n\}",
        app_js,
        re.S,
    )
    assert ensure_authenticated is not None
    assert logout is not None

    script = textwrap.dedent(
        """
        const assert = require("node:assert/strict");

        class FakeElement {
          constructor(value = "") {
            this.value = value;
            this.checked = false;
            this.textContent = "";
            this.disabled = false;
            this.listeners = new Map();
            this.classList = {
              values: new Set(),
              toggle: (name, enabled) => {
                if (enabled) this.classList.values.add(name);
                else this.classList.values.delete(name);
              },
            };
          }

          addEventListener(type, listener) {
            this.listeners.set(type, listener);
          }

          dispatch(type) {
            return this.listeners.get(type)?.({ target: this });
          }
        }

        const elements = {
          interfaceRadiusSlider: new FakeElement("4"),
          interfaceRadiusValue: new FakeElement(),
          interfaceLineChartStyleStraight: new FakeElement("straight"),
          interfaceLineChartStyleSmooth: new FakeElement("smooth"),
          interfaceSettingsStatus: new FakeElement(),
          saveInterfaceSettingsBtn: new FakeElement(),
          resetInterfaceSettingsBtn: new FakeElement(),
        };
        const cssVariables = new Map();
        globalThis.document = {
          getElementById: (id) => elements[id] || null,
          documentElement: {
            dataset: {},
            style: {
              setProperty: (name, value) => cssVariables.set(name, value),
              getPropertyValue: (name) => cssVariables.get(name) || "",
            },
          },
        };

        let nextTimerId = 1;
        const timers = new Map();
        globalThis.setTimeout = (callback, delay) => {
          const id = nextTimerId++;
          timers.set(id, { callback, delay });
          return id;
        };
        globalThis.clearTimeout = (id) => timers.delete(id);
        const runAllTimers = () => {
          const pending = [...timers.values()];
          timers.clear();
          pending.forEach(({ callback }) => callback());
        };

        const apiCalls = [];
        const toasts = [];
        const sessionStorageRemovals = [];
        const storageValues = new Map();
        const storageWrites = [];
        const authState = { csrfToken: "old-token", user: { id: "old-user", role: "admin" } };
        const USER_AVATAR_SESSION_KEY = "avatar-key";
        globalThis.window = {
          location: { href: "/" },
        };
        globalThis.sessionStorage = {
          removeItem: (key) => sessionStorageRemovals.push(key),
        };
        globalThis.localStorage = {
          getItem: (key) => storageValues.get(key) ?? null,
          setItem: (key, value) => {
            storageValues.set(key, String(value));
            storageWrites.push({ key, value: String(value) });
          },
          removeItem: (key) => storageValues.delete(key),
        };
        let apiImpl = async () => ({ settings: { radius_px: 4, line_chart_style: "straight" } });
        let fetchImpl = async () => ({
          ok: true,
          json: async () => ({
            authenticated: true,
            csrf_token: "new-token",
            user: { id: "new-user", role: "user" },
          }),
        });
        let confirmResult = true;
        let revealCount = 0;
        let adaptLegacyRadiusResponses = true;
        function legacyRadiusResponse(payload) {
          const settings = payload?.settings;
          if (!settings || !Object.hasOwn(settings, "radius_px")) return payload;
          return {
            ...payload,
            settings: {
              line_chart_style: "straight",
              ...settings,
            },
          };
        }
        async function api(path, options = {}) {
          apiCalls.push({ path, options });
          const payload = await apiImpl(path, options);
          return adaptLegacyRadiusResponses ? legacyRadiusResponse(payload) : payload;
        }
        async function fetch(path, options = {}) {
          return fetchImpl(path, options);
        }
        function showToast(message, type = "info") {
          toasts.push({ message, type });
        }
        async function showConfirm() {
          return confirmResult;
        }
        function activateThemeUserStorage() {}
        function applySavedUserTheme() {}
        function updateCurrentUsername() {}
        function applyRoleAccess() {}
        function clearReportNavigationCache() {}
        function revealAuthenticatedApp() {
          revealCount += 1;
        }
        function resetSystemThemeColorsForAuthChange() { return 1; }
        async function loadSystemThemeColors() { return true; }
        function captureSystemThemeColors() { return {}; }
        function restoreSystemThemeColors() { return true; }
        const systemThemeColorState = { authRevision: 1 };
        function deferred() {
          let resolve;
          let reject;
          const promise = new Promise((res, rej) => {
            resolve = res;
            reject = rej;
          });
          return { promise, resolve, reject };
        }
        async function flushMicrotasks() {
          for (let index = 0; index < 6; index += 1) await Promise.resolve();
        }

        __INTERFACE_RADIUS_BLOCK__
        __ENSURE_AUTHENTICATED__
        __LOGOUT__

        const radiusHarness = {
          state: interfaceRadiusState,
          load: loadInterfaceRadiusPreference,
          save: saveInterfaceRadiusPreference,
          discard: discardUnsavedInterfaceRadius,
          resetAuth: resetInterfaceRadiusForAuthChange,
          useStrictResponses: () => { adaptLegacyRadiusResponses = false; },
          ensureAuthenticated,
          logout,
          authState,
          elements,
          cssVariables,
          apiCalls,
          toasts,
          storageValues,
          storageWrites,
          runAllTimers,
          timerCount: () => timers.size,
          revealCount: () => revealCount,
        };

        (async () => {
        __SCENARIO__
        })().catch((error) => {
          console.error(error.stack || error);
          process.exitCode = 1;
        });
        """
    ).replace("__INTERFACE_RADIUS_BLOCK__", block.group(0)).replace(
        "__ENSURE_AUTHENTICATED__",
        ensure_authenticated.group(0),
    ).replace(
        "__LOGOUT__",
        logout.group(0),
    ).replace(
        "__SCENARIO__",
        textwrap.indent(textwrap.dedent(scenario_source).strip(), "  "),
    )
    script_path = tmp_path / "interface_radius_state_machine.cjs"
    script_path.write_text(script, encoding="utf-8")
    subprocess.run(["node", str(script_path)], check=True, cwd=ROOT)


def _run_theme_color_node_scenario(tmp_path: Path, scenario_source: str) -> None:
    app_js = _read(APP_JS)
    block = re.search(
        r"// Theme color runtime start.*?// Theme color runtime end",
        app_js,
        re.S,
    )
    assert block is not None
    script = textwrap.dedent(
        """
        const assert = require("node:assert/strict");
        const cssVariables = new Map();
        const attributes = new Map([
          ["data-theme", "space-tech"],
          ["data-color-mode", "light"],
        ]);
        globalThis.document = {
          documentElement: {
            getAttribute: (name) => attributes.get(name) || null,
            setAttribute: (name, value) => attributes.set(name, String(value)),
            style: {
              setProperty: (name, value) => cssVariables.set(name, value),
              getPropertyValue: (name) => cssVariables.get(name) || "",
            },
          },
        };

        __THEME_COLOR_BLOCK__

        (async () => {
        __SCENARIO__
        })().catch((error) => {
          console.error(error.stack || error);
          process.exitCode = 1;
        });
        """
    ).replace("__THEME_COLOR_BLOCK__", block.group(0)).replace(
        "__SCENARIO__",
        textwrap.indent(textwrap.dedent(scenario_source).strip(), "  "),
    )
    script_path = tmp_path / "theme_color_runtime.cjs"
    script_path.write_text(script, encoding="utf-8")
    subprocess.run(["node", str(script_path)], check=True, cwd=ROOT)


def _run_system_theme_color_node_scenario(tmp_path: Path, scenario_source: str) -> None:
    app_js = _read(APP_JS)
    runtime_block = re.search(
        r"// Theme color runtime start.*?// Theme color runtime end",
        app_js,
        re.S,
    )
    state_block = re.search(
        r"// System theme colors start.*?// System theme colors end",
        app_js,
        re.S,
    )
    assert runtime_block is not None
    assert state_block is not None

    script = textwrap.dedent(
        """
        const assert = require("node:assert/strict");

        class FakeElement {
          constructor(value = "") {
            this.value = value;
            this.textContent = "";
            this.hidden = false;
            this.disabled = false;
            this.style = {};
            this.listeners = new Map();
            this.classList = {
              values: new Set(),
              toggle: (name, enabled) => {
                if (enabled) this.classList.values.add(name);
                else this.classList.values.delete(name);
              },
            };
          }

          addEventListener(type, listener) {
            this.listeners.set(type, listener);
          }

          dispatch(type) {
            return this.listeners.get(type)?.({ target: this });
          }
        }

        const elements = {
          systemThemeColorsSection: new FakeElement(),
          systemVitalityThemeColor: new FakeElement("#3F6FAF"),
          systemCalmThemeColor: new FakeElement("#355F63"),
          systemVitalityThemeColorSwatch: new FakeElement(),
          systemCalmThemeColorSwatch: new FakeElement(),
          systemVitalityThemeColorError: new FakeElement(),
          systemCalmThemeColorError: new FakeElement(),
          systemThemeColorsStatus: new FakeElement(),
          saveSystemThemeColorsBtn: new FakeElement(),
          resetSystemThemeColorsBtn: new FakeElement(),
        };
        elements.systemThemeColorsSection.hidden = true;
        const cssVariables = new Map();
        const attributes = new Map([
          ["data-theme", "space-tech"],
          ["data-color-mode", "light"],
        ]);
        globalThis.document = {
          getElementById: (id) => elements[id] || null,
          documentElement: {
            getAttribute: (name) => attributes.get(name) || null,
            setAttribute: (name, value) => attributes.set(name, String(value)),
            style: {
              setProperty: (name, value) => cssVariables.set(name, value),
              getPropertyValue: (name) => cssVariables.get(name) || "",
            },
          },
        };

        const storageValues = new Map();
        const storageWrites = [];
        globalThis.localStorage = {
          getItem: (key) => storageValues.get(key) ?? null,
          setItem: (key, value) => {
            storageValues.set(key, String(value));
            storageWrites.push({ key, value: String(value) });
          },
          removeItem: (key) => storageValues.delete(key),
        };
        const apiCalls = [];
        const toasts = [];
        let apiImpl = async () => ({
          colors: {
            system: { vitality: "#3F6FAF", calm: "#355F63" },
            personal: { vitality: null, calm: null },
            effective: { vitality: "#3F6FAF", calm: "#355F63" },
          },
          capabilities: { can_manage_system_theme_colors: true },
        });
        async function api(path, options = {}) {
          apiCalls.push({ path, options });
          return apiImpl(path, options);
        }
        function showToast(message, type = "info") {
          toasts.push({ message, type });
        }
        function deferred() {
          let resolve;
          let reject;
          const promise = new Promise((res, rej) => {
            resolve = res;
            reject = rej;
          });
          return { promise, resolve, reject };
        }
        async function flushMicrotasks() {
          for (let index = 0; index < 6; index += 1) await Promise.resolve();
        }

        __THEME_RUNTIME_BLOCK__
        __SYSTEM_THEME_COLOR_BLOCK__

        const themeHarness = {
          state: systemThemeColorState,
          load: loadSystemThemeColors,
          save: saveSystemThemeColors,
          reset: resetSystemThemeColorDraft,
          discard: discardUnsavedSystemThemeColors,
          resetAuth: resetSystemThemeColorsForAuthChange,
          capture: captureSystemThemeColors,
          restore: restoreSystemThemeColors,
          elements,
          cssVariables,
          storageValues,
          storageWrites,
          apiCalls,
          toasts,
        };

        (async () => {
        __SCENARIO__
        })().catch((error) => {
          console.error(error.stack || error);
          process.exitCode = 1;
        });
        """
    ).replace("__THEME_RUNTIME_BLOCK__", runtime_block.group(0)).replace(
        "__SYSTEM_THEME_COLOR_BLOCK__",
        state_block.group(0),
    ).replace(
        "__SCENARIO__",
        textwrap.indent(textwrap.dedent(scenario_source).strip(), "  "),
    )
    script_path = tmp_path / "system_theme_color_state_machine.cjs"
    script_path.write_text(script, encoding="utf-8")
    subprocess.run(["node", str(script_path)], check=True, cwd=ROOT)


def test_configurable_theme_gradient_contract_is_absent_and_fixed_logo_gradient_is_present():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)
    frontend = "\n".join((html, app_js, css))

    for canceled_token in (
        "interfaceThemeGradientToggle",
        "themeGradientEnabled",
        "theme_gradient_enabled",
        "data-theme-gradient",
        "themeGradient",
        "autoCheckLastInterfaceThemeGradient",
    ):
        assert canceled_token not in frontend
    for fixed_token in (
        "#3466D9",
        "#6AA4FF",
        "--theme-accent-gradient",
    ):
        assert fixed_token in frontend


def test_fixed_logo_theme_palette_runtime_normalizes_and_guarantees_contrast(tmp_path):
    _run_theme_color_node_scenario(
        tmp_path,
        r"""
        assert.equal(normalizeThemeHex("#3466d9"), "#3466D9");
        assert.equal(normalizeThemeHex("#fff"), null);
        assert.equal(normalizeThemeHex("rgb(63, 111, 175)"), null);

        const lightPalette = deriveThemePalette("#3466D9", "light");
        assert.equal(lightPalette.accent, "#3466D9");
        assert.equal(lightPalette.gradientEnd, "#6AA4FF");
        assert.ok(["#000000", "#FFFFFF"].includes(lightPalette.onAccent));
        assert.ok(
          contrastRatio(lightPalette.accent, lightPalette.onAccent)
            >= contrastRatio(lightPalette.accent, lightPalette.onAccent === "#000000" ? "#FFFFFF" : "#000000")
        );
        assert.ok(contrastRatio(lightPalette.readableAccent, "#F7FAFC") >= 4.5);
        assert.match(lightPalette.focusRing, /^rgba\(52, 102, 217, 0\.\d+\)$/);

        const darkPalette = deriveThemePalette("#3466D9", "dark");
        assert.equal(darkPalette.accent, "#3466D9");
        assert.ok(contrastRatio(darkPalette.readableAccent, "#121318") >= 4.5);

        const applied = applyEffectiveThemeColors({ vitality: "#FFFFFF", calm: "#000000" });
        assert.deepEqual(applied.colors, {
          vitality: "#3466D9",
          calm: "#355F63",
        });
        assert.equal(cssVariables.get("--theme-accent"), "#3466D9");
        assert.equal(cssVariables.get("--theme-accent-gradient-end"), "#6AA4FF");
        assert.equal(cssVariables.get("--theme-on-accent"), applied.palette.onAccent);
        assert.equal(cssVariables.get("--theme-accent-readable"), applied.palette.readableAccent);
        assert.equal(cssVariables.get("--theme-focus-ring"), applied.palette.focusRing);

        attributes.set("data-theme", "light");
        attributes.set("data-color-mode", "dark");
        const calmApplied = applyEffectiveThemeColors(applied.colors);
        assert.equal(calmApplied.palette.accent, "#3466D9");
        assert.ok(contrastRatio(calmApplied.palette.readableAccent, "#121318") >= 4.5);
        """,
    )


def test_theme_and_dark_mode_helpers_force_the_single_light_theme():
    app_js = _read(APP_JS)
    commit_theme = re.search(
        r"function commitTheme\(theme\) \{(?P<body>.*?)\n\}",
        app_js,
        re.S,
    )
    apply_dark_mode = re.search(
        r"function applyDarkMode\(darkMode\) \{(?P<body>.*?)\n\}",
        app_js,
        re.S,
    )
    assert commit_theme is not None
    assert apply_dark_mode is not None
    assert 'document.documentElement.setAttribute("data-theme", "space-tech");' in commit_theme.group("body")
    assert 'document.documentElement.setAttribute("data-color-mode", "light");' in apply_dark_mode.group("body")
    for body in (commit_theme.group("body"), apply_dark_mode.group("body")):
        assert "applyEffectiveThemeColors(effectiveThemeColors);" in body


def test_fixed_theme_runtime_derives_the_page_background(tmp_path):
    _run_theme_color_node_scenario(
        tmp_path,
        r"""
        const light = applyEffectiveThemeColors({
          vitality: "#3466D9",
          calm: "#355F63",
        });
        assert.match(light.palette.pageBackground, /^#[0-9A-F]{6}$/);
        assert.equal(cssVariables.get("--theme-page-background"), light.palette.pageBackground);
        assert.notEqual(light.palette.pageBackground, light.palette.accent);

        assert.equal(light.colors.vitality, "#3466D9");
        assert.equal(light.palette.gradientEnd, "#6AA4FF");
        """,
    )


def test_theme_emphasis_surfaces_use_fixed_gradient_without_header_leakage():
    css = _read(STYLES_CSS)

    for selector in (
        ".top-nav-item.active,",
        ".top-nav-group.active > .top-nav-group-toggle",
        ".nav-item.active",
        ".btn-primary",
        ".trend-quick-btn.active",
        "#page-local-storage .local-storage-tab.active,",
        "#page-settings .card-icon-blue,",
    ):
        assert selector in css

    solid_contract_start = css.index("/* Solid theme emphasis surfaces: start */")
    solid_contract_end = css.index("/* Solid theme emphasis surfaces: end */", solid_contract_start)
    solid_contract = css[solid_contract_start:solid_contract_end]
    assert "background: var(--theme-accent);" in solid_contract
    assert "color: var(--theme-on-accent);" in solid_contract
    assert "radial-gradient" not in solid_contract
    for protected_selector in (" th,", " th {", "\nth {", "thead", ".table-header", ".app-modal-shell"):
        assert protected_selector not in solid_contract


def test_theme_forms_and_calendar_use_solid_tokens_and_neutral_surfaces():
    css = _read(STYLES_CSS)
    contract_start = css.index("/* Solid theme form and calendar controls: start */")
    contract_end = css.index("/* Solid theme form and calendar controls: end */", contract_start)
    contract = css[contract_start:contract_end]

    for selector in (
        ".main-content input:not([type=\"checkbox\"])",
        ".app-modal-shell input:not([type=\"checkbox\"])",
        ".custom-input-shell:focus-within input.custom-input-native",
        ".custom-select-trigger::after",
        ".custom-select-option.active",
        ".custom-date-shell::after",
        ".custom-date-head strong",
        ".custom-date-day.active",
        ".custom-date-actions button",
    ):
        assert selector in contract

    for token in (
        "caret-color: var(--theme-accent-readable)",
        "border-color: var(--theme-accent-readable)",
        "box-shadow: 0 0 0 3px var(--theme-focus-ring)",
        "background: var(--surface-container-lowest)",
        "background: var(--theme-accent)",
        "color: var(--theme-on-accent)",
    ):
        assert token in contract
    assert "linear-gradient" not in contract
    assert "radial-gradient" not in contract
    for protected_selector in (" th,", " th {", "\nth {", "thead", ".table-header"):
        assert protected_selector not in contract


def test_theme_surface_contract_does_not_change_protected_header_or_modal_rules():
    css = _read(STYLES_CSS)
    theme_contract = css[
        css.index("/* Solid theme emphasis surfaces: start */"):
        css.index("/* Solid theme form and calendar controls: end */")
    ]
    for protected_surface in (
        ".app-modal-shell > .app-modal-header",
        ".app-modal-shell > .app-modal-body",
        ".app-modal-shell > .app-modal-footer",
        ".modal-info",
        ".pbc-modal",
        ".user-modal",
    ):
        assert protected_surface not in theme_contract
    assert " th," not in theme_contract
    assert " th {" not in theme_contract
    assert "\nth {" not in theme_contract
    assert "thead" not in theme_contract


def test_removed_system_theme_color_controls_stay_absent_from_settings_runtime():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    frontend = html + "\n" + app_js
    for removed_token in (
        "systemThemeColorsSection",
        "systemVitalityThemeColor",
        "systemCalmThemeColor",
        "saveSystemThemeColorsBtn",
        "resetSystemThemeColorsBtn",
        "loadSystemThemeColors",
        "saveSystemThemeColors",
        "can_manage_system_theme_colors",
        "System theme colors start",
    ):
        assert removed_token not in frontend
    assert 'vitality: "#3466D9"' in app_js
    assert 'gradientEnd: accent === "#3466D9" ? "#6AA4FF"' in app_js


def _obsolete_system_theme_color_valid_preview_normalization_and_atomic_save(tmp_path):
    _run_system_theme_color_node_scenario(
        tmp_path,
        r"""
        const h = themeHarness;
        assert.equal(h.elements.systemThemeColorsSection.hidden, true);
        assert.equal(await h.load({ silent: true }), true);
        assert.equal(h.elements.systemThemeColorsSection.hidden, false);
        assert.deepEqual(h.state.savedColors, {
          vitality: "#3F6FAF",
          calm: "#355F63",
        });
        assert.deepEqual(h.state.draftColors, h.state.savedColors);
        assert.deepEqual(h.state.lastValidDraft, h.state.savedColors);
        assert.equal(h.cssVariables.get("--theme-accent"), "#3F6FAF");

        h.elements.systemVitalityThemeColor.value = "#abcdef";
        h.elements.systemVitalityThemeColor.dispatch("input");
        assert.equal(h.state.rawInputs.vitality, "#abcdef");
        assert.equal(h.state.draftColors.vitality, "#ABCDEF");
        assert.equal(h.state.lastValidDraft.vitality, "#ABCDEF");
        assert.equal(h.cssVariables.get("--theme-accent"), "#ABCDEF");
        assert.equal(h.elements.systemVitalityThemeColorError.hidden, true);
        h.elements.systemVitalityThemeColor.dispatch("blur");
        assert.equal(h.elements.systemVitalityThemeColor.value, "#ABCDEF");

        h.elements.systemCalmThemeColor.value = "#102030";
        h.elements.systemCalmThemeColor.dispatch("input");
        assert.equal(h.state.dirty, true);
        assert.equal(h.elements.saveSystemThemeColorsBtn.disabled, false);
        let postBody = null;
        apiImpl = async (path, options) => {
          assert.equal(path, "/api/settings/interface/theme-colors");
          assert.equal(options.method, "POST");
          postBody = JSON.parse(options.body);
          return {
            colors: {
              system: { vitality: "#ABCDEF", calm: "#102030" },
              personal: { vitality: null, calm: null },
              effective: { vitality: "#ABCDEF", calm: "#102030" },
            },
            capabilities: { can_manage_system_theme_colors: true },
          };
        };
        assert.equal(await h.save(), true);
        assert.deepEqual(postBody, {
          vitality_theme_color: "#ABCDEF",
          calm_theme_color: "#102030",
        });
        assert.deepEqual(h.state.savedColors, postBody && {
          vitality: postBody.vitality_theme_color,
          calm: postBody.calm_theme_color,
        });
        assert.equal(h.state.dirty, false);
        assert.equal(h.elements.systemThemeColorsStatus.textContent, "保存成功");
        assert.deepEqual(
          JSON.parse(h.storageValues.get("autoCheckLastEffectiveThemeColors")),
          { vitality: "#ABCDEF", calm: "#102030" },
        );
        """,
    )


def _obsolete_system_theme_color_invalid_input_reset_failure_and_leave_discard(tmp_path):
    _run_system_theme_color_node_scenario(
        tmp_path,
        r"""
        const h = themeHarness;
        apiImpl = async () => ({
          colors: {
            system: { vitality: "#112233", calm: "#445566" },
            personal: { vitality: null, calm: null },
            effective: { vitality: "#112233", calm: "#445566" },
          },
          capabilities: { can_manage_system_theme_colors: true },
        });
        assert.equal(await h.load({ silent: true }), true);
        const appliedBeforeInvalid = h.cssVariables.get("--theme-accent");

        h.elements.systemVitalityThemeColor.value = "#123";
        h.elements.systemVitalityThemeColor.dispatch("input");
        assert.equal(h.state.rawInputs.vitality, "#123");
        assert.equal(h.state.lastValidDraft.vitality, "#112233");
        assert.equal(h.cssVariables.get("--theme-accent"), appliedBeforeInvalid);
        assert.equal(h.elements.systemVitalityThemeColorError.hidden, false);
        assert.equal(h.elements.saveSystemThemeColorsBtn.disabled, true);
        assert.equal(await h.save(), false);

        const callsBeforeReset = h.apiCalls.length;
        h.reset();
        assert.deepEqual(h.state.draftColors, {
          vitality: "#3F6FAF",
          calm: "#355F63",
        });
        assert.equal(h.apiCalls.length, callsBeforeReset);
        assert.equal(h.state.dirty, true);
        assert.equal(h.cssVariables.get("--theme-accent"), "#3F6FAF");

        apiImpl = async () => { throw new Error("save failed"); };
        assert.equal(await h.save(), false);
        assert.deepEqual(h.state.savedColors, {
          vitality: "#112233",
          calm: "#445566",
        });
        assert.deepEqual(h.state.draftColors, {
          vitality: "#3F6FAF",
          calm: "#355F63",
        });
        assert.equal(h.cssVariables.get("--theme-accent"), "#3F6FAF");
        assert.equal(h.state.dirty, true);
        assert.equal(h.elements.systemThemeColorsStatus.textContent, "保存失败");
        assert.equal(h.toasts.length, 1);

        assert.equal(h.discard(), true);
        assert.deepEqual(h.state.draftColors, h.state.savedColors);
        assert.equal(h.cssVariables.get("--theme-accent"), "#112233");
        assert.equal(h.state.dirty, false);
        """,
    )


def _obsolete_system_theme_color_state_rejects_stale_get_post_and_auth_revisions(tmp_path):
    _run_system_theme_color_node_scenario(
        tmp_path,
        r"""
        const h = themeHarness;
        const oldGet = deferred();
        apiImpl = async () => oldGet.promise;
        const oldLoading = h.load({ silent: true });
        await flushMicrotasks();
        h.resetAuth();
        oldGet.resolve({
          colors: {
            system: { vitality: "#AAAAAA", calm: "#BBBBBB" },
            personal: { vitality: null, calm: null },
            effective: { vitality: "#AAAAAA", calm: "#BBBBBB" },
          },
          capabilities: { can_manage_system_theme_colors: true },
        });
        assert.equal(await oldLoading, false);
        assert.deepEqual(h.state.savedColors, {
          vitality: "#3F6FAF",
          calm: "#355F63",
        });

        apiImpl = async () => ({
          colors: {
            system: { vitality: "#102030", calm: "#405060" },
            personal: { vitality: null, calm: null },
            effective: { vitality: "#102030", calm: "#405060" },
          },
          capabilities: { can_manage_system_theme_colors: true },
        });
        assert.equal(await h.load({ silent: true }), true);

        const getDuringMutation = deferred();
        const currentPost = deferred();
        apiImpl = async (_path, options) => (
          options.method === "POST" ? currentPost.promise : getDuringMutation.promise
        );
        const loading = h.load({ silent: true });
        await flushMicrotasks();
        h.elements.systemVitalityThemeColor.value = "#708090";
        h.elements.systemVitalityThemeColor.dispatch("input");
        const saving = h.save();
        await flushMicrotasks();
        currentPost.resolve({
          colors: {
            system: { vitality: "#708090", calm: "#405060" },
            personal: { vitality: null, calm: null },
            effective: { vitality: "#708090", calm: "#405060" },
          },
          capabilities: { can_manage_system_theme_colors: true },
        });
        assert.equal(await saving, true);
        getDuringMutation.resolve({
          colors: {
            system: { vitality: "#111111", calm: "#222222" },
            personal: { vitality: null, calm: null },
            effective: { vitality: "#111111", calm: "#222222" },
          },
          capabilities: { can_manage_system_theme_colors: true },
        });
        assert.equal(await loading, false);
        assert.equal(h.state.savedColors.vitality, "#708090");

        h.elements.systemCalmThemeColor.value = "#A0B0C0";
        h.elements.systemCalmThemeColor.dispatch("input");
        const stalePost = deferred();
        apiImpl = async () => stalePost.promise;
        const staleSaving = h.save();
        await flushMicrotasks();
        h.resetAuth();
        stalePost.resolve({
          colors: {
            system: { vitality: "#708090", calm: "#A0B0C0" },
            personal: { vitality: null, calm: null },
            effective: { vitality: "#708090", calm: "#A0B0C0" },
          },
          capabilities: { can_manage_system_theme_colors: true },
        });
        assert.equal(await staleSaving, false);
        assert.deepEqual(h.state.savedColors, {
          vitality: "#3F6FAF",
          calm: "#355F63",
        });
        assert.equal(h.state.saving, false);
        assert.equal(h.elements.systemThemeColorsSection.hidden, true);
        """,
    )


def _obsolete_system_theme_colors_integrate_with_settings_navigation_and_auth_boundaries():
    app_js = _read(APP_JS)
    ensure_authenticated = re.search(
        r"async function ensureAuthenticated\(\) \{(?P<body>.*?)\n\}",
        app_js,
        re.S,
    )
    switch_page = re.search(
        r"async function switchPage\(name, options = \{\}\) \{(?P<body>.*?)\n\}",
        app_js,
        re.S,
    )
    logout = re.search(
        r"async function logout\(\) \{(?P<body>.*?)\n\}",
        app_js,
        re.S,
    )
    assert ensure_authenticated is not None
    assert switch_page is not None
    assert logout is not None
    assert "resetSystemThemeColorsForAuthChange();" in ensure_authenticated.group("body")
    assert "loadSystemThemeColors({ silent: true })" in ensure_authenticated.group("body")
    assert "discardUnsavedSystemThemeColors();" in switch_page.group("body")
    assert "captureSystemThemeColors();" in logout.group("body")
    assert "resetSystemThemeColorsForAuthChange();" in logout.group("body")
    assert "restoreSystemThemeColors(" in logout.group("body")
    assert 'loadPageSection("全局主题色", () => loadSystemThemeColors({ silent: false }))' in app_js


def test_interface_preferences_expose_complete_accessible_wysiwyg_controls_and_state_contract(tmp_path):
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert (
        '<div id="interfaceLineChartStyle" role="radiogroup" aria-label="折线图风格">'
        in html
    )
    assert re.search(
        r'<input id="interfaceLineChartStyleStraight" type="radio" '
        r'name="interfaceLineChartStyle" value="straight" checked>',
        html,
    )
    assert re.search(
        r'<input id="interfaceLineChartStyleSmooth" type="radio" '
        r'name="interfaceLineChartStyle" value="smooth">',
        html,
    )
    assert re.findall(
        r'<input[^>]+name="interfaceLineChartStyle"[^>]+value="([^"]+)"',
        html,
    ) == ["straight", "smooth"]
    assert html.index('value="straight" checked') < html.index('value="smooth"')
    assert 'type="color"' not in html
    assert "HEX" not in html
    assert "tension" not in html.lower()

    for selector in (
        "#page-settings #interfaceLineChartStyle",
        "#page-settings #interfaceLineChartStyle label:has(input:focus-visible)",
        "#page-settings #interfaceLineChartStyle label:has(input:disabled)",
    ):
        assert selector in css

    assert "统一控制系统折线图的数据点连接方式" not in html

    assert "const DEFAULT_INTERFACE_PREFERENCES = Object.freeze({" in app_js
    for field in (
        "radiusPx: 4",
        'lineChartStyle: "straight"',
        "savedPreferences:",
        "draftPreferences:",
        "function readInterfacePreferencesPayload(payload)",
        "function applyInterfacePreferences(preferences)",
        "function cacheAuthenticatedInterfacePreferences(preferences)",
    ):
        assert field in app_js
    assert "Object.defineProperties(interfaceRadiusState" not in app_js

    _run_interface_radius_node_scenario(
        tmp_path,
        """
        const h = radiusHarness;
        assert.deepEqual(h.state.savedPreferences, {
          radiusPx: 4,
          lineChartStyle: "straight",
        });
        assert.deepEqual(h.state.draftPreferences, h.state.savedPreferences);

        h.elements.interfaceLineChartStyleSmooth.checked = true;
        h.elements.interfaceLineChartStyleSmooth.dispatch("input");
        assert.equal(h.state.draftPreferences.lineChartStyle, "smooth");
        assert.equal(h.state.statusText, "正在预览，尚未保存");

        let postBody = null;
        apiImpl = async (_path, options) => {
          postBody = JSON.parse(options.body);
          return { settings: {
            radius_px: 4,
            line_chart_style: "smooth",
          } };
        };
        assert.equal(await h.save(), true);
        assert.deepEqual(postBody, {
          radius_px: 4,
          line_chart_style: "smooth",
        });
        assert.deepEqual(h.state.savedPreferences, h.state.draftPreferences);

        h.elements.interfaceLineChartStyleStraight.checked = true;
        h.elements.interfaceLineChartStyleStraight.dispatch("input");
        h.elements.interfaceRadiusSlider.value = "8";
        h.elements.interfaceRadiusSlider.dispatch("input");
        apiImpl = async () => ({ settings: {} });
        assert.equal(await h.save(), false);
        assert.deepEqual(h.state.savedPreferences, {
          radiusPx: 4,
          lineChartStyle: "smooth",
        });
        assert.deepEqual(h.state.draftPreferences, {
          radiusPx: 8,
          lineChartStyle: "straight",
        });
        assert.equal(h.state.statusText, "保存失败");

        assert.equal(h.discard(), true);
        assert.deepEqual(h.state.draftPreferences, h.state.savedPreferences);
        """,
    )


def test_interface_preferences_strict_payloads_and_success_only_display_cache(tmp_path):
    _run_interface_radius_node_scenario(
        tmp_path,
        """
        const h = radiusHarness;
        const radiusCacheKey = "autoCheckLastInterfaceRadius";
        h.useStrictResponses();

        apiImpl = async () => ({ settings: {
          radius_px: 7,
          line_chart_style: "smooth",
        } });
        assert.equal(await h.load({ silent: true }), true);
        assert.deepEqual(h.state.savedPreferences, {
          radiusPx: 7,
          lineChartStyle: "smooth",
        });
        assert.equal(h.storageValues.get(radiusCacheKey), "7");
        assert.equal(h.storageWrites.length, 1);

        const invalidPayloads = [
          { radius_px: 7, line_chart_style: "curve" },
          { radius_px: 7 },
          { radius_px: 0, line_chart_style: "smooth" },
        ];
        for (const settings of invalidPayloads) {
          h.resetAuth();
          apiImpl = async () => ({ settings });
          assert.equal(await h.load({ silent: true }), false);
          assert.deepEqual(h.state.savedPreferences, {
            radiusPx: 4,
            lineChartStyle: "straight",
          });
          assert.equal(h.storageValues.get(radiusCacheKey), "7");
          assert.equal(h.storageWrites.length, 1);
        }

        apiImpl = async () => { throw new Error("load failed"); };
        assert.equal(await h.load({ silent: true }), false);
        assert.equal(h.storageWrites.length, 1);

        h.elements.interfaceLineChartStyleSmooth.checked = true;
        h.elements.interfaceLineChartStyleSmooth.dispatch("input");
        apiImpl = async () => ({ settings: {
          radius_px: 4,
          line_chart_style: "curve",
        } });
        assert.equal(await h.save(), false);
        assert.equal(h.state.savedPreferences.lineChartStyle, "straight");
        assert.equal(h.state.draftPreferences.lineChartStyle, "smooth");
        assert.equal(h.storageWrites.length, 1);

        fetchImpl = async () => ({
          ok: false,
          json: async () => ({ authenticated: false }),
        });
        await assert.rejects(() => h.ensureAuthenticated(), /login required/);
        assert.equal(h.storageWrites.length, 1);
        """,
    )


def test_reason_filter_contains_all_current_reasons():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)

    assert '<option value="">全部差异类型</option>' in html
    assert "全部预估差异原因" not in html

    for reason in [
        "资产缺失",
        "资产重复",
        "资产差异",
        "负债及权益科目差异",
        "负债及权益科目缺失",
        "负债及权益科目重复",
        "实收本金差异",
        "实收本金缺失",
        "实收本金重复",
        "暂无法确定",
    ]:
        assert f'<option value="{reason}">{reason}</option>' in html

    for detail_reason in [
        "AM标的缺失",
        "FA与AM标的不一致",
        "合同投融资余额为0但FA科目余额不为0",
        "实收信托有误",
    ]:
        assert f'<option value="{detail_reason}">{detail_reason}</option>' not in html

    assert "function differenceReasonMatchesFilter(differenceReason, selectedReason)" in app_js
    assert ".split(/\\s*\\+\\s*/)" in app_js
    assert "function resultMatchesReasonFilter(item, selectedReason)" in app_js
    assert "differenceReasonMatchesFilter(item.difference_reason, selectedReason)" in app_js
    assert "resultMatchesReasonFilter(item, reason)" in app_js
    assert "item.difference_reason === reason" not in app_js


def test_result_list_and_export_do_not_show_difference_direction():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)

    result_table = re.search(r'<table class="result-table">(?P<table>.*?)</table>', html, re.S)
    assert result_table is not None
    assert "差异类型" in result_table.group("table")
    assert "预估差异原因" not in result_table.group("table")
    assert "差异方向" not in result_table.group("table")
    assert "col-direction" not in result_table.group("table")

    render_results = re.search(r"function renderResults\(\) \{(?P<body>.*?)function renderDetails", app_js, re.S)
    assert render_results is not None
    assert "item.direction" not in render_results.group("body")

    export_to_excel = re.search(r"function exportToExcel\(\) \{(?P<body>.*?)exportBtn.addEventListener", app_js, re.S)
    assert export_to_excel is not None
    assert "差异方向" not in export_to_excel.group("body")
    assert "item.direction" not in export_to_excel.group("body")


def test_result_list_and_export_show_valuation_asset_total_before_asset_total():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)

    result_table = re.search(r'<table class="result-table">(?P<table>.*?)</table>', html, re.S)
    assert result_table is not None
    table = result_table.group("table")
    assert table.index("资产合计（估值表）") < table.index("资产合计 (元)")
    assert '<tr><td colspan="9" class="empty">暂无结果</td></tr>' in table

    render_results = re.search(r"function renderResults\(\) \{(?P<body>.*?)function renderDetails", app_js, re.S)
    assert render_results is not None
    render_body = render_results.group("body")
    assert "item.valuation_asset_total" in render_body
    assert render_body.index("item.valuation_asset_total") < render_body.index("item.asset_total")
    assert 'colspan="9"' in render_body

    export_to_excel = re.search(r"function exportToExcel\(\) \{(?P<body>.*?)exportBtn.addEventListener", app_js, re.S)
    assert export_to_excel is not None
    export_body = export_to_excel.group("body")
    assert app_js.index('header: "资产合计（估值表）"') < app_js.index('header: "资产合计"')
    assert "item.valuation_asset_total" in export_body
    assert export_body.index("item.valuation_asset_total") < export_body.index("item.asset_total")


def test_auto_check_no_source_data_empty_state_and_date_box_layout():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert "run-bar-icon" not in html
    assert ".run-bar-icon" not in css
    assert "RESULT_EMPTY_SOURCE" in app_js
    assert "resultEmptyState = noSourceData ? RESULT_EMPTY_SOURCE : \"\";" in app_js
    assert "Boolean(h.no_source_data)" in app_js
    assert "let hideLastRunTimeForNoSourceData = false;" in app_js
    assert "hideLastRunTimeForNoSourceData = true;" in app_js
    assert "if (lastRunTime) lastRunTime.hidden = true;" in app_js
    assert "hideLastRunTimeForNoSourceData = false;" in app_js
    assert "lastRunTime.hidden = Boolean(resultRestoreHistoryMeta) || hideLastRunTimeForNoSourceData;" in app_js
    assert "lastRunTime.textContent && !resultRestoreHistoryMeta && !hideLastRunTimeForNoSourceData" in app_js
    assert "renderSourceNoDataRow" in app_js
    assert "报表对应日期无数据" in app_js
    assert "appendRunLog(noDataMessage)" not in app_js
    assert ".no-source-panel" in css
    assert "@keyframes noSourceScan" in css


def test_cards_hover_glow_tracks_theme_palette_instead_of_dark_shadow():
    css = _read(STYLES_CSS)
    readme = _read(README_MD)

    root_rule = re.search(r"(?m)^:root\s*\{(?P<body>.*?)\}", css, re.S)
    space_theme_rule = re.search(r'(?m)^\[data-theme="space-tech"\]\s*\{(?P<body>.*?)\}', css, re.S)
    assert root_rule is not None
    assert space_theme_rule is not None
    assert "--card-hover-glow: #166534;" in root_rule.group("body")
    assert "--card-hover-shadow: rgba(22, 101, 52" in root_rule.group("body")
    assert "--card-hover-glow: #38bdf8;" in space_theme_rule.group("body")
    assert "--card-hover-shadow: rgba(14, 116, 144" in space_theme_rule.group("body")

    card_hover = re.search(r"(?m)^\.card:hover\s*\{(?P<body>.*?)\}", css, re.S)
    assert card_hover is not None
    card_hover_body = card_hover.group("body")
    assert "border-color: color-mix(in srgb, var(--card-hover-glow)" in card_hover_body
    assert "0 0 0 1px color-mix(in srgb, var(--card-hover-glow)" in card_hover_body
    assert "0 0 24px color-mix(in srgb, var(--card-hover-glow)" in card_hover_body
    assert "var(--card-hover-shadow" in card_hover_body
    assert "#38bdf8" not in card_hover_body
    assert "rgba(0, 0, 0" not in card_hover_body

    settings_card_hover = re.search(r"(?m)^#page-settings \.card:hover\s*\{(?P<body>.*?)\}", css, re.S)
    assert settings_card_hover is not None
    settings_hover_body = settings_card_hover.group("body")
    assert "border-color: color-mix(in srgb, var(--card-hover-glow)" in settings_hover_body
    assert "0 0 24px color-mix(in srgb, var(--card-hover-glow)" in settings_hover_body
    assert "var(--card-hover-shadow-soft)" in settings_hover_body
    assert "transform: none;" in settings_hover_body
    assert "#38bdf8" not in settings_hover_body
    assert "rgba(0, 0, 0" not in settings_hover_body

    for selector in [
        r'\[data-theme="space-tech"\] \.home-stat-card:hover',
        r'\[data-theme="space-tech"\]\[data-color-mode="dark"\] \.home-stat-card:hover',
        r"\.glass-card:hover",
        r'\[data-color-mode="dark"\] \.glass-card:hover',
        r"\.home-stat-card:hover",
        r"\.home-analysis-card:hover",
        r"\.glass-stat-card:hover",
        r"\.tool-card:hover",
        r"\.data-manage-item:hover",
        r"#page-settings \.settings-dashboard-card:hover",
        r'\[data-theme="space-tech"\] #page-settings \.settings-dashboard-card:hover',
        r'\[data-theme="space-tech"\]\[data-color-mode="dark"\] #page-settings \.settings-dashboard-card:hover',
    ]:
        rule = re.search(rf"(?m)^{selector}\s*\{{(?P<body>.*?)\}}", css, re.S)
        assert rule is not None
        rule_body = rule.group("body")
        assert "var(--card-hover-glow)" in rule_body
        assert "#38bdf8" not in rule_body
        assert "rgba(0, 0, 0" not in rule_body
    assert "悬停光晕" not in readme


def test_export_to_excel_includes_processing_script_column_after_detail():
    app_js = _read(APP_JS)

    export_to_excel = re.search(r"function exportToExcel\(\) \{(?P<body>.*?)exportBtn.addEventListener", app_js, re.S)
    assert export_to_excel is not None
    export_body = export_to_excel.group("body")
    assert 'header: "处理脚本"' in app_js
    assert app_js.index('header: "差异原因详情"') < app_js.index('header: "处理脚本"')
    assert "buildProcessingScriptText(item)" in export_body
    assert "function buildProcessingScriptText(item)" in app_js
    assert "window.buildProcessingScript(item)" in app_js
    assert "escapeExcelSingleLineText(buildProcessingScriptText(item))" in export_body


def test_export_to_excel_includes_remark_for_combined_difference_reason():
    app_js = _read(APP_JS)

    export_to_excel = re.search(r"function exportToExcel\(\) \{(?P<body>.*?)exportBtn.addEventListener", app_js, re.S)
    assert export_to_excel is not None
    export_body = export_to_excel.group("body")
    assert '{ header: "备注", width: 220, style: "SpecificReason", type: "string" }' in app_js
    assert "function remarkText(item)" in app_js
    assert "资产端存在多组候选资产均可解释差额" in app_js
    assert "负债及权益端存在多组候选科目均可解释差额" in app_js
    assert "资产端差额已由" in app_js
    assert "资产端差额可归入" in app_js
    assert "解释实收本金部分，剩余部分由" in app_js
    assert 'remarkText(item) || "无"' in export_body
    assert app_js.index('header: "差异类型"') < app_js.index('header: "具体原因"')
    assert app_js.index('header: "处理脚本"') < app_js.index('header: "备注"')


def test_export_button_shows_progress_and_failure_feedback():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert 'data-export-label>导出</span>' in html
    assert 'class="export-progress" id="exportProgress" hidden' in html
    assert 'id="exportProgressText">准备导出...' in html
    assert "const exportBtnLabel = exportBtn?.querySelector(\"[data-export-label]\");" in app_js
    assert "function setExportState(exporting, message = \"\")" in app_js
    assert "function updateExportProgress(message)" in app_js
    assert "async function exportToExcel()" in app_js
    assert "await waitForExportUiFrame();" in app_js
    assert "showToast(`已导出 ${data.length} 条结果`, \"success\");" in app_js
    assert "showToast(`导出失败：${message}`, \"error\");" in app_js
    assert 'showToast("无数据可导出", "warning");' in app_js
    assert ".export-progress" in css
    assert ".export-progress-spinner" in css


def test_export_to_excel_groups_rows_by_difference_reason_stably():
    app_js = _read(APP_JS)

    assert "function exportRowsForExcel(data)" in app_js
    export_rows = re.search(r"function exportRowsForExcel\(data\) \{(?P<body>.*?)\n\}", app_js, re.S)
    assert export_rows is not None
    export_body = export_rows.group("body")
    assert "reasonOrder" in export_body
    assert "difference_reason" in export_body
    assert "originalIndex" in export_body
    assert "reasonDelta || left.originalIndex - right.originalIndex" in export_body

    export_to_excel = re.search(r"function exportToExcel\(\) \{(?P<body>.*?)exportBtn.addEventListener", app_js, re.S)
    assert export_to_excel is not None
    assert "const data = exportRowsForExcel(filteredResults());" in export_to_excel.group("body")


def test_export_to_excel_only_wraps_project_name_and_blocks_overflow():
    app_js = _read(APP_JS)

    assert "const EXPORT_COLUMNS" in app_js
    assert '{ header: "项目名称", width: 180, style: "ProjectName", type: "string" }' in app_js
    assert '{ header: "差异金额", width: 110, style: "Money", type: "number" }' in app_js
    assert '{ header: "差异类型", width: 180, style: "Text", type: "string" }' in app_js
    assert '{ header: "具体原因", width: 180, style: "SpecificReason", type: "string" }' in app_js
    assert '{ header: "差异原因详情", width: 360, style: "Detail", type: "string" }' in app_js
    assert '{ header: "处理脚本", width: 360, style: "Script", type: "string" }' in app_js
    assert '{ header: "备注", width: 220, style: "SpecificReason", type: "string" }' in app_js
    assert "预估差异原因" not in app_js
    assert app_js.index('header: "差异类型"') < app_js.index('header: "具体原因"')
    assert app_js.index('header: "具体原因"') < app_js.index('header: "匹配状态"')
    assert app_js.index('header: "处理脚本"') < app_js.index('header: "备注"')
    assert 'wrapText="1"' in app_js
    assert app_js.count('wrapText="1"') == 1
    assert 'shrinkToFit="1"' not in app_js
    assert 'ht="20" customHeight="1"' not in app_js
    assert "white-space:pre" not in app_js
    assert "function wrapExcelDetailLine(value)" not in app_js
    assert "EXPORT_OVERFLOW_GUARD_CELL" not in app_js

    export_to_excel = re.search(r"function exportToExcel\(\) \{(?P<body>.*?)exportBtn.addEventListener", app_js, re.S)
    assert export_to_excel is not None
    export_body = export_to_excel.group("body")
    assert "buildExcelWorkbookBlob(rows" in export_body
    assert 'type: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"' in app_js
    assert "a.download = `自动对数_${ds}_${ts}.xlsx`;" in export_body


def test_export_to_excel_keeps_detail_cells_single_line():
    app_js = _read(APP_JS)

    export_to_excel = re.search(r"function exportToExcel\(\) \{(?P<body>.*?)exportBtn.addEventListener", app_js, re.S)
    assert export_to_excel is not None
    export_body = export_to_excel.group("body")
    assert "buildDetailText(item)" in export_body
    assert "escapeExcelSingleLineText(buildDetailText(item))" not in export_body
    assert "function escapeExcelDetailText(value)" not in app_js
    assert "const detail = escapeExcelSingleLineText(buildDetailText(item));" not in export_body
    assert "escapeExcelText(buildDetailText(item))" not in export_body


def test_export_to_xlsx_preserves_numeric_and_detail_formatting(tmp_path):
    app_js = _read(APP_JS)
    export_section = app_js[app_js.index("const EXPORT_COLUMNS"):app_js.index("function exportRowsForExcel")]
    output_path = tmp_path / "result.xlsx"
    script_path = tmp_path / "build_export_xlsx.js"
    script_path.write_text(
        export_section
        + textwrap.dedent(
            f"""
            const fs = require("fs");
            globalThis.Blob = require("buffer").Blob;
            (async () => {{
              const rows = [[
                "P001",
                "项目一",
                "1000",
                900,
                "800.5",
                -100,
                "资产缺失 + 负债及权益科目差异",
                "FA与AM标的不一致",
                "已解释",
                "第一行\\n第二行\\n第三行",
                "select 1; update t set a = 1;",
                "资产端差额已由“资产缺失”解释，修正资产端后仍存在剩余差额，剩余部分由“负债及权益科目差异”解释，因此展示为组合差异类型。"
              ]];
              const blob = buildExcelWorkbookBlob(rows, "自动对数结果 — 2026-05-31");
              const buffer = Buffer.from(await blob.arrayBuffer());
              fs.writeFileSync({str(output_path)!r}, buffer);
            }})();
            """
        ),
        encoding="utf-8",
    )

    subprocess.run(["node", str(script_path)], check=True, cwd=ROOT)

    with ZipFile(output_path) as archive:
        assert "xl/worksheets/sheet1.xml" in archive.namelist()
        worksheet_xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
        assert '<row r="2" ht="30" customHeight="1">' in worksheet_xml
        assert '<row r="3"' in worksheet_xml
        assert '<c r="C3" s="5"><v>1000</v></c>' in worksheet_xml

        assert '<c r="F3" s="5"><v>-100</v></c>' in worksheet_xml
        assert "第一行\n第二行\n第三行" in worksheet_xml

    workbook = load_workbook(output_path)
    sheet = workbook.active
    assert sheet.row_dimensions[2].height == 30
    assert sheet["C3"].value == 1000
    assert sheet["C3"].data_type == "n"
    assert sheet["F3"].value == -100
    assert sheet["F3"].data_type == "n"
    assert sheet["G2"].value == "差异类型"
    assert sheet["H2"].value == "具体原因"
    assert sheet["I2"].value == "匹配状态"
    assert sheet["J2"].value == "差异原因详情"
    assert sheet["K2"].value == "处理脚本"
    assert sheet["L2"].value == "备注"
    assert sheet["G3"].value == "资产缺失 + 负债及权益科目差异"
    assert sheet["H3"].value == "FA与AM标的不一致"
    assert sheet["H3"].alignment.horizontal == "left"
    assert sheet["J3"].value == "第一行\n第二行\n第三行"
    assert sheet["J3"].alignment.wrap_text is None
    assert sheet["J3"].alignment.shrink_to_fit is None
    assert sheet["J3"].font.sz == 11
    assert sheet["K3"].value == "select 1; update t set a = 1;"
    assert sheet["K3"].alignment.wrap_text is None
    assert sheet["K3"].alignment.shrink_to_fit is None
    assert sheet["K3"].font.sz == 11
    assert sheet["L3"].value == "资产端差额已由“资产缺失”解释，修正资产端后仍存在剩余差额，剩余部分由“负债及权益科目差异”解释，因此展示为组合差异类型。"
    assert sheet["L3"].alignment.wrap_text is None
    assert sheet["L3"].alignment.shrink_to_fit is None
    assert sheet["L3"].font.sz == 11


def test_home_target_code_mismatch_count_prefers_refinement_rows_without_double_counting(tmp_path):
    app_js = _read(APP_JS)
    normalize_fn = app_js[
        app_js.index("function normalizeHomeReasonText"):
        app_js.index("function homeSpecificReasonMatchesPaidIn")
    ]
    text_match_fn = app_js[
        app_js.index("function homeTargetCodeMismatchTextMatches"):
        app_js.index("function homeTargetCodeMismatchCount")
    ]
    count_fn = app_js[
        app_js.index("function homeTargetCodeMismatchCount"):
        app_js.index("function homeReasonCategoryFromItem")
    ]
    script_path = tmp_path / "home_target_code_count.js"
    script_path.write_text(
        normalize_fn
        + text_match_fn
        + count_fn
        + textwrap.dedent(
            """
            const mixedDetails = {
              details: [
                {
                  kind: "fa_am",
                  data: { specific_reason: "FA与AM标的不一致" },
                },
                {
                  kind: "asset_missing_refinement",
                  data: {
                    rows: [
                      { check_result: "FA和AM标的不一致", pact_id: "PACT_A" },
                      { check_result: "FA和AM标的不一致", pact_id: "PACT_B" },
                    ],
                  },
                },
              ],
            };
            const legacyFaAmOnly = {
              details: [
                {
                  kind: "fa_am",
                  data: { specific_reason: "FA与AM标的不一致" },
                },
              ],
            };
            const counts = [
              homeTargetCodeMismatchCount(mixedDetails),
              homeTargetCodeMismatchCount(legacyFaAmOnly),
            ];
            if (counts[0] !== 2 || counts[1] !== 1) {
              throw new Error(`unexpected counts: ${counts.join(",")}`);
            }
            """
        ),
        encoding="utf-8",
    )

    subprocess.run(["node", str(script_path)], check=True, cwd=ROOT)


def test_candidate_ambiguous_status_is_available_in_result_filters_and_badge():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)

    assert '<option value="候选不唯一">候选不唯一</option>' in html
    assert 'if (s === "候选不唯一") return `<span class="status-badge status-badge--warn">候选不唯一</span>`;' in app_js


def test_export_to_excel_keeps_processing_script_single_line():
    app_js = _read(APP_JS)

    export_to_excel = re.search(r"function exportToExcel\(\) \{(?P<body>.*?)exportBtn.addEventListener", app_js, re.S)
    assert export_to_excel is not None
    export_body = export_to_excel.group("body")
    assert "EXPORT_SCRIPT_MAX_CHARS" not in app_js
    assert "escapeExcelTruncatedSingleLineText" not in app_js
    assert "replace(/\\s*(?:\\r\\n?|\\n)\\s*/g, \" \")" in app_js
    assert ".slice(0, maxLength)" not in app_js
    assert "escapeExcelSingleLineText(buildProcessingScriptText(item))" in export_body
    assert "escapeExcelDetailText(buildProcessingScriptText(item))" not in export_body
    assert "escapeExcelText(buildProcessingScriptText(item))" not in export_body


def test_result_list_shows_loading_animation_when_returning_to_auto_check_page():
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert "let resultListLoadingTimer" in app_js
    assert "function renderResultListLoading()" in app_js
    assert "function showResultListReturnLoading()" in app_js
    loading_fn = re.search(r"function renderResultListLoading\(\) \{(?P<body>.*?)\n\}", app_js, re.S)
    assert loading_fn is not None
    assert "success-panel" not in loading_fn.group("body")
    assert "launchConfetti" not in loading_fn.group("body")
    assert "if (!resultBody || !results.length) return;" in app_js
    assert "renderResultListLoading();" in app_js
    assert "resultListLoadingTimer = setTimeout(() => {" in app_js
    assert "renderResults();" in app_js

    switch_page = re.search(r"function switchPage\(name, options = \{\}\) \{(?P<body>.*?)\n\}", app_js, re.S)
    assert switch_page is not None
    assert 'const previousPage = document.documentElement.getAttribute("data-page") || "";' in switch_page.group("body")
    assert 'if (name === "auto-check" && previousPage !== "auto-check") showResultListReturnLoading();' in switch_page.group("body")

    run_handler = re.search(r"runBtn.addEventListener\(\"click\", async \(\) => \{(?P<body>.*?)\n\}\);", app_js, re.S)
    assert run_handler is not None
    assert "renderResultListLoading();" not in run_handler.group("body")

    assert "openResultDetailRow" not in app_js
    assert "result-detail-loading" not in app_js
    assert 'class="result-loading-row"' in app_js
    assert 'class="loading-spinner result-loading-spinner"' in app_js
    assert "正在加载执行结果列表..." in app_js

    assert ".result-loading-row td" in css
    assert ".result-loading-spinner" in css
    assert "@keyframes resultListLoadingSweep" in css
    assert ".result-detail-loading" not in css
    assert ".detail-row.is-loading" not in css


def test_result_detail_expansion_is_single_open():
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    handler = re.search(r"resultBody\.addEventListener\(\"click\", \(e\) => \{(?P<body>.*?)\n\}\);", app_js, re.S)
    assert handler is not None
    body = handler.group("body")
    assert 'class="result-main-row" data-result-index="${gi}"' in app_js
    assert '<td class="result-project-code">${escapeHtml(item.project_code)}</td>' in app_js
    assert 'style="font-family:Consolas;color:#061623;"' not in app_js
    assert "function hasSelectedResultText()" in app_js
    assert 'const mainRow = e.target.closest(".result-main-row");' in body
    assert "if (!btn && hasSelectedResultText()) return;" in body
    assert 'resultBody.querySelectorAll(".detail-row").forEach((detailRow) => {' in body
    assert "if (detailRow === row) return;" in body
    assert "detailRow.hidden = true;" in body
    assert 'resultBody.querySelectorAll(".result-main-row").forEach((otherRow) => {' in body
    assert 'otherRow.classList.remove("is-expanded");' in body
    assert 'resultBody.querySelectorAll(".expand-btn").forEach((otherBtn) => {' in body
    assert "if (otherBtn === currentBtn) return;" in body
    assert 'otherBtn.textContent = "+";' in body
    assert "row.hidden = !wasHidden;" in body
    assert 'mainRow?.classList.toggle("is-expanded", wasHidden);' in body
    assert ".result-table tbody tr.result-main-row" in css
    assert "cursor: pointer;" in css
    assert ".result-table tbody tr.result-main-row td" in css
    assert "user-select: text;" in css
    assert ".result-project-code" in css
    assert '[data-color-mode="dark"] .result-project-code' in css
    assert '[data-color-mode="dark"] .result-table tbody td' in css
    assert '[data-color-mode="dark"] .result-table tbody tr:hover' in css
    assert ".result-table tbody tr.result-main-row.is-expanded" in css
    assert ".detail-row td {\n  padding: 10px 18px 14px 58px;" in css


def test_result_detail_uses_report_asset_total_label_everywhere():
    app_js = _read(APP_JS)
    export_detail_js = _read(EXPORT_DETAIL_JS)
    server_py = _read(ROOT / "src" / "auto_check" / "app" / "server.py")
    readme = _read(README_MD)

    assert '"label": "资负报表资产合计"' in server_py
    assert 'displayDetailLabel(r.label)' in app_js
    assert 'return label === "zf_detail 资产合计" ? "资负报表资产合计" : label;' in app_js
    assert 'rowValueAny(rows, ["资负报表资产合计", "zf_detail 资产合计"])' in export_detail_js
    assert "资产核对：资负报表资产=" in export_detail_js
    assert "结果列表支持点击项目所在行展开或收回详情" in readme


def test_space_tech_result_detail_title_icon_uses_solid_theme_color():
    css = _read(STYLES_CSS)

    result_icon = re.search(r'(?m)^\[data-theme="space-tech"\] \.result-card \.card-title-icon\s*\{(?P<body>.*?)\}', css, re.S)
    assert result_icon is not None
    body = result_icon.group("body")
    assert "color: var(--theme-accent-readable)" in body
    assert "background: none" in body
    assert "background-clip: text" not in body
    assert "-webkit-text-fill-color: currentColor" in body
    assert "drop-shadow" in body


def test_difference_direction_is_kept_in_detail_payload():
    server_py = _read(SERVER_PY)

    assert '{"label": "差异方向", "value": result.direction}' in server_py


def test_no_runtime_reference_to_removed_balance_checker_directory():
    for path in [
        ROOT / "pyproject.toml",
        ROOT / "src" / "auto_check" / "__main__.py",
        ROOT / "src" / "auto_check" / "app" / "server.py",
    ]:
        assert "balance_checker" not in _read(path)


def test_session_expire_setting_replaces_default_run_date_setting():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)

    assert 'id="sessionExpireHours"' in html
    assert "会话过期时间" in html
    assert "默认运行日期" not in html
    assert 'id="defaultRunDate"' not in html
    assert "sessionExpireHours" in app_js
    assert "session_expire_hours" in app_js
    assert "defaultRunDate" not in app_js
    assert "default_run_date: normalized" not in app_js
    assert 'runDate.value = d.default_run_date || settingsPayload?.api_default_run_date || "";' in app_js


def test_home_auto_refresh_setting_controls_chart_reload():
    app_js = _read(APP_JS)

    assert "function shouldAutoRefreshHome()" in app_js
    assert "function syncDefaultSettingsControls()" in app_js
    assert '["visualEffects", "autoRefreshHome"].forEach((id)' in app_js
    assert 'const refreshData = options.forceHomeRefresh || shouldAutoRefreshHome();' in app_js
    assert 'else if (homeChartsNeedThemeRefresh)' in app_js
    assert 'redrawHomeChartsFromCache();' in app_js
    assert "homeChartsNeedThemeRefresh = false;" in app_js
    assert "renderHomeStats(); renderChart(); renderTrendChart();" in app_js
    assert 'switchPage(savedPage, { forceHomeRefresh: savedPage === "home" })' in app_js
    assert 'switchPage("report-navigation")' in app_js


def test_multilevel_navigation_groups_reconcile_pages_and_renames_labels():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert html.count('data-page="report-navigation"') == 2
    assert html.count('data-nav-group="smart-reconcile"') == 2
    assert html.count('data-nav-group-toggle="smart-reconcile"') == 2
    assert html.count('aria-expanded="false"') >= 2

    for page, label in [
        ("home", "对数总览"),
        ("auto-check", "对数执行"),
        ("history", "对数历史"),
    ]:
        assert html.count(f'data-page="{page}"') == 2
        assert html.count(f">{label}<") >= 1

    assert "智能核数" in html
    assert "const smartReconcilePages" in app_js
    assert "function syncNavGroupState" in app_js
    assert 'item.classList.toggle("active", item.dataset.page === name)' in app_js
    assert '.nav-group[data-nav-group="smart-reconcile"]' in css
    assert ".top-nav-group" in css
    assert ".nav-subitem" in css
    assert ".top-nav-submenu" in css
    assert "max-height: 0;" in css
    assert ".nav-group.open .nav-submenu" in css
    assert "max-height: 140px;" in css


def test_smart_reconcile_parent_uses_theme_specific_toggle_and_hover_behavior():
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)
    readme = _read(README_MD)

    for selector in [
        ".top-nav-group:hover .top-nav-submenu",
        ".top-nav-group:focus-within .top-nav-submenu",
        ".top-nav-group::after",
    ]:
        assert selector in css
    assert ".nav-group:hover .nav-submenu" not in css
    assert ".nav-group:focus-within .nav-submenu" not in css

    toggle_handler = re.search(
        r"navGroupToggles\.forEach\(\(toggle\) => \{(?P<body>.*?)\n\}\);\n\ndocument\.addEventListener\(\"click\"",
        app_js,
        re.S,
    )
    assert toggle_handler is not None
    handler_body = toggle_handler.group("body")
    assert 'event.preventDefault();' in handler_body
    assert 'group.classList.contains("nav-group")' in handler_body
    assert 'setNavGroupOpen(group, !group.classList.contains("open"));' in handler_body
    assert 'group.classList.contains("top-nav-group")' in handler_body
    assert 'switchPage("home");' in handler_body
    assert "event.detail > 0" in handler_body
    assert "toggle.blur();" in handler_body
    assert 'if (group.classList.contains("nav-group")) setNavGroupOpen(group, active);' not in app_js
    assert app_js.count("group.contains(document.activeElement)") >= 2
    assert app_js.count("document.activeElement.blur();") >= 2
    assert "当前唯一启用的浅色主题中，悬浮父菜单显示二级菜单，点击父菜单进入“对数总览”" in readme
    assert "系统优化及BUG修复。" in app_js


def test_top_nav_submenu_pointer_selection_releases_focus_for_mouse_leave_close():
    app_js = _read(APP_JS)

    item_handler = re.search(
        r"\[\.\.\.navItems, \.\.\.topNavItems\]\.forEach\(\(item\) => \{(?P<body>.*?)\n\}\);\n\nnavGroupToggles",
        app_js,
        re.S,
    )
    assert item_handler is not None
    handler_body = item_handler.group("body")
    assert 'item.classList.contains("top-nav-subitem")' in handler_body
    assert "e.detail > 0" in handler_body
    assert "item.blur();" in handler_body


def test_report_navigation_is_default_route_and_preserves_home_dashboard_hash():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert 'id="page-report-navigation"' in html
    assert 'id="page-home"' in html
    assert ':root[data-page="report-navigation"] #page-report-navigation' in css
    assert ':root[data-page="home"] #page-home' in css
    assert 'switchPage("report-navigation")' in app_js
    assert 'name = "report-navigation";' in app_js
    assert 'if (name === "home") {' in app_js
    assert 'const refreshData = options.forceHomeRefresh || shouldAutoRefreshHome();' in app_js


def test_report_navigation_page_uses_monthly_and_period_statistics_scopes():
    html = _read(INDEX_HTML)

    page = re.search(
        r'<section class="page" id="page-report-navigation"[^>]*>(?P<body>.*?)</section>\s*<!-- 首页 -->',
        html,
        re.S,
    )
    assert page is not None
    body = page.group("body")

    assert body.index('class="report-nav-stats-toolbar"') < body.index('id="reportNavMonth"')
    assert 'class="report-nav-overview-group"' in body
    assert '<h3 class="report-nav-overview-title">报送概览</h3>' in body
    overview_label_rule = re.search(
        r"#page-report-navigation \.report-nav-overview-label::before\s*\{(?P<body>[^}]*)\}",
        _read(STYLES_CSS),
    )
    assert overview_label_rule is not None
    assert "width: 4px;" in overview_label_rule.group("body")
    assert "height: 19px;" in overview_label_rule.group("body")
    assert "background-image: var(--theme-accent-gradient);" in overview_label_rule.group("body")
    report_period_rule = re.search(
        r"#page-report-navigation \.report-nav-report-period\s*\{(?P<body>[^}]*)\}",
        _read(STYLES_CSS),
    )
    assert report_period_rule is not None
    assert "margin-left: -2px;" in report_period_rule.group("body")
    assert "padding-left: 0;" in report_period_rule.group("body")
    assert 'class="report-nav-report-month"' in body
    assert 'id="reportNavMonthlyStat"' in body
    assert 'id="reportNavPeriodStats"' in body
    assert 'id="reportNavStats"' in body
    assert '<span class="report-nav-task-label">任务统计</span>' in body
    assert 'id="reportNavPeriodSelect" tabindex="-1"' in body
    for value in ["week", "month", "quarter", "year"]:
        assert f'value="{value}"' in body
        assert f'data-report-nav-period="{value}"' in body
    assert body.index('id="reportNavPeriodSelect"') < body.index('id="reportNavMonthlyStat"')
    assert body.index('id="reportNavMonthlyStat"') < body.index('id="reportNavPeriodStats"')
    assert 'id="reportNavLastRun"' in body
    assert 'id="reportNavRefreshButton"' in body


def test_report_navigation_page_uses_readonly_panorama_details_and_compact_todo_rows():
    html = _read(INDEX_HTML)

    page = re.search(
        r'<section class="page" id="page-report-navigation"[^>]*>(?P<body>.*?)</section>\s*<!-- 首页 -->',
        html,
        re.S,
    )
    assert page is not None
    body = page.group("body")

    for element_id in ["reportNavFishbone", "reportNavFishboneSpine", "reportNavBranches"]:
        assert f'id="{element_id}"' in body
    assert 'id="reportNavProcessDetails" aria-live="polite" hidden' in body
    assert "报送流程进度" in body
    assert 'class="report-nav-flow-legend"' not in body
    assert 'class="report-nav-legend-item' not in body
    assert 'class="report-nav-fishbone"' in body
    assert 'class="report-nav-schedule-layout"' in body
    assert 'id="reportNavTodoTitle"' in body
    assert "我的待办" in body
    assert 'id="reportNavTodoCount"' in body
    assert 'class="report-nav-todo-all"' in body
    assert 'id="reportNavTodoAllBtn"' in body
    assert 'id="reportNavTodoAllModal"' in _read(INDEX_HTML)
    assert 'id="reportNavTodoAllList"' in _read(INDEX_HTML)
    assert 'id="reportNavHistoryModal"' in _read(INDEX_HTML)
    assert 'id="reportNavHistoryList"' in _read(INDEX_HTML)
    assert 'id="reportNavHistoryBtn"' in _read(INDEX_HTML)
    assert ">处理记录</button>" in _read(INDEX_HTML)
    assert 'id="reportNavTodoAllPagination"' in _read(INDEX_HTML)
    assert 'id="reportNavHistoryPagination"' in _read(INDEX_HTML)
    assert '<button type="button" class="report-nav-todo-all" id="reportNavTodoAllBtn"' in body
    assert "注意事项" not in body
    assert 'class="report-nav-filter-chips"' not in body
    for title in ["数据治理流程", "报表特殊治理", "源系统输出确认"]:
        assert title not in body
    assert 'id="reportNavTodoList"' in body or 'class="report-nav-todo-list"' in body
    assert body.count('class="report-nav-todo-primary"') == 0
    assert body.count('class="report-nav-todo-action"') == 0
    assert body.count('class="report-nav-todo-deadline"') == 0
    assert body.count('type="button" class="report-nav-todo-action"') == 0
    assert body.count(">处理</button>") == 0
    assert "（3）" not in body
    todo_all = re.search(
        r"#page-report-navigation \.report-nav-todo-all\s*\{(?P<body>.*?)\}",
        _read(STYLES_CSS),
        re.S,
    )
    assert todo_all is not None
    assert "color: var(--outline);" in todo_all.group("body")
    redesign_css = _read(STYLES_CSS).split(
        "/* ===== Report navigation: restrained read-only panorama ===== */",
        1,
    )[1]
    todo_dot = re.search(
        r"#page-report-navigation \.report-nav-todo > i\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert todo_dot is not None
    assert "background: var(--report-nav-danger);" in todo_dot.group("body")
    todo_mid_dot = re.search(
        r"#page-report-navigation \.report-nav-todo\.mid > i\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert todo_mid_dot is not None
    assert "background: var(--report-nav-warning);" in todo_mid_dot.group("body")
    todo_action = re.search(
        r"#page-report-navigation \.report-nav-todo-action\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert todo_action is not None
    for declaration in [
        "border: 1px solid var(--theme-accent);",
        "border-radius: var(--ui-radius);",
        "color: var(--theme-accent-readable);",
        "background: transparent;",
    ]:
        assert declaration in todo_action.group("body")
    todo_title = re.search(
        r"#page-report-navigation \.report-nav-todo h3\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert todo_title is not None
    assert "font-size: 13px;" in todo_title.group("body")
    assert "font-weight: 600;" in todo_title.group("body")
    assert "立即处理" not in body
    assert "查看</button>" not in body
    assert 'id="reportNavSchedules"' not in body


def test_report_navigation_frontend_preserves_snapshot_period_refresh_and_card_maintenance_logic():
    app_js = _read(APP_JS)

    assert 'api(`/api/report-navigation/dashboard?period=${encodeURIComponent(period)}`)' in app_js
    assert 'async function loadReportNavigation' in app_js
    assert 'function renderReportNavigation' in app_js
    assert 'function renderReportNavigationCards' in app_js
    assert 'function renderReportNavigationProcesses' in app_js
    assert 'function renderReportNavTodos' in app_js
    assert 'renderReportNavTodos(payload.todos || [])' in app_js
    assert "REPORT_NAV_TODO_PREVIEW_LIMIT = 5" in app_js
    assert "REPORT_NAV_TODO_PAGE_SIZE = 10" in app_js
    assert "items.slice(0, REPORT_NAV_TODO_PREVIEW_LIMIT)" in app_js
    assert "openReportNavTodoAllModal" in app_js
    assert "closeReportNavTodoAllModal" in app_js
    assert "openReportNavHistoryModal" in app_js
    assert "loadReportNavHistory" in app_js
    assert "processing-history" in app_js
    assert "发起人：" in app_js
    assert "includeInitiator: true" in app_js
    assert "includeInitiator = false" in app_js
    assert "处理时间：" in app_js
    assert "AutoCheckModuleHost?.openDetailOverlay" in app_js
    assert 'openMode === "detail"' in app_js
    assert "auto-check:report-navigation-refresh" in app_js
    assert "暂无待办" in app_js
    assert "暂无处理记录" in app_js
    assert "buildReportNavTodoHash" in app_js
    assert "handleReportNavTodoAction" in app_js
    assert "AutoCheckModuleHost?.openConfirmOverlay" in app_js
    assert 'openMode === "confirm"' in app_js
    assert "openConfirmOverlay" in _read(MODULE_HOST_JS)
    assert "openDetailOverlay" in _read(MODULE_HOST_JS)
    assert 'if (name === "report-navigation") await loadReportNavigation();' in app_js
    assert 'reportNavPeriodSelect?.addEventListener("change"' in app_js
    assert 'process.process_code === "five_articles"' in app_js
    assert "[1, 4, 7, 10].includes" in app_js
    assert 'reportNavStats?.addEventListener("click"' in app_js
    assert "function openReportNavigationCardMaintenance(cardCode)" in app_js
    assert '/api/report-navigation/cards/${encodeURIComponent(cardCode)}' in app_js
    assert 'id="reportNavCardMaintenanceModal"' in _read(INDEX_HTML)
    assert 'id="reportNavCardMaintenanceSave"' in _read(INDEX_HTML)
    assert "本周" in app_js and "本月" in app_js and "本季度" in app_js and "本年" in app_js


def test_report_nav_todo_initiator_and_processing_history_modal():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert 'id="reportNavHistoryModal"' in html
    assert ">处理记录</button>" in html
    assert "我的处理记录" in html
    assert "REPORT_NAV_TODO_PAGE_SIZE = 10" in app_js
    assert "processing-history" in app_js
    assert "openDetailOverlay" in app_js
    assert "发起人：" in app_js
    assert "includeInitiator: true" in app_js
    assert "includeInitiator = false" in app_js
    assert "处理时间：" in app_js
    assert "发起时间：" in app_js
    assert "openReportNavHistoryModal" in app_js
    assert "buildReportNavHistoryItemHtml" in app_js
    assert re.search(r"#reportNavHistoryModal\s*\{\s*z-index:\s*3100;", css)
    assert re.search(r"#reportNavTodoAllModal\s*\{\s*z-index:\s*3000;", css)
    assert "我的待办发起人与处理记录。" in app_js
    assert "function isReportNavModuleOverlayOpen()" in app_js
    assert "if (isReportNavModuleOverlayOpen()) return;" in app_js
    action_body = app_js.split("async function handleReportNavTodoAction", 1)[1].split(
        "function buildReportNavInitiatorSuffix", 1
    )[0]
    assert 'if (openMode === "detail")' in action_body
    assert 'showToast("详情弹窗打开失败", "error");' in action_body
    detail_block = action_body.split('if (openMode === "detail")', 1)[1].split(
        'if (openMode === "confirm")', 1
    )[0]
    assert "return;" in detail_block
    assert "buildReportNavTodoHash" not in detail_block
    assert "activate(" not in detail_block
    assert 'if (isReportNavHistoryModalOpen() || isReportNavTodoAllModalOpen())' in action_body
    assert 'showToast("确认弹窗打开失败", "error");' in action_body
    drawer_js = _read(
        ROOT
        / "src"
        / "auto_check"
        / "modules"
        / "report_special_processing"
        / "web"
        / "components"
        / "record_drawer.js"
    )
    assert "event.stopPropagation();" in drawer_js


def test_report_nav_processing_history_load_distinguishes_error_from_empty():
    app_js = _read(APP_JS)
    load_history = re.search(
        r"async function loadReportNavHistory\(page = reportNavHistoryPage\) \{(?P<body>.*?)\n\}",
        app_js,
        re.S,
    )
    assert load_history is not None
    body = load_history.group("body")
    assert "REPORT_NAV_HISTORY_EMPTY_MESSAGE" in app_js
    assert "REPORT_NAV_HISTORY_ERROR_MESSAGE" in app_js
    assert "reportNavHistoryCache" in body
    assert "renderReportNavHistoryList(items)" in body
    assert "renderReportNavHistoryList([], { emptyMessage: REPORT_NAV_HISTORY_ERROR_MESSAGE })" in body
    assert "renderReportNavHistoryList(reportNavHistoryCache)" in body
    assert 'showToast(`处理记录加载失败：${error.message}`, "error")' in body
    assert "renderReportNavHistoryList([]);" not in body.split("} catch (error) {", 1)[1].split("} finally {", 1)[0]

    render_history = re.search(
        r"function renderReportNavHistoryList\(items = \[\], options = \{\}\) \{(?P<body>.*?)\n\}",
        app_js,
        re.S,
    )
    assert render_history is not None
    assert "options.emptyMessage || REPORT_NAV_HISTORY_EMPTY_MESSAGE" in render_history.group("body")
    assert "暂无处理记录" in app_js
    assert "处理记录加载失败" in app_js


def test_report_navigation_schedule_save_updates_locally_before_background_refresh():
    app_js = _read(APP_JS)
    body = app_js.split(
        "async function editReportNavigationSchedule(processCode)", 1
    )[1].split("function renderReportNavigationProcesses", 1)[0]

    assert "const result = await api(" in body
    assert "process.report_date = savedDate;" in body
    assert "renderReportNavigationProcesses(payload);" in body
    assert "renderReportNavigationSchedule(payload);" in body
    assert "writeReportNavigationCache(" in body
    assert 'showToast("截止日期已更新", "success");' in body
    assert "void loadReportNavigation();" in body
    assert "await loadReportNavigation();" not in body


def test_report_navigation_processes_sort_by_report_date_ascending_and_stably(tmp_path):
    app_js = _read(APP_JS)
    start_marker = "// Report navigation process ordering start"
    end_marker = "// Report navigation process ordering end"

    assert start_marker in app_js
    assert end_marker in app_js
    ordering_code = app_js.split(start_marker, 1)[1].split(end_marker, 1)[0]
    script = textwrap.dedent(
        f"""
        const assert = require("node:assert/strict");
        {ordering_code}

        const processes = [
          {{ process_code: "late", report_date: "2026-07-20" }},
          {{ process_code: "same-a", report_date: "2026-07-05" }},
          {{ process_code: "missing" }},
          {{ process_code: "early", report_date: "2026-07-01T08:00:00" }},
          {{ process_code: "same-b", report_date: "2026-07-05" }},
          {{ process_code: "invalid-day", report_date: "2026-02-30" }},
          {{ process_code: "invalid-month", report_date: "2026-13-01" }},
          {{ process_code: "invalid", report_date: "not-a-date" }},
        ];
        const originalOrder = processes.map((item) => item.process_code);
        const sorted = sortReportNavigationProcessesByDate(processes);

        assert.deepEqual(
          sorted.map((item) => item.process_code),
          [
            "early", "same-a", "same-b", "late",
            "missing", "invalid-day", "invalid-month", "invalid",
          ],
        );
        assert.deepEqual(processes.map((item) => item.process_code), originalOrder);
        """
    )
    script_path = tmp_path / "report_navigation_process_ordering_test.cjs"
    script_path.write_text(script, encoding="utf-8")
    subprocess.run(["node", str(script_path)], check=True, cwd=ROOT)

    display_body = app_js.split(
        "function reportNavigationDisplayProcesses(payload = {})", 1
    )[1].split("function reportNavigationSpineProgress", 1)[0]
    assert "sortReportNavigationProcessesByDate" in display_body


def test_report_navigation_reorder_uses_immediate_lightweight_flip_animation():
    app_js = _read(APP_JS)
    readme = _read(README_MD)
    css = _read(STYLES_CSS)

    assert "function captureReportNavigationSortPositions(container)" in app_js
    assert "function animateReportNavigationSort(container, previousPositions)" in app_js
    animation_body = app_js.split(
        "function animateReportNavigationSort(container, previousPositions)", 1
    )[1].split("function animateReportNavigationScheduleCardHeight", 1)[0]
    assert "getBoundingClientRect()" in animation_body
    assert "if (!visualEffectsEnabled()) return;" in animation_body
    assert 'window.matchMedia("(prefers-reduced-motion: reduce)").matches' in animation_body
    assert "element.animate(" in animation_body
    assert "translate:" in animation_body
    assert "duration: 220" in animation_body
    assert 'easing: "cubic-bezier(0.22, 1, 0.36, 1)"' in animation_body

    flow_body = app_js.split(
        "function renderReportNavigationProcesses(payload)", 1
    )[1].split("function renderReportNavigation(payload", 1)[0]
    assert "const previousPositions = captureReportNavigationSortPositions(reportNavBranches);" in flow_body
    assert '<div class="report-nav-branch ${side} ${shift}${done ? " done" : " running"}">' in flow_body
    assert 'class="report-nav-process-card${selected ? " selected" : ""}" data-report-nav-sort-key="flow:${escapeHtml(process.process_code || "")}"' in flow_body
    assert "animateReportNavigationSort(reportNavBranches, previousPositions);" in flow_body
    assert flow_body.index("captureReportNavigationSortPositions") < flow_body.index("reportNavBranches.innerHTML")
    assert flow_body.index("reportNavBranches.innerHTML") < flow_body.index("animateReportNavigationSort")

    schedule_body = app_js.split(
        "function renderReportNavigationSchedule(payload = {})", 1
    )[1].split("function selectReportNavigationScheduleProcess", 1)[0]
    assert "const previousPositions = captureReportNavigationSortPositions(reportNavScheduleTable);" in schedule_body
    assert 'data-report-nav-sort-key="schedule:${processCode}"' in schedule_body
    assert 'data-report-nav-sort-key="schedule-detail:${processCode}"' in schedule_body
    assert 'const detailReordering = previousPositions.has(`schedule-detail:${processCode}`);' in schedule_body
    assert 'class="report-nav-schedule-detail ${state.code}${detailReordering ? " reordering" : ""}"' in schedule_body
    assert "animateReportNavigationSort(reportNavScheduleTable, previousPositions);" in schedule_body
    assert schedule_body.index("captureReportNavigationSortPositions") < schedule_body.index("reportNavScheduleTable.innerHTML")
    assert schedule_body.index("reportNavScheduleTable.innerHTML") < schedule_body.index("animateReportNavigationSort")
    detail_reordering_rule = re.search(
        r"#page-report-navigation \.report-nav-schedule-detail\.reordering\s*\{(?P<body>[^}]*)\}",
        css,
    )
    assert detail_reordering_rule is not None
    assert "animation: none;" in detail_reordering_rule.group("body")
    assert "鱼骨图从左到右、报送日程从上到下按报送日期升序排列" in readme
    assert "保存截止日期后立即以轻量位移动画完成重排" in readme


def test_report_navigation_browser_refresh_restores_scoped_session_cache(tmp_path):
    app_js = _read(APP_JS)
    block = re.search(
        r"// Report navigation session cache start(?P<body>.*?)// Report navigation session cache end",
        app_js,
        re.S,
    )
    assert block is not None

    script = textwrap.dedent(
        f"""
        const assert = require("node:assert/strict");
        const values = new Map();
        const sessionStorage = {{
          get length() {{ return values.size; }},
          key: (index) => [...values.keys()][index] ?? null,
          getItem: (key) => values.get(key) ?? null,
          setItem: (key, value) => values.set(key, String(value)),
          removeItem: (key) => values.delete(key),
        }};
        const authState = {{ user: {{ id: "user-a" }} }};
        {block.group("body")}

        const now = new Date(2026, 6, 24, 12, 0, 0);
        const payload = {{
          period: "month",
          business_report_date: "2026-06-30",
          processes: [{{ process_code: "cached-process" }}],
        }};

        writeReportNavigationCache("month", payload, now);
        assert.deepEqual(readReportNavigationCache("month", now)?.payload, payload);
        assert.equal(readReportNavigationCache("week", now), null);

        authState.user = {{ id: "user-b" }};
        assert.equal(readReportNavigationCache("month", now), null);

        authState.user = {{ id: "user-a" }};
        assert.equal(readReportNavigationCache("month", new Date(2026, 7, 1, 12, 0, 0)), null);

        sessionStorage.setItem("unrelated-key", "keep");
        clearReportNavigationCache();
        assert.equal(readReportNavigationCache("month", now), null);
        assert.equal(sessionStorage.getItem("unrelated-key"), "keep");
        """
    )
    script_path = tmp_path / "report_navigation_cache_test.cjs"
    script_path.write_text(script, encoding="utf-8")
    subprocess.run(["node", str(script_path)], check=True, cwd=ROOT)


def test_report_navigation_browser_refresh_keeps_cached_content_on_request_failure():
    app_js = _read(APP_JS)
    start = app_js.index("async function loadReportNavigation")
    end = app_js.index("function syncReportNavigationPeriodTabs", start)
    body = app_js[start:end]

    assert "const cached = readReportNavigationCache(period);" in body
    assert 'let restoredFromCache = String(reportNavigationPayload?.period || "") === period;' in body
    assert "restoredFromCache = true;" in body
    assert "if (!restoredFromCache)" in body
    assert "renderReportNavigation({})" not in body
    assert "writeReportNavigationCache(period, payload);" in body


def test_report_navigation_has_first_load_placeholder_and_accessible_loading_state():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    page = re.search(r'<section class="page" id="page-report-navigation"[^>]*>', html)
    assert page is not None
    assert 'data-loading-state="initial-loading"' in page.group(0)
    assert 'aria-busy="true"' in page.group(0)
    assert 'id="reportNavInitialLoading"' in html
    assert 'class="report-nav-initial-loading"' in html
    assert 'role="status"' in html
    assert "function setReportNavigationLoadingState(state)" in app_js
    assert 'reportNavPage.dataset.loadingState = state;' in app_js
    assert 'reportNavPage.setAttribute("aria-busy"' in app_js
    assert 'const wasHidden = previous === "initial-loading" || previous === "error-empty";' in app_js
    assert "window.requestAnimationFrame(refreshReportNavigationScheduleLayout)" in app_js
    assert 'setReportNavigationLoadingState(restoredFromCache ? "refreshing-with-cache" : "initial-loading");' in app_js
    assert 'setReportNavigationLoadingState("ready");' in app_js
    assert 'setReportNavigationLoadingState(restoredFromCache ? "error-with-cache" : "error-empty");' in app_js
    assert '#page-report-navigation[data-loading-state="initial-loading"] > :not(.report-nav-initial-loading)' in css
    assert '#page-report-navigation[data-loading-state="error-empty"] > :not(.report-nav-initial-loading)' in css
    assert "@keyframes report-nav-loading-spin" in css


def test_report_navigation_refresh_optimization_is_documented():
    readme = _read(README_MD)
    app_js = _read(APP_JS)

    assert "浏览器刷新时优先恢复当前用户和统计周期的最近成功画面" in readme
    assert "<li>系统优化及BUG修复。</li>" in app_js


def test_report_navigation_statistics_keep_the_existing_four_icon_colors():
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert 'report_forms: {\n    color: "blue"' in app_js
    assert 'supplement_tasks: {\n    color: "green"' in app_js
    assert 'data_governance: {\n    color: "orange"' in app_js
    assert 'label: "数据治理"' in app_js
    assert 'special_governance: {\n    color: "red"' in app_js
    assert 'escapeHtml(style.label || card.name || "")' in app_js
    assert 'class="report-nav-stat-body"' in app_js
    assert '<path d="M14 2H6a2 2 0 0 0-2 2v16' in app_js
    assert '<path d="M20 6L9 17l-5-5"' in app_js
    assert '<path d="M21 12a9 9 0 1 1-6.22-8.56"' in app_js
    assert '<path d="M10.29 3.86L1.82 18' in app_js
    for color in ["blue", "green", "orange", "red"]:
        assert f"#page-report-navigation .report-nav-stat-card.{color} .report-nav-stat-icon" in css


def test_report_navigation_display_only_step_is_visually_neutral():
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert "function reportNavigationStepIsDisplayOnly(step = {})" in app_js
    assert 'Boolean(step.display_only || step.status === "display_only")' in app_js
    assert '? "display-only"' in app_js
    assert 'const statusMarkup = displayOnly ? ""' in app_js
    assert '? "仅展示"' not in app_js
    assert 'class="report-nav-step-display-only"' not in app_js
    assert ".report-nav-step-row.display-only .report-nav-step-index" in css
    assert ".report-nav-schedule-step-list > li.display-only > i" in css


def test_completed_report_navigation_schedule_detail_shows_completion_time():
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert "function reportNavigationCompletionTimeText(process = {})" in app_js
    assert 'process.status === "completed" && process.completed_at' in app_js
    assert '`完成时间：${reportNavigationTimestampText(process.completed_at)}`' in app_js
    assert 'class="report-nav-schedule-detail-completed-at"' in app_js
    assert 'completionTimeValue.hidden = !completionTimeText;' in app_js
    assert ".report-nav-schedule-detail-completed-at" in css
    assert (
        "#page-report-navigation .report-nav-schedule-detail "
        "small.report-nav-schedule-detail-completed-at"
    ) in css
    assert "color: var(--on-surface);" in css
    assert "#page-report-navigation .report-nav-schedule-detail-completed-at[hidden]" in css
    assert "全部完成时在展开详情的“下一步”区域显示“完成时间：YYYY-MM-DD HH:mm:ss”" in _read(README_MD)


def test_report_navigation_overdue_state_is_real_and_current_period_only():
    readme = _read(README_MD)
    app = _read(APP_JS)
    design = _read(
        ROOT
        / "docs"
        / "superpowers"
        / "specs"
        / "2026-07-15-report-navigation-statistics-design.md"
    )

    assert "未完成流程在当前报告期持续保持逾期状态" in readme
    assert "进入下个月后按新的报告期重新开始，不结转上月逾期状态" in readme
    assert "报送导航报告期固定为当前月份上一个自然月的最后一天" in readme
    assert "8 月 1 日对应 7 月 31 日" in readme
    assert "按最新报送日期兜底" not in readme
    assert "节点完成时间固定为该报送日期当天 `20:00:00`" not in readme
    assert "当期逾期跟踪" in app
    assert "报送日期到期完成兜底" not in app
    assert "当前月份上一个自然月的最后一天" in design
    assert "进入下个月后创建新的报告期快照" in design
    assert "报送日期当天 `20:00:00`" not in design
    assert "MAX(COALESCE(update_date, create_date))" not in design
    assert "完成时间仅取匹配记录的 `MAX(create_date)`" in design
    assert "步骤 6 仍是最终完成判断节点" in design


def test_report_navigation_process_cards_select_readonly_step_details():
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)
    redesign_css = css.split("/* ===== Report navigation: restrained read-only panorama ===== */", 1)[1]

    assert "const REPORT_NAV_CHECK_ICON =" in app_js
    assert '${completed ? REPORT_NAV_CHECK_ICON : ""}' in app_js
    assert 'class="report-nav-step-order">${index + 1}、</span>' in app_js
    assert '${done ? REPORT_NAV_CHECK_ICON : ""}' in app_js
    assert 'class="report-nav-process-detail-head ${done ? "completed" : "running"}"' in app_js
    assert "function renderReportNavigationProcessDetails(process)" in app_js
    assert "function animateReportNavigationFlowCardHeight(startHeight)" in app_js
    assert "reportNavigationFlowCardAnimation" in app_js
    assert "reportNavFlowCard.animate(" in app_js
    assert 'height: `${startHeight}px`' in app_js
    assert 'height: `${endHeight}px`' in app_js
    assert 'overflow: "hidden"' in app_js
    assert "duration: 180" in app_js
    assert app_js.count("animateReportNavigationFlowCardHeight(startHeight);") == 2
    assert "reportNavProcessDetails.animate(" not in app_js
    assert "clipPath:" not in app_js
    assert "function selectReportNavigationProcess(processCode)" in app_js
    assert 'data-report-nav-process="${escapeHtml(process.process_code || "")}"' in app_js
    assert 'aria-pressed="${selected ? "true" : "false"}"' in app_js
    assert 'event.target.closest("[data-report-nav-process]")' in app_js
    assert 'event.key !== "Enter" && event.key !== " "' in app_js
    assert "function clearReportNavigationProcessSelection()" in app_js
    assert 'reportNavFlowCard?.addEventListener("click"' in app_js
    assert 'event.target.closest("[data-report-nav-process], #reportNavProcessDetails")' in app_js
    assert 'reportNavProcessDetails.hidden = !process;' in app_js
    assert 'reportNavFlowCard?.classList.toggle("has-selection", Boolean(process));' in app_js
    assert 'class="report-nav-step-row ${stateClass}"' in app_js
    assert "截止日期" in app_js
    assert "完成于" in app_js
    assert 'const completionText = done\n      ? (process.completed_at ? escapeHtml(reportNavigationTimestampText(process.completed_at)) : "--")\n      : "进行中";' in app_js
    assert "暂无可展示步骤" in app_js
    detail_renderer = re.search(
        r"function renderReportNavigationProcessDetails\(process\) \{(?P<body>.*?)\n\}",
        app_js,
        re.S,
    )
    assert detail_renderer is not None
    assert "status_message" not in detail_renderer.group("body")
    assert "error_message" not in detail_renderer.group("body")
    assert "report-nav-step-state" not in detail_renderer.group("body")
    assert "reportNavigationTimestampText(process.completed_at)" not in detail_renderer.group("body")
    process_icon = re.search(
        r"#page-report-navigation \.report-nav-process-state-icon\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert process_icon is not None
    assert "width: 16px;" in process_icon.group("body")
    assert "height: 16px;" in process_icon.group("body")
    assert "border: 1px solid var(--outline);" in process_icon.group("body")
    running_state = re.search(
        r"#page-report-navigation \.report-nav-branch\.running \.report-nav-process-meta\.completion time\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert running_state is not None
    assert "color: var(--on-surface-variant);" in running_state.group("body")
    assert "font-size:" not in running_state.group("body")
    step_icon = re.search(
        r"#page-report-navigation \.report-nav-step-index\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert step_icon is not None
    assert "width: 14px;" in step_icon.group("body")
    assert "height: 14px;" in step_icon.group("body")
    assert "align-self: center;" in step_icon.group("body")
    step_content = re.search(
        r"#page-report-navigation \.report-nav-step-content\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert step_content is not None
    assert "align-items: center;" in step_content.group("body")
    step_svg = re.search(
        r"#page-report-navigation \.report-nav-step-index svg\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert step_svg is not None
    assert "display: block;" in step_svg.group("body")
    flow_renderer = re.search(
        r"function renderReportNavigationProcesses\(payload\) \{(?P<body>.*?)\n\}",
        app_js,
        re.S,
    )
    assert flow_renderer is not None
    flow_body = flow_renderer.group("body")
    assert "data-manual-action" not in flow_body
    assert "manual-complete" not in flow_body
    assert "manual-cancel" not in flow_body
    assert "errorProcess" not in flow_body
    assert "runningProcess" not in flow_body
    assert "processes[0]" not in flow_body
    assert "reportNavigationPanelWidth" not in app_js


def test_report_navigation_uses_bent_connectors_below_details_and_admin_schedule_editing():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)
    redesign_css = css.split("/* ===== Report navigation: restrained read-only panorama ===== */", 1)[1]

    assert 'class="report-nav-flow-legend"' not in html
    assert 'class="report-nav-legend-item' not in html
    assert 'const shift = index === 0' in app_js
    assert 'class="report-nav-branch ${side} ${shift}' in app_js
    assert ".report-nav-branch-line::after" in redesign_css
    assert "height: 16px;" in redesign_css
    assert "width: 23px;" in redesign_css
    assert "rotate(-52deg)" in redesign_css
    assert "rotate(128deg)" in redesign_css
    assert "bottom: calc(50% + 34px);" in redesign_css
    assert "top: calc(50% + 34px);" in redesign_css
    assert "padding: 0 136px;" in redesign_css
    assert "var(--report-nav-card-shift)" in redesign_css
    assert "#page-report-navigation .report-nav-flow-card.has-selection" not in redesign_css
    details = re.search(
        r"#page-report-navigation \.report-nav-process-details\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert details is not None
    assert "background: transparent;" in details.group("body")
    assert "margin: 14px 20px 0;" in details.group("body")
    assert 'data-report-nav-date="${escapeHtml(process.process_code || "")}"' in app_js
    assert '<time${dateInteraction}>${escapeHtml(reportNavigationDateText(process.report_date))}</time>' in app_js
    assert 'title="右击修改截止日期"' in app_js
    deadline_edit_rule = re.search(
        r"#page-report-navigation \.report-nav-process-deadline time\.editable\s*\{(?P<body>[^}]*)\}",
        redesign_css,
    )
    assert deadline_edit_rule is not None
    assert "cursor: pointer;" in deadline_edit_rule.group("body")
    for declaration in ["padding:", "margin:", "background:", "outline-color:"]:
        assert declaration not in deadline_edit_rule.group("body")
    assert "outline: none;" in deadline_edit_rule.group("body")
    deadline_edit_hover_rule = re.search(
        r"#page-report-navigation \.report-nav-process-deadline time\.editable:hover,\s*"
        r"#page-report-navigation \.report-nav-process-deadline time\.editable:focus-visible\s*\{(?P<body>[^}]*)\}",
        redesign_css,
    )
    assert deadline_edit_hover_rule is not None
    assert "text-decoration: underline;" in deadline_edit_hover_rule.group("body")
    assert "background:" not in deadline_edit_hover_rule.group("body")
    assert "outline-color:" not in deadline_edit_hover_rule.group("body")
    assert 'reportNavBranches?.addEventListener("contextmenu"' in app_js
    assert 'authState.user?.role !== "admin"' in app_js
    assert '"修改截止日期"' in app_js
    assert '{ type: "date", defaultValue: currentDate }' in app_js
    assert '/api/report-navigation/schedules/${encodeURIComponent(processCode)}' in app_js
    assert "report_month: reportMonth, report_date: nextDate" in app_js


def test_report_navigation_spine_colors_completed_prefix_and_keeps_future_nodes_gray():
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)
    redesign_css = css.split("/* ===== Report navigation: restrained read-only panorama ===== */", 1)[1]

    assert "const allProcessesCompleted = processes.length > 0" in app_js
    assert 'reportNavFishbone?.classList.toggle("all-done", allProcessesCompleted)' in app_js
    assert 'const reportNavFishboneSpine = document.getElementById("reportNavFishboneSpine")' in app_js
    assert "function reportNavigationSpineProgress(processes = [])" in app_js
    assert 'if (process.status !== "completed") break;' in app_js
    assert "((completedPrefixCount - 0.5) / processes.length) * 100" in app_js
    assert '"--report-nav-spine-progress"' in app_js
    assert "`${reportNavigationSpineProgress(processes)}%`" in app_js

    progress_rule = re.search(
        r"#page-report-navigation \.report-nav-fishbone-spine::before\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert progress_rule is not None
    assert "width: var(--report-nav-spine-progress, 0%);" in progress_rule.group("body")
    assert "background: var(--theme-accent);" in progress_rule.group("body")
    assert "background-image: var(--theme-accent-gradient);" in progress_rule.group("body")
    assert "box-shadow: none;" in progress_rule.group("body")
    assert "transform: none;" in progress_rule.group("body")

    assert ".report-nav-branch.done.top .report-nav-branch-line" in redesign_css
    assert ".report-nav-branch.done.bottom .report-nav-branch-line" in redesign_css
    assert ".report-nav-branch.done .report-nav-branch-line::after" in redesign_css
    done_node = re.search(
        r"#page-report-navigation \.report-nav-branch\.done \.report-nav-branch-node\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert done_node is not None
    assert "border-color: var(--theme-accent);" in done_node.group("body")
    assert "background: var(--theme-accent);" in done_node.group("body")

    assert "#page-report-navigation .report-nav-fishbone.all-done .report-nav-fishbone-spine" in css
    assert "#page-report-navigation .report-nav-fishbone.all-done .report-nav-fishbone-tail" in css
    spine_matches = re.findall(
        r"#page-report-navigation \.report-nav-fishbone-spine\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert spine_matches
    base_spine = next(
        body for body in reversed(spine_matches)
        if "background:" in body
    )
    assert "var(--outline)" in base_spine or "var(--outline-variant)" in base_spine
    assert "var(--theme-accent)" not in base_spine
    all_done_tail_matches = re.findall(
        r"#page-report-navigation \.report-nav-fishbone\.all-done \.report-nav-fishbone-tail\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert all_done_tail_matches
    final_all_done_tail = all_done_tail_matches[-1]
    assert "background: var(--theme-accent);" in final_all_done_tail
    assert "background-image: var(--theme-accent-gradient);" in final_all_done_tail


def test_report_navigation_styles_follow_theme_radius_and_vertical_mobile_timeline():
    css = _read(STYLES_CSS)
    redesign_css = css.split("/* ===== Report navigation: restrained read-only panorama ===== */", 1)[1]

    fishbone_terminals = re.search(
        r"#page-report-navigation \.report-nav-fishbone-head,\s*"
        r"#page-report-navigation \.report-nav-fishbone-tail\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert fishbone_terminals is not None
    assert "border-radius: 999px !important;" in fishbone_terminals.group("body")
    assert "var(--ui-radius)" not in fishbone_terminals.group("body")

    for selector in [
        "#page-report-navigation .report-nav-period-bar",
        "#page-report-navigation .report-nav-stats-layout",
        "#page-report-navigation .report-nav-overview-group",
        "#page-report-navigation .report-nav-period-tabs",
        "#page-report-navigation .report-nav-stat-card",
        "#page-report-navigation .report-nav-process-card",
        "#page-report-navigation .report-nav-process-details",
        "#page-report-navigation .report-nav-step-row",
        "#page-report-navigation .report-nav-todo",
    ]:
        assert selector in css
    assert "border-radius: var(--ui-radius)" in css
    assert "var(--theme-accent)" in css
    assert "@media (max-width: 760px)" in css
    assert "grid-template-columns: 28px minmax(0, 1fr);" in css
    assert '[data-color-mode="dark"] #page-report-navigation' not in redesign_css
    assert "width: min(220px, calc(200% - 12px));" in redesign_css
    assert "grid-template-columns: repeat(4, minmax(0, 1fr));" in redesign_css
    assert "min-height: 260px;" in redesign_css
    assert "overflow-y: hidden;" in redesign_css
    assert "#page-report-navigation .report-nav-flow-card.has-selection" not in redesign_css
    assert "#page-report-navigation .report-nav-process-details[hidden]" in redesign_css
    todo_content = re.search(
        r"#page-report-navigation \.report-nav-todo > \.report-nav-todo-main\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert todo_content is not None
    assert "display: grid;" in todo_content.group("body")
    assert "gap: 3px;" in todo_content.group("body")
    todo_bar = re.search(
        r"#page-report-navigation \.report-nav-todo > i\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert todo_bar is not None
    assert "width: 5px;" in todo_bar.group("body")
    assert "height: 5px;" in todo_bar.group("body")
    assert "border-radius: 50%;" in todo_bar.group("body")
    assert "flex-wrap: wrap;" in redesign_css
    assert "#page-report-navigation .report-nav-stat-card::before" in redesign_css
    assert "display: none;" in re.search(
        r"#page-report-navigation \.report-nav-stat-card::before\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    ).group("body")
    details_style = re.search(
        r"#page-report-navigation \.report-nav-process-details\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert details_style is not None
    assert "background: transparent;" in details_style.group("body")
    assert "display: block !important;" in re.search(
        r"#page-report-navigation \.report-nav-page-title\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    ).group("body")
    assert "#page-report-navigation .report-nav-branch.done .report-nav-branch-node" in redesign_css
    assert "box-shadow: 0 0 0 3px var(--surface-container-lowest);" in redesign_css
    assert ".report-nav-card-maintenance-grid" in css


def test_report_navigation_headings_and_fishbone_spacing_match_the_requested_layout():
    css = _read(STYLES_CSS)
    redesign_css = css.split("/* ===== Report navigation: restrained read-only panorama ===== */", 1)[1]

    heading = re.search(
        r"#page-report-navigation \.report-nav-card-head h2\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert heading is not None
    assert "font-size: 16px;" in heading.group("body")

    for selector in ["report-nav-fishbone", "report-nav-branches"]:
        rule = re.search(
            rf"#page-report-navigation \.{selector}\s*\{{(?P<body>.*?)\}}",
            redesign_css,
            re.S,
        )
        assert rule is not None
        assert "min-height: 260px;" in rule.group("body")

    fishbone = re.search(
        r"#page-report-navigation \.report-nav-fishbone\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert fishbone is not None
    assert "padding: 0 136px;" in fishbone.group("body")
    assert "left: 20px;" in redesign_css
    assert "right: 20px;" in redesign_css

    spine = re.search(
        r"#page-report-navigation \.report-nav-fishbone-spine\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert spine is not None
    assert "left: 116px;" in spine.group("body")
    assert "right: 116px;" in spine.group("body")

    terminals = re.search(
        r"#page-report-navigation \.report-nav-fishbone-head,\s*"
        r"#page-report-navigation \.report-nav-fishbone-tail\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert terminals is not None
    assert "min-height: 34px;" in terminals.group("body")
    assert "padding: 8px 14px;" in terminals.group("body")

    process_card = re.search(
        r"#page-report-navigation \.report-nav-process-card\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert process_card is not None
    assert "min-height: 88px;" in process_card.group("body")
    assert "padding: 10px 12px;" in process_card.group("body")
    assert "margin-bottom: 6px;" in redesign_css
    assert "line-height: 1.6;" in redesign_css


def test_report_navigation_matches_the_selected_desktop_design_proportions():
    css = _read(STYLES_CSS)
    redesign_css = css.split("/* ===== Report navigation: restrained read-only panorama ===== */", 1)[1]

    stats_layout = re.search(
        r"#page-report-navigation \.report-nav-stats-layout\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert stats_layout is not None
    assert "grid-template-columns: repeat(4, minmax(0, 1fr));" in stats_layout.group("body")
    assert "column-gap: 16px;" in stats_layout.group("body")
    content_area_hover = re.search(
        r"#page-report-navigation \.report-nav-stats-layout:hover,\s*"
        r"#page-report-navigation \.report-nav-flow-card:hover,\s*"
        r"#page-report-navigation \.report-nav-schedule-card:hover,\s*"
        r"#page-report-navigation \.report-nav-attention-card:hover\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert content_area_hover is not None
    assert "border-color: color-mix(in srgb, var(--theme-accent) 50%, var(--outline-variant));" in content_area_hover.group("body")
    assert "transform: translateY(-2px);" in content_area_hover.group("body")
    assert "0 8px 18px rgba(15, 23, 42, 0.08);" in content_area_hover.group("body")

    stat_groups = re.search(
        r"#page-report-navigation \.report-nav-overview-group,\s*"
        r"#page-report-navigation \.report-nav-period-stat-group\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert stat_groups is not None
    assert "min-height: 0;" in stat_groups.group("body")
    assert "padding: 0;" in stat_groups.group("body")
    assert "border: 0;" in stat_groups.group("body")
    assert "background: transparent;" in stat_groups.group("body")

    flow_card = re.search(
        r"#page-report-navigation \.report-nav-flow-card\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert flow_card is not None
    assert "padding: 16px 22px 18px;" in flow_card.group("body")

    fishbone = re.search(
        r"#page-report-navigation \.report-nav-fishbone\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert fishbone is not None
    assert "min-height: 260px;" in fishbone.group("body")
    assert "padding: 0 136px;" in fishbone.group("body")

    process_card = re.search(
        r"#page-report-navigation \.report-nav-process-card\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert process_card is not None
    assert "width: min(220px, calc(200% - 12px));" in process_card.group("body")
    assert "min-height: 88px;" in process_card.group("body")
    assert "padding: 10px 12px;" in process_card.group("body")
    assert "border-radius: var(--ui-radius);" in process_card.group("body")

    stat_card = re.search(
        r"#page-report-navigation \.report-nav-stat-card\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert stat_card is not None
    assert "display: flex;" in stat_card.group("body")
    assert "align-items: flex-start;" in stat_card.group("body")
    assert "gap: 16px;" in stat_card.group("body")

    process_details = re.search(
        r"#page-report-navigation \.report-nav-process-details\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert process_details is not None
    assert "padding: 12px 18px;" in process_details.group("body")
    assert "border-radius: var(--ui-radius);" in process_details.group("body")

    attention = re.search(
        r"#page-report-navigation \.report-nav-attention-card\s*\{(?P<body>[^}]*padding: 16px;[^}]*)\}",
        redesign_css,
        re.S,
    )
    assert attention is not None
    assert "padding: 16px;" in attention.group("body")
    assert "align-self: flex-start;" in attention.group("body")
    assert "align-self: stretch;" not in attention.group("body")
    assert "display: flex;" in attention.group("body")
    assert "flex-direction: column;" in attention.group("body")

    todo_heading = re.search(
        r"#page-report-navigation \.report-nav-attention-card \.report-nav-card-head h2\s*\{(?P<body>.*?)\}",
        redesign_css,
        re.S,
    )
    assert todo_heading is not None
    assert "gap: 6px;" in todo_heading.group("body")
    assert "font-size: 16px;" in todo_heading.group("body")
    assert "font-weight: 600;" in todo_heading.group("body")


def test_space_tech_uses_gap_safe_panel_shadows_across_all_pages():
    css = _read(STYLES_CSS)

    for selector in [
        '[data-theme="space-tech"] .card',
        '[data-theme="space-tech"] .home-stat-card',
        '[data-theme="space-tech"] #page-report-navigation .report-nav-stat-card',
        '[data-theme="space-tech"] #page-report-navigation .report-nav-card',
        '[data-theme="space-tech"] #page-settings .settings-dashboard-card',
        '[data-theme="space-tech"] #page-users .user-stat-card',
        '[data-theme="space-tech"] #page-users .user-filter-bar',
        '[data-theme="space-tech"] #page-users .user-table-card',
        '[data-theme="space-tech"] #page-local-storage .local-storage-metric',
        '[data-theme="space-tech"] #page-local-storage .local-storage-table-panel',
        '[data-theme="space-tech"] #page-local-storage .local-storage-detail-panel',
    ]:
        assert selector in css

    assert "--space-panel-shadow: inset 0 1px 0 rgba(255, 255, 255, 0.56);" in css
    assert "--space-panel-hover-shadow:" in css
    assert "box-shadow: var(--space-panel-shadow);" in css
    assert "box-shadow: var(--space-panel-hover-shadow);" in css
    assert "--space-panel-shadow: inset 0 1px 0 rgba(255, 255, 255, 0.08);" in css


def test_report_navigation_month_label_uses_snapshot_while_period_only_reloads_statistics_cards():
    app_js = _read(APP_JS)

    assert "function reportNavigationMonthText(value)" in app_js
    assert 'const businessPeriod = String(payload.business_report_date || payload.report_month || "").slice(0, 7);' in app_js
    assert "reportNavMonth.textContent = reportNavigationMonthText(businessPeriod);" in app_js
    assert 'const monthlyCard = cards.find((card) => card.card_code === "report_forms")' in app_js
    assert 'const periodCards = cards.filter((card) => card.card_code !== "report_forms")' in app_js
    assert 'const period = reportNavPeriodSelect?.value || "month";' in app_js
    assert 'dashboard?period=${encodeURIComponent(period)}' in app_js
    assert "report_month=" not in app_js


def test_report_navigation_schedule_extends_through_overdue_completion_date():
    app_js = _read(APP_JS)

    assert "const completionDates = processes" in app_js
    assert 'process.status === "completed"' in app_js
    assert ".map((process) => reportNavigationDateOnly(process.completed_at))" in app_js
    assert "const latestCompletion = completionDates.length" in app_js
    assert "latestCompletion.getTime()," in app_js
    assert "hasOverdueIncomplete ? today.getTime() : latestDeadline.getTime()," in app_js


def test_report_navigation_schedule_timeline_expands_with_hover_step_preview():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    flow_position = html.index('class="report-nav-card report-nav-flow-card"')
    schedule_position = html.index('id="reportNavScheduleCard"')
    layout_position = html.index('class="report-nav-schedule-layout"')
    attention_position = html.index('class="report-nav-card report-nav-attention-card"')
    assert flow_position < schedule_position < attention_position
    assert flow_position < layout_position <= schedule_position
    for element_id in ["reportNavScheduleCard", "reportNavScheduleRange", "reportNavScheduleTable"]:
        assert f'id="{element_id}"' in html
    assert "报送日程" in html
    assert "我的待办" in html
    for process_code, process_name in [
        ("pbc_template", "资管产品模板、逐笔报送"),
        ("full_elements", "全要素报送"),
        ("east5", "EAST5.0报送"),
    ]:
        assert f'{process_code}: "{process_name}"' in app_js
    assert "function reportNavigationProcessDisplayName(process = {})" in app_js
    assert 'east5_1: "归档并上传 EAST5.0 报送"' in app_js
    assert "function reportNavigationStepDisplayName(step = {})" in app_js
    assert "step_name: reportNavigationStepDisplayName(step)" in app_js
    assert "const REPORT_NAV_FISHBONE_PROCESS_NAMES = {" in app_js
    assert 'pbc_template: "资管产品模板、逐笔"' in app_js
    assert "function reportNavigationFishboneProcessName(process = {})" in app_js
    assert "const fishboneProcessName = reportNavigationFishboneProcessName(process);" in app_js
    assert "<strong>${escapeHtml(fishboneProcessName)}</strong>" in app_js
    assert '<div class="report-nav-schedule-summary"><strong>${escapeHtml(process.process_name || "")}</strong>' in app_js
    schedule_layout_rule = re.search(
        r"#page-report-navigation \.report-nav-schedule-layout\s*\{(?P<body>[^}]*)\}",
        css,
    )
    assert schedule_layout_rule is not None
    assert "display: flex;" in schedule_layout_rule.group("body")
    assert "flex-wrap: wrap;" in schedule_layout_rule.group("body")
    assert "flex: 999 1 var(--report-nav-schedule-min-width, 1148px);" in css
    assert "min-width: min(100%, var(--report-nav-schedule-min-width, 1148px));" in css
    assert "flex: 1 1 320px;" in css
    assert "grid-template-columns: 172px minmax(0, 1fr);" in css
    assert "function reportNavigationScheduleMinWidth(dates = [])" in app_js
    assert "return 172 + (dayCount * 38) + 64;" in app_js
    assert "function refreshReportNavigationScheduleLayout()" in app_js
    assert 'reportNavScheduleCard?.style.setProperty("--report-nav-schedule-min-width", `${scheduleMinWidth}px`);' in app_js
    assert 'const reportNavTodoCard = document.querySelector("#page-report-navigation .report-nav-attention-card");' in app_js
    assert "function syncReportNavigationTodoCardHeight()" in app_js
    assert "Math.abs(reportNavScheduleCard.offsetTop - reportNavTodoCard.offsetTop) < 2" in app_js
    assert "Math.abs(scheduleRect.top - todoRect.top) < 2" not in app_js
    assert "reportNavTodoLockedHeightPx" in app_js
    assert "scheduleHeight - detailHeight" not in app_js
    assert 'reportNavTodoCard.style.removeProperty("height");' in app_js
    assert "syncReportNavigationTodoCardHeight();" in app_js
    assert "首次展开前锁死待办高度" in app_js or "reportNavTodoLockedHeightPx = Math.round(startHeight)" in app_js
    assert 'class="report-nav-todo-summary"' in app_js
    assert "#page-report-navigation .report-nav-todo-summary" in css
    assert "text-overflow: ellipsis" in css
    todo_summary = re.search(
        r"#page-report-navigation \.report-nav-todo-summary\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert todo_summary is not None
    assert "overflow: hidden" in todo_summary.group("body")
    assert "white-space: nowrap" in todo_summary.group("body")
    assert "grid-column: 1" in todo_summary.group("body")
    assert 'class="report-nav-todo-initiator"' in app_js
    assert "#page-report-navigation .report-nav-todo-initiator" in css
    assert "#reportNavTodoAllModal .report-nav-todo-initiator" in css
    # Card + modal share the same timestamp/initiator blue.
    assert (
        "#page-report-navigation .report-nav-todo time,\n"
        "#page-report-navigation .report-nav-todo-initiator,\n"
        "#reportNavTodoAllModal .report-nav-todo time,\n"
        "#reportNavTodoAllModal .report-nav-todo-initiator,\n"
        "#reportNavHistoryModal .report-nav-todo time,\n"
        "#reportNavHistoryModal .report-nav-todo-initiator"
    ) in css
    shared_time = re.search(
        r"#page-report-navigation \.report-nav-todo time,\s*\n"
        r"#page-report-navigation \.report-nav-todo-initiator,\s*\n"
        r"#reportNavTodoAllModal \.report-nav-todo time,\s*\n"
        r"#reportNavTodoAllModal \.report-nav-todo-initiator,\s*\n"
        r"#reportNavHistoryModal \.report-nav-todo time,\s*\n"
        r"#reportNavHistoryModal \.report-nav-todo-initiator\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert shared_time is not None
    assert "color: var(--theme-accent-readable)" in shared_time.group("body")
    assert "-webkit-text-fill-color: var(--theme-accent-readable)" in shared_time.group("body")
    assert "font-family: var(--report-nav-font-mono)" in shared_time.group("body")
    assert "#reportNavTodoAllModal,\n#reportNavHistoryModal" in css
    assert '--report-nav-font-sans: "PingFang SC"' in css
    assert 'font-family: var(--report-nav-font-sans);' in css
    assert 'window.addEventListener("resize", refreshReportNavigationScheduleLayout);' in app_js
    reveal_fn = re.search(
        r"function revealAuthenticatedApp\(\) \{(?P<body>.*?)\n\}",
        app_js,
        re.S,
    )
    assert reveal_fn is not None
    reveal_body = reveal_fn.group("body")
    assert "refreshReportNavigationScheduleLayout();" in reveal_body
    assert "window.requestAnimationFrame(refreshReportNavigationScheduleLayout)" in reveal_body
    assert 'class="report-nav-schedule-header-label" aria-hidden="true"></div>' in app_js
    assert ">流程 / 进度</div>" not in app_js

    for function_name in [
        "reportNavigationScheduleDates",
        "reportNavigationScheduleState",
        "renderReportNavigationSchedule",
        "selectReportNavigationScheduleProcess",
        "animateReportNavigationScheduleCardHeight",
        "editReportNavigationScheduleOwner",
        "openReportNavigationScheduleStepsPreview",
        "closeReportNavigationScheduleStepsPreview",
    ]:
        assert f"function {function_name}" in app_js or f"async function {function_name}" in app_js
    assert 'data-report-nav-schedule-process="${processCode}"' in app_js
    assert 'data-report-nav-schedule-date="${processCode}"' in app_js
    assert '<span>负责人</span>' not in app_js
    assert 'data-report-nav-schedule-owner="${processCode}"' not in app_js
    assert 'class="report-nav-schedule-steps-preview"' in app_js
    assert 'role="button" aria-label="查看步骤"' in app_js
    assert '<span>查看步骤</span>' in app_js
    assert 'class="report-nav-schedule-steps-popover"' in app_js
    assert 'popover="manual"' in app_js
    assert 'class="report-nav-schedule-step-list"' in app_js
    assert 'class="report-nav-schedule-steps-arrow" aria-hidden="true"' in app_js
    assert 'class="report-nav-schedule-detail-spacer" aria-hidden="true"' in app_js
    assert "`${index + 1}、${escapeHtml(reportNavigationStepDisplayName(step))}`" in app_js
    assert "const firstIncompleteStepIndex = processSteps.findIndex(" in app_js
    assert '(step) => !reportNavigationStepIsDisplayOnly(step) && step.status !== "completed"' in app_js
    assert 'index === firstIncompleteStepIndex ? "running" : "waiting"' in app_js
    assert 'stepState === "completed" ? "已完成" : (stepState === "running" ? "进行中" : "未完成")' in app_js
    assert "function positionReportNavigationScheduleStepsPopover" in app_js
    assert 'data-report-nav-schedule-view="${processCode}"' not in app_js
    assert 'class="report-nav-schedule-view-steps"' not in app_js
    assert 'class="report-nav-schedule-today"' not in app_js
    assert "reportNavigationScheduleDates(payload.report_month, processes)" in app_js
    assert "const deadlines = processes" in app_js
    assert "const latestDeadline = new Date(Math.max(...deadlines.map((item) => item.getTime())));" in app_js
    assert "const hasOverdueIncomplete = processes.some((process) => {" in app_js
    assert 'process.status !== "completed" && deadline && today > deadline' in app_js
    assert "cursor <= scheduleEndDate" in app_js
    assert "const deadlinePosition = ((deadlineIndex + 0.5) / dates.length) * 100;" in app_js
    assert "const completionOffsetDays" in app_js
    assert 'code: "early-completed", label: "提前完成"' in app_js
    assert "const todayPosition = ((todayIndex + 0.5) / dates.length) * 100;" in app_js
    assert 'const endpointPosition = ["early-completed", "overdue-completed"].includes(state.code)' in app_js
    assert '? completionPosition' in app_js
    assert ': deadlinePosition;' in app_js
    assert 'const fillPosition = ["completed", "early-completed", "overdue-completed"].includes(state.code)' in app_js
    assert ': todayPosition;' in app_js
    assert "const completionText = reportNavigationDateText(completedKey);" in app_js
    assert '`提前${state.earlyDays}天，${completionText}完成`' in app_js
    assert 'state.code === "completed"' in app_js
    assert '`${completionText}按时完成`' in app_js
    assert 'state.code === "running" ? "" : deadlineText' in app_js
    assert '`逾期${state.overdueDays}天 · 原截止${deadlineText}`' not in app_js
    assert "const todayText = reportNavigationDateText(todayKey);" not in app_js
    assert '`已逾期${state.overdueDays}天`' in app_js
    assert app_js.count('`逾期${state.overdueDays}天完成`') >= 2
    assert 'class="report-nav-schedule-deadline-warning"' in app_js
    assert 'aria-label="原截止日期：${escapeHtml(deadlineText)}">!</span>' in app_js
    assert "const adjacentOverdueCompletionLabels" not in app_js
    assert '--report-nav-schedule-deadline:${deadlinePosition}%' in app_js
    assert 'class="report-nav-schedule-early-tail"' not in app_js
    assert 'class="report-nav-schedule-early-deadline"' not in app_js
    assert "const scheduleEdgePosition = 50 / dates.length;" in app_js
    assert "const scheduleEndPosition" not in app_js
    assert "const earlyTailEndPosition" not in app_js
    assert "const earlyDeadlineOffsetPosition" not in app_js
    assert "const overdueStopPosition" in app_js
    assert "--report-nav-schedule-early-tail-end" not in app_js
    assert "--report-nav-schedule-early-deadline-offset" not in app_js
    assert '--report-nav-schedule-overdue-stop:${overdueStopPosition}%' in app_js
    assert "const dotPosition = ((index + 0.5) / dates.length) * 100;" in app_js
    assert "dotPosition <= fillPosition" in app_js
    assert 'class="reached"' in app_js
    assert 'class="early-target"' not in app_js
    assert 'class="early-original-deadline"' in app_js
    assert 'class="after-early-completion"' not in app_js
    assert 'class="report-nav-schedule-original-deadline-label' not in app_js
    assert 'class="reached before-deadline"' in app_js
    assert 'class="reached after-deadline"' in app_js
    assert "dates.push(new Date(cursor));" in app_js
    assert "payload.work_calendar?.holidays" in app_js
    assert "payload.work_calendar?.adjusted_workdays" in app_js
    assert "const holiday = holidayKeys.has(itemKey) || (weekend && !adjustedWorkday);" in app_js
    assert '${holiday ? " holiday" : ""}' in app_js
    assert "overdue-completed" in app_js
    assert 'api(`/api/report-navigation/schedule-owners/${encodeURIComponent(processCode)}`' in app_js
    assert 'return { code: "running", label: "进行中", overdueDays: 0, earlyDays: 0 };' in app_js
    assert 'return { code: "pending", label: "待开始"' not in app_js
    assert "function viewReportNavigationScheduleSteps" not in app_js
    assert "duration: 180" in app_js
    assert 'reportNavScheduleCard?.addEventListener("click"' in app_js
    assert 'reportNavScheduleTable?.addEventListener("contextmenu"' in app_js
    assert 'reportNavScheduleTable?.addEventListener("pointerover"' in app_js
    assert 'reportNavScheduleTable?.addEventListener("pointerout"' in app_js
    assert 'reportNavScheduleTable?.addEventListener("focusin"' in app_js
    assert 'reportNavScheduleTable?.addEventListener("focusout"' in app_js
    assert "const REPORT_NAV_SCHEDULE_STEPS_SHOW_DELAY = 120;" in app_js
    assert "const REPORT_NAV_SCHEDULE_STEPS_HIDE_DELAY = 140;" in app_js

    for selector in [
        ".report-nav-schedule-card",
        ".report-nav-schedule-scroll",
        ".report-nav-schedule-header",
        ".report-nav-schedule-summary",
        ".report-nav-schedule-dates",
        ".report-nav-schedule-track",
        ".report-nav-schedule-detail",
        ".report-nav-schedule-steps-preview",
        ".report-nav-schedule-steps-popover",
    ]:
        assert f"#page-report-navigation {selector}" in css
    assert "position: sticky;" in css
    assert "margin: 14px 0 0;" in css
    assert "overflow-x: hidden;" in css
    schedule_header_rule = re.search(
        r"#page-report-navigation \.report-nav-schedule-header\s*\{(?P<body>[^}]*)\}",
        css,
    )
    assert schedule_header_rule is not None
    assert "min-height: 58px;" in schedule_header_rule.group("body")
    assert ".report-nav-schedule-date-head.today {" not in css
    assert "#page-report-navigation .report-nav-schedule-today" not in css
    today_circle_rule = re.search(
        r"#page-report-navigation \.report-nav-schedule-date-head\.today b\s*\{(?P<body>[^}]*)\}",
        css,
    )
    assert today_circle_rule is not None
    assert "border-radius: 50%;" in today_circle_rule.group("body")
    assert "border: 1px solid var(--theme-accent);" in today_circle_rule.group("body")
    date_head_rule = re.search(
        r"#page-report-navigation \.report-nav-schedule-date-head\s*\{(?P<body>[^}]*)\}",
        css,
    )
    assert date_head_rule is not None
    assert "grid-template-rows: 28px 14px;" in date_head_rule.group("body")
    for selector in [
        ".report-nav-schedule-date-head b",
        ".report-nav-schedule-date-head em",
    ]:
        date_text_rule = re.search(
            rf"#page-report-navigation \{selector}\s*\{{(?P<body>[^}}]*)\}}",
            css,
        )
        assert date_text_rule is not None
        assert "display: grid;" in date_text_rule.group("body")
        assert "place-items: center;" in date_text_rule.group("body")
    assert "grid-template-columns: repeat(var(--report-nav-schedule-day-count), minmax(0, 1fr));" in css
    schedule_grid_rule = re.search(
        r"#page-report-navigation \.report-nav-schedule-header,\s*"
        r"#page-report-navigation \.report-nav-schedule-row\s*\{(?P<body>[^}]*)\}",
        css,
    )
    assert schedule_grid_rule is not None
    assert "gap: 0;" in schedule_grid_rule.group("body")
    assert "align-items: stretch;" in schedule_grid_rule.group("body")
    schedule_summary_title_rule = re.search(
        r"#page-report-navigation \.report-nav-schedule-summary strong\s*\{(?P<body>[^}]*)\}",
        css,
    )
    assert schedule_summary_title_rule is not None
    assert "font-size: 12px;" in schedule_summary_title_rule.group("body")
    for declaration in [
        "overflow: visible;",
        "overflow-wrap: anywhere;",
        "text-overflow: clip;",
        "white-space: normal;",
    ]:
        assert declaration in schedule_summary_title_rule.group("body")
    schedule_summary_percent_rule = re.search(
        r"#page-report-navigation \.report-nav-schedule-summary span\s*\{(?P<body>[^}]*)\}",
        css,
    )
    assert schedule_summary_percent_rule is not None
    assert "color: var(--outline);" in schedule_summary_percent_rule.group("body")
    assert "--report-nav-schedule-baseline-color: var(--outline-variant);" in css
    schedule_hover_rule = re.search(
        r"#page-report-navigation \.report-nav-schedule-row:hover:not\(\.selected\),\s*"
        r"#page-report-navigation \.report-nav-schedule-row:focus-visible:not\(\.selected\)\s*\{(?P<body>[^}]*)\}",
        css,
    )
    assert schedule_hover_rule is not None
    assert "--report-nav-schedule-baseline-color: var(--outline-variant);" in schedule_hover_rule.group("body")
    assert "background: linear-gradient(" in schedule_hover_rule.group("body")
    assert "color-mix(in srgb, var(--theme-accent) 12%, var(--surface-container-lowest))" in schedule_hover_rule.group("body")
    assert "background: var(--theme-focus-ring);" not in schedule_hover_rule.group("body")
    assert "box-shadow: none;" in schedule_hover_rule.group("body")
    assert ".report-nav-schedule-row:hover:not(.selected)::after" in css
    assert ".report-nav-schedule-row:focus-visible:not(.selected)::after" in css
    assert "width: 3px;" in css
    assert ".report-nav-schedule-row:hover:not(.selected) .report-nav-schedule-summary strong" in css
    schedule_hover_dot_rule = re.search(
        r"#page-report-navigation \.report-nav-schedule-row:hover:not\(\.selected\) \.report-nav-schedule-dots > i\.reached,\s*"
        r"#page-report-navigation \.report-nav-schedule-row:focus-visible:not\(\.selected\) \.report-nav-schedule-dots > i\.reached\s*\{(?P<body>[^}]*)\}",
        css,
    )
    assert schedule_hover_dot_rule is not None
    assert "transform: scale(1.35);" in schedule_hover_dot_rule.group("body")
    assert "box-shadow" not in schedule_hover_dot_rule.group("body")
    assert "color-mix(in srgb, var(--theme-accent) 14%, transparent)" not in schedule_hover_dot_rule.group("body")
    schedule_selected_rule = re.search(
        r"#page-report-navigation \.report-nav-schedule-row\.selected\s*\{(?P<body>[^}]*)\}",
        css,
    )
    assert schedule_selected_rule is not None
    assert "--report-nav-schedule-baseline-color: var(--outline-variant);" in schedule_selected_rule.group("body")
    assert "background: transparent;" in schedule_selected_rule.group("body")
    schedule_selected_outline_rule = re.search(
        r"#page-report-navigation \.report-nav-schedule-row\.selected::after\s*\{(?P<body>[^}]*)\}",
        css,
    )
    assert schedule_selected_outline_rule is not None
    assert "border: 1px solid var(--theme-accent);" in schedule_selected_outline_rule.group("body")
    assert "border-bottom: 0;" in schedule_selected_outline_rule.group("body")
    assert "transform: scale(1.35);" in css
    assert "--report-nav-schedule-edge: calc(50% / var(--report-nav-schedule-day-count));" in css
    assert "left: var(--report-nav-schedule-fill);" in css
    assert "left: var(--report-nav-schedule-edge);" in css
    assert "right: var(--report-nav-schedule-edge);" in css
    assert "calc(var(--report-nav-schedule-fill) - var(--report-nav-schedule-edge))" in css
    schedule_line_rule = re.search(
        r"#page-report-navigation \.report-nav-schedule-baseline,\s*"
        r"#page-report-navigation \.report-nav-schedule-fill\s*\{(?P<body>[^}]*)\}",
        css,
    )
    assert schedule_line_rule is not None
    assert "height: 3px;" in schedule_line_rule.group("body")
    assert "var(--report-nav-schedule-baseline-color) 0 5px" in css
    assert "var(--report-nav-schedule-early-tail-end)" not in css
    assert "#page-report-navigation .report-nav-schedule-dots > i.reached" in css
    assert "background: var(--report-nav-schedule-state);" in css
    assert ".report-nav-schedule-row.pending," not in css
    assert "--report-nav-schedule-state: var(--theme-accent);" in css
    assert ".report-nav-schedule-row.early-completed" in css
    assert ".report-nav-schedule-row.overdue-completed," in css
    assert "left: var(--report-nav-schedule-endpoint);" in css
    running_deadline_rule = re.search(
        r"#page-report-navigation \.report-nav-schedule-row\.running "
        r"\.report-nav-schedule-endpoint\s*\{(?P<body>[^}]*)\}",
        css,
    )
    assert running_deadline_rule is not None
    assert "border-color: var(--report-nav-danger);" in running_deadline_rule.group("body")
    assert "const endpointAtLastDate = endpointIndex >= dates.length - 1;" in app_js
    assert "const endpointAtFirstDate = endpointIndex <= 0;" in app_js
    assert "const endpointNearLastDate = endpointIndex === dates.length - 2;" in app_js
    assert "const endpointNearFirstDate = endpointIndex === 1;" in app_js
    assert "const deadlineAtLastDate" not in app_js
    assert "const deadlineAtFirstDate" not in app_js
    assert "const deadlineNearLastDate" not in app_js
    assert "const deadlineNearFirstDate" not in app_js
    assert '${endpointNearLastDate ? " near-last" : ""}' in app_js
    assert "#page-report-navigation .report-nav-schedule-endpoint.at-first em" in css
    assert "#page-report-navigation .report-nav-schedule-endpoint.at-last em" in css
    assert "#page-report-navigation .report-nav-schedule-endpoint.near-first em" in css
    assert "#page-report-navigation .report-nav-schedule-endpoint.near-last em" in css
    assert "translateX(calc(-50% - 14px));" in css
    assert "right: 50%;" in css
    assert "text-align: right;" in css
    assert "#page-report-navigation .report-nav-schedule-deadline-warning" in css
    assert "left: var(--report-nav-schedule-deadline);" in css
    schedule_deadline_warning_rule = re.search(
        r"#page-report-navigation \.report-nav-schedule-deadline-warning\s*\{(?P<body>[^}]*)\}",
        css,
    )
    assert schedule_deadline_warning_rule is not None
    assert "width: 14px;" in schedule_deadline_warning_rule.group("body")
    assert "height: 14px;" in schedule_deadline_warning_rule.group("body")
    assert "#page-report-navigation .report-nav-schedule-deadline-warning em" not in css
    assert "#page-report-navigation .report-nav-schedule-row.overdue .report-nav-schedule-endpoint" in css
    overdue_endpoint_size_rule = re.search(
        r"#page-report-navigation \.report-nav-schedule-endpoint\s*\{(?P<body>[^}]*)\}",
        css,
    )
    assert overdue_endpoint_size_rule is not None
    assert "width: 14px;" in overdue_endpoint_size_rule.group("body")
    assert "height: 14px;" in overdue_endpoint_size_rule.group("body")
    overdue_endpoint_color_rule = re.search(
        r"#page-report-navigation \.report-nav-schedule-row\.overdue "
        r"\.report-nav-schedule-endpoint,\s*"
        r"#page-report-navigation \.report-nav-schedule-row\.risk "
        r"\.report-nav-schedule-endpoint\s*\{(?P<body>[^}]*)\}",
        css,
    )
    assert overdue_endpoint_color_rule is not None
    assert "width:" not in overdue_endpoint_color_rule.group("body")
    assert "height:" not in overdue_endpoint_color_rule.group("body")
    assert "#page-report-navigation .report-nav-schedule-early-tail" not in css
    assert "#page-report-navigation .report-nav-schedule-early-deadline" not in css
    assert ".report-nav-schedule-row.early-completed .report-nav-schedule-baseline" not in css
    assert "#page-report-navigation .report-nav-schedule-dots > i.after-early-completion" not in css
    early_original_deadline_rule = re.search(
        r"#page-report-navigation \.report-nav-schedule-dots > i\.early-original-deadline\s*"
        r"\{(?P<body>[^}]*)\}",
        css,
    )
    assert early_original_deadline_rule is not None
    assert "width: 14px;" in early_original_deadline_rule.group("body")
    assert "height: 14px;" in early_original_deadline_rule.group("body")
    assert "border: 2px solid var(--report-nav-danger);" in early_original_deadline_rule.group("body")
    assert "#page-report-navigation .report-nav-schedule-original-deadline-label" not in css
    assert ".report-nav-schedule-row.overdue .report-nav-schedule-fill" in css
    assert ".report-nav-schedule-row.overdue-completed .report-nav-schedule-fill" in css
    assert "var(--theme-accent) 0 var(--report-nav-schedule-overdue-stop)" in css
    assert "var(--report-nav-danger) var(--report-nav-schedule-overdue-stop) 100%" in css
    assert ".report-nav-schedule-dots > i.reached.before-deadline" in css
    assert ".report-nav-schedule-dots > i.reached.after-deadline" in css
    assert "#page-report-navigation .report-nav-schedule-row.completed .report-nav-schedule-endpoint em" in css
    schedule_endpoint_label_rule = re.search(
        r"#page-report-navigation \.report-nav-schedule-endpoint em\s*\{(?P<body>[^}]*)\}",
        css,
    )
    assert schedule_endpoint_label_rule is not None
    for declaration in [
        "bottom: calc(100% + 5px);",
        "left: 50%;",
        "text-align: center;",
        "white-space: nowrap;",
        "transform: translateX(-50%);",
    ]:
        assert declaration in schedule_endpoint_label_rule.group("body")
    assert "left: 18px;" not in schedule_endpoint_label_rule.group("body")
    report_navigation_main_spacing_rule = re.search(
        r':root\[data-page="report-navigation"\] \.main-content\s*\{(?P<body>[^}]*)\}',
        css,
    )
    assert report_navigation_main_spacing_rule is not None
    assert "padding-bottom: 14px;" in report_navigation_main_spacing_rule.group("body")
    report_navigation_page_spacing_rule = re.search(
        r':root\[data-page="report-navigation"\] #page-report-navigation\s*\{(?P<body>[^}]*)\}',
        css,
    )
    assert report_navigation_page_spacing_rule is not None
    assert "padding-bottom: 0;" in report_navigation_page_spacing_rule.group("body")
    schedule_detail_rule = re.search(
        r"#page-report-navigation \.report-nav-schedule-detail\s*\{(?P<body>[^}]*)\}",
        css,
    )
    assert schedule_detail_rule is not None
    assert "grid-template-columns: 172px 150px 300px minmax(300px, 1fr) auto;" in schedule_detail_rule.group("body")
    assert "min-height: 64px;" in schedule_detail_rule.group("body")
    assert "padding: 8px 18px;" in schedule_detail_rule.group("body")
    assert "background: transparent;" in schedule_detail_rule.group("body")
    assert "animation: report-nav-schedule-detail-expand 160ms ease-out both;" in schedule_detail_rule.group("body")
    assert "@keyframes report-nav-schedule-detail-expand" in css
    assert "clip-path: inset(0 0 100% 0);" in css
    assert "clip-path: inset(0);" in css
    assert "#page-report-navigation .report-nav-schedule-row.selected::after" in css
    assert "#page-report-navigation .report-nav-schedule-detail::before" in css
    assert "border: 1px solid var(--theme-accent);" in css
    schedule_detail_label_rule = re.search(
        r"#page-report-navigation \.report-nav-schedule-detail > div > span\s*\{(?P<body>[^}]*)\}",
        css,
    )
    assert schedule_detail_label_rule is not None
    assert "color: var(--outline);" in schedule_detail_label_rule.group("body")
    assert "font-size: 12px;" in schedule_detail_label_rule.group("body")
    assert "font-weight: 400;" in schedule_detail_label_rule.group("body")
    schedule_detail_text_rule = re.search(
        r"#page-report-navigation \.report-nav-schedule-detail > div > strong\s*\{(?P<body>[^}]*)\}",
        css,
    )
    assert schedule_detail_text_rule is not None
    assert "color: var(--on-surface);" in schedule_detail_text_rule.group("body")
    assert "font-size: 12px;" in schedule_detail_text_rule.group("body")
    assert "font-weight: 400;" in schedule_detail_text_rule.group("body")
    assert ".report-nav-schedule-steps-preview.open .report-nav-schedule-steps-popover" in css
    schedule_steps_trigger_rule = re.search(
        r"#page-report-navigation \.report-nav-schedule-steps-preview\s*\{(?P<body>[^}]*)\}",
        css,
    )
    assert schedule_steps_trigger_rule is not None
    assert "border: 1px solid var(--outline-variant);" in schedule_steps_trigger_rule.group("body")
    assert "border: 1px solid var(--theme-accent);" not in schedule_steps_trigger_rule.group("body")
    assert "background: var(--surface-container-lowest);" in schedule_steps_trigger_rule.group("body")
    assert "box-shadow:" in schedule_steps_trigger_rule.group("body")
    assert ".report-nav-schedule-steps-preview.open::before" in css
    assert "z-index: 2147483000;" in css
    assert "pointer-events: auto;" in css
    assert "@keyframes report-nav-schedule-step-ping" in css
    assert "position: fixed;" in re.search(
        r"#page-report-navigation \.report-nav-schedule-steps-popover\s*\{(?P<body>[^}]*)\}",
        css,
    ).group("body")
    assert "鼠标悬浮“查看步骤”以状态面板展示该流程的全部步骤" in _read(README_MD)
    assert "提前完成使用绿色并在实际完成日结束" in _read(README_MD)
    assert "实际完成日后保留原灰色虚线和灰色圆点" in _read(README_MD)
    assert "原截止日仅显示与完成勾同尺寸的红色空心圈" in _read(README_MD)
    assert "截止日及之前保持主题蓝" in _read(README_MD)
    assert "超过截止日后改为红色" in _read(README_MD)
    assert "不显示今天文字标签，保留当天日期圆圈" in _read(README_MD)
    assert "鼠标悬浮日程行时采用参考页 01 的浅蓝横向渐变底" in _read(README_MD)
    assert "--report-nav-schedule-table-width" not in app_js
    assert "border-color: var(--theme-accent);" in css
    assert ".report-nav-schedule-date-head.holiday" in css
    assert "color: var(--report-nav-danger);" in css


def test_report_navigation_step_popover_uses_clickable_status_for_manual_confirmation():
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert 'class="report-nav-schedule-step-status-action"' in app_js
    assert 'data-report-nav-step-action="${manualAction}"' in app_js
    assert 'data-report-nav-step-code="${escapeHtml(String(step.step_code || ""))}"' in app_js
    assert 'data-report-nav-step-row="${escapeHtml(String(step.step_code || ""))}"' in app_js
    assert 'step.manual_completion_allowed && (!completed || step.manual_completed)' in app_js
    assert "setReportNavigationManualStepState" in app_js
    assert "await loadReportNavigation({ preserveSchedule: true });" in app_js
    assert "function updateReportNavigationScheduleStepsPopover" in app_js
    update_body = app_js.split(
        "function updateReportNavigationScheduleStepsPopover", 1
    )[1].split("function flushDeferredReportNavigationScheduleRender", 1)[0]
    assert "list.innerHTML" not in update_body
    assert "preview.focus" not in update_body
    assert "statusControl.textContent = stepStatusText;" in update_body
    assert "row.className = stepState;" in update_body
    assert "function updateReportNavigationScheduleProcessSummary" in app_js
    assert "summaryPercent.textContent = `${percent}%`;" in app_js
    assert "progressValue.textContent = `${percent}%（${completedSteps}/${totalSteps}）`;" in app_js
    assert 'nextStepValue.textContent = nextStep?.step_name || "全部步骤已完成";' in app_js
    action_body = app_js.split(
        "async function setReportNavigationManualStepState", 1
    )[1].split("async function editReportNavigationScheduleOwner", 1)[0]
    assert "control.disabled = true;" not in action_body
    assert 'control?.dataset.reportNavStepPending === "true"' in action_body
    assert 'control.dataset.reportNavStepPending = "true";' in action_body
    assert 'delete control.dataset.reportNavStepPending;' in action_body
    assert "updateReportNavigationScheduleProcessSummary(processCode, payload);" in action_body
    assert "if (!preserveSchedule) renderReportNavigationSchedule(payload);" in app_js
    assert "reportNavigationScheduleRenderDeferred = true;" in app_js
    assert "flushDeferredReportNavigationScheduleRender();" in app_js
    assert "reopenReportNavigationScheduleStepsPreview" not in app_js
    assert "reportNavigationScheduleStepsKeepOpenProcessCode" not in app_js
    assert "#page-report-navigation .report-nav-schedule-step-status-action" in css
    assert "text-decoration: underline;" in css
    assert "cursor: pointer;" in css


def test_report_navigation_manual_refresh_has_icon_cooldown_and_error_feedback():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert 'id="reportNavRefreshButton"' in html
    assert 'id="reportNavRefreshCountdown"' in html
    assert 'class="report-nav-refresh-icon"' in html
    assert 'aria-label="立即刷新报送导航统计"' in html
    assert 'api("/api/report-navigation/refresh", {' in app_js
    assert 'method: "POST"' in app_js
    assert "const REPORT_NAV_REFRESH_COOLDOWN_SECONDS = 300;" in app_js
    assert "function setReportNavigationRefreshCooldown(seconds)" in app_js
    assert "result.retry_after_seconds ?? result.cooldown_seconds" in app_js
    assert "reportNavRefreshButton?.addEventListener(\"click\"" in app_js
    assert 'reportNavRefreshButton.classList.add("refreshing")' in app_js
    assert "error.payload?.retry_after_seconds" in app_js
    assert "报送导航刷新失败" in app_js
    assert "reportNavRefreshCountdown.textContent" in app_js
    assert 'String(minutes).padStart(2, "0")' in app_js
    assert "function renderReportNavigationRefreshIssues(issues = [])" in app_js
    assert 'showInfo("报送导航统计异常"' in app_js
    assert "result.issues || []" in app_js
    assert "普通用户 5 分钟可见倒计时、管理员免冷却" in app_js
    assert "管理员不受冷却限制" in _read(README_MD)
    assert "#page-report-navigation .report-nav-refresh-button" in css
    assert "#page-report-navigation .report-nav-refresh-countdown" in css
    assert "\n.report-nav-refresh-issues {" in css
    assert "#page-report-navigation .report-nav-refresh-button.refreshing svg" in css
    assert "@keyframes report-nav-refresh-spin" in css


def test_report_navigation_docs_changelog_and_page_titles_are_updated():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    readme = _read(README_MD)

    for title in ["对数总览", "对数执行", "对数历史"]:
        assert f"<h2>{title}</h2>" in html

    assert "- 报送导航：" in readme
    assert "- 智能核数：" in readme
    assert "默认进入“报送导航”" in readme

    current_changelog = re.search(
        r'<span class="changelog-version">v1\.2\.13</span>(?P<body>.*?)</div>\s*<div class="changelog-item">',
        app_js,
        re.S,
    )
    assert current_changelog is not None
    assert "新增报送导航状态定时统计" in current_changelog.group("body")
    assert "新增智能核数多级菜单" in current_changelog.group("body")
    assert "系统优化及BUG修复。" in current_changelog.group("body")


def test_home_dashboard_uses_clickable_reconcile_stats_and_keeps_line_charts():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)
    readme = _read(README_MD)

    for label, stat_id, stat_key in [
        ("总差异数", "homeStatTotalDiff", "total"),
        ("未解释", "homeStatUnresolved", "unresolved"),
        ("已解释", "homeStatExplained", "explained"),
        ("实收本金不一致", "homeStatPaidIn", "paidIn"),
        ("标的代码不一致", "homeStatTargetCode", "targetCode"),
        ("报告期", "homeStatReportPeriod", "reportPeriod"),
    ]:
        assert label in html
        assert f'id="{stat_id}"' in html
        assert f'data-home-stat="{stat_key}"' in html

    assert 'id="homeStatReportRunAt"' in html
    for delta_id in [
        "homeStatTotalDiffDelta",
        "homeStatUnresolvedDelta",
        "homeStatExplainedDelta",
        "homeStatPaidInDelta",
        "homeStatTargetCodeDelta",
    ]:
        assert f'id="{delta_id}"' in html
    assert html.index('data-home-stat="reportPeriod"') < html.index('data-home-stat="total"')
    assert 'id="homeQualityScope"' not in html
    assert 'id="homeReasonScope"' not in html
    assert 'id="homeFocusScope"' in html
    assert 'id="homeQualityScore"' not in html
    assert 'id="homeQualityTag"' not in html
    assert "质量分" not in html
    assert "home-quality-ring" not in html
    assert "home-quality-tag" not in html
    assert 'class="home-quality-body">\n                <div class="home-quality-bars"' in html
    assert '"homeQualityScore"' not in app_js
    assert "periodExplainedPct" not in app_js
    assert ".home-quality-ring" not in css
    assert ".home-quality-tag" not in css
    assert 'id="chartCanvas"' in html
    assert 'id="trendCanvas"' in html
    assert "执行趋势" in html
    assert "最新趋势" not in html
    assert "多指标统计" not in html
    assert 'card-title card-title--trend' in html
    assert 'class="trend-legend" aria-label="指标颜色说明"' in html
    assert "trend-legend-dot--first-run" in html
    assert "trend-legend-dot--avg" not in html
    assert "trend-legend-dot--runs" in html
    assert "近12期" in html
    assert "差异类型分布" in html
    assert "高频差异项目" in html
    assert "重点差异项目" not in html
    assert "drawGlassChart(canvas, values, labels, renderChartAnimId, true, tooltipItems)" in app_js
    assert "tooltipItems = []" in app_js
    assert "差异数: ${formatMoney(r.total_count || 0)}" in app_js
    assert "function drawGlassMultiMetricChart" in app_js
    assert "drawGlassMultiMetricChart(canvas, [" in app_js
    assert "function trendFirstRunMetricStyle()" in app_js
    assert "const palette = canvasThemePalette();" in app_js
    assert 'color: firstRunStyle.color' in app_js
    assert 'shadow: firstRunStyle.shadow' in app_js
    assert "function refreshHomeChartsForTheme()" in app_js
    assert "refreshHomeChartsForTheme();" in app_js
    assert "const sharedMaxVal = Math.max(...series.flatMap((metric) => metric.values), 1);" in app_js
    assert "series.forEach((metric) => { metric.maxVal = sharedMaxVal; });" in app_js
    assert "function drawLegend" not in app_js
    assert "drawLegend();" not in app_js
    assert "每期差异个数" in html
    assert "每期首次执行差异个数" not in html
    assert "每期平均差异个数" not in html
    assert "每期差异个数" in app_js
    assert "每期首次执行差异个数" not in app_js
    assert "每期平均差异个数" not in app_js
    assert "每期执行次数" in app_js
    assert "firstRunDiff" in app_js
    assert "averageDiff" not in app_js
    assert "executionCount" in app_js
    assert "totalDiff / executionCount" not in app_js
    assert "const firstRun = [...dateRuns].sort(compareHomeRunTimeAsc)[0];" in app_js
    assert "firstRunDiff: Number(firstRun?.total_count || 0)" in app_js
    assert "function formatMetricChartNumber(metric = {}, value)" in app_js
    assert "if (metric.integerValues) return String(Math.round(Number(value || 0)));" in app_js
    assert "const firstRunValues = filtered.map((item) => Math.round(item.firstRunDiff));" in app_js
    assert "integerValues: true" in app_js
    assert "function formatChartRunAtLabel(runAt = \"\", fallbackDate = \"\")" in app_js
    assert "const labels = dateRuns.map((r) => formatChartRunAtLabel(r.run_at, targetDate));" in app_js
    assert "bezierCurveTo" in app_js
    assert "chart-bar" not in html

    assert "const HOME_REASON_DEFS" in app_js
    assert "function homeDetailReasonText(details = [])" in app_js
    assert "function homeDisplayDetailReasonText(displayDetails = [])" in app_js
    assert "homeDetailReasonText(item.details)" in app_js
    assert "homeDisplayDetailReasonText(item.display_details)" in app_js
    assert '"specific_reason", "reason", "check_result", "reason_text", "basis"' in app_js
    assert "function homeSpecificReasonMatchesPaidIn(item = {})" in app_js
    assert 'text.includes("4001与c1000存在差异")' in app_js
    assert 'text.includes("4001-c1000差额正好解释主差异")' in app_js
    assert "function homeSpecificReasonMatchesTargetCode(item = {})" in app_js
    assert 'text.includes("fa/am标的不一致")' in app_js
    assert 'text.includes("fa与am标的不一致")' in app_js
    assert 'text.includes("fa和am标的不一致")' in app_js
    assert "function homeTargetCodeMismatchCount(item = {})" in app_js
    assert 'detail?.kind === "fa_am"' in app_js
    assert "summary.targetCode += homeTargetCodeMismatchCount(item);" in app_js
    assert "function homeReasonCategoryFromItem(item = {})" in app_js
    assert "function homeResultCountsAsUnresolved(item = {})" in app_js
    assert '["未解释", "候选不唯一"].includes(String(item.match_status || ""))' in app_js
    assert "const HOME_STATUS_ORDER" in app_js
    assert "buildHomeResultGroups(results)" in app_js
    assert "recentHomePeriodDates(runs, 12)" in app_js
    assert "homeRunsForPeriodDates(runs, recentPeriodDates)" in app_js
    assert "Promise.all(recentPeriodSummaries.map((run) => loadHomeRunDetail(run)))" in app_js
    assert "homeDifferenceTypeParts" in app_js
    assert "function firstHomeRunsForPeriodDates(runs = [], dates = [])" in app_js
    assert "compareHomeRunTimeAsc(run, current) < 0" in app_js
    assert "function aggregateHomeSummaryForRuns(runs = [], summaryBuilder = () => ({}))" not in app_js
    assert "const currentReportFirstRun = firstHomeRunsForPeriodDates(recentPeriodRuns, [latestRun.run_date])[0] || latestRun;" in app_js
    assert "const currentReportStatusCounts = homeStatusCountsForRun(currentReportFirstRun);" in app_js
    assert "const currentReportTypeSummary = homeDifferenceTypeSummaryForRun(currentReportFirstRun);" in app_js
    assert "averageHomeSummaryByPeriod" not in app_js
    assert "formatHomeRoundedCount(item.count)" in app_js
    assert "renderHomeQualityRows(\n      currentReportStatusCounts," in app_js
    assert "renderHomeReasonList(\n      currentReportTypeSummary," in app_js
    assert "buildHomeFrequencyItems(recentPeriodRuns, recentPeriodDates)" in app_js
    assert "renderHomeFrequencyList(frequencyItems, recentPeriodDates.length)" in app_js
    assert "const totalText = `近${periodCount}期 ${item.periodCount}次`;" in app_js
    assert "按报告期去重累计 ${item.periodCount} 次" in app_js
    assert "至少 2 期后分析高频项目" in app_js
    assert "等待首次核对后生成质量分布" in app_js
    assert "等待首次核对后统计差异类型" in app_js
    assert "const periodScopeText = `近${recentPeriodDates.length}期`;" not in app_js
    assert 'document.getElementById("homeQualityScope")' not in app_js
    assert 'document.getElementById("homeReasonScope")' not in app_js
    assert 'data-home-stat="periodExplained"' not in html
    assert 'data-home-stat="periodUnresolved"' not in html
    assert 'data-home-stat="frequent"' not in html
    assert "homeStatsState = {" in app_js
    assert "function findHomeStatsBaselineRun" in app_js
    assert "return { run: samePeriodRuns[currentIndex - 1], label: \"较上次\" };" in app_js
    assert "const previousDate = [...runs]" in app_js
    assert "String(date) < currentDate" in app_js
    assert "return previousPeriodRun ? { run: previousPeriodRun, label: \"较上期\" } : { run: null, label: \"较上期\" };" in app_js
    assert "const deltaText = delta >= 0 ? `+${delta}` : String(delta);" in app_js
    assert "el.hidden = true;" in app_js
    assert "el.textContent = \"\";" in app_js
    assert "el.hidden = false;" in app_js
    assert "暂无对比" not in app_js
    assert 'el.innerHTML = `${escapeHtml(label)} <span class="home-stat-delta-value">${escapeHtml(deltaText)}</span>`;' in app_js
    assert "renderHomeStatDeltas(" in app_js
    assert "counts: {" in app_js
    assert "showHomeStatResults" in app_js
    assert "reportRuns: reportPeriodRuns" in app_js
    assert "run.run_date === latestRun.run_date" in app_js
    assert "function summarizeHomeRunForReport(run = {})" in app_js
    assert "function renderHomeReportPeriodTable(periodRuns = [])" in app_js
    assert '<th class="col-run-at">执行时间</th>' in app_js
    assert '<th class="col-total">差异数</th>' in app_js
    assert '<th class="col-paid-in">实收本金不一致</th>' in app_js
    assert '<th class="col-target-code">标的代码不一致</th>' in app_js
    assert 'const isReportPeriod = key === "reportPeriod";' in app_js
    assert 'const modalTitle = isReportPeriod ? "报送期差异数详情" : `${label}项目明细`;' in app_js
    assert "报送期 ${run?.run_date || \"--\"}，共 ${reportRuns.length} 次执行，按执行时间倒序。" in app_js
    assert 'id="infoDetailAction"' in html
    assert "HOME_TOP_STAT_KEYS" in app_js
    assert "openHomeStatResultList" in app_js
    assert 'trigger.closest(".home-stats-row")' in app_js
    assert "detailActionLabel: \"查看明细\"" in app_js
    assert 'const reasonValue = `home-category:${key}`' in app_js
    assert 'const reasonValue = "home-status:unresolved";' in app_js
    assert 'ensureSelectOption(reasonFilter, reasonValue, "未解释/候选不唯一");' in app_js
    assert 'String(selectedReason) === "home-status:unresolved"' in app_js
    assert '["paidIn", "targetCode"].includes(key)' in app_js
    assert "resultMatchesReasonFilter(item, reason)" in app_js
    assert "resultFilterHint" in app_js
    assert "结果列表已筛选" in app_js
    assert "clearHomeResultFilter" in app_js
    assert "resultRestoreHistoryMeta" in app_js
    assert "结果列表已恢复到历史数据：报告期" in app_js
    assert "restoreLatestResults" in app_js
    assert "回到最新结果" in app_js
    assert "function restoreLatestResultsToResultList()" in app_js
    assert 'setStatus("结果列表已还原到最新结果")' in app_js
    assert 'showToast("结果列表已还原到最新结果", "success")' in app_js
    assert "function hasActiveResultListFilter()" in app_js
    assert "const hadExistingFilter = hasActiveResultListFilter();" in app_js
    assert 'applyHomeResultListFilter(key, { hadExistingFilter });' in app_js
    assert 'key === "total"' in app_js
    assert 'homeResultListFilterLabel = hadExistingFilter ? (HOME_STAT_LABELS[key] || "") : "";' in app_js
    assert '<span class="result-filter-hint" id="resultFilterHint" hidden></span>' in html
    assert 'home-focus-item home-stat-click-target' not in app_js
    assert 'data-home-stat="${escapeHtml(statKey)}"' not in app_js
    assert "const fullName" in app_js
    assert 'title="${escapeHtml(itemTitle)}"' in app_js
    assert "home-focus-title-row" in app_js
    assert "home-focus-total" in app_js
    assert 'const totalText = `近${periodCount}期 ${item.periodCount}次`;' in app_js
    assert 'const detailText = `${periodText} · 最近类型：${reason}`;' in app_js
    assert '<td class="col-name" title="${escapeHtml(nameText)}">' in app_js
    assert '<th class="col-code">' in app_js
    assert '<th class="col-asset">' in app_js
    assert '<th class="col-liability">' in app_js
    assert '<th class="col-specific">' not in app_js
    assert '<td class="col-code" title="${escapeHtml(codeText)}">' in app_js
    assert "home-stat-modal-table" in css
    assert "table-layout: fixed;" in css
    assert ".home-stat-modal-table .col-asset" in css
    assert ".home-stat-modal-table .col-liability" in css
    assert ".home-stat-modal-table .col-run-at" in css
    assert ".home-stat-modal-table .col-target-code" in css
    assert ".home-stat-modal-table .col-specific" not in css
    assert ".home-stat-modal-table td.num {\n  text-align: left;" in css
    assert '[data-color-mode="dark"] .home-stat-modal-table-wrap' in css
    assert '[data-color-mode="dark"] .home-stat-modal-table th' in css
    assert '[data-color-mode="dark"] .home-stat-modal-table td' in css
    assert ".info-detail-action" in css
    assert ".result-filter-hint" in css
    assert ".result-filter-clear" in css
    assert ".home-stat-subvalue" in css
    assert ".home-stat-delta" in css
    assert ".home-stat-delta-value" in css
    assert ".home-stat-delta--up" in css
    assert ".home-stat-delta--down" in css
    assert ".result-card .card-title-left {\n  display: flex;" in css
    assert ".trend-legend" in css
    assert ".card-title--trend .card-title-icon" in css
    assert "display: none;" in css
    assert ".trend-legend-dot--first-run" in css
    assert ".trend-legend-dot--avg" not in css
    assert "background: linear-gradient(90deg, var(--secondary), var(--on-secondary-container))" in css
    assert '[data-theme="space-tech"] .trend-legend-dot--first-run' in css
    assert "background: linear-gradient(90deg, #3b82f6, #06b6d4)" in css
    assert ".trend-legend-dot--runs" in css
    assert "grid-template-columns: minmax(0, 0.92fr) minmax(0, 1fr) minmax(0, 1.08fr)" in css
    assert ".home-quality-track--sys" in css
    assert "grid-template-columns: 22px minmax(0, 1fr)" in css
    assert ".home-focus-title-row" in css
    assert ".home-focus-total" in css
    assert ".home-focus-detail" in css
    assert ".home-analysis-card:hover" in css
    assert "transform: translateY(-2px);" in css
    home_content_surface = re.search(
        r"#page-home \.glass-card,\s*"
        r"#page-home \.glass-stat-card\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert home_content_surface is not None
    assert "background: var(--surface-container-lowest) !important;" in home_content_surface.group("body")
    assert "border: 1px solid var(--outline-variant) !important;" in home_content_surface.group("body")
    assert ':root[data-page="home"] body' in css
    assert "grid-template-columns: repeat(6, minmax(0, 1fr))" in css
    page_home_rule = re.search(r"(?m)^#page-home\s*\{(?P<body>.*?)\}", css, re.S)
    home_grid_rule = re.search(r"(?m)^\.home-grid\s*\{(?P<body>.*?)\}", css, re.S)
    assert page_home_rule is not None
    assert home_grid_rule is not None
    assert "--home-card-glow-gutter: 12px;" in page_home_rule.group("body")
    assert "overflow: visible;" in page_home_rule.group("body")
    assert "padding: var(--home-card-glow-gutter);" in home_grid_rule.group("body")
    assert "margin: calc(-1 * var(--home-card-glow-gutter));" in home_grid_rule.group("body")
    assert "overflow: visible;" in home_grid_rule.group("body")
    home_glass_shadow = re.search(
        r'\[data-theme="space-tech"\] #page-home \.glass-card,\s*'
        r'\[data-theme="space-tech"\] #page-home \.glass-stat-card\s*'
        r'\{(?P<body>.*?)\}',
        css,
        re.S,
    )
    assert home_glass_shadow is not None
    assert "box-shadow: var(--space-panel-shadow) !important;" in home_glass_shadow.group("body")
    home_glass_hover_shadow = re.search(
        r'\[data-theme="space-tech"\] #page-home \.glass-card:hover,\s*'
        r'\[data-theme="space-tech"\] #page-home \.glass-stat-card:hover\s*'
        r'\{(?P<body>.*?)\}',
        css,
        re.S,
    )
    assert home_glass_hover_shadow is not None
    assert "box-shadow: var(--space-panel-shadow) !important;" in home_glass_hover_shadow.group("body")
    assert "首页调整为自动对数概览工作台" in readme
    assert "对数质量和差异类型分布不再展示统计期数" in readme
    assert "高频差异项目继续展示实际统计期数" in readme
    assert "同报告期第二次及以后执行对比同报告期上一次执行" in readme
    assert "当期首次执行且存在上一报告期时显示“较上期”" in readme
    assert "首页执行趋势改用“月/日 时:分”展示横轴时间并支持悬浮查看执行时间和差异数" in readme
    assert "无任何对比基准时隐藏小字" in readme
    assert "多指标统计按核对日期展示每期差异个数" in readme
    assert "按每期首次执行记录取数" in readme
    assert "多指标统计按核对日期真实计算每期平均差异个数" not in readme
    assert "首页高频差异项目恢复连续/出现期数小字" in readme
    assert "项目名称右侧展示近 X 期按报告期去重后的出现次数" in readme
    assert "首页“实收本金不一致”仅统计具体原因中 `4001 - c1000` 差额正好解释主差异的项目" in readme
    assert "“标的代码不一致”仅统计具体原因中包含 FA/AM 标的不一致的项目" in readme
    assert "首页统计中“候选不唯一”归入“未解释”" in readme
    assert "不影响结果列表原始状态、历史和导出" in readme
    assert "首页报送期统计弹框改为展示该报送期内全部执行记录" in readme
    assert "顶部首位展示报告期统计卡片" in readme
    assert "沉稳主题下差异个数线和图例使用沉稳主题色" in readme
    assert "在“结果详情”标题后同行展示筛选说明" in readme
    assert "总差异数跳转时若结果列表原本已是全部差异则不显示筛选说明" in readme
    assert "对数质量和差异类型分布只统计当前报告期第一次执行" in readme
    assert "高频差异项目仍最多统计近 12 个报告期" in readme
    assert "对数质量移除质量分、圆环和评价标签" in readme
    assert "每期全部执行次数先取平均后汇总" not in readme
    assert "同期前序执行中出现过的差异仍纳入平均统计" not in readme
    assert "按每期首次执行记录取数并按整数展示" in readme
    assert "顶部统计项可查看项目明细并跳转到自动对数结果列表自动筛选" in readme
    assert "长项目名称在边界内省略并支持鼠标悬浮查看全称" in readme


def test_readme_documents_bounded_reconcile_matching_and_reference_codes():
    readme = _read(README_MD)
    app_js = _read(APP_JS)

    assert "50～100 行候选池仍可快速匹配 2～5 条组合" in readme
    assert "偶发 20～30 条共同命中" in readme
    assert "同一项目组合计算最长 60 秒" in readme
    assert "JS0508-2.51" in readme
    assert "空串和纯空格" in readme
    assert "系统优化及BUG修复。" in app_js


def _extract_home_baseline_functions(app_js: str) -> str:
    body = []
    for name in (
        "compareHomeRunTimeAsc",
        "homeRunsForPeriodDates",
        "firstHomeRunsForPeriodDates",
        "compareHomeRunsAsc",
        "isSameHomeRun",
        "findHomeStatsBaselineRun",
    ):
        block = re.search(rf"function {name}\(.*?\n\}}\n", app_js, re.S)
        assert block is not None, name
        body.append(block.group(0))
    return "\n".join(body)


def test_home_stats_baseline_uses_first_run_of_previous_period(tmp_path):
    app_js = _read(APP_JS)
    script = textwrap.dedent(
        """
        const assert = require("node:assert/strict");
        __HOME_BASELINE_FUNCS__
        const runs = [
          { id: "r1", run_date: "2026-05-31", run_at: "09:00:00" },
          { id: "r2", run_date: "2026-05-31", run_at: "18:00:00" },
          { id: "r3", run_date: "2026-06-30", run_at: "09:00:00" },
          { id: "r4", run_date: "2026-06-30", run_at: "18:00:00" },
        ];
        const firstRunBaseline = findHomeStatsBaselineRun(
          runs,
          { id: "r3", run_date: "2026-06-30", run_at: "09:00:00" }
        );
        assert.equal(firstRunBaseline.label, "较上期");
        assert.equal(firstRunBaseline.run.id, "r1");
        const laterRunBaseline = findHomeStatsBaselineRun(
          runs,
          { id: "r4", run_date: "2026-06-30", run_at: "18:00:00" }
        );
        assert.equal(laterRunBaseline.label, "较上次");
        assert.equal(laterRunBaseline.run.id, "r3");
        const noBaseline = findHomeStatsBaselineRun(runs, { id: "r5", run_date: "2026-04-30", run_at: "09:00:00" });
        assert.equal(noBaseline.label, "较上期");
        assert.equal(noBaseline.run, null);
        """
    ).replace("__HOME_BASELINE_FUNCS__", _extract_home_baseline_functions(app_js))
    script_path = tmp_path / "home_stats_baseline.cjs"
    script_path.write_text(script, encoding="utf-8")
    subprocess.run(["node", str(script_path)], check=True, cwd=ROOT)


def test_home_report_period_stat_card_fits_scale_ratio_changes():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert "home-stat-card--report-period" in html
    assert html.index("home-stat-card--report-period") < html.index('data-home-stat="total"')
    assert ".home-stat-card--report-period {\n  min-width: 0;\n}" in css
    assert ".home-stat-card--report-period .home-stat-value" in css
    assert "font-size: var(--home-report-period-font-size, 24px);" in css
    assert "const HOME_REPORT_PERIOD_MIN_FONT_SIZE = 16;" in app_js
    assert "const HOME_REPORT_PERIOD_MAX_FONT_SIZE = 25;" in app_js
    assert "function fitHomeReportPeriodValue()" in app_js
    assert 'value.style.setProperty("--home-report-period-font-size", `${size}px`);' in app_js
    assert "value.scrollWidth > value.clientWidth + 1" in app_js
    assert 'if (id === "homeStatReportPeriod") fitHomeReportPeriodValue();' in app_js
    assert 'window.addEventListener("resize", fitHomeReportPeriodValue);' in app_js
    assert "\u9996\u9875\u62a5\u544a\u671f\u7edf\u8ba1\u5361\u7247\u6309\u5b9e\u9645\u663e\u793a\u6bd4\u4f8b\u81ea\u9002\u5e94\u5b57\u53f7" in _read(README_MD)


def test_outer_content_panels_share_one_glow_free_hover_rule():
    css = _read(STYLES_CSS)
    readme = _read(README_MD)

    shared_base_rule = re.search(
        r"#page-home \.glass-card,\s*"
        r"#page-home \.glass-stat-card,\s*"
        r"#page-report-navigation \.report-nav-stats-layout,\s*"
        r"#page-report-navigation \.report-nav-flow-card,\s*"
        r"#page-report-navigation \.report-nav-schedule-card,\s*"
        r"#page-report-navigation \.report-nav-attention-card,\s*"
        r"#page-auto-check > \.card,\s*"
        r"#page-history > \.card,\s*"
        r"#page-settings \.settings-dashboard-card,\s*"
        r"#page-users \.user-stat-card,\s*"
        r"#page-users \.user-filter-bar,\s*"
        r"#page-users \.user-table-card\s*"
        r"\{(?P<body>[^}]*)\}",
        css,
    )
    assert shared_base_rule is not None
    assert "box-shadow: inset 0 1px 0 rgba(255, 255, 255, 0.56) !important;" in shared_base_rule.group("body")

    shared_hover_rule = re.search(
        r"#page-home \.glass-card:hover,\s*"
        r"#page-home \.glass-stat-card:hover,\s*"
        r"#page-report-navigation \.report-nav-stats-layout:hover,\s*"
        r"#page-report-navigation \.report-nav-flow-card:hover,\s*"
        r"#page-report-navigation \.report-nav-schedule-card:hover,\s*"
        r"#page-report-navigation \.report-nav-attention-card:hover,\s*"
        r"#page-auto-check > \.card:hover,\s*"
        r"#page-history > \.card:hover,\s*"
        r"#page-settings \.settings-dashboard-card:hover,\s*"
        r"#page-users \.user-stat-card:hover,\s*"
        r"#page-users \.user-filter-bar:hover,\s*"
        r"#page-users \.user-table-card:hover\s*"
        r"\{(?P<body>[^}]*)\}",
        css,
    )
    assert shared_hover_rule is not None
    hover_body = shared_hover_rule.group("body")
    assert "border-width: 1px !important;" in hover_body
    assert "border-color: color-mix(in srgb, var(--theme-accent) 36%, var(--outline-variant)) !important;" in hover_body
    assert "box-shadow: inset 0 1px 0 rgba(255, 255, 255, 0.56) !important;" in hover_body
    assert "transform: translateY(-2px) !important;" in hover_body
    assert "0 0" not in hover_body
    assert "rgba(15, 23, 42" not in hover_body
    assert "外层内容模块悬浮统一使用对数总览的 1px 浅主题描边" in readme
    assert "内容模块之间的缝隙不显示外投影" in readme


def test_home_charts_rerender_after_scale_ratio_changes():
    app_js = _read(APP_JS)
    readme = _read(README_MD)

    assert "let homeChartsResizeTimer = null;" in app_js
    assert "const HOME_CHARTS_RESIZE_DEBOUNCE_MS = 160;" in app_js
    assert "function scheduleHomeChartsResize()" in app_js
    assert 'document.documentElement.getAttribute("data-page") !== "home"' in app_js
    assert "window.clearTimeout(homeChartsResizeTimer);" in app_js
    assert "homeChartsResizeTimer = window.setTimeout(() => {" in app_js
    assert "renderChart();" in app_js
    assert "renderTrendChart();" in app_js
    assert 'window.addEventListener("resize", scheduleHomeChartsResize);' in app_js
    assert "canvas" in readme


def test_home_analysis_cards_keep_height_in_short_scale_ratio_viewports():
    css = _read(STYLES_CSS)
    readme = _read(README_MD)

    assert "@media (max-width: 1200px) and (max-height: 700px)" in css
    assert ":root[data-page=\"home\"] body" in css
    assert ":root[data-page=\"home\"] .main-content" in css
    assert "#page-home,\n  #page-home .home-grid" in css
    assert ".home-charts-row,\n  .home-analysis-row {\n    flex: none;" in css
    assert ".home-analysis-card {\n    height: auto;\n    min-height: 160px;" in css
    assert "1366" in readme
    assert "125%" in readme


def test_saving_page_size_immediately_rerenders_results():
    app_js = _read(APP_JS)
    save_settings = re.search(r"function saveSettings\(\) \{(?P<body>.*?)function resetSettings", app_js, re.S)

    assert save_settings is not None
    assert "currentPage = 1" in save_settings.group("body")
    assert "renderResults()" in save_settings.group("body")


def test_run_page_has_stop_button_logs_and_background_job_polling():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert 'id="stopRunBtn"' in html
    assert 'class="stop-icon"' in html
    assert 'id="runLogPanel"' in html
    assert 'id="runLogList"' in html
    assert 'id="runLogToggleBtn"' in html
    assert 'api("/api/run/start"' in app_js
    assert 'api(`/api/run/status/${encodeURIComponent(jobId)}`)' in app_js
    assert 'api("/api/run/cancel"' in app_js
    assert "function renderRunLogs" in app_js
    assert "function buildRunCompletionNotice" in app_js
    assert "function buildRunCompletionLogMessage" in app_js
    assert 'formatHistoryDiffCount(history, "added_count")' in app_js
    assert 'formatHistoryDiffCount(history, "removed_count")' in app_js
    assert "上次执行时间 ${baselineRunAt || \"无\"}" in app_js
    assert "appendRunLog(buildRunCompletionLogMessage(completionNotice, h));" in app_js
    assert "runLogPanel.hidden = true" in app_js
    assert "if (!logs.length && runLogPanel.hidden) return;" in app_js
    assert 'runLogPanel.classList.toggle("collapsed")' in app_js
    assert ".btn-stop" in css
    assert ".run-log-panel" in css
    assert ".run-log-panel.collapsed .run-log-list" in css
    assert "display: none" not in re.search(r"\.run-log-panel\.collapsed \.run-log-list\s*\{(?P<body>.*?)\}", css, re.S).group("body")
    assert "max-height" in re.search(r"\.run-log-list\s*\{(?P<body>.*?)\}", css, re.S).group("body")
    assert "opacity" in re.search(r"\.run-log-panel\.collapsed \.run-log-list\s*\{(?P<body>.*?)\}", css, re.S).group("body")
    assert "transform" in re.search(r"\.run-log-panel\.collapsed \.run-log-list\s*\{(?P<body>.*?)\}", css, re.S).group("body")
    assert "@keyframes runLogTogglePop" in css
    assert "@keyframes runLogLineFloatIn" in css
    assert "@media (prefers-reduced-motion: reduce)" in css


def test_combination_limit_is_configurable_in_default_settings():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)

    assert 'id="combinationLimit"' in html
    assert "组合候选阈值" in html
    assert "function getCombinationLimit()" in app_js
    assert "defaultSettings.combinationLimit || \"50\"" in app_js
    assert 'api("/api/settings/defaults"' in app_js
    assert 'localStorage.removeItem("autoCheckSettings")' in app_js
    assert "max_combination_rows: getCombinationLimit()" in app_js


def test_visual_effects_setting_replaces_page_size_control():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert 'id="pageSize"' not in html
    assert "每页显示条数" not in html
    assert 'id="visualEffects"' in html
    assert "动画效果" in html
    assert "visualEffects" in app_js
    assert "visual_effects" in app_js
    assert "function visualEffectsEnabled()" in app_js
    assert "function applyVisualEffectsSetting()" in app_js
    assert 'dataVisualEffects' not in app_js
    assert 'dataset.visualEffects' in app_js
    assert '[data-visual-effects="off"]' in css


def test_latest_history_results_load_by_default_and_last_run_time_is_retained():
    app_js = _read(APP_JS)

    assert "function formatDisplayTime(value)" in app_js
    assert 'return String(value)' in app_js
    assert r'.replace(/\.\d+(?=(?:Z|[+-]\d{2}:?\d{2})?$)/, "")' in app_js
    assert 'const displayTime = formatDisplayTime(value || "");' in app_js
    assert "if (!displayTime) return;" in app_js
    assert "latestRunAt = displayTime;" in app_js
    assert "async function loadLatestHistoryResults()" in app_js
    assert "await loadLatestHistoryResults()" in app_js
    assert 'if (lastRunTime.textContent && !resultRestoreHistoryMeta && !hideLastRunTimeForNoSourceData) lastRunTime.hidden = false;' in app_js
    assert "lastRunTime.hidden = Boolean(resultRestoreHistoryMeta) || hideLastRunTimeForNoSourceData;" in app_js
    set_last_run = re.search(r"function setLastRunTime\(value, executorName = \"\"\) \{(?P<body>.*?)\n\}", app_js, re.S)
    assert set_last_run is not None
    assert "hideLastRunTimeForNoSourceData = false;" in set_last_run.group("body")
    assert "setLastRunTime(latestHistory.run_at, historyExecutorName(latestHistory))" in app_js
    assert "normalizeExecutorDisplayName" in app_js
    assert 'executorName: normalizeExecutorDisplayName(extra.executorName || latestRunExecutor, "")' in app_js
    assert "if (!latestRunAt) setLastRunTime(formatLastRunTime())" not in app_js


def test_frontend_dates_use_beijing_time_helpers():
    app_js = _read(APP_JS)

    assert 'const BEIJING_TIME_ZONE = "Asia/Shanghai";' in app_js
    assert "function formatBeijingDate(" in app_js
    assert "function formatBeijingDateTime(" in app_js
    assert "function shiftBeijingDate(" in app_js
    assert 'trendDateStart = shiftBeijingDate({ months: -6 });' in app_js
    assert "formatClockTime()" in app_js
    assert "return formatBeijingTime();" in app_js
    assert 'const displayTime = formatDisplayTime(value || "");' in app_js
    assert "latestRunAt = displayTime;" in app_js
    assert "return formatBeijingDateTime();" in app_js
    assert "savedAt: formatBeijingDateTime()" in app_js
    assert "`users-${formatBeijingDate()}.csv`" in app_js
    assert "`auto-check-configs-${formatBeijingDate()}.json`" in app_js
    assert "new Date().toISOString().slice(0, 10)" not in app_js


def test_latest_history_results_uses_history_sort_order():
    app_js = _read(APP_JS)

    assert "function compareHistoryRunsDesc(a, b)" in app_js
    assert "function compareHistoryRunsByRunAtDesc(a, b)" in app_js
    assert "const sorted = getFilteredHistoryRuns().sort(compareHistoryRunsDesc);" in app_js
    assert "const latest = [...runs].sort(compareHistoryRunsByRunAtDesc)[0];" in app_js


def test_history_page_has_pagination_controls_and_logic():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert 'id="historyReportFilter"' in html
    assert 'id="historyExecutorFilter"' in html
    assert 'class="history-toolbar"' in html
    assert 'type="date"' in html
    assert '<select id="historyExecutorFilter" class="filter-select history-filter-input">' in html
    assert '<option value="">全部执行人</option>' in html
    assert 'id="clearHistoryReportFilter"' in html
    assert 'id="clearHistoryExecutorFilter"' in html
    assert 'id="clearKeywordFilter"' in html
    assert 'id="clearReasonFilter"' in html
    assert 'id="clearStatusFilter"' in html
    assert 'class="filter-clear-shell filter-clear-shell--select result-reason-filter-shell"' in html
    assert html.index('id="historyReportFilter"') < html.index('id="historyRefreshBtn"')
    assert html.index('id="historyExecutorFilter"') < html.index('id="historyRefreshBtn"')
    assert 'id="historyPageInfo"' in html
    assert 'id="historyPrevPage"' in html
    assert 'id="historyNextPage"' in html
    assert 'id="historyPageCurrent"' in html
    assert 'id="historyJumpPage"' in html
    assert '跳至 <input id="historyJumpPage" type="number" min="1" /> 页' in html
    history_pagination = re.search(r'id="historyPagination"(?P<body>.*?)</div>\s*</section>', html, re.S)
    assert history_pagination is not None
    assert "sysInfoFeedback" not in history_pagination.group("body")
    assert "let historyCurrentPage = 1;" in app_js
    assert "function getHistoryFilterValues()" in app_js
    assert "function getFilteredHistoryRuns()" in app_js
    assert "function updateHistoryExecutorOptions()" in app_js
    assert "executors.set(name.toLowerCase(), name);" in app_js
    assert "updateHistoryExecutorOptions();" in app_js
    assert "run.run_date !== filters.reportDate" in app_js
    assert 'executorText !== filters.executor' in app_js
    assert "暂无符合条件的历史记录" in app_js
    assert 'historyReportFilter?.addEventListener("change"' in app_js
    assert 'historyExecutorFilter?.addEventListener("change"' in app_js
    assert "clearHistoryFilterControl(historyReportFilter)" in app_js
    assert "clearHistoryFilterControl(historyExecutorFilter)" in app_js
    assert "clearResultFilterControl(keywordFilter)" in app_js
    assert "clearResultFilterControl(reasonFilter)" in app_js
    assert "clearResultFilterControl(statusFilter)" in app_js
    assert "function updateFilterClearButtons()" in app_js
    assert "function getHistoryPageItems()" in app_js
    assert "function updateHistoryPagination()" in app_js
    assert "historyPrevPageBtn?.addEventListener" in app_js
    assert "historyNextPageBtn?.addEventListener" in app_js
    assert "historyJumpPage?.addEventListener" in app_js
    assert ".history-toolbar" in css
    assert ".history-filter-field" in css
    assert ".history-filter-input" in css
    assert ".filter-clear-shell" in css
    assert ".filter-clear-button" in css
    assert ".filter-clear-button::before" in css
    assert ".filter-clear-button::after" in css
    result_reason_shell = re.search(
        r"(?m)^\.result-card \.filters-row \.result-reason-filter-shell\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert result_reason_shell is not None
    assert "width: 220px" in result_reason_shell.group("body")
    assert ".result-card .filters-row .filter-clear-shell--text" not in css
    assert ".result-card .filters-row .filter-clear-shell--select" not in css
    history_field = re.search(r"(?m)^\.history-filter-field\s*\{(?P<body>.*?)\}", css, re.S)
    assert history_field is not None
    assert "white-space: nowrap" in history_field.group("body")
    assert ".history-filter-field > span:first-child" in css
    assert "flex: 0 0 150px" in css
    history_date_shell = re.search(
        r"(?m)^\.history-filter-clear-shell\.filter-clear-shell--date\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert history_date_shell is not None
    assert "flex: 0 0 176px" in history_date_shell.group("body")
    assert "width: 176px" in history_date_shell.group("body")
    assert ".filter-clear-shell:hover .filter-clear-button.is-visible" in css
    assert ".filter-clear-shell--select .custom-select-trigger" in css


def test_system_info_actions_show_running_and_completion_feedback():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert 'id="sysInfoFeedback"' in html
    assert "async function runSystemInfoAction" in app_js
    assert "setSystemInfoFeedback" in app_js
    assert 'id="testAllConnBtn"' not in html
    assert "testAllConnBtn" not in app_js
    assert 'refreshInfoBtn?.addEventListener("click"' in app_js
    assert "button.textContent = pendingText" in app_js
    assert "button.disabled = true" in app_js
    assert ".sys-info-feedback" in css
    assert ".sys-info-feedback--running" in css
    assert ".sys-info-feedback--success" in css
    assert ".sys-info-feedback--error" in css


def test_changelog_documents_latest_setting_and_cleanup_changes():
    app_js = _read(APP_JS)

    for text in [
        "v1.2.2",
        "新增后台执行、停止执行和执行日志。",
        "新增默认设置持久化、历史分页和系统信息操作反馈。",
        "新增估值表资产合计列、导出详情和 1541 财产权核对。",
        "v1.2.1",
        "2026-06-01",
        "系统优化及BUG修复。",
        "v1.2.0",
        "v1.1.0",
        "新增系统设置、业务设置、主题和图表能力。",
        "v1.0.0",
        "初始版本：自动对数、历史记录、多数据源和 Excel 导出。",
    ]:
        assert text in app_js

    for verbose_text in [
        "执行过程中新增后台控制台日志和可折叠页面执行日志，细化到",
        "资产缺失/重复新增 1541 财产权合同投融资核对：",
        "统一使用 fa_valuationreport_dws.c_projcode",
        "后台执行、停止执行和执行日志优化。",
        "优化工作台布局、进度条和结果详情区域。",
        "资产科目匹配、名称匹配和组合候选规则优化。",
    ]:
        assert verbose_text not in app_js


def test_v21_changelog_documents_interface_radius_concisely():
    app_js = _read(APP_JS)
    changelog = re.search(
        r'<span class="changelog-version">v1\.2\.13</span>(?P<body>.*?)<div class="changelog-item">',
        app_js,
        re.S,
    )

    assert changelog is not None
    body = changelog.group("body")
    assert "<li>新增界面圆角个性化设置。</li>" in body
    assert "<li>系统优化及BUG修复。</li>" in body
    assert body.count("新增界面圆角个性化设置") == 1
    assert "1–15px" not in body
    assert "导航、卡片、按钮" not in body
    for verbose_theme_detail in [
        "全局纯色主题",
        "主题色",
        "折线图风格",
        "#3F6FAF",
        "#355F63",
        "#RRGGBB",
    ]:
        assert verbose_theme_detail not in body
    assert 'const DEFAULT_VERSION = "V1.2";' in app_js


def test_balanced_modal_refresh_is_documented_with_concise_in_app_changelog():
    readme = _read(README_MD)
    app_js = _read(APP_JS)

    for text in [
        "系统弹窗统一为轻量平衡风格",
        "白色表面、细分隔线、克制阴影、统一标题栏、独立滚动内容区和固定操作区",
        "确认、输入、信息、用户、数据源、人行导入与校验、流程工具弹窗保留各自适配业务内容的尺寸",
        "历史详情按完整结果、新增差异、减少差异分组",
        "主题蓝、红、绿色条区分",
        "表头和内容继续保持居中",
        "中性浅灰表头、透明描边状态",
        "隐藏滚动条但保留滚动",
        "恢复按钮与状态列居中",
        "弹窗圆角继续跟随当前用户的界面设置",
        "当前唯一启用的浅色主题",
    ]:
        assert text in readme
    assert "`v1.2.13` (2026-07-18) 主要变化：" in readme

    current = re.search(
        r'<span class="changelog-version">v1\.2\.13</span>(?P<body>.*?)<div class="changelog-item">',
        app_js,
        re.S,
    )
    assert current is not None
    assert '<span class="changelog-date">2026-07-18</span>' in current.group("body")
    assert "系统优化及BUG修复。" in current.group("body")
    assert "弹窗" not in current.group("body")
    assert "历史详情" not in current.group("body")


def test_changelog_and_readme_document_pbc_import_and_space_nav_updates():
    app_js = _read(APP_JS)
    readme = _read(README_MD)

    for text in [
        "新增工具页面与人行全量产品一键导入能力。",
        "系统优化及BUG修复。",
    ]:
        assert text in app_js

    for text in [
        "人行全量产品一键导入增强",
        "上传支持 zip/rar/7z/xlsx/xls/csv",
        "字段映射区布局优化",
        "太空主题滚动效果优化",
        "内容越过导航栏时增加高透明内容模糊遮罩",
    ]:
        assert text in readme


def test_version_206_documents_db_validation_engine_update():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    readme = _read(README_MD)

    assert 'const DEFAULT_VERSION = "V1.2";' in app_js
    assert 'id="statusText">V1.2</span>' in html
    assert 'id="topNavStatus" title="V1.2">V1.2</span>' in html

    for text in [
        "v1.2.13",
        "v1.2.12",
        "新增监管智核品牌名称和系统 Logo。",
        "首页组合图表指标改为每期差异个数。",
        "v1.2.11",
        "新增流程链配置、手工执行及执行记录查看。",
        "流程链配置支持从流程表选择流程。",
        "自动对数新增3001共同类科目与实收本金多次重复识别。",
        "v1.2.10",
        "新增人行逐笔校验引擎公开信息校验、模板校验和规则说明能力。",
        "自动对数差异原因调整为固定基础分类，细分原因在详情展示。",
        "自动对数资产缺失细分新增多资产格式化具体原因和详情表格。",
        "自动对数资产重复新增私募产品细分原因和详情表格。",
        "自动对数资产差异新增贷款及财产权合同细分原因和详情表格。",
        "自动对数资产端组合候选过多时支持科目分组组合，并新增债券DM证券余额差异细分。",
        "自动对数资产差异和负债权益科目差异新增逆/正回购金额比对。",
        "自动对数资产差异解释后支持继续核对剩余差额并展示组合差异类型。",
        "自动对数差异类型筛选支持组合差异类型匹配。",
        "自动对数资产端和负债权益主差异多组候选时展示候选不唯一。",
        "自动对数资产缺失候选不唯一时支持AM复核确认候选组，实收本金缺失/重复新增c1000防误判判断。",
        "自动对数导出Excel新增组合差异备注列。",
        "自动对数处理脚本支持多个FA/AM标的不一致生成。",
        "自动对数负债权益正回购差异新增具体原因。",
        "自动对数结果列表和导出字段改为差异类型，并新增具体原因列。",
        "自动对数历史详情同步展示具体原因。",
        "自动对数结果详情改为单行展开查看。",
        "自动对数实收本金差异与负债权益混合场景支持剩余差额核对。",
        "自动对数实收本金差异新增TA差异细分原因。",
        "自动对数负债权益和实收本金新增格式化具体原因和详情表格。",
        "系统优化及BUG修复。",
    ]:
        assert text in app_js

    for text in [
        "v1.2.11",
        "自动对数资产端候选新增 `3001.XX` 正数共同类科目",
        "自动对数实收本金重复支持多次重复计入识别",
        "新增流程执行工具",
        "流程链支持弹窗新增/编辑",
        "流程顺序可从申报平台流程表中选择",
        "流程链配置弹框点击空白遮罩不再关闭",
        "sp_task.end_time",
        "仅支持手工执行",
        "流程功能独立于自动对数",
        "v1.2.10",
        "人行逐笔校验引擎新增数据库逐笔校验能力",
        "新增公开信息交叉校验",
        "新增模板交叉校验",
        "baseinfo.table_name_zh",
        "baseinfo.template_json",
        "支持按 `baseinfo` 英文表名生成和读取 30 张物理模板表",
        "ZG09/ZG10 模板交叉校验对齐旧程序口径",
        "`cpkj=1` 分别对比 `balance_sheet_info`、`balance_sheet_info2`",
        "`cpkj=2` 分别对比 `balance_sheet_info_zcglxt`、`balance_sheet_info2_zcglxt`",
        "历史改为按真实执行时间倒序展示",
        "规则说明同步更新为最新代码口径",
        "自动对数差异原因调整为固定基础分类",
        "资产端解释后仍有剩余差额时可组合展示多个差异类型",
        "自动对数资产缺失细分扩展为",
        "自动对数资产重复细分扩展为",
        "自动对数资产差异新增贷款/财产权合同逐一核对",
        "自动对数资产差异新增逆回购金额比对",
        "`subcode LIKE '7%'` 的 `buyback_money + expenses`",
        "资产差异细分",
        "命中 `1101.05.06.01*` 时核查 AM 标的表 `c_spv_type` 和 `c_assettype`",
        "特定目的载体范围扩展至信托、银行理财、保险理财、场外证券理财产品、场外基金理财产品和期货",
        "自动对数导出处理脚本支持从“资产缺失细分”表中识别多条 FA/AM 标的不一致记录",
        "特定目的载体、债券、股票、公募基金、私募基金、逆回购、贷款、股权投资、信托计划收益权、资产收益权",
        "资产缺失细分",
        "正回购差异",
        "自动对数负债及权益科目差异新增正回购金额比对",
        "匹配状态为 `候选不唯一`",
        "资产缺失方向支持继续用 AM 标的和合同投融资余额复核唯一确认候选组",
        "实收本金缺失/重复新增 `c1000` 防误判闸门",
        "自动对数导出 Excel 新增“备注”列",
        "`subcode LIKE '8%'` 的 `buyback_money - expenses`",
        "负债及权益科目细分",
        "导出 Excel 在“差异类型”后新增“具体原因”列",
        "结果页展开详情和导出 Excel 同步展示“具体原因”",
        "自动对数结果详情改为单行展开查看",
        "合同投融资余额非 0",
        "继续核查 SPV DM 表和报表明细",
        "实收本金差异与负债权益混合场景",
        "`a0001-d0000-(4001-c1000)` 的剩余差额",
        "自动对数实收本金缺失、重复、差异的具体原因改为",
        "实收本金细分",
        "`currency_report_24.currency_detail_project_2_1_8`",
        "自动对数“实收本金差异”新增 TA 细分原因",
        "DM TA 表与 DWS TA 表份额余额+待结转收益汇总",
        "客户类型依赖字段为空记录",
        "load-local-pg-20260601-formatted-reason-scenarios.ps1",
        "load-local-pg-20260614-16-reconcile-scenarios.ps1",
        "seed_current_reconcile_20260614_16.py",
        "seed_current_home_frequency_reports.py",
        "seed_current_history_delta_20260622.py",
        "高频差异项目",
        "按报告期和执行人筛选",
        "执行人下拉按现有历史记录去重生成",
        "悬浮小叉快速清除",
        "保留输入框、日期框和下拉框的粒子悬浮效果",
        "结果列表仅加宽差异类型筛选框",
        "避免执行人文字被日期图标区遮挡",
        "历史详情弹窗移除内容区边缘线",
        "底部“恢复到结果页”按钮区域固定在弹窗底部",
        "报表对应日期无数据时隐藏顶部“最近执行”提示",
        "新增差异 10 条、减少差异 10 条",
        "DELTA20260622",
        "实收本金不一致",
        "标的代码不一致",
        "HFJST2026",
        "AC20260614",
        "AC20260615",
        "AC20260616",
        "3001 共同类资产/负债",
        "债券 DM 余额差异",
    ]:
        assert text in readme


def test_version_208_documents_regulatory_intelligence_core_brand_update():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    login_html = _read(ROOT / "src" / "auto_check" / "web" / "login.html")
    logo = _read(ROOT / "src" / "auto_check" / "web" / "assets" / "logo-full.svg")
    login_logo = _read(ROOT / "src" / "auto_check" / "web" / "assets" / "logo-login.svg")
    login_dark_logo = _read(ROOT / "src" / "auto_check" / "web" / "assets" / "logo-login-dark.svg")
    favicon_asset = _read(ROOT / "src" / "auto_check" / "web" / "assets" / "favicon-64x64.svg")
    readme = _read(README_MD)

    assert "<title>监管智核</title>" in html
    assert "<title>监管智核</title>" in login_html
    assert 'href="/assets/favicon-64x64.svg?v=1.2.12-regulatory-intelligence-core"' in html
    assert 'href="/assets/favicon-64x64.svg?v=1.2.12-regulatory-intelligence-core"' in login_html
    assert 'class="brand-wordmark-main">监管智核</span>' in html
    assert 'class="brand-wordmark-sub">监管报送核验平台</span>' in html
    assert 'src="/assets/logo-login.svg?v=1.2.12-regulatory-intelligence-core-horizontal" alt="监管智核"' in login_html
    assert '"/assets/logo-login-dark.svg?v=1.2.12-regulatory-intelligence-core-horizontal"' not in login_html
    assert 'alt="监管智核 Logo"' in html
    assert "准星" not in html
    assert "准星" not in login_html
    assert 'viewBox="0 0 520 160"' in login_logo
    assert "监管智核横向标志" in login_logo
    assert "ric-horizontal" in login_logo
    assert "ric-stacked" not in login_logo
    assert 'viewBox="0 0 520 160"' in login_dark_logo
    assert "监管智核深色横向标志" in login_dark_logo
    assert "ric-horizontal-dark" in login_dark_logo
    assert "ric-stacked" not in login_dark_logo

    for asset in [logo, favicon_asset, login_logo, login_dark_logo]:
        assert "ric-" in asset
        assert "监管智核" in asset
        assert any(color in asset for color in ("#3466d9", "#4f7cff"))
        assert any(color in asset for color in ("#ffbd38", "#f0a12b"))
        assert "scheme-a" not in asset
        assert "scheme-a-zx-grid-hit" not in asset
        assert "A compact ZX monogram" not in asset

    for text in [
        "v1.2.12",
        "新增监管智核品牌名称和系统 Logo。",
        "新增点击 Logo 切换主题能力。",
        "新增登录进入主界面动效。",
        "系统优化及BUG修复。",
    ]:
        assert text in app_js

    for text in [
        "v1.2.12",
        "系统对外名称更新为“监管智核”",
        "使用 `logo/regulatory-intelligence-core` 资源包中的双环对勾设计替换系统 Logo",
        "主应用、登录页、关于系统和浏览器页签品牌文案同步调整",
        "点击侧边栏或顶部导航 Logo 切换活力/沉稳主题",
        "侧边栏与顶部导航之间的衔接过渡动画",
        "登录成功进入主界面时新增一次性入场动画",
        "浅色登录页品牌区同步加入暗色模式浮动圆形动效",
        "前端静态测试同步更新监管智核品牌和 Logo 资源断言",
    ]:
        assert text in readme


def test_version_21_documents_reconcile_schema_and_flow_updates():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    readme = _read(README_MD)

    assert 'const DEFAULT_VERSION = "V1.2";' in app_js
    assert 'id="statusText">V1.2</span>' in html
    assert 'id="topNavStatus" title="V1.2">V1.2</span>' in html
    assert "- 应用界面版本：`V1.2`" in readme
    assert 'id="sysVersion">V1.2</span>' in html

    change_items = [
        "人行逐笔校验执行历史新增执行人展示",
        "自动对数 AM 复核在名称无法匹配时新增兜底",
        "首页最新趋势横轴改为展示每次自动对数的执行日期和时间",
        "兜底明细与资产缺失细分明细重复计数",
        "流程链配置可选流程列表保留 500 条初始展示上限",
        "流程链停止按钮改为按后台任务",
        "自动对数导出处理脚本按 AM 合同来源判断",
        "自动对数仓储查询支持在系统设置页面通过表单维护表名、字段名和表级数据源",
        "表字段配置保存失败弹框按缺失字段逐行展示",
        "自动对账表字段配置新增“标准中文名”输入框",
    ]
    assert "`v1.2.13` (2026-07-18) 主要变化：" in readme
    assert '<span class="changelog-version">v1.2.13</span>' in app_js
    assert '<span class="changelog-date">2026-07-18</span>' in app_js
    for text in change_items:
        assert text in readme

    for text in [
        "应用自身配置、用户和历史记录改为保存到 MySQL `auto_check` 应用库",
        "config.json 仅保留 `app_database` 启动连接信息",
        "sql/app_storage/mysql/001_init_schema.sql",
        "scripts/export_sqlite_to_mysql.py",
        "本地数据查询页面及入口已隐藏",
        "删除旧 SQLite `auto-check.db` 后应用仍应只依赖 MySQL 应用库运行",
    ]:
        assert text in readme

    for text in [
        "重复启动本地服务时检测默认端口占用",
        "系统设置和工具页的配置加载改为模块间互不阻塞",
        "系统信息改用轻量统计接口",
        "避免切页或刷新时拉取全量历史记录",
        "逐笔字段映射加载结果少于系统内置表单",
        "历史记录写入当前登录用户的姓名、账号和用户 ID",
    ]:
        assert text in readme
    assert "系统优化及BUG修复。" in app_js
    assert "本地数据查询" not in app_js
    assert "auto-check.db" not in app_js


def test_version_205_documents_scheme_a_logo_update():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    logo = _read(ROOT / "src" / "auto_check" / "web" / "assets" / "logo-full.svg")
    favicon_asset = _read(ROOT / "src" / "auto_check" / "web" / "assets" / "favicon-64x64.svg")
    readme = _read(README_MD)

    assert 'const DEFAULT_VERSION = "V1.2";' in app_js
    assert 'id="statusText">V1.2</span>' in html
    assert 'id="topNavStatus" title="V1.2">V1.2</span>' in html

    for text in [
        "v1.2.9",
        "系统优化及BUG修复。",
    ]:
        assert text in app_js

    for text in [
        "v1.2.9",
        "系统 Logo 采用方案 A",
        "ZX 数据网格设计",
        "绿色命中点",
        "favicon 同步改为方案 A 的小尺寸简化版本",
        "修复数据源测试连接长时间未返回后关闭弹窗",
        "人行逐笔校验引擎保存配置后自动刷新字段映射",
        "全站下拉框统一为方案 5 粒子悬浮风格",
        "修复首页趋势日期下拉菜单滚动时被关闭的问题",
        "适当加宽日期下拉框",
        "系统输入框同步加入粒子悬浮与毛玻璃聚焦效果",
        "默认太空主题和暗色模式风格一致",
        "用户创建/编辑弹框移除多余隐藏下拉框",
        "关于系统品牌区仅保留 Logo",
        "日期选择组件采用太空粒子风格",
        "日期选择弹层同步替换为自定义方案 B 日历面板",
        "统一运行日期和人行逐笔报告期",
        "修复用户创建/编辑弹框输入框和人行逐笔报告期日期组件",
        "调整系统 Logo 与 favicon 图案重心",
        "浏览器页签 favicon 引用增加版本参数",
        "沉稳主题下输入框、下拉框和日期选择组件描边统一使用沉稳主题色",
        "活力主题保持蓝青紫粒子风格",
        "侧边栏品牌区移除“精准核对 · 合规报送”标语",
        "自动对数原因“缺失资产在AM信息中正常，需排除生成数据SQL”调整为“缺失资产在投资端信息无异常，请核查报表是否正常生成”",
        "自动对数导出 Excel 改为 `.xlsx` 工作簿格式",
        "表头行高固定为 30",
        "金额列写入为数值格式",
        "差异原因详情”点击公式栏查看时保留原始换行格式",
        "自动对数导出按钮增加导出中进度反馈和成功/失败提示",
        "活力主题下结果详情标题图标改为蓝青紫渐变",
        "顶部导航和侧边栏用户头像改为风格 B 彩色首字母头像",
        "系统默认设置中的“默认运行日期”替换为“会话过期时间”",
        "前端静态测试同步增加方案 A Logo 资源断言",
    ]:
        assert text in readme


def test_home_chart_date_select_keeps_scrollable_wider_dropdown():
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    chart_select_rule = re.search(r"(?m)^\.chart-date-select\s*\{(?P<body>.*?)\}", css, re.S)
    assert chart_select_rule is not None
    assert "min-width: 150px" in chart_select_rule.group("body")

    assert 'target.closest(".custom-select-dropdown")' in app_js
    assert "const dropdownWidth = compactRoleSelect" in app_js
    assert "const openAbove = availableBelow < 160 && availableAbove > availableBelow;" in app_js


def test_home_chart_date_select_keeps_fixed_width_after_custom_select_enhancement():
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)
    readme = _read(README_MD)

    chart_select_rule = re.search(r"(?m)^\.chart-date-select\s*\{(?P<body>.*?)\}", css, re.S)
    chart_shell_rule = re.search(r"(?m)^\.custom-select-shell\.chart-date-select\s*\{(?P<body>.*?)\}", css, re.S)
    assert chart_select_rule is not None
    assert chart_shell_rule is not None
    assert "width: 150px;" in chart_select_rule.group("body")
    assert "width: 150px;" in chart_shell_rule.group("body")
    assert "flex: 0 0 150px;" in chart_shell_rule.group("body")
    assert "function scheduleHomeChartDateSelectMeasure()" not in app_js
    assert "对数总览执行趋势日期下拉框固定为 150px" in readme


def test_version_204_documents_tab_and_brand_hierarchy_update():
    app_js = _read(APP_JS)
    readme = _read(README_MD)

    for text in [
        "v1.2.8",
        "新增浏览器页签品牌精简显示。",
        "系统优化及BUG修复。",
    ]:
        assert text in app_js

    for text in [
        "v1.2.8",
        "浏览器页签标题精简为“准星”",
        "开启您的智能工作台",
        "小字监管报送助手",
        "关于系统品牌区仅保留 Logo",
        "移除中点分隔",
    ]:
        assert text in readme


def test_version_203_documents_brand_logo_update():
    app_js = _read(APP_JS)
    readme = _read(README_MD)

    for text in [
        "v1.2.7",
        "新增准星·监管报送助手品牌名称和系统 Logo。",
        "系统优化及BUG修复。",
    ]:
        assert text in app_js

    for text in [
        "v1.2.7",
        "系统对外名称更新为“准星·监管报送助手”",
        "使用 `logo/scheme-D-zx-grid` 资源包中的 ZX 数据网格设计替换系统 Logo",
        "登录页功能卡片中的旧闪电符号替换为校验符号",
        "前端静态测试同步更新品牌、版本和 Logo 资源断言",
    ]:
        assert text in readme


def test_version_202_documents_security_login_update():
    app_js = _read(APP_JS)
    readme = _read(README_MD)

    for text in [
        "v1.2.6",
        "导出 Excel 新增处理脚本列。",
        "新增用户姓名，导航用户按钮、用户列表和执行历史优先显示姓名。",
        "新增对数任务全局互斥和一键导入同表冲突提示。",
        "系统优化及BUG修复。",
    ]:
        assert text in app_js

    for text in [
        "v1.2.6",
        "自动对数导出 Excel 新增“处理脚本”列",
        "系统时间统一按北京时间生成和展示",
        "核对历史按核对日期倒序、同日按执行时间倒序排列",
        "登录页页签标题与主应用保持英文一致",
        "登录和用户管理密码规则调整为至少 6 位且包含字母",
        "首页趋势和核对历史展示全部对数记录",
        "MySQL 数据源隐藏 Schema 输入",
        "一键导入上传解析支持自动跳过模板标题区",
        "用户列表进入页面时使用稳定骨架加载态",
    ]:
        assert text in readme

    for verbose_text in [
        "FA 与 AM 标的不一致时生成修正 SQL",
        "执行日志显示失败原因、正在执行用户和可再次执行提示",
        "登录页浅色和暗色模式按钮、品牌图标、特性图标统一",
        "沉稳主题下系统设置栏目图标背景统一",
        "导出、历史详情、执行日志和登录体验优化。",
        "默认数据源切换增加即时反馈，用户列表加载更稳定。",
        "默认数据源切换先在本地置顶并播放切换动画",
    ]:
        assert verbose_text not in app_js


def test_version_201_documents_confirm_button_update():
    app_js = _read(APP_JS)
    readme = _read(README_MD)

    for text in [
        "v1.2.5",
        "系统优化及BUG修复。",
    ]:
        assert text in app_js

    for verbose_text in [
        "导入确认弹窗、系统设置和业务设置布局优化。",
        "主题滚动层次和暗色模式对比度优化。",
    ]:
        assert verbose_text not in app_js

    for text in [
        "v1.2.5",
        "导入确认弹窗按钮更清晰",
        "系统设置改为太空科技三列卡片布局",
        "太空主题导航上下内容模糊感加重",
    ]:
        assert text in readme


def test_business_settings_displays_current_table_field_mapping():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)
    readme = _read(README_MD)

    assert 'id="businessSettingsBody"' in html
    assert 'id="businessSettingsContent"' in html
    assert 'id="initReconcileSchemaFromFileBtn"' in html
    assert 'id="reconcileSchemaForm"' in html
    assert html.index('id="initReconcileSchemaFromFileBtn"') < html.index('id="reconcileSchemaForm"')
    business_card = re.search(
        r'<section class="card settings-dashboard-card card-business admin-only">(?P<body>.*?)</section>',
        html,
        re.S,
    )
    assert business_card is not None
    business_header = re.search(
        r'<div class="card-header">(?P<body>.*?)</div>\s*<div id="businessSettingsBody"',
        business_card.group("body"),
        re.S,
    )
    assert business_header is not None
    assert 'id="initReconcileSchemaFromFileBtn"' in business_header.group("body")
    assert 'id="saveReconcileSchemaBtn"' in business_header.group("body")
    schema_panel_head = business_card.group("body")[
        business_card.group("body").index('<div class="reconcile-settings-panel reconcile-schema-panel">'):
        business_card.group("body").index('<div id="reconcileSchemaForm"')
    ]
    assert 'id="initReconcileSchemaFromFileBtn"' not in schema_panel_head
    assert 'id="saveReconcileSchemaBtn"' not in schema_panel_head
    assert 'id="reconcileSchemaEditor"' not in html
    assert "function renderBusinessSettings()" in app_js
    assert "function loadReconcileSchemaSettings()" in app_js
    assert "function renderReconcileSchemaForm(" in app_js
    assert "function readReconcileSchemaForm()" in app_js
    assert "function loadReconcileTableColumns(" in app_js
    assert "function currentBusinessFieldGroups()" in app_js
    assert "function filterReconcileColumnOptions(" in app_js
    assert "function renderReconcileFieldOptions(" in app_js
    assert "function splitFrontendReconcileSchemaMissingItems(" in app_js
    assert "function parseReconcileSchemaErrorItem(" in app_js
    assert "function formatReconcileSchemaSaveErrors(" in app_js
    assert "function showReconcileSchemaSaveError(" in app_js
    assert "function readTrimmedControlValue(" in app_js
    assert "const fallback = optional ? tableConfig.fields : tableConfig.optional_fields;" in app_js
    assert "function validateReconcileSchemaRequiredFields(" in app_js
    assert "function markReconcileSchemaRequiredError(" in app_js
    assert "function clearReconcileSchemaRequiredErrors(" in app_js
    assert "function expandReconcileSchemaTable(" in app_js
    assert "function reconcileSchemaVisibleControl(" in app_js
    assert "function selectReconcileSchemaFieldOption(" in app_js
    assert "function reconcileSchemaFieldOptionsOpen(" in app_js
    assert "function openReconcileFieldOptionsForInput(" in app_js
    assert "reconcile-schema-required-error" in app_js
    assert "reconcile-schema-error-message" in app_js
    assert "v1.2.13" in app_js
    assert "scrollIntoView" in app_js
    assert 'querySelector("select.reconcile-schema-source")' in app_js
    assert 'querySelector("input.reconcile-schema-display-name")' in app_js
    assert 'querySelector("input.reconcile-schema-table-name")' in app_js
    assert 'input.reconcile-schema-field-search' in app_js
    assert 'tableEl.querySelector(".reconcile-schema-source")' not in app_js
    assert 'tableEl.querySelector(".reconcile-schema-display-name")' not in app_js
    assert 'tableEl.querySelector(".reconcile-schema-table-name")' not in app_js
    assert 'loadReconcileTableColumns(key, { openCombo: combo' in app_js
    assert "reconcile-schema-load-columns" not in app_js
    assert ">读取字段</button>" not in app_js
    assert "loadReconcileTableColumns(key, { openCombo: combo" in app_js
    assert "reconcile-schema-field-combobox" in app_js
    assert "reconcile-schema-display-name" in app_js
    assert "标准中文名" in app_js
    assert "function currentReconcileTableDisplayName(" in app_js
    assert "displayName: currentReconcileTableDisplayName(primaryKey, group.table)" in app_js
    assert "<strong>${escapeHtml(group.displayName || group.table)}</strong>" in app_js
    assert "display_name: displayName" in app_js
    assert "reconcile-schema-field-search" in app_js
    assert '<div class="reconcile-schema-field-row">' in app_js
    assert '<label class="reconcile-schema-field-row">' not in app_js
    assert "reconcile-schema-field-option-name" in app_js
    assert "reconcile-schema-field-option-comment" in app_js
    assert 'optionalFields: [["data_source"' not in app_js
    assert 'optionalFields: [["contract_start_date"' not in app_js
    assert "delete optionalFields[fieldKey];" in app_js
    assert 'reconcileSchemaForm?.addEventListener("mousedown"' in app_js
    assert app_js.index('reconcileSchemaForm?.addEventListener("mousedown"') < app_js.index('reconcileSchemaForm?.addEventListener("focusin"')
    assert 'if (event.target.closest(".reconcile-schema-field-option")) return;' in app_js
    assert 'event.target.closest("input.reconcile-schema-field-search")' in app_js
    assert "_reconcileSuppressNextInputClick" in app_js
    assert "openReconcileFieldOptionsForInput(fieldInput)" in app_js
    assert 'reconcileSchemaForm?.addEventListener("keydown"' in app_js
    assert 'if (event.key === "Escape") closeReconcileFieldOptions(reconcileSchemaForm);' in app_js
    assert 'closeReconcileFieldOptions(reconcileSchemaForm);' in app_js[app_js.index('const toggle = event.target.closest(".reconcile-schema-toggle");'):app_js.index('reconcileSchemaForm?.addEventListener("input"')]
    assert "event.preventDefault();" in app_js[app_js.index('const option = event.target.closest(".reconcile-schema-field-option[data-value]");'):app_js.index('const toggle = event.target.closest(".reconcile-schema-toggle");')]
    assert "showInfo(title, `" in app_js
    assert "modal-info--reconcile-schema-error" in app_js
    assert "reconcile-schema-save-error" in app_js
    assert "reconcile-schema-save-error-table" in app_js
    assert "reconcile-schema-field-select" not in app_js
    assert "<datalist" not in app_js
    assert 'querySelector(`.reconcile-schema-field-combobox[data-field-key="${fieldKey}"] .reconcile-schema-field-search`)?.value.trim()' not in app_js
    assert "grid-template-columns: repeat(2, minmax(360px, 1fr));" in css
    assert "text-overflow: ellipsis;" in css
    assert "#page-settings .reconcile-schema-required-error" in css
    assert "#page-settings .reconcile-schema-error-message" in css
    assert ".modal-info--reconcile-schema-error" in css
    assert ".reconcile-schema-save-error pre" in css
    assert ".reconcile-schema-save-error-table" in css
    assert "white-space: normal;" in css
    assert "overflow-wrap: anywhere;" in css
    assert "#page-settings .card-header-actions" in css
    assert "向数据库校验表和字段是否真实存在" in readme
    assert "自动对数后台失败日志展示真实错误摘要" in readme
    assert "表字段配置保存失败弹框按缺失字段逐行展示" in readme
    assert "标准中文名" in readme
    assert "/api/settings/reconcile-schema/init-from-file" in app_js
    assert "/api/settings/reconcile-schema/columns" in app_js
    assert "/api/settings/reconcile-schema" in app_js
    assert 'setupSettingsDashboardCollapsible("businessSettingsToggle", "businessSettingsBody")' not in app_js

    for text in [
        "zf_detail_2024",
        "fa_valuationreport_dws",
        "c_projcode",
        "fa_accountbalance_dws",
        "dm.ta_pact_survamt_day_zgxg_dm",
        "ta_pact_detail_dws",
        "am_pactasset_dws",
        "am_projinvest_dws",
        "dm.fa_security_balance_zgxg_dm",
        "dm.am_projinvest_zgxg_dm",
        "dm.am_projinvest_spv_zgxg_dm",
        "zgxg_zhbs.ccqxx",
        "ass_man_reg.ex_pledge_back",
        "currency_report_24.currency_detail_project_2_1_*",
        "currency_report_duration",
        "projinnercode",
        "a0001",
        "f_marketvalue",
        "tpm_clientkind_tusp",
        "tpm_clientkindex",
        "tpm_spvtype",
        "f_alltincom",
        "c_stockcode",
        "c_spv_type",
        "c_assettype",
        "c_datasource",
        "f_acbalance",
        "d_bdate",
        "sbm_seclas_h2024",
        "sbm_gpgqtype_h",
        "sbm_fundtype",
        "pin_gqtype_h",
        "svd_assettype",
        "buyback_money",
        "expenses",
        "每次表名或字段调整时，需要同步更新此业务设置",
    ]:
        assert text in app_js

    assert "c_procode" not in app_js


def test_business_settings_body_scrolls_inside_fixed_dashboard_card():
    css = _read(STYLES_CSS)
    html = _read(INDEX_HTML)

    rule = re.search(
        r"#page-settings \.card-business \.settings-business-scroll\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert rule is not None
    assert "overflow-y: auto" in rule.group("body")
    assert 'id="businessSettingsBody" class="card-body settings-business-scroll"' in html
    assert "#businessSettingsBody:not(.collapsed)" not in css


def test_settings_page_uses_space_tech_dashboard_layout_without_extra_theme_modes():
    html = _read(INDEX_HTML)
    css = _read(STYLES_CSS)
    app_js = _read(APP_JS)
    readme = _read(README_MD)
    settings_section = re.search(
        r'<section class="page" id="page-settings">(?P<body>.*?)\n      </section>\n\n      <!-- 确认弹窗 -->',
        html,
        re.S,
    )
    assert settings_section is not None
    settings_html = settings_section.group("body")

    for text in [
        'class="settings-container"',
        'class="page-header settings-page-header"',
        'class="dashboard-grid settings-dashboard-grid"',
        'class="card settings-dashboard-card card-system-info"',
        'class="card settings-dashboard-card card-interface"',
        'class="card settings-dashboard-card card-default"',
        'class="card settings-dashboard-card card-db-validation admin-only"',
        'class="card settings-dashboard-card card-flow admin-only"',
        'class="card settings-dashboard-card card-business"',
        'class="card settings-dashboard-card card-datasource"',
        'class="card settings-dashboard-card card-about"',
        'id="sysInfoBody"',
        'id="defaultSettingsBody"',
        'id="dbValidationSettingsBody"',
        'id="businessSettingsBody"',
        'id="configBody"',
    ]:
        assert text in settings_html

    assert "card-theme" not in settings_html
    assert "主题设置" not in settings_html
    assert 'id="themeBody"' not in settings_html
    system_info_pos = settings_html.index('class="card settings-dashboard-card card-system-info"')
    interface_pos = settings_html.index('class="card settings-dashboard-card card-interface"')
    default_pos = settings_html.index('class="card settings-dashboard-card card-default admin-only"')
    db_validation_pos = settings_html.index('class="card settings-dashboard-card card-db-validation admin-only"')
    flow_pos = settings_html.index('class="card settings-dashboard-card card-flow admin-only"')
    datasource_pos = settings_html.index('class="card settings-dashboard-card card-datasource admin-only"')
    business_pos = settings_html.index('class="card settings-dashboard-card card-business admin-only"')
    about_pos = settings_html.index('class="card settings-dashboard-card card-about"')
    assert system_info_pos < interface_pos < default_pos < db_validation_pos < flow_pos < datasource_pos < business_pos < about_pos
    for removed_data_management_markup in (
        'class="card settings-dashboard-card card-data admin-only"',
        'id="dataManageToggle"',
        'id="dataManageBody"',
        'id="clearHistoryBtn"',
        'id="exportConfigBtn"',
        'id="importConfigBtn"',
        'id="importConfigFile"',
    ):
        assert removed_data_management_markup not in settings_html
    for retained_data_management_handler in (
        'setupCollapsible("dataManageToggle", "dataManageBody", "dataManageArrow");',
        'document.getElementById("clearHistoryBtn")?.addEventListener("click", async () => {',
        'document.getElementById("exportConfigBtn")?.addEventListener("click", async () => {',
        'document.getElementById("importConfigBtn")?.addEventListener("click", () => {',
        'document.getElementById("importConfigFile")?.addEventListener("change", async (e) => {',
    ):
        assert retained_data_management_handler in app_js
    assert 'id="businessSettingsBody" class="card-body settings-business-scroll"' in settings_html
    assert "settings-collapsed-card" not in settings_html
    assert "settings-collapsible-body" not in settings_html
    assert "<h2>系统设置</h2>" in settings_html
    about_card = re.search(
        r'<section class="card settings-dashboard-card card-about">(?P<body>.*?)</section>',
        settings_html,
        re.S,
    )
    assert about_card is not None
    assert "about-features" in about_card.group("body")
    assert "about-tech" in about_card.group("body")
    assert "主要功能" in about_card.group("body")
    assert "技术栈" in about_card.group("body")

    assert 'name="theme"' not in html
    assert 'id="themeToggle"' not in html
    assert 'value="dark"' not in html
    assert 'value="auto"' not in html

    dashboard_grid_rule = re.search(
        r"#page-settings \.settings-dashboard-grid\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert dashboard_grid_rule is not None
    assert "grid-template-columns: repeat(3, minmax(0, 1fr))" in dashboard_grid_rule.group("body")
    assert "align-items: stretch" in dashboard_grid_rule.group("body")
    space_container_rule = re.search(
        r"\[data-theme=\"space-tech\"\] #page-settings \.settings-container\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert space_container_rule is not None
    assert "max-width: none" in space_container_rule.group("body")
    assert "[data-theme=\"space-tech\"] #page-settings .settings-page-header" in css
    assert "display: none" in re.search(
        r"\[data-theme=\"space-tech\"\] #page-settings \.settings-page-header\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    ).group("body")

    for pattern in [
        r"#page-settings \.settings-container\s*\{",
        r"#page-settings \.settings-dashboard-card\s*\{",
        r"#page-settings \.card-business \.settings-business-scroll\s*\{",
        r"#page-settings \.card-business,\s*#page-settings \.card-about\s*\{",
        r"#page-settings \.config-item-name\s*\{",
        r"#page-settings \.config-item-info\s*\{",
        r"#page-settings \.config-item-actions\s*\{",
        r"#page-settings \.reconcile-settings-panel\s*\{",
        r"#page-settings \.db-validation-source-row\s*\{",
        r"\[data-theme=\"space-tech\"\] #page-settings \.settings-dashboard-card\s*\{",
        r"\[data-theme=\"space-tech\"\]\[data-color-mode=\"dark\"\] #page-settings \.settings-dashboard-card\s*\{",
    ]:
        assert re.search(pattern, css) is not None

    assert "对账业务设置" in html
    assert "function getReconcileBusinessSourceName()" in app_js
    assert "function loadReconcileSchemaSettings()" in app_js
    assert "/api/settings/reconcile-schema/init-from-file" in app_js
    assert "filterRunsByReconcileBusinessSource" not in app_js
    assert "全部数据源" in app_js
    assert "defaultConfigSwitchAnimationName" not in app_js
    assert 'config-item--default-switched' not in app_js
    assert "function applyDefaultConfigLocally" not in app_js
    assert "设为默认" not in html
    assert "modalSetDefault" not in html

    span_one_rule = re.search(
        r"#page-settings \.card-system-info,\s*#page-settings \.card-data,\s*#page-settings \.card-default,\s*#page-settings \.card-flow,\s*#page-settings \.card-about\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    span_two_rule = re.search(
        r"#page-settings \.card-datasource,\s*#page-settings \.card-business\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    db_validation_rule = re.search(
        r"#page-settings \.card-db-validation\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert span_one_rule is not None
    assert span_two_rule is not None
    assert db_validation_rule is not None
    assert "grid-column: span 1" in span_one_rule.group("body")
    assert "grid-column: span 2" in span_two_rule.group("body")
    assert "grid-column: 1 / -1" in db_validation_rule.group("body")
    assert "db-validation-settings-grid" in settings_html
    assert settings_html.count('class="db-validation-source-row') == 4
    assert settings_html.index("报表信息配置数据源") < settings_html.index("逐笔数据源")
    assert "字段匹配数据源" not in settings_html
    assert settings_html.count('<span class="setting-label">报送子系统编号</span>') == 3
    assert settings_html.count('<span class="setting-label">分类编号</span>') == 3
    assert settings_html.count("填写报送子系统编号，多个用;分隔") == 3
    assert settings_html.count("填写分类编号，多个用;分隔") == 3
    assert "对账报表库数据源" not in settings_html
    assert "对账业务库数据源" not in settings_html
    assert '<span class="setting-label">sys_manage_id</span>' not in settings_html
    assert '<span class="setting-label">classification_id</span>' not in settings_html
    assert 'id="dbValidationBaseinfoTable" type="hidden"' in settings_html
    assert 'id="dbValidationFieldInfoTable" type="hidden"' in settings_html
    assert 'id="dbValidationPublicInfoTable"' not in settings_html
    assert '<span class="setting-label">公开信息表</span>' not in settings_html
    assert "minmax(240px, 320px)" in css
    assert "minmax(260px, 1fr)" in css
    assert ".db-validation-filter-input:last-of-type" in css
    assert "人行逐笔校验引擎独占整行" in readme
    assert "数据管理移动至原主题设置位置" in readme
    assert "对账业务设置前移至数据源配置后" in readme
    assert "对账业务设置隐藏旧版全局对账数据源选择" in readme
    assert "自动对数执行以表字段配置中的表级数据源为准" in readme
    assert "系统信息改为展示历史核对次数、登录用户和首页自动刷新等运行指标" in readme
    assert "系统设置中的“对账业务设置”" in readme
    equal_height_rule = re.search(
        r"#page-settings \.card-business,\s*#page-settings \.card-about\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert equal_height_rule is not None
    assert "height: 800px" in equal_height_rule.group("body")
    assert "overflow-y: auto" in re.search(
        r"#page-settings \.card-business \.settings-business-scroll\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    ).group("body")
    about_body_rule = re.search(
        r"#page-settings \.card-about \.card-body\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert about_body_rule is not None
    assert "display: flex" in about_body_rule.group("body")
    assert "flex-direction: column" in about_body_rule.group("body")
    assert "overflow: hidden" in about_body_rule.group("body")
    about_content_rule = re.search(
        r"#page-settings \.card-about \.about-content\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert about_content_rule is not None
    assert "overflow-y: auto" in about_content_rule.group("body")
    about_links_rule = re.search(
        r"#page-settings \.card-about \.about-links\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert about_links_rule is not None
    assert "flex-shrink: 0" in about_links_rule.group("body")
    assert 'id="aboutHelp"' in html
    assert 'id="aboutChangelog"' in html
    assert re.search(
        r'class="about-content">.*?class="about-links"',
        html,
        re.S,
    ) is not None
    assert 'setupSettingsDashboardCollapsible("businessSettingsToggle", "businessSettingsBody")' not in app_js
    assert "settings-expanded-card" not in css
    assert "settings-expanded-card" not in app_js
    assert "<h4>主要功能</h4>" in app_js
    assert "<h4>技术栈</h4>" in app_js

    assert 'class="settings-header"' not in html
    assert ".settings-header" not in css
    assert 'if (!toggle.classList.contains("collapsible")) return;' in app_js
    assert 'if (!configToggle.classList.contains("collapsible")) return;' in app_js


def test_interface_settings_card_is_shared_and_has_one_exact_radius_slider():
    html = _read(INDEX_HTML)
    settings_section = re.search(
        r'<section class="page" id="page-settings">(?P<body>.*?)\n      </section>\n\n      <!-- 确认弹窗 -->',
        html,
        re.S,
    )
    assert settings_section is not None
    settings_html = settings_section.group("body")

    interface_card = re.search(
        r'<section class="card settings-dashboard-card card-interface">(?P<body>.*?)</section>',
        settings_html,
        re.S,
    )
    assert interface_card is not None
    card_html = interface_card.group("body")
    assert "admin-only" not in interface_card.group(0)
    assert "<h3>界面设置</h3>" in card_html
    assert "<p>配置系统圆角和折线图风格</p>" in card_html
    assert '<input id="interfaceRadiusSlider" type="range" min="1" max="15" step="1" value="4" />' in card_html
    assert '<output id="interfaceRadiusValue">4px</output>' in card_html
    assert '<span id="interfaceSettingsStatus" role="status">已保存</span>' in card_html
    assert '<button id="saveInterfaceSettingsBtn" type="button" class="btn-outline btn-sm" data-action-tone="primary" data-action-variant="weak">保存界面设置</button>' in card_html
    assert '<button id="resetInterfaceSettingsBtn" type="button" class="btn-outline btn-sm" data-action-tone="warning" data-action-variant="weak">恢复默认</button>' in card_html
    assert "导航、卡片、弹窗、矩形按钮和输入选择将统一使用该圆角" not in card_html
    assert html.count('id="interfaceRadiusSlider"') == 1
    assert len(re.findall(r'<input[^>]+id="interfaceRadiusSlider"[^>]*>', html)) == 1

    system_info_pos = settings_html.index('class="card settings-dashboard-card card-system-info"')
    interface_pos = settings_html.index('class="card settings-dashboard-card card-interface"')
    default_pos = settings_html.index('class="card settings-dashboard-card card-default admin-only"')
    assert system_info_pos < interface_pos < default_pos


def test_user_radius_override_is_semantic_and_border_radius_only():
    css = _read(STYLES_CSS)
    start_marker = "/* User interface radius preference: start */"
    end_marker = "/* User interface radius preference: end */"

    assert css.count(start_marker) == 1
    assert css.count(end_marker) == 1
    override = css.split(start_marker, 1)[1].split(end_marker, 1)[0]
    blocks = re.findall(r"(?P<selectors>[^{}]+)\{(?P<body>[^{}]+)\}", override, re.S)
    assert blocks
    selectors = {
        selector.strip()
        for selector_group, _body in blocks
        for selector in selector_group.split(",")
        if selector.strip()
    }

    for selector in (
        ".nav-item",
        ".nav-group-toggle",
        ".nav-submenu",
        '[data-theme="space-tech"] .top-nav-item',
        '[data-theme="space-tech"] .top-nav-group-toggle',
        '[data-theme="space-tech"] .top-nav-submenu',
        '[data-theme="space-tech"] .top-nav-subitem',
        ".card",
        ".home-stat-card",
        ".home-analysis-card",
        "#page-home .glass-card",
        "#page-home .glass-stat-card",
        ".tool-card",
        ".run-log-panel",
        "#page-report-navigation .report-nav-stat-card",
        "#page-report-navigation .report-nav-card",
        "#page-report-navigation .report-nav-branch-panel",
        "#page-report-navigation .report-nav-batch",
        "#page-report-navigation .report-nav-todo",
        "#page-report-navigation .report-nav-load-state",
        ".toast",
        ".flow-toast",
        ".top-nav-status",
        ".dark-mode-toggle",
        ".sidebar-footer .status",
        "#page-report-navigation .report-nav-refresh-button",
        ".history-summary-item",
        ".history-section",
        ".detail-block",
        ".detail-item",
        ".status-badge",
        ".flow-chain-list",
        ".flow-chain-selection-summary",
        ".flow-run-panel:last-child #flowLog",
        ".flow-history-table-wrap",
        ".db-validation-history-table-wrap",
        ".pbc-upload-area",
        "#page-settings .metric-item",
        "#page-settings .db-validation-source-row",
        "#page-settings .reconcile-schema-table",
        "#page-settings .business-settings-note",
        "#page-settings .business-field-group",
        ".flow-chain-config",
        ".config-item",
        "#page-settings .about-description",
        "#page-settings .about-features",
        "#page-settings .about-tech",
        "#page-settings .settings-dashboard-card",
        "#page-users .user-stat-card",
        "#page-users .user-filter-bar",
        "#page-users .user-table-card",
        ".reconcile-settings-panel",
        ".db-validation-panel",
        ".flow-run-panel",
        ".flow-step-builder-panel",
        "#page-local-storage .local-storage-metric",
        "#page-local-storage .local-storage-table-panel",
        "#page-local-storage .local-storage-detail-panel",
        ".modal",
        ".pbc-modal",
        ".user-modal",
        ".btn-primary",
        ".btn-outline",
        ".btn-danger",
        ".btn-stop",
        ".btn-confirm-primary",
        ".btn-close",
        ".page-btn",
        ".trend-quick-btn",
        ".trend-quick-btns",
        ".pbc-btn",
        ".report-nav-action-button",
        ".info-detail-action",
        ".flow-toast-action",
        ".user-menu-trigger",
        ".user-menu-panel",
        ".user-menu-logout",
        ".filter-input",
        ".filter-select",
        ".chart-date-select",
        ".setting-input",
        ".prompt-input",
        ".user-form-control",
        '.main-content input:not([type="checkbox"]):not([type="radio"]):not([type="range"]):not([type="file"]):not([type="hidden"])',
        ".main-content select",
        ".main-content textarea",
        '.modal-field input:not([type="checkbox"]):not([type="radio"]):not([type="range"])',
        ".modal-field select",
        '.modal input:not([type="checkbox"]):not([type="radio"]):not([type="range"]):not([type="file"]):not([type="hidden"])',
        ".modal select",
        ".modal textarea",
        ".custom-input-shell",
        ".custom-select-shell",
        ".custom-select-trigger",
        ".custom-select-dropdown",
        ".custom-date-shell",
        ".custom-date-dropdown",
        "#page-report-navigation .report-nav-filter-chips span",
    ):
        assert selector in selectors

    for forbidden_selector in (
        ".login-card",
        ".login-form",
        ".login-input",
        ".brand-theme-toggle",
        ".user-initial-avatar",
        ".user-menu-icon",
        ".status-dot",
        ".filter-clear-button",
        ".expand-btn",
        ".custom-date-nav",
        ".custom-date-day",
        ".user-filter-pill",
        ".user-icon-action",
        ".user-enable-switch",
        ".flow-toast-close",
        ".pbc-file-remove-btn",
        ".pbc-mapping-action",
        ".tool-card-badge",
        ".badge",
        ".tag",
        ".progress-bar",
        ".progress-fill",
        ".checkbox",
        ".radio",
        ".range-track",
        ".range-thumb",
        "svg",
        "overlay",
        "[class*=card]",
    ):
        assert forbidden_selector not in selectors

    for selector in selectors:
        assert "*" not in selector
        assert re.match(r"^(?:button|input|select|textarea)(?:$|[.#:\[])", selector) is None

    for _selector_group, body in blocks:
        declarations = [item.strip() for item in body.split(";") if item.strip()]
        assert declarations == ["border-radius: var(--ui-radius) !important"]


def test_remaining_user_modal_home_stat_validation_flow_and_report_radius_overrides_are_scoped():
    css = _read(STYLES_CSS)
    start_marker = "/* User interface radius preference: start */"
    end_marker = "/* User interface radius preference: end */"
    override = css.split(start_marker, 1)[1].split(end_marker, 1)[0]
    blocks = re.findall(r"(?P<selectors>[^{}]+)\{(?P<body>[^{}]+)\}", override, re.S)
    selectors = {
        selector.strip()
        for selector_group, _body in blocks
        for selector in selector_group.split(",")
        if selector.strip()
    }

    required_selectors = {
        "#page-users .user-filter-pill",
        "#page-users .user-avatar",
        "#page-users .role-badge",
        "#page-users .user-status-badge",
        ".user-modal .user-role-card",
        ".user-modal .user-role-card-icon",
        ".user-modal .user-enable-row",
        "#configModal .modal-section",
        "#infoModal .home-stat-modal-table-wrap",
        "#dbValidationModal .db-validation-table-item",
        "#dbValidationModal #dbValidationLog",
        ".flow-chain-editor-overlay .flow-definition-table",
        ".flow-chain-editor-overlay .flow-selected-step",
        ".flow-chain-editor-overlay .flow-selected-step-actions .btn-icon",
        "#page-report-navigation .report-nav-done-meta",
        "#page-report-navigation .report-nav-no-panel-done-meta",
    }
    assert len(required_selectors) == 16
    assert required_selectors <= selectors

    for excluded_selector in (
        ".user-avatar-status",
        ".current-user-badge",
        ".user-enable-switch",
        ".user-enable-switch .switch-track",
        ".user-enable-switch .switch-thumb",
        ".flow-selected-step-actions .btn-icon",
    ):
        assert excluded_selector not in selectors


def test_readme_documents_expanded_interface_radius_surface_coverage():
    readme = _read(README_MD)

    for text in (
        "鱼骨详情卡",
        "报送日期分组卡",
        "统计周期分段选择器",
        "注意事项卡",
        "注意事项筛选标签",
        "统计失败提示条",
        "右上角通知",
        "顶部版本标签",
        "流程浮动通知",
        "历史详情与执行历史表格外框",
        "结果详情分区",
        "结果状态标签",
        "一键导入上传区",
        "系统信息指标卡",
        "数据源与流程链配置行",
        "关于系统内容卡",
        "对账表字段配置卡",
        "对账业务维护说明条",
        "对账业务字段表格分组",
    ):
        assert text in readme


def test_pbc_close_uses_shared_radius_instead_of_diamond():
    css = _read(STYLES_CSS)

    for selector in ("tool-card", "pbc-modal"):
        rule = re.search(
            rf"(?m)^\.{selector}\s*\{{(?P<body>.*?)\}}",
            css,
            re.S,
        )
        assert rule is not None
        assert re.search(r"clip-path\s*:\s*polygon\(", rule.group("body")) is None

    close_rule = re.search(
        r"(?m)^\.pbc-modal-close\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert close_rule is not None
    close_body = close_rule.group("body")
    assert "clip-path" not in close_body
    assert "rotate(90deg)" not in close_body
    assert "rotate(45deg)" not in close_body

    override = css.split("/* User interface radius preference: start */", 1)[1].split(
        "/* User interface radius preference: end */",
        1,
    )[0]
    assert ".pbc-modal-close" in override


def test_all_system_modals_use_balanced_shared_shell():
    html = _read(INDEX_HTML)
    css = _read(STYLES_CSS)

    overlay_ids = (
        "pbcModalOverlay",
        "dbValidationModalOverlay",
        "dbValidationHistoryOverlay",
        "dbValidationMappingOverlay",
        "crossTableMappingPromptModal",
        "flowModalOverlay",
        "flowHistoryOverlay",
        "flowChainEditorOverlay",
        "confirmModal",
        "promptModal",
        "infoModal",
        "reportNavTodoAllModal",
        "reportNavHistoryModal",
        "reportNavCardMaintenanceModal",
        "userModal",
        "configModal",
        "rolePermissionsModal",
        "roleDefinitionModal",
    )
    for overlay_id in overlay_ids:
        opening = re.search(
            rf'<div class="(?P<classes>[^"]+)" id="{overlay_id}"',
            html,
        )
        assert opening is not None
        assert "app-modal-overlay" in opening.group("classes").split()

    assert html.count("app-modal-shell") == len(overlay_ids)
    assert html.count("app-modal-header") == len(overlay_ids)
    assert "pbc-modal-icon" not in html
    assert "user-modal-icon" not in html

    for selector in (
        ".app-modal-overlay",
        ".app-modal-shell",
        ".app-modal-header",
        ".app-modal-body",
        ".app-modal-footer",
        ".app-modal-close",
    ):
        assert selector in css
    assert '[data-color-mode="dark"] .app-modal-overlay' in css
    assert '[data-color-mode="dark"] .app-modal-shell' in css


def test_all_modal_footers_use_the_neutral_shared_surface():
    html = _read(INDEX_HTML)
    css = _read(STYLES_CSS)

    footer_classes = [
        classes
        for classes in re.findall(r'<div class="(?P<classes>[^"]+)"', html)
        if {"modal-footer", "pbc-modal-footer"}.intersection(classes.split())
    ]
    assert footer_classes
    assert all("app-modal-footer" in classes.split() for classes in footer_classes)

    shared_footer = re.search(
        r"(?m)^\.app-modal-shell \.app-modal-footer\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert shared_footer is not None
    assert "background: var(--surface-container-lowest)" in shared_footer.group("body")
    assert "border-top: 1px solid var(--outline-variant)" in shared_footer.group("body")
    for layout_declaration in ("display:", "height:", "min-height:", "padding:", "margin:"):
        assert layout_declaration not in shared_footer.group("body")

    confirm_footer = re.search(
        r"(?m)^\.modal-confirm \.modal-footer\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert confirm_footer is not None
    assert "surface-container-high" in confirm_footer.group("body")
    assert css.index(".modal-confirm .modal-footer") < css.index(
        ".app-modal-shell .app-modal-footer"
    )

    dark_footer = re.search(
        r'(?m)^\[data-color-mode="dark"\] \.app-modal-shell \.app-modal-footer\s*'
        r"\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert dark_footer is not None
    assert "background: #0f172a" in dark_footer.group("body")


def test_modal_surfaces_are_solid_and_tool_gray_is_limited_to_progress_tracks():
    css = _read(STYLES_CSS)

    start = css.index(
        "/* Modal surfaces stay solid; neutral gray is reserved for tool progress tracks. */"
    )
    end = css.index("/* User interface radius preference: start */", start)
    surface_contract = css[start:end]

    for selector in (
        ".app-modal-shell > .app-modal-header",
        ".app-modal-shell > .app-modal-body",
        ".app-modal-shell > .app-modal-footer",
        "#pbcModal :is(",
        "#dbValidationModal :is(",
        "#flowModal :is(",
        ".db-validation-panel",
        ".db-validation-table-item",
        ".flow-run-panel",
        ".flow-chain-list",
        ".pbc-import-log",
        "#flowLog",
    ):
        assert selector in surface_contract

    assert surface_contract.count(
        "background: var(--surface-container-lowest);"
    ) >= 2
    assert surface_contract.count("background: var(--surface-container);") == 1
    for progress_selector in (
        ".pbc-upload-progress-track",
        ".pbc-progress-bar-track",
    ):
        assert progress_selector in surface_contract
    assert ".pbc-step-num" not in surface_contract
    step_state = re.search(
        r"(?m)^\.pbc-step--active \.pbc-step-num,\s*\n"
        r"\.pbc-step--done \.pbc-step-num\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert step_state is not None
    assert "background: var(--primary)" in step_state.group("body")
def test_modal_table_headers_match_history_tokens_without_layout_overrides():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    global_headers = re.search(r"(?m)^th\s*\{(?P<body>.*?)\}", css, re.S)
    assert global_headers is not None
    global_body = global_headers.group("body")
    for preserved_declaration in (
        "position: sticky",
        "top: 0",
        "z-index: 2",
        "font-size: 12px",
    ):
        assert preserved_declaration in global_body
    for visual_declaration in (
        "color: var(--on-surface-variant)",
        "font-weight: 600",
        "background: var(--surface-container-low)",
        "border-bottom: 1px solid color-mix(in srgb, var(--outline-variant) 32%, var(--surface-container-lowest))",
    ):
        assert visual_declaration in global_body

    global_dark_headers = re.search(
        r'(?m)^\[data-color-mode="dark"\] th\s*\{(?P<body>.*?)\}',
        css,
        re.S,
    )
    assert global_dark_headers is not None
    assert "color: #cbd5e1" in global_dark_headers.group("body")
    assert "background: rgba(30, 41, 59, 0.94)" in global_dark_headers.group("body")

    shared_headers = re.search(
        r"(?m)^\.app-modal-shell table th,\s*\n"
        r"\.app-modal-shell \.pbc-file-list-header,\s*\n"
        r"\.app-modal-shell \.db-validation-table-header,\s*\n"
        r"\.app-modal-shell \.flow-def-header\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert shared_headers is not None
    shared_body = shared_headers.group("body")
    for declaration in (
        "color: var(--on-surface-variant)",
        "font-weight: 600",
        "background: var(--surface-container-low)",
        "border-bottom: 1px solid color-mix(in srgb, var(--outline-variant) 32%, var(--surface-container-lowest))",
    ):
        assert declaration in shared_body
    for layout_property in (
        "text-align:",
        "position:",
        "top:",
        "z-index:",
        "padding:",
        "width:",
        "height:",
        "display:",
        "grid-template-columns:",
        "!important",
    ):
        assert layout_property not in shared_body

    shared_rule_start = css.index(".app-modal-shell table th,")
    assert '<div class="app-modal-shell modal modal-info">' in html
    assert 'class="home-stat-modal-table"' in app_js
    assert ".app-modal-shell table th" in shared_headers.group(0)
    assert css.index(".home-stat-modal-table th") < shared_rule_start
    for modal_specific_selector in (
        ".db-validation-history-table th",
        ".flow-def-header th",
    ):
        assert css.index(modal_specific_selector) < shared_rule_start

    dark_headers = re.search(
        r'(?m)^\[data-color-mode="dark"\] \.app-modal-shell table th,\s*\n'
        r'\[data-color-mode="dark"\] \.app-modal-shell \.pbc-file-list-header,\s*\n'
        r'\[data-color-mode="dark"\] \.app-modal-shell \.db-validation-table-header,\s*\n'
        r'\[data-color-mode="dark"\] \.app-modal-shell \.flow-def-header\s*'
        r"\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert dark_headers is not None
    dark_body = dark_headers.group("body")
    assert "color: #cbd5e1" in dark_body
    assert "background: rgba(30, 41, 59, 0.94)" in dark_body
    assert (
        "border-bottom: 1px solid color-mix(in srgb, var(--outline-variant) 32%, "
        "var(--surface-container-lowest))"
    ) in dark_body

    user_header = re.search(r"(?m)^\.user-table th\s*\{(?P<body>.*?)\}", css, re.S)
    assert user_header is not None
    assert "font-size: 12px" in user_header.group("body")
    for duplicate_visual in ("background:", "color:", "font-weight:"):
        assert duplicate_visual not in user_header.group("body")


def test_main_result_and_history_headers_match_user_table_height_only():
    css = _read(STYLES_CSS)
    main_headers = re.search(
        r"(?m)^\.result-card > \.table-wrap > \.result-table > thead > tr > th\s*"
        r"\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert main_headers is not None
    assert "padding-top: 14px" in main_headers.group("body")
    assert "padding-bottom: 14px" in main_headers.group("body")
    assert "height:" not in main_headers.group("body")
    assert ".app-modal-shell" not in main_headers.group(0)


def test_result_detail_labels_use_table_header_color_and_values_are_transparent():
    css = _read(STYLES_CSS)

    detail_item = re.search(r"(?m)^\.detail-item\s*\{(?P<body>.*?)\}", css, re.S)
    assert detail_item is not None
    assert "background: transparent" in detail_item.group("body")

    detail_label = re.search(r"(?m)^\.detail-item span\s*\{(?P<body>.*?)\}", css, re.S)
    assert detail_label is not None
    assert "background: var(--surface-container-low)" in detail_label.group("body")
    assert "surface-container-high" not in detail_label.group("body")


def test_section_headers_use_the_shared_table_header_color():
    css = _read(STYLES_CSS)

    for selector in (
        ".modal-section-header",
        ".business-field-header",
        ".db-validation-panel",
        ".flow-chain-selected-count",
        ".flow-chain-list",
        ".flow-chain-selection-summary",
        ".pbc-import-log",
        ".flow-run-panel:last-child #flowLog",
    ):
        section_header = re.search(
            rf"(?m)^{re.escape(selector)}\s*\{{(?P<body>.*?)\}}",
            css,
            re.S,
        )
        assert section_header is not None
        assert "background: var(--surface-container-low)" in section_header.group("body")
        assert "surface-container-high" not in section_header.group("body")


def test_primary_tool_modals_preserve_their_pre_shared_layout_contract():
    html = _read(INDEX_HTML)
    css = _read(STYLES_CSS)

    content_markers = {
        "pbcModal": "<!-- Steps indicator -->",
        "dbValidationModal": '<div class="db-validation-grid">',
        "flowModal": '<div class="flow-run-grid">',
    }
    for modal_id, content_marker in content_markers.items():
        modal = re.search(
            rf'<div class="app-modal-shell [^"]+" id="{modal_id}">(?P<body>.*?)'
            r'<div class="app-modal-footer pbc-modal-footer"',
            html,
            re.S,
        )
        assert modal is not None
        assert content_marker in modal.group("body")
        assert '<div class="app-modal-body pbc-modal-body">' not in modal.group("body")

    for auxiliary_body_class in (
        "db-validation-history-table-wrap",
        "flow-history-table-wrap",
        "flow-chain-editor-body",
    ):
        assert re.search(
            r'<div class="app-modal-body pbc-modal-body">\s*'
            rf'<div class="{auxiliary_body_class}">',
            html,
        )

    shared_shell = re.search(
        r"(?m)^\.app-modal-shell\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert shared_shell is not None
    for layout_declaration in (
        "display: flex",
        "flex-direction: column",
        "padding: 0",
        "overflow: hidden",
    ):
        assert layout_declaration not in shared_shell.group("body")

    primary_safe_scope = (
        ".app-modal-shell:not(#pbcModal):not(#dbValidationModal):not(#flowModal)"
    )
    for selector_suffix in (
        "",
        " > .app-modal-header",
        " > .app-modal-body",
        " .app-modal-close",
    ):
        assert f"{primary_safe_scope}{selector_suffix}" in css
    assert f"{primary_safe_scope}:not(.modal-info--history-detail) > .app-modal-footer" in css

    pbc_modal = re.search(r"(?m)^\.pbc-modal\s*\{(?P<body>.*?)\}", css, re.S)
    pbc_header = re.search(r"(?m)^\.pbc-modal-header\s*\{(?P<body>.*?)\}", css, re.S)
    pbc_footer = re.search(r"(?m)^\.pbc-modal-footer\s*\{(?P<body>.*?)\}", css, re.S)
    flow_modal = re.search(r"(?m)^\.flow-modal\s*\{(?P<body>.*?)\}", css, re.S)
    assert pbc_modal is not None
    assert "padding: 32px 32px 24px" in pbc_modal.group("body")
    assert "overflow-y: auto" in pbc_modal.group("body")
    assert pbc_header is not None
    assert "margin-bottom: 24px" in pbc_header.group("body")
    assert "padding-bottom: 16px" in pbc_header.group("body")
    assert pbc_footer is not None
    assert "margin-top: 24px" in pbc_footer.group("body")
    assert "padding-top: 16px" in pbc_footer.group("body")
    assert flow_modal is not None
    assert "display: flex" in flow_modal.group("body")
    assert "overflow: hidden" in flow_modal.group("body")

    shared_close = re.search(
        r"(?m)^\.app-modal-close\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert shared_close is not None
    for layout_declaration in ("top:", "right:", "width:", "height:"):
        assert layout_declaration not in shared_close.group("body")

    pbc_close = re.search(r"(?m)^\.pbc-modal-close\s*\{(?P<body>.*?)\}", css, re.S)
    assert pbc_close is not None
    assert "top: 16px; right: 16px" in pbc_close.group("body")
    assert "width: 36px; height: 36px" in pbc_close.group("body")


def test_primary_tool_modal_header_margins_are_not_overridden_by_shared_visuals():
    css = _read(STYLES_CSS)

    shared_headings = re.search(
        r"(?m)^\.app-modal-header h2,\s*\n\.app-modal-header h3\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert shared_headings is not None
    assert "margin:" not in shared_headings.group("body")

    shared_paragraph = re.search(
        r"(?m)^\.app-modal-header p\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert shared_paragraph is not None
    assert "margin:" not in shared_paragraph.group("body")

    primary_safe_scope = (
        ".app-modal-shell:not(#pbcModal):not(#dbValidationModal):not(#flowModal)"
    )
    safe_headings = re.search(
        rf"(?m)^{re.escape(primary_safe_scope)} > \.app-modal-header h2,\s*\n"
        rf"{re.escape(primary_safe_scope)} > \.app-modal-header h3\s*\{{(?P<body>.*?)\}}",
        css,
        re.S,
    )
    assert safe_headings is not None
    assert "margin: 0" in safe_headings.group("body")

    safe_paragraph = re.search(
        rf"(?m)^{re.escape(primary_safe_scope)} > \.app-modal-header p\s*\{{(?P<body>.*?)\}}",
        css,
        re.S,
    )
    assert safe_paragraph is not None
    assert "margin: 4px 0 0" in safe_paragraph.group("body")


def test_interface_radius_has_default_and_regular_user_three_card_responsive_layout():
    css = _read(STYLES_CSS)

    root = re.search(r"(?m)^:root\s*\{(?P<body>.*?)\n\}", css, re.S)
    assert root is not None
    assert "--ui-radius: 4px;" in root.group("body")

    interface_body = re.search(
        r"(?m)^#page-settings \.card-interface \.card-body\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert interface_body is not None
    assert "display: flex" in interface_body.group("body")
    assert "flex-direction: column" in interface_body.group("body")
    assert "flex: 1" in interface_body.group("body")

    interface_control = re.search(
        r"(?m)^#page-settings \.card-interface \.setting-item\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert interface_control is not None
    assert "display: grid" in interface_control.group("body")
    assert "grid-template-columns:" in interface_control.group("body")

    slider = re.search(
        r"(?m)^#page-settings #interfaceRadiusSlider\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert slider is not None
    assert "width: 100%" in slider.group("body")
    assert "accent-color: var(--secondary)" in slider.group("body")

    value = re.search(
        r"(?m)^#page-settings #interfaceRadiusValue\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert value is not None
    assert "color: var(--on-surface)" in value.group("body")

    status = re.search(
        r"(?m)^#page-settings #interfaceSettingsStatus\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert status is not None
    assert "color: var(--on-surface-variant)" in status.group("body")

    desktop_user_grid = re.search(
        r'(?m)^\[data-role="user"\] #page-settings \.settings-dashboard-grid\s*\{(?P<body>.*?)\}',
        css,
        re.S,
    )
    assert desktop_user_grid is not None
    assert "grid-template-columns: repeat(3, minmax(0, 1fr))" in desktop_user_grid.group("body")
    assert "align-items: stretch" in desktop_user_grid.group("body")

    user_cards = re.search(
        r'\[data-role="user"\] #page-settings \.card-system-info,\s*'
        r'\[data-role="user"\] #page-settings \.card-interface,\s*'
        r'\[data-role="user"\] #page-settings \.card-about\s*\{(?P<body>.*?)\}',
        css,
        re.S,
    )
    assert user_cards is not None
    assert "grid-column: span 1" in user_cards.group("body")
    assert "height: 100%" in user_cards.group("body")

    tablet_css = css[css.index("@media (max-width: 1200px)") :]
    tablet_user_grid = re.search(
        r'\[data-role="user"\] #page-settings \.settings-dashboard-grid\s*\{(?P<body>.*?)\}',
        tablet_css,
        re.S,
    )
    assert tablet_user_grid is not None
    assert "grid-template-columns: repeat(2, minmax(0, 1fr))" in tablet_user_grid.group("body")

    mobile_css = css[css.index("@media (max-width: 760px)") :]
    mobile_user_grid = re.search(
        r'\[data-role="user"\] #page-settings \.settings-dashboard-grid\s*\{(?P<body>.*?)\}',
        mobile_css,
        re.S,
    )
    assert mobile_user_grid is not None
    assert "grid-template-columns: 1fr" in mobile_user_grid.group("body")

    admin_grid = re.search(
        r"(?m)^#page-settings \.settings-dashboard-grid\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert admin_grid is not None
    assert "grid-template-columns: repeat(3, minmax(0, 1fr))" in admin_grid.group("body")


def test_interface_radius_loads_before_theme_and_auth_reveal_with_internal_fallback():
    app_js = _read(APP_JS)

    ensure_auth = re.search(
        r"async function ensureAuthenticated\(\) \{(?P<body>.*?)\n\}",
        app_js,
        re.S,
    )
    assert ensure_auth is not None
    auth_body = ensure_auth.group("body")
    reset_call = "resetInterfaceRadiusForAuthChange();"
    load_call = "await loadInterfaceRadiusPreference({ silent: true });"
    assert reset_call in auth_body
    assert load_call in auth_body
    assert auth_body.index(reset_call) < auth_body.index('authState.csrfToken = payload.csrf_token || "";')
    assert auth_body.index(reset_call) < auth_body.index("authState.user = payload.user || null;")
    assert auth_body.index("authState.user = payload.user || null;") < auth_body.index(load_call)
    assert auth_body.index("document.documentElement.dataset.role") < auth_body.index(load_call)
    assert auth_body.index(load_call) < auth_body.index("applySavedUserTheme();")
    assert auth_body.index(load_call) < auth_body.index("revealAuthenticatedApp();")

    loader = re.search(
        r"async function loadInterfaceRadiusPreference\(\{ silent = false \} = \{\}\) \{(?P<body>.*?)\n\}",
        app_js,
        re.S,
    )
    assert loader is not None
    load_body = loader.group("body")
    assert 'api("/api/settings/interface", { signal: abortController.signal })' in load_body
    assert "const requestId = ++interfaceRadiusState.loadRequestId;" in load_body
    assert "const editRevision = interfaceRadiusState.editRevision;" in load_body
    assert "const mutationRevision = interfaceRadiusState.serverMutationRevision;" in load_body
    assert "const abortController = new AbortController();" in load_body
    assert "setTimeout(() => abortController.abort(), INTERFACE_RADIUS_LOAD_TIMEOUT_MS)" in load_body
    assert "clearTimeout(timeoutId);" in load_body
    assert "requestId !== interfaceRadiusState.loadRequestId" in load_body
    assert "mutationRevision !== interfaceRadiusState.serverMutationRevision" in load_body
    assert "editRevision === interfaceRadiusState.editRevision" in load_body
    assert "const preferences = readInterfacePreferencesPayload(payload);" in load_body
    assert "} catch (error) {" in load_body
    assert "if (!interfaceRadiusState.loaded)" in load_body
    assert "interfaceRadiusState.savedPreferences = { ...DEFAULT_INTERFACE_PREFERENCES };" in load_body
    assert "interfaceRadiusState.draftPreferences = { ...DEFAULT_INTERFACE_PREFERENCES };" in load_body
    assert "applyInterfacePreferences(DEFAULT_INTERFACE_PREFERENCES);" in load_body
    assert "interfaceRadiusState.loadFailed = true;" in load_body
    assert 'interfaceRadiusState.statusText = "加载失败，当前使用默认 4px";' in load_body
    assert "if (!silent)" in load_body
    assert "showToast(" in load_body
    assert "return false;" in load_body


def test_interface_radius_state_normalization_rendering_and_api_boundary():
    app_js = _read(APP_JS)
    block = re.search(
        r"// Interface radius start(?P<body>.*?)// Interface radius end",
        app_js,
        re.S,
    )
    assert block is not None
    body = block.group("body")

    assert "const DEFAULT_INTERFACE_PREFERENCES = Object.freeze({" in body
    assert "const MIN_INTERFACE_RADIUS_PX = 1;" in body
    assert "const MAX_INTERFACE_RADIUS_PX = 15;" in body
    assert "const INTERFACE_RADIUS_LOAD_TIMEOUT_MS = 2500;" in body
    for state_line in [
        "savedPreferences: { ...DEFAULT_INTERFACE_PREFERENCES }",
        "draftPreferences: { ...DEFAULT_INTERFACE_PREFERENCES }",
        "loaded: false",
        "loadFailed: false",
        "saving: false",
        'statusText: "已保存"',
        "loadRequestId: 0",
        "editRevision: 0",
        "serverMutationRevision: 0",
    ]:
        assert state_line in body

    normalize = re.search(
        r"function normalizeInterfaceRadius\(radiusPx\) \{(?P<body>.*?)\n\}",
        body,
        re.S,
    )
    assert normalize is not None
    normalize_body = normalize.group("body")
    assert "Number.isInteger(radiusPx)" in normalize_body
    assert "radiusPx >= MIN_INTERFACE_RADIUS_PX" in normalize_body
    assert "radiusPx <= MAX_INTERFACE_RADIUS_PX" in normalize_body
    assert "return DEFAULT_INTERFACE_RADIUS_PX;" in normalize_body

    apply_radius = re.search(
        r"function applyInterfaceRadius\(radiusPx\) \{(?P<body>.*?)\n\}",
        body,
        re.S,
    )
    assert apply_radius is not None
    apply_body = apply_radius.group("body")
    assert 'document.documentElement.style.setProperty("--ui-radius", `${normalizedRadiusPx}px`);' in apply_body
    assert "return normalizedRadiusPx;" in apply_body
    assert apply_body.count("setProperty(") == 1
    assert "api(" not in apply_body
    assert "localStorage" not in apply_body

    strict_payload = re.search(
        r"function readInterfacePreferencesPayload\(payload\) \{(?P<body>.*?)\n\}",
        body,
        re.S,
    )
    assert strict_payload is not None
    strict_body = strict_payload.group("body")
    assert "payload?.settings?.radius_px" in strict_body
    assert "payload?.settings?.line_chart_style" in strict_body
    assert "!Number.isInteger(radiusPx)" in strict_body
    assert "radiusPx < MIN_INTERFACE_RADIUS_PX" in strict_body
    assert "radiusPx > MAX_INTERFACE_RADIUS_PX" in strict_body
    assert "throw new Error(" in strict_body

    render = re.search(
        r"function renderInterfaceRadiusPreference\(\) \{(?P<body>.*?)\n\}",
        body,
        re.S,
    )
    assert render is not None
    render_body = render.group("body")
    assert "interfaceRadiusSlider.value = String(interfaceRadiusState.draftPreferences.radiusPx);" in render_body
    assert "interfaceRadiusValue.textContent = `${interfaceRadiusState.draftPreferences.radiusPx}px`;" in render_body
    assert 'interfaceLineChartStyleStraight.checked = interfaceRadiusState.draftPreferences.lineChartStyle === "straight";' in render_body
    assert "interfaceSettingsStatus.textContent = interfaceRadiusState.statusText;" in render_body
    assert "const saving = interfaceRadiusState.saving;" in render_body
    assert "interfaceRadiusSlider.disabled = saving;" in render_body
    assert "saveInterfaceSettingsBtn.disabled = saving;" in render_body
    assert 'saveInterfaceSettingsBtn.classList.toggle("loading", saving);' in render_body
    assert "resetInterfaceSettingsBtn.disabled = saving;" in render_body
    assert "interfaceRadiusState.statusText =" not in render_body

    api_paths = set(re.findall(r'["\'](/api/[^"\']+)["\']', body))
    assert api_paths == {"/api/settings/interface"}
    assert body.count('api("/api/settings/interface"') == 2

    assert "已保存" in body
    assert "正在预览，尚未保存" in body
    assert "保存成功" in body
    assert "保存失败" in body
    assert "加载失败，当前使用默认 4px" in body

    loader = re.search(
        r"async function loadInterfaceRadiusPreference\(\{ silent = false \} = \{\}\) \{(?P<body>.*?)\n\}",
        body,
        re.S,
    )
    assert loader is not None
    load_body = loader.group("body")
    load_success_body, load_catch_body = load_body.split("} catch (error) {", 1)
    assert "cacheAuthenticatedInterfacePreferences(preferences);" in load_success_body
    assert "cacheAuthenticatedInterfacePreferences(" not in load_catch_body


def test_interface_radius_preview_reset_save_and_discard_are_draft_safe():
    app_js = _read(APP_JS)
    block = re.search(
        r"// Interface radius start(?P<body>.*?)// Interface radius end",
        app_js,
        re.S,
    )
    assert block is not None
    body = block.group("body")

    slider_handler = re.search(
        r'interfaceRadiusSlider\?\.addEventListener\("input", \(\) => \{(?P<body>.*?)\n\}\);',
        body,
        re.S,
    )
    assert slider_handler is not None
    slider_body = slider_handler.group("body")
    assert "updateInterfacePreferenceDraft({" in slider_body
    assert "normalizeInterfaceRadius(Number(interfaceRadiusSlider.value))" in slider_body
    assert "api(" not in slider_body
    assert "POST" not in slider_body
    assert "cacheAuthenticatedInterfacePreferences(" not in slider_body

    draft_updater = re.search(
        r"function updateInterfacePreferenceDraft\(change\) \{(?P<body>.*?)\n\}",
        body,
        re.S,
    )
    assert draft_updater is not None
    draft_body = draft_updater.group("body")
    assert "if (interfaceRadiusState.saving) return;" in draft_body
    assert "interfaceRadiusState.editRevision += 1;" in draft_body
    assert "interfaceRadiusState.draftPreferences = {" in draft_body
    assert "applyInterfacePreferences(interfaceRadiusState.draftPreferences);" in draft_body
    assert "syncInterfaceRadiusDirtyStatus();" in draft_body

    reset_handler = re.search(
        r'resetInterfaceSettingsBtn\?\.addEventListener\("click", \(\) => \{(?P<body>.*?)\n\}\);',
        body,
        re.S,
    )
    assert reset_handler is not None
    reset_body = reset_handler.group("body")
    assert "if (interfaceRadiusState.saving) return;" in reset_body
    assert "interfaceRadiusState.editRevision += 1;" in reset_body
    assert "interfaceRadiusState.draftPreferences = { ...DEFAULT_INTERFACE_PREFERENCES };" in reset_body
    assert "applyInterfacePreferences(interfaceRadiusState.draftPreferences);" in reset_body
    assert "syncInterfaceRadiusDirtyStatus();" in reset_body
    assert "interfaceRadiusState.savedPreferences" not in reset_body
    assert "api(" not in reset_body
    assert "POST" not in reset_body
    assert "cacheAuthenticatedInterfacePreferences(" not in reset_body

    save = re.search(
        r"async function saveInterfaceRadiusPreference\(\) \{(?P<body>.*?)\n\}",
        body,
        re.S,
    )
    assert save is not None
    save_body = save.group("body")
    assert "if (interfaceRadiusState.saving) return false;" in save_body
    assert "interfaceRadiusState.serverMutationRevision += 1;" in save_body
    assert "interfaceRadiusState.saving = true;" in save_body
    assert 'api("/api/settings/interface", {' in save_body
    assert 'method: "POST"' in save_body
    for payload_field in (
        "radius_px: interfaceRadiusState.draftPreferences.radiusPx",
        "line_chart_style: interfaceRadiusState.draftPreferences.lineChartStyle",
    ):
        assert payload_field in save_body
    assert "const savedPreferences = readInterfacePreferencesPayload(payload);" in save_body
    assert "interfaceRadiusState.savedPreferences = copyInterfacePreferences(savedPreferences);" in save_body
    assert "interfaceRadiusState.draftPreferences = copyInterfacePreferences(savedPreferences);" in save_body
    assert 'interfaceRadiusState.statusText = "保存成功";' in save_body
    assert "applyInterfacePreferences(savedPreferences);" in save_body
    assert "cacheAuthenticatedInterfacePreferences(savedPreferences);" in save_body
    assert "} catch (error) {" in save_body
    assert "} finally {" in save_body
    catch_body = save_body.split("} catch (error) {", 1)[1].split("} finally {", 1)[0]
    assert 'interfaceRadiusState.statusText = "保存失败";' in catch_body
    assert "interfaceRadiusState.savedPreferences =" not in catch_body
    assert "interfaceRadiusState.draftPreferences =" not in catch_body
    assert "applyInterfacePreferences(" not in catch_body
    assert "cacheAuthenticatedInterfacePreferences(" not in catch_body
    finally_body = save_body.split("} finally {", 1)[1]
    assert "interfaceRadiusState.saving = false;" in finally_body
    assert "renderInterfaceRadiusPreference();" in finally_body

    discard = re.search(
        r"function discardUnsavedInterfaceRadius\(\) \{(?P<body>.*?)\n\}",
        body,
        re.S,
    )
    assert discard is not None
    discard_body = discard.group("body")
    assert "!interfacePreferencesMatch(" in discard_body
    assert "interfaceRadiusState.draftPreferences = copyInterfacePreferences(interfaceRadiusState.savedPreferences);" in discard_body
    assert "applyInterfacePreferences(interfaceRadiusState.savedPreferences);" in discard_body
    assert "syncInterfaceRadiusDirtyStatus();" in discard_body
    assert "renderInterfaceRadiusPreference();" in discard_body
    assert "api(" not in discard_body

    switch_page = re.search(
        r"async function switchPage\(name, options = \{\}\) \{(?P<body>.*?)\n\}",
        app_js,
        re.S,
    )
    assert switch_page is not None
    switch_body = switch_page.group("body")
    assert 'if (previousPage === "settings" && name !== "settings") {' in switch_body
    assert "discardUnsavedInterfaceRadius();" in switch_body
    assert "discardUnsavedSystemThemeColors();" not in switch_body
    assert switch_body.index("discardUnsavedInterfaceRadius();") < switch_body.index(
        "document.documentElement.setAttribute('data-page', name);"
    )


def test_login_uses_last_authenticated_interface_radius_display_cache():
    login_html = _read(ROOT / "src" / "auto_check" / "web" / "login.html")

    assert "--ui-radius: 4px;" in login_html
    assert 'const LAST_INTERFACE_RADIUS_CACHE_KEY = "autoCheckLastInterfaceRadius";' in login_html
    assert "function normalizeLoginInterfaceRadius(value)" in login_html
    assert "Number.isInteger(parsed)" in login_html
    assert "parsed >= 1 && parsed <= 15" in login_html
    assert "localStorage.getItem(LAST_INTERFACE_RADIUS_CACHE_KEY)" in login_html
    assert 'document.documentElement.style.setProperty("--ui-radius", `${radiusPx}px`);' in login_html
    assert login_html.index('id="initialInterfaceRadiusScript"') < login_html.index("<style>")

    for selector in (
        ".right-panel",
        ".form-input",
        ".login-btn",
    ):
        assert selector in login_html


def test_interface_radius_settings_use_server_authority_with_login_display_cache():
    app_js = _read(APP_JS)

    settings_loader = re.search(
        r"async function loadSettingsPageData\(\) \{(?P<body>.*?)\n\}",
        app_js,
        re.S,
    )
    assert settings_loader is not None
    settings_body = settings_loader.group("body")
    assert "Promise.all" in settings_body
    assert 'loadPageSection("界面设置", () => loadInterfaceRadiusPreference({ silent: false }))' in settings_body

    block = re.search(
        r"// Interface radius start(?P<body>.*?)// Interface radius end",
        app_js,
        re.S,
    )
    assert block is not None
    body = block.group("body")

    assert 'const LAST_INTERFACE_RADIUS_CACHE_KEY = "autoCheckLastInterfaceRadius";' in body
    assert "function cacheAuthenticatedInterfacePreferences(preferences)" in body
    assert "localStorage.setItem(LAST_INTERFACE_RADIUS_CACHE_KEY, String(normalizedRadiusPx));" in body
    assert "localStorage.getItem(LAST_INTERFACE_RADIUS_CACHE_KEY)" not in app_js
    assert "localStorage.removeItem(LAST_INTERFACE_RADIUS_CACHE_KEY)" not in app_js
    assert app_js.count("localStorage.setItem(LAST_INTERFACE_RADIUS_CACHE_KEY") == 1
    assert "autoCheckRadius" not in app_js


def test_interface_radius_node_keeps_new_draft_when_older_get_finishes(tmp_path):
    _run_interface_radius_node_scenario(
        tmp_path,
        """
        const h = radiusHarness;
        const getRequest = deferred();
        apiImpl = async () => getRequest.promise;

        const loading = h.load({ silent: false });
        await flushMicrotasks();
        h.elements.interfaceRadiusSlider.value = "9";
        h.elements.interfaceRadiusSlider.dispatch("input");
        assert.equal(h.state.draftPreferences.radiusPx, 9);
        assert.equal(h.cssVariables.get("--ui-radius"), "9px");

        getRequest.resolve({ settings: { radius_px: 4 } });
        assert.equal(await loading, true);
        assert.equal(h.state.savedPreferences.radiusPx, 4);
        assert.equal(h.state.draftPreferences.radiusPx, 9);
        assert.equal(h.cssVariables.get("--ui-radius"), "9px");
        assert.equal(h.state.statusText, "正在预览，尚未保存");
        assert.equal(h.elements.interfaceSettingsStatus.textContent, "正在预览，尚未保存");
        """,
    )


def test_interface_radius_node_discard_invalidates_pending_get_and_allows_reload(tmp_path):
    _run_interface_radius_node_scenario(
        tmp_path,
        """
        const h = radiusHarness;
        const oldSuccess = deferred();
        const oldFailure = deferred();
        let requestCount = 0;
        apiImpl = async () => {
          requestCount += 1;
          if (requestCount === 1) return oldSuccess.promise;
          if (requestCount === 2) return { settings: { radius_px: 6 } };
          if (requestCount === 3) return oldFailure.promise;
          return { settings: { radius_px: 7 } };
        };

        assert.equal(h.state.loaded, false);
        assert.equal(h.state.savedPreferences.radiusPx, 4);
        assert.equal(h.state.draftPreferences.radiusPx, 4);
        const staleSuccessLoading = h.load({ silent: false });
        await flushMicrotasks();
        h.elements.interfaceRadiusSlider.value = "9";
        h.elements.interfaceRadiusSlider.dispatch("input");
        assert.equal(h.discard(), true);
        assert.equal(h.state.savedPreferences.radiusPx, 4);
        assert.equal(h.state.draftPreferences.radiusPx, 4);
        assert.equal(h.state.loaded, false);
        assert.equal(h.state.loadFailed, false);
        assert.equal(h.state.statusText, "已保存");
        assert.equal(h.cssVariables.get("--ui-radius"), "4px");

        oldSuccess.resolve({ settings: { radius_px: 6 } });
        assert.equal(await staleSuccessLoading, false);
        assert.equal(h.state.savedPreferences.radiusPx, 4);
        assert.equal(h.state.draftPreferences.radiusPx, 4);
        assert.equal(h.state.loaded, false);
        assert.equal(h.state.loadFailed, false);
        assert.equal(h.state.statusText, "已保存");
        assert.equal(h.cssVariables.get("--ui-radius"), "4px");
        assert.equal(h.toasts.length, 0);

        assert.equal(await h.load({ silent: false }), true);
        assert.equal(h.state.savedPreferences.radiusPx, 6);
        assert.equal(h.state.draftPreferences.radiusPx, 6);
        assert.equal(h.state.loaded, true);
        assert.equal(h.cssVariables.get("--ui-radius"), "6px");

        const staleFailureLoading = h.load({ silent: false });
        await flushMicrotasks();
        h.elements.interfaceRadiusSlider.value = "9";
        h.elements.interfaceRadiusSlider.dispatch("input");
        assert.equal(h.discard(), true);
        oldFailure.reject(new Error("stale load failed"));
        assert.equal(await staleFailureLoading, false);
        assert.equal(h.state.savedPreferences.radiusPx, 6);
        assert.equal(h.state.draftPreferences.radiusPx, 6);
        assert.equal(h.state.loaded, true);
        assert.equal(h.state.loadFailed, false);
        assert.equal(h.state.statusText, "已保存");
        assert.equal(h.cssVariables.get("--ui-radius"), "6px");
        assert.equal(h.toasts.length, 0);

        assert.equal(await h.load({ silent: false }), true);
        assert.equal(h.state.savedPreferences.radiusPx, 7);
        assert.equal(h.state.draftPreferences.radiusPx, 7);
        assert.equal(h.state.loaded, true);
        assert.equal(h.cssVariables.get("--ui-radius"), "7px");
        assert.equal(h.toasts.length, 0);
        """,
    )


def test_interface_radius_node_ignores_get_that_predates_successful_save(tmp_path):
    _run_interface_radius_node_scenario(
        tmp_path,
        """
        const h = radiusHarness;
        const getRequest = deferred();
        const postRequest = deferred();
        apiImpl = async (_path, options) => (
          options.method === "POST" ? postRequest.promise : getRequest.promise
        );

        const loading = h.load({ silent: false });
        await flushMicrotasks();
        h.elements.interfaceRadiusSlider.value = "9";
        h.elements.interfaceRadiusSlider.dispatch("input");
        const saving = h.save();
        postRequest.resolve({ settings: { radius_px: 9 } });
        assert.equal(await saving, true);

        getRequest.resolve({ settings: { radius_px: 4 } });
        await loading;
        assert.equal(h.state.savedPreferences.radiusPx, 9);
        assert.equal(h.state.draftPreferences.radiusPx, 9);
        assert.equal(h.cssVariables.get("--ui-radius"), "9px");
        assert.equal(h.state.statusText, "保存成功");
        """,
    )


def test_interface_radius_node_disables_and_guards_draft_controls_while_saving(tmp_path):
    _run_interface_radius_node_scenario(
        tmp_path,
        """
        const h = radiusHarness;
        const postRequest = deferred();
        apiImpl = async () => postRequest.promise;

        h.elements.interfaceRadiusSlider.value = "9";
        h.elements.interfaceRadiusSlider.dispatch("input");
        const saving = h.save();
        assert.equal(h.elements.interfaceRadiusSlider.disabled, true);
        assert.equal(h.elements.interfaceLineChartStyleStraight.disabled, true);
        assert.equal(h.elements.interfaceLineChartStyleSmooth.disabled, true);
        assert.equal(h.elements.resetInterfaceSettingsBtn.disabled, true);

        h.elements.interfaceRadiusSlider.value = "8";
        h.elements.interfaceRadiusSlider.dispatch("input");
        h.elements.resetInterfaceSettingsBtn.dispatch("click");
        assert.equal(h.state.draftPreferences.radiusPx, 9);
        assert.equal(h.cssVariables.get("--ui-radius"), "9px");

        postRequest.resolve({ settings: { radius_px: 9 } });
        assert.equal(await saving, true);
        assert.equal(h.state.savedPreferences.radiusPx, 9);
        assert.equal(h.state.draftPreferences.radiusPx, 9);
        assert.equal(h.elements.interfaceRadiusSlider.disabled, false);
        assert.equal(h.elements.interfaceLineChartStyleStraight.disabled, false);
        assert.equal(h.elements.interfaceLineChartStyleSmooth.disabled, false);
        assert.equal(h.elements.resetInterfaceSettingsBtn.disabled, false);
        """,
    )


def test_interface_radius_node_derives_saved_status_when_draft_returns_to_baseline(tmp_path):
    _run_interface_radius_node_scenario(
        tmp_path,
        """
        const h = radiusHarness;

        h.elements.resetInterfaceSettingsBtn.dispatch("click");
        assert.equal(h.state.draftPreferences.radiusPx, 4);
        assert.equal(h.state.statusText, "已保存");
        assert.equal(h.elements.interfaceSettingsStatus.textContent, "已保存");

        h.elements.interfaceRadiusSlider.value = "9";
        h.elements.interfaceRadiusSlider.dispatch("input");
        assert.equal(h.state.statusText, "正在预览，尚未保存");
        h.elements.interfaceRadiusSlider.value = "4";
        h.elements.interfaceRadiusSlider.dispatch("input");
        assert.equal(h.state.draftPreferences.radiusPx, 4);
        assert.equal(h.state.statusText, "已保存");
        assert.equal(h.elements.interfaceSettingsStatus.textContent, "已保存");
        """,
    )


def test_interface_radius_node_times_out_get_without_real_waiting(tmp_path):
    _run_interface_radius_node_scenario(
        tmp_path,
        """
        const h = radiusHarness;
        let capturedOptions = null;
        apiImpl = async (_path, options) => {
          capturedOptions = options;
          return new Promise((_resolve, reject) => {
            options.signal?.addEventListener("abort", () => {
              const error = new Error("aborted");
              error.name = "AbortError";
              reject(error);
            }, { once: true });
          });
        };

        const loading = h.load({ silent: false });
        await flushMicrotasks();
        assert.ok(capturedOptions.signal);
        assert.equal(h.timerCount(), 1);
        let settled = false;
        let result = null;
        loading.then((value) => {
          settled = true;
          result = value;
        });
        h.runAllTimers();
        await flushMicrotasks();

        assert.equal(settled, true);
        assert.equal(result, false);
        assert.equal(h.timerCount(), 0);
        assert.equal(h.state.savedPreferences.radiusPx, 4);
        assert.equal(h.state.draftPreferences.radiusPx, 4);
        assert.equal(h.cssVariables.get("--ui-radius"), "4px");
        assert.equal(h.state.statusText, "加载失败，当前使用默认 4px");
        assert.equal(h.toasts.length, 1);
        """,
    )


def test_interface_radius_node_preserves_new_draft_when_current_get_fails(tmp_path):
    _run_interface_radius_node_scenario(
        tmp_path,
        """
        const h = radiusHarness;
        const getRequest = deferred();
        apiImpl = async () => getRequest.promise;

        const loading = h.load({ silent: false });
        await flushMicrotasks();
        h.elements.interfaceRadiusSlider.value = "9";
        h.elements.interfaceRadiusSlider.dispatch("input");
        getRequest.reject(new Error("network failed"));

        assert.equal(await loading, false);
        assert.equal(h.state.loaded, false);
        assert.equal(h.state.loadFailed, true);
        assert.equal(h.state.savedPreferences.radiusPx, 4);
        assert.equal(h.state.draftPreferences.radiusPx, 9);
        assert.equal(h.cssVariables.get("--ui-radius"), "9px");
        assert.equal(h.state.statusText, "正在预览，尚未保存");
        assert.equal(h.toasts.length, 1);
        """,
    )


def test_interface_radius_node_rejects_invalid_get_payload_as_load_failure(tmp_path):
    _run_interface_radius_node_scenario(
        tmp_path,
        """
        const h = radiusHarness;
        apiImpl = async () => ({ settings: {} });

        const result = await h.load({ silent: true });
        assert.equal(result, false);
        assert.equal(h.state.loaded, false);
        assert.equal(h.state.loadFailed, true);
        assert.equal(h.state.savedPreferences.radiusPx, 4);
        assert.equal(h.state.draftPreferences.radiusPx, 4);
        assert.equal(h.cssVariables.get("--ui-radius"), "4px");
        assert.equal(h.state.statusText, "加载失败，当前使用默认 4px");
        assert.equal(h.toasts.length, 0);
        """,
    )


def test_interface_radius_node_rejects_invalid_post_without_losing_draft(tmp_path):
    _run_interface_radius_node_scenario(
        tmp_path,
        """
        const h = radiusHarness;
        apiImpl = async () => ({ settings: { radius_px: 6 } });
        assert.equal(await h.load({ silent: true }), true);
        h.elements.interfaceRadiusSlider.value = "9";
        h.elements.interfaceRadiusSlider.dispatch("input");
        apiImpl = async () => ({ settings: {} });

        const result = await h.save();
        assert.equal(result, false);
        assert.equal(h.state.savedPreferences.radiusPx, 6);
        assert.equal(h.state.draftPreferences.radiusPx, 9);
        assert.equal(h.cssVariables.get("--ui-radius"), "9px");
        assert.equal(h.state.statusText, "保存失败");
        assert.equal(h.elements.interfaceSettingsStatus.textContent, "保存失败");
        assert.equal(h.toasts.length, 1);
        """,
    )


def test_interface_radius_auth_boundary_logout_resets_immediately_and_invalidates_get(tmp_path):
    _run_interface_radius_node_scenario(
        tmp_path,
        """
        const h = radiusHarness;
        apiImpl = async () => ({ settings: { radius_px: 8 } });
        assert.equal(await h.load({ silent: true }), true);
        assert.equal(h.cssVariables.get("--ui-radius"), "8px");

        const oldGet = deferred();
        const logoutRequest = deferred();
        apiImpl = async (path) => {
          if (path === "/api/settings/interface") return oldGet.promise;
          if (path === "/api/auth/logout") return logoutRequest.promise;
          throw new Error(`unexpected API path: ${path}`);
        };
        const oldLoading = h.load({ silent: false });
        await flushMicrotasks();
        const loggingOut = h.logout();
        await flushMicrotasks();

        assert.equal(h.state.savedPreferences.radiusPx, 4);
        assert.equal(h.state.draftPreferences.radiusPx, 4);
        assert.equal(h.state.loaded, false);
        assert.equal(h.state.loadFailed, false);
        assert.equal(h.state.saving, false);
        assert.equal(h.state.statusText, "已保存");
        assert.equal(h.cssVariables.get("--ui-radius"), "4px");

        oldGet.resolve({ settings: { radius_px: 12 } });
        assert.equal(await oldLoading, false);
        assert.equal(h.state.savedPreferences.radiusPx, 4);
        assert.equal(h.state.draftPreferences.radiusPx, 4);
        assert.equal(h.cssVariables.get("--ui-radius"), "4px");

        logoutRequest.resolve({});
        await loggingOut;
        assert.equal(h.authState.csrfToken, "");
        assert.equal(window.location.href, "/login.html");
        """,
    )


def test_interface_radius_auth_boundary_resets_before_loading_new_user(tmp_path):
    _run_interface_radius_node_scenario(
        tmp_path,
        """
        const h = radiusHarness;
        h.useStrictResponses();
        apiImpl = async () => ({ settings: {
          radius_px: 8,
          line_chart_style: "smooth",
        } });
        assert.equal(await h.load({ silent: true }), true);
        assert.equal(h.cssVariables.get("--ui-radius"), "8px");
        assert.equal(h.storageWrites.length, 1);

        const oldGet = deferred();
        const newGet = deferred();
        let preferenceRequestCount = 0;
        apiImpl = async (path) => {
          assert.equal(path, "/api/settings/interface");
          preferenceRequestCount += 1;
          return preferenceRequestCount === 1 ? oldGet.promise : newGet.promise;
        };
        const oldLoading = h.load({ silent: false });
        await flushMicrotasks();
        const authenticating = h.ensureAuthenticated();
        await flushMicrotasks();

        assert.equal(h.authState.user.id, "new-user");
        assert.equal(h.state.savedPreferences.radiusPx, 4);
        assert.equal(h.state.draftPreferences.radiusPx, 4);
        assert.equal(h.state.loaded, false);
        assert.equal(h.cssVariables.get("--ui-radius"), "4px");
        assert.equal(h.revealCount(), 0);

        oldGet.resolve({ settings: {
          radius_px: 12,
          line_chart_style: "smooth",
        } });
        assert.equal(await oldLoading, false);
        assert.equal(h.state.savedPreferences.radiusPx, 4);
        assert.equal(h.cssVariables.get("--ui-radius"), "4px");
        assert.equal(h.storageWrites.length, 1);

        newGet.resolve({ settings: {
          radius_px: 6,
          line_chart_style: "straight",
        } });
        await authenticating;
        assert.deepEqual(h.state.savedPreferences, {
          radiusPx: 6,
          lineChartStyle: "straight",
        });
        assert.deepEqual(h.state.draftPreferences, h.state.savedPreferences);
        assert.equal(h.cssVariables.get("--ui-radius"), "6px");
        assert.equal(h.storageWrites.length, 2);
        assert.equal(h.revealCount(), 1);
        """,
    )


def test_interface_radius_node_keeps_dirty_draft_when_get_starts_after_edit(tmp_path):
    _run_interface_radius_node_scenario(
        tmp_path,
        """
        const h = radiusHarness;
        h.elements.interfaceRadiusSlider.value = "9";
        h.elements.interfaceRadiusSlider.dispatch("input");
        assert.equal(h.state.savedPreferences.radiusPx, 4);
        assert.equal(h.state.draftPreferences.radiusPx, 9);

        apiImpl = async () => ({ settings: { radius_px: 6 } });
        assert.equal(await h.load({ silent: false }), true);
        assert.equal(h.state.savedPreferences.radiusPx, 6);
        assert.equal(h.state.draftPreferences.radiusPx, 9);
        assert.equal(h.cssVariables.get("--ui-radius"), "9px");
        assert.equal(h.state.statusText, "正在预览，尚未保存");
        """,
    )


def test_interface_radius_node_rejects_get_started_while_post_is_pending(tmp_path):
    _run_interface_radius_node_scenario(
        tmp_path,
        """
        const h = radiusHarness;
        apiImpl = async () => ({ settings: { radius_px: 6 } });
        assert.equal(await h.load({ silent: true }), true);
        h.elements.interfaceRadiusSlider.value = "9";
        h.elements.interfaceRadiusSlider.dispatch("input");

        const postRequest = deferred();
        const getRequest = deferred();
        let getRequestCount = 0;
        apiImpl = async (_path, options) => {
          if (options.method === "POST") return postRequest.promise;
          getRequestCount += 1;
          return getRequest.promise;
        };
        const saving = h.save();
        await flushMicrotasks();
        const loading = h.load({ silent: false });
        await flushMicrotasks();
        getRequest.resolve({ settings: { radius_px: 4 } });
        assert.equal(await loading, false);
        assert.equal(getRequestCount, 0);

        postRequest.reject(new Error("save failed"));
        assert.equal(await saving, false);
        assert.equal(h.state.savedPreferences.radiusPx, 6);
        assert.equal(h.state.draftPreferences.radiusPx, 9);
        assert.equal(h.cssVariables.get("--ui-radius"), "9px");
        assert.equal(h.state.statusText, "保存失败");
        """,
    )


def test_interface_radius_auth_reset_invalidates_old_post_success_and_finally(tmp_path):
    _run_interface_radius_node_scenario(
        tmp_path,
        """
        const h = radiusHarness;
        apiImpl = async () => ({ settings: { radius_px: 8 } });
        assert.equal(await h.load({ silent: true }), true);
        h.elements.interfaceRadiusSlider.value = "9";
        h.elements.interfaceRadiusSlider.dispatch("input");

        const oldPost = deferred();
        const newGet = deferred();
        const newPost = deferred();
        let postRequestCount = 0;
        apiImpl = async (_path, options) => {
          if (options.method === "POST") {
            postRequestCount += 1;
            return postRequestCount === 1 ? oldPost.promise : newPost.promise;
          }
          return newGet.promise;
        };
        const oldSaving = h.save();
        await flushMicrotasks();
        const authenticating = h.ensureAuthenticated();
        await flushMicrotasks();
        newGet.resolve({ settings: { radius_px: 6 } });
        await authenticating;

        h.elements.interfaceRadiusSlider.value = "7";
        h.elements.interfaceRadiusSlider.dispatch("input");
        const newSaving = h.save();
        await flushMicrotasks();
        assert.equal(h.state.saving, true);

        oldPost.resolve({ settings: { radius_px: 9 } });
        assert.equal(await oldSaving, false);
        assert.equal(h.state.savedPreferences.radiusPx, 6);
        assert.equal(h.state.draftPreferences.radiusPx, 7);
        assert.equal(h.cssVariables.get("--ui-radius"), "7px");
        assert.equal(h.state.saving, true);
        assert.equal(h.toasts.length, 0);

        newPost.resolve({ settings: { radius_px: 7 } });
        assert.equal(await newSaving, true);
        assert.equal(h.state.savedPreferences.radiusPx, 7);
        assert.equal(h.state.draftPreferences.radiusPx, 7);
        assert.equal(h.state.saving, false);
        """,
    )


def test_interface_radius_auth_reset_invalidates_old_post_failure(tmp_path):
    _run_interface_radius_node_scenario(
        tmp_path,
        """
        const h = radiusHarness;
        apiImpl = async () => ({ settings: { radius_px: 8 } });
        assert.equal(await h.load({ silent: true }), true);
        h.elements.interfaceRadiusSlider.value = "9";
        h.elements.interfaceRadiusSlider.dispatch("input");

        const oldPost = deferred();
        const newGet = deferred();
        apiImpl = async (_path, options) => (
          options.method === "POST" ? oldPost.promise : newGet.promise
        );
        const oldSaving = h.save();
        await flushMicrotasks();
        const authenticating = h.ensureAuthenticated();
        await flushMicrotasks();
        newGet.resolve({ settings: { radius_px: 6 } });
        await authenticating;

        oldPost.reject(new Error("old user save failed"));
        assert.equal(await oldSaving, false);
        assert.equal(h.state.savedPreferences.radiusPx, 6);
        assert.equal(h.state.draftPreferences.radiusPx, 6);
        assert.equal(h.cssVariables.get("--ui-radius"), "6px");
        assert.equal(h.state.statusText, "已保存");
        assert.equal(h.state.saving, false);
        assert.equal(h.toasts.length, 0);
        """,
    )


def test_interface_radius_logout_failure_restores_dirty_snapshot(tmp_path):
    _run_interface_radius_node_scenario(
        tmp_path,
        """
        const h = radiusHarness;
        h.useStrictResponses();
        apiImpl = async () => ({ settings: {
          radius_px: 8,
          line_chart_style: "smooth",
        } });
        assert.equal(await h.load({ silent: true }), true);
        h.elements.interfaceRadiusSlider.value = "9";
        h.elements.interfaceRadiusSlider.dispatch("input");
        h.elements.interfaceLineChartStyleStraight.checked = true;
        h.elements.interfaceLineChartStyleStraight.dispatch("input");

        const logoutRequest = deferred();
        apiImpl = async (path) => {
          assert.equal(path, "/api/auth/logout");
          return logoutRequest.promise;
        };
        const loggingOut = h.logout();
        await flushMicrotasks();
        assert.equal(h.state.savedPreferences.radiusPx, 4);
        assert.equal(h.state.draftPreferences.radiusPx, 4);
        assert.equal(h.cssVariables.get("--ui-radius"), "4px");

        logoutRequest.reject(new Error("logout failed"));
        await loggingOut;
        assert.deepEqual(h.state.savedPreferences, {
          radiusPx: 8,
          lineChartStyle: "smooth",
        });
        assert.deepEqual(h.state.draftPreferences, {
          radiusPx: 9,
          lineChartStyle: "straight",
        });
        assert.equal(h.state.loaded, true);
        assert.equal(h.state.loadFailed, false);
        assert.equal(h.state.saving, false);
        assert.equal(h.state.statusText, "正在预览，尚未保存");
        assert.equal(h.cssVariables.get("--ui-radius"), "9px");
        assert.equal(h.toasts.length, 1);
        assert.equal(window.location.href, "/");
        """,
    )


def test_settings_dark_mode_keeps_business_codes_and_about_links_readable():
    css = _read(STYLES_CSS)

    assert '[data-theme="space-tech"][data-color-mode="dark"] #page-settings .business-field-table code' in css
    assert '[data-theme="space-tech"][data-color-mode="dark"] #page-settings .about-links a' in css
    dark_code_rule = re.search(
        r'\[data-theme="space-tech"\]\[data-color-mode="dark"\] #page-settings \.business-field-table code\s*\{(?P<body>.*?)\}',
        css,
        re.S,
    )
    dark_link_rule = re.search(
        r'\[data-theme="space-tech"\]\[data-color-mode="dark"\] #page-settings \.about-links a\s*\{(?P<body>.*?)\}',
        css,
        re.S,
    )
    assert dark_code_rule is not None
    assert dark_link_rule is not None
    assert "color: #dbeafe" in dark_code_rule.group("body")
    assert "color: #7dd3fc" in dark_link_rule.group("body")


def test_home_chart_empty_state_keeps_centered_chart_structure():
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert "function setChartEmptyState" in app_js
    for function_name in ["renderChart", "renderTrendChart"]:
        body = re.search(
            rf"async function {function_name}\([^)]*\) \{{(?P<body>.*?)\n\}}",
            app_js,
            re.S,
        )
        assert body is not None
        assert "container.innerHTML" not in body.group("body")
        assert "setChartEmptyState" in body.group("body")

    placeholder_rule = re.search(
        r"\.chart-container \.placeholder-text\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert placeholder_rule is not None
    rule_body = placeholder_rule.group("body")
    assert "position: absolute" in rule_body
    assert "inset: 0" in rule_body
    assert "align-items: center" in rule_body
    assert "justify-content: center" in rule_body


def test_home_chart_loading_state_is_animated():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert 'class="chart-loading-indicator"' in html
    assert "function setChartLoadingState" in app_js
    assert 'setChartLoadingState(container, true);' in app_js
    assert 'setChartLoadingState(container, false);' in app_js
    assert ".chart-loading-indicator" in css
    assert "@keyframes chartLoadingPulse" in css


def test_home_trend_curve_control_points_stay_inside_plot_area():
    app_js = _read(APP_JS)
    smooth_curve = re.search(
        r"function smoothCurveThrough\(ctx, pts, tension = 0\.35, bounds = null\) \{(?P<body>.*?)\n\}",
        app_js,
        re.S,
    )

    assert "function clampNumber" in app_js
    assert smooth_curve is not None
    assert "clampNumber(cp1y" in smooth_curve.group("body")
    assert "clampNumber(cp2y" in smooth_curve.group("body")
    assert "{ top: pad.top, bottom: pad.top + ph }" in app_js


def test_home_chart_line_style_normalization_and_geometry_are_shared(tmp_path):
    app_js = _read(APP_JS)
    normalize_style = re.search(
        r"function normalizeLineChartStyle\(value\) \{.*?\n\}",
        app_js,
        re.S,
    )
    smooth_curve = re.search(
        r"function smoothCurveThrough\(ctx, pts, tension = 0\.35, bounds = null\) \{.*?\n\}",
        app_js,
        re.S,
    )
    trace_line = re.search(
        r"function traceChartLine\(ctx, points, style, bounds = null\) \{.*?\n\}",
        app_js,
        re.S,
    )
    assert normalize_style is not None
    assert smooth_curve is not None
    assert trace_line is not None

    script_path = tmp_path / "chart-line-style.cjs"
    script_path.write_text(
        textwrap.dedent(
            f"""
            const assert = require("node:assert/strict");
            function clampNumber(value, min, max) {{
              return Math.min(Math.max(value, min), max);
            }}
            {normalize_style.group(0)}
            {smooth_curve.group(0)}
            {trace_line.group(0)}

            assert.equal(normalizeLineChartStyle(), "straight");
            assert.equal(normalizeLineChartStyle("straight"), "straight");
            assert.equal(normalizeLineChartStyle("smooth"), "smooth");
            assert.equal(normalizeLineChartStyle("SMOOTH"), "straight");
            assert.equal(normalizeLineChartStyle("curve"), "straight");

            const points = [{{ x: 1, y: 4 }}, {{ x: 3, y: 2 }}, {{ x: 7, y: 6 }}];
            const calls = [];
            const ctx = {{
              beginPath: () => calls.push(["beginPath"]),
              moveTo: (...args) => calls.push(["moveTo", ...args]),
              lineTo: (...args) => calls.push(["lineTo", ...args]),
              bezierCurveTo: (...args) => calls.push(["bezierCurveTo", ...args]),
            }};
            traceChartLine(ctx, points, "straight");
            assert.deepEqual(calls.map((call) => call[0]), [
              "beginPath", "moveTo", "lineTo", "lineTo",
            ]);

            calls.length = 0;
            traceChartLine(ctx, points, "smooth", {{ top: 0, bottom: 10 }});
            assert.equal(calls[0][0], "beginPath");
            assert.equal(calls[1][0], "moveTo");
            assert.equal(calls.filter((call) => call[0] === "bezierCurveTo").length, 2);
            """
        ),
        encoding="utf-8",
    )
    subprocess.run(["node", str(script_path)], check=True, cwd=ROOT)


def test_home_chart_renderers_use_logo_gradients_and_style_points():
    app_js = _read(APP_JS)
    single_chart = app_js[
        app_js.index("function drawGlassChart(") :
        app_js.index("function drawGlassMultiMetricChart(")
    ]
    multi_chart = app_js[
        app_js.index("function drawGlassMultiMetricChart(") :
        app_js.index("function cssRootValue(")
    ]

    assert "function canvasThemePalette()" in app_js
    assert 'cssRootValue("--theme-accent"' in app_js
    assert 'cssRootValue("--theme-accent-readable"' in app_js
    for body in (single_chart, multi_chart):
        assert "const palette = canvasThemePalette();" in body
        assert "const lineStyle = currentLineChartStyle();" in body
        assert "traceChartLine(ctx, drawPts, lineStyle, curveBounds);" in body
        assert 'if (lineStyle === "smooth"' in body
        assert "createLinearGradient" in body

    assert "ctx.fillStyle = palette.areaFill;" in single_chart
    assert "lineGradient.addColorStop(0, palette.primary);" in single_chart
    assert "lineGradient.addColorStop(1, palette.gradientEnd);" in single_chart
    assert "if (metric.gradientEnd)" in multi_chart
    assert "metricPaint.addColorStop(0, metric.strokeColor);" in multi_chart
    assert "metricPaint.addColorStop(1, metric.gradientEnd);" in multi_chart
    assert 'gradientEnd: "#FFBD38"' in app_js


def test_home_chart_theme_and_line_style_redraw_cached_data_without_refetch():
    app_js = _read(APP_JS)
    redraw = re.search(
        r"function redrawHomeChartsFromCache\(\) \{(?P<body>.*?)\n\}",
        app_js,
        re.S,
    )
    refresh_theme = re.search(
        r"function refreshHomeChartsForTheme\(\) \{(?P<body>.*?)\n\}",
        app_js,
        re.S,
    )
    apply_preferences = re.search(
        r"function applyInterfacePreferences\(preferences\) \{(?P<body>.*?)\n\}",
        app_js,
        re.S,
    )
    load_history = re.search(
        r"async function loadHomeChartHistory\(\{ refreshData = true \} = \{\}\) \{(?P<body>.*?)\n\}",
        app_js,
        re.S,
    )

    assert redraw is not None
    assert "renderChart({ refreshData: false })" in redraw.group("body")
    assert "renderTrendChart({ refreshData: false })" in redraw.group("body")
    assert refresh_theme is not None
    assert "redrawHomeChartsFromCache();" in refresh_theme.group("body")
    assert apply_preferences is not None
    assert "normalizeLineChartStyle(preferences.lineChartStyle)" in apply_preferences.group("body")
    assert "refreshHomeChartsForTheme();" in apply_preferences.group("body")
    assert load_history is not None
    assert "homeChartHistoryCache" in load_history.group("body")
    assert 'api("/api/history")' in load_history.group("body")

    schedule_resize = re.search(
        r"function scheduleHomeChartsResize\(\) \{(?P<body>.*?)\n\}",
        app_js,
        re.S,
    )
    assert schedule_resize is not None
    assert "redrawHomeChartsFromCache();" in schedule_resize.group("body")


def test_space_tech_theme_has_structural_top_navigation_and_switching():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert 'class="top-nav"' in html
    for page in ["home", "auto-check", "history", "settings"]:
        assert re.search(rf'class="[^"]*\btop-nav-item\b[^"]*" data-page="{page}"', html)
    assert 'data-theme-toggle-logo' not in html
    assert 'aria-label="切换主题"' not in html
    assert 'name="theme"' not in html
    assert 'document.documentElement.setAttribute("data-theme", "space-tech")' in app_js
    assert 'document.documentElement.setAttribute("data-color-mode", "light")' in app_js
    assert "function syncNavState" in app_js
    assert "topNavItems" in app_js
    assert '[data-theme="space-tech"] .top-nav' in css
    assert '[data-theme="space-tech"] .main-content' in css

    top_nav_brand = re.search(r'\[data-theme="space-tech"\] \.top-nav-brand\s*\{(?P<body>.*?)\}', css, re.S)
    assert top_nav_brand is not None
    assert "width: 100%" in top_nav_brand.group("body")
    assert "justify-self: start" in top_nav_brand.group("body")

    top_nav_tabs = re.search(r'\[data-theme="space-tech"\] \.top-nav-tabs\s*\{(?P<body>.*?)\}', css, re.S)
    assert top_nav_tabs is not None
    assert "justify-self: center" in top_nav_tabs.group("body")

    top_status = re.search(r'\[data-theme="space-tech"\] \.top-nav-status\s*\{(?P<body>.*?)\}', css, re.S)
    assert top_status is not None
    assert "flex: 0 0 auto" in top_status.group("body")
    assert "max-width: 200px" in top_status.group("body")

    top_notice_status = re.search(
        r'\[data-theme="space-tech"\] \.top-nav-status\.top-nav-status--notice\s*\{(?P<body>.*?)\}',
        css,
        re.S,
    )
    assert top_notice_status is not None
    assert "flex: 0 1 auto" in top_notice_status.group("body")
    assert "max-width: min(680px, 42vw)" in top_notice_status.group("body")
    assert 'topNavStatus.classList.toggle("top-nav-status--notice", nextText !== DEFAULT_VERSION);' in app_js
    assert 'topNavStatus.classList.remove("top-nav-status--notice");' in app_js


def test_space_tech_top_navigation_centers_pages_and_keeps_actions_right():
    html = _read(INDEX_HTML)
    css = _read(STYLES_CSS)
    app_js = _read(APP_JS)
    readme = _read(README_MD)

    top_nav = re.search(r'<header class="top-nav">(?P<body>.*?)</header>', html, re.S)
    assert top_nav is not None
    top_nav_body = top_nav.group("body")

    tabs = re.search(r'<nav class="top-nav-tabs">(?P<body>.*?)</nav>', top_nav_body, re.S)
    actions = re.search(r'<div class="top-nav-actions">(?P<body>.*?)</div>\s*</header>', html, re.S)
    assert tabs is not None
    assert actions is not None
    assert 'data-page="report-navigation"' in tabs.group("body")
    assert 'id="topDarkModeToggle"' not in tabs.group("body")
    assert 'id="topUserMenu"' not in tabs.group("body")
    assert 'id="topDarkModeToggle"' not in actions.group("body")
    assert 'id="topUserMenu"' in actions.group("body")

    top_nav_rule = re.search(r'\[data-theme="space-tech"\] \.top-nav\s*\{(?P<body>.*?)\}', css, re.S)
    tabs_rule = re.search(r'\[data-theme="space-tech"\] \.top-nav-tabs\s*\{(?P<body>.*?)\}', css, re.S)
    actions_rule = re.search(r'\.top-nav-actions\s*\{(?P<body>.*?)\}', css, re.S)
    assert top_nav_rule is not None
    assert tabs_rule is not None
    assert actions_rule is not None
    assert "display: grid" in top_nav_rule.group("body")
    assert "grid-template-columns: minmax(0, 1fr) auto minmax(0, 1fr)" in top_nav_rule.group("body")
    assert "justify-self: center" in tabs_rule.group("body")
    assert "justify-self: end" in actions_rule.group("body")
    assert "当前唯一启用的亮色活力主题" in readme
    assert "系统优化及BUG修复。" in app_js


def test_space_tech_theme_uses_reference_light_palette():
    css = _read(STYLES_CSS)
    theme = re.search(r"\[data-theme=\"space-tech\"\]\s*\{(?P<body>.*?)\}", css, re.S)

    assert theme is not None
    body = theme.group("body")
    for text in [
        "color-scheme: light",
        "--surface: #f8fafc",
        "--surface-container-lowest: #ffffff",
        "--on-surface: #0f172a",
        "--secondary: #3b82f6",
        "background: var(--theme-page-background)",
    ]:
        assert text in body
    assert "--space-gradient-primary" not in body
    assert '[data-theme="space-tech"] body' in css
    assert "rgba(255, 255, 255, 0.72)" in css


def test_space_tech_theme_hides_in_page_heading_text_and_tightens_gap():
    css = _read(STYLES_CSS)

    assert '[data-theme="space-tech"] .page-header h2' in css
    assert '[data-theme="space-tech"] .page-header' in css
    # Edge-stuck shell: content hugs the nav without floating offsets.
    assert "padding: 12px 32px 32px" in css
    assert "margin: 12px 32px 0" not in css
    assert "padding: 64px 0 32px" not in css
    assert "margin: 8px 14px 0" not in css
    assert "padding: 78px 0 18px" not in css


def test_theme_is_saved_per_user_without_updating_global_defaults():
    app_js = _read(APP_JS)

    for text in [
        'theme: "space-tech"',
        'darkMode: "false"',
        "function normalizeDarkMode",
        "darkMode === true",
        "const THEME_ACTIVE_USER_KEY",
        "function themeStorageKey(baseKey)",
        "function saveUserThemePreference(keyBase, value)",
        "function withSavedUserTheme(settings = {})",
        "let serverDefaultSettings",
        "activateThemeUserStorage();",
        "applySavedUserTheme();",
        "theme: settings.theme",
        "darkMode: normalizeDarkMode(settings.darkMode)",
        "theme: normalized.theme",
        "dark_mode: normalized.darkMode",
        "async function saveAndApplyTheme",
        "async function saveAndApplyDarkMode",
        "defaultSettings.theme = theme",
        "defaultSettings.darkMode = darkMode",
        "theme: serverDefaultSettings.theme",
        "darkMode: serverDefaultSettings.darkMode",
    ]:
        assert text in app_js

    save_theme = re.search(r"async function saveAndApplyTheme\(theme, options = \{\}\) \{(?P<body>.*?)\n\}", app_js, re.S)
    save_dark = re.search(r"async function saveAndApplyDarkMode\(darkMode\) \{(?P<body>.*?)\n\}", app_js, re.S)
    assert save_theme is not None
    assert save_dark is not None
    assert 'api("/api/settings/defaults"' not in save_theme.group("body")
    assert 'api("/api/settings/defaults"' not in save_dark.group("body")


def test_fixed_theme_is_applied_before_stylesheet_and_radius_cache_stays_out_of_boot_script():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)

    assert 'id="initialThemeScript"' in html
    stylesheet_ref = re.search(r'href="/styles\.css[^"]*"', html)
    assert stylesheet_ref is not None
    assert html.index('id="initialThemeScript"') < html.index(stylesheet_ref.group(0))
    for text in ["data-theme", "data-color-mode"]:
        assert text in html
    initial_script = html[html.index('id="initialThemeScript"'):html.index('</script>', html.index('id="initialThemeScript"'))]
    assert 'setAttribute("data-theme", "space-tech")' in initial_script
    assert 'setAttribute("data-color-mode", "light")' in initial_script
    assert "autoCheckTheme" not in initial_script
    assert "autoCheckDarkMode" not in initial_script
    assert "autoCheckLastInterfaceRadius" not in initial_script
    assert 'const LAST_INTERFACE_RADIUS_CACHE_KEY = "autoCheckLastInterfaceRadius";' in app_js


def test_latest_result_detail_list_is_restored_from_local_snapshot_before_history_fetch():
    app_js = _read(APP_JS)

    for text in [
        "const LATEST_RESULTS_SNAPSHOT_KEY",
        "function saveLatestResultsSnapshot",
        "function restoreLatestResultsSnapshot",
        "function clearLatestResultsSnapshot",
        "autoCheckLatestResults",
        "saveLatestResultsSnapshot(",
    ]:
        assert text in app_js

    initial_load = re.search(r"// Initial load(?P<body>.*?loadSystemInfo\(\);)", app_js, re.S)
    assert initial_load is not None
    body = initial_load.group("body")
    assert body.index("restoreLatestResultsSnapshot()") < body.index("await loadLatestHistoryResults()")


def test_tool_and_settings_page_loaders_are_isolated():
    app_js = _read(APP_JS)

    assert "async function loadPageSection" in app_js
    assert "async function loadToolsPageData" in app_js
    assert "async function loadSettingsPageData" in app_js

    tools_loader = re.search(r"async function loadToolsPageData\(\) \{(?P<body>.*?)\n\}", app_js, re.S)
    assert tools_loader is not None
    tools_body = tools_loader.group("body")
    assert "Promise.all" in tools_body
    assert 'loadPageSection("PBC导入配置", loadPbcImportSettings)' in tools_body
    assert 'loadPageSection("逐笔校验配置", loadDbValidationSettings)' in tools_body
    assert 'loadPageSection("流程执行配置", loadFlowSettings)' in tools_body

    settings_loader = re.search(r"async function loadSettingsPageData\(\) \{(?P<body>.*?)\n\}", app_js, re.S)
    assert settings_loader is not None
    settings_body = settings_loader.group("body")
    assert "Promise.all" in settings_body
    assert 'loadPageSection("系统信息", loadSystemInfo)' in settings_body
    assert 'loadPageSection("数据源配置", loadConfigList)' in settings_body
    assert 'loadPageSection("逐笔校验配置", loadDbValidationSettings)' in settings_body
    assert 'loadPageSection("流程执行配置", loadFlowSettings)' in settings_body
    assert 'loadPageSection("业务字段配置", loadReconcileSchemaSettings)' in settings_body
    assert "applySettingsRoleAccess();" in settings_body

    switch_page = re.search(r"async function switchPage\(name, options = \{\}\) \{(?P<body>.*?)\n\}", app_js, re.S)
    assert switch_page is not None
    switch_body = switch_page.group("body")
    assert 'if (name === "tools") loadToolsPageData();' in switch_body
    assert "void loadSettingsPageData();" in switch_body
    assert 'if (name === "settings") await loadSettingsPageData();' not in switch_body
    assert 'if (name === "tools") await loadToolsPageData();' not in switch_body
    assert 'await loadPbcImportSettings(); await loadDbValidationSettings(); await loadFlowSettings();' not in switch_body
    assert "function clearTopNavGroupFocus(group)" in app_js
    assert 'group.addEventListener("pointerleave"' in app_js
    assert "submenu-dismissed" not in app_js
    assert "submenu-dismissed" not in _read(STYLES_CSS)


def test_system_info_uses_lightweight_summary_api():
    app_js = _read(APP_JS)

    load_system_info = re.search(r"async function loadSystemInfo\(\) \{(?P<body>.*?)\n\}", app_js, re.S)
    assert load_system_info is not None
    body = load_system_info.group("body")
    assert 'api("/api/system-info")' in body
    assert 'api("/api/history")' not in body


def test_tool_settings_load_failures_render_visible_placeholders():
    app_js = _read(APP_JS)

    assert "function renderDbValidationSettingsLoading" in app_js
    assert "function renderDbValidationSettingsError" in app_js
    assert "正在加载逐笔校验配置" in app_js
    assert "逐笔校验配置加载失败" in app_js

    assert "function renderFlowSettingsLoadError" in app_js
    assert "流程链配置加载失败" in app_js
    assert "flowStartBtn.disabled = true" in app_js


def test_db_validation_field_mapping_warns_when_metadata_is_partial():
    app_js = _read(APP_JS)

    assert "少于系统内置表单" in app_js
    assert "请检查字段映射数据源、baseinfo/field_info 或筛选条件" in app_js


def test_pbc_import_tool_is_exposed_in_tools_page():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)

    assert 'data-page="tools"' in html
    assert 'id="page-tools"' in html
    assert 'id="toolCardPbc"' in html
    assert 'id="pbcModalOverlay"' in html
    for element_id in [
        "pbcZipFile",
        "pbcUploadArea",
        "pbcFileList",
        "pbcDataSource",
        "pbcTargetTable",
        "pbcMappingList",
        "pbcImportLog",
        "pbcProgressFill",
        "pbcProgressPercent",
        "pbcNextBtn",
        "pbcFinishBtn",
    ]:
        assert f'id="{element_id}"' in html

    for endpoint in [
        "/api/tools/pbc-import/settings",
        "/api/tools/pbc-import/upload",
        "/api/tools/pbc-import/columns",
        "/api/tools/pbc-import/start",
        "/api/tools/pbc-import/status/",
    ]:
        assert endpoint in app_js


def test_pbc_import_tool_is_generic_one_click_import_with_upload_progress():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert "一键导入工具" in html
    assert "支持任何数据批量导入，自动校验数据完整性与格式合规性，快速完成导入流程。" in html
    assert "支持人行产品数据批量导入" not in html
    assert "人行全量产品一键导入" not in html
    assert 'id="pbcUploadProgress"' in html
    assert 'id="pbcUploadProgressFill"' in html
    assert "function setPbcUploadState" in app_js
    assert "function uploadPbcFileWithProgress" in app_js
    assert "xhr.upload.onprogress" in app_js
    assert "setPbcUploadState(true" in app_js
    assert ".pbc-upload-area.uploading" in css
    assert ".pbc-upload-progress" in css


def test_pbc_target_table_defaults_to_recent():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)

    assert 'placeholder=""' in html
    assert "function getPbcTargetTable" in app_js
    assert "pbcTargetTable.placeholder" not in app_js
    assert "target_table: getPbcTargetTable()" in app_js
    mapping_handler = re.search(
        r'pbcLoadMappingsBtn\?\.addEventListener\("click", async \(\) => \{(?P<body>.*?)\n\}\);',
        app_js,
        re.S,
    )
    assert mapping_handler is not None
    assert "请填写目标表名" not in mapping_handler.group("body")


def test_pbc_import_modal_flow_and_defaults_do_not_skip_mapping_step():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert '<option value="replace" selected>清空后导入</option>' in html
    assert 'pbcModalOverlay?.addEventListener("click"' not in app_js
    assert "pbcNextBtn.onclick" not in app_js
    assert 'if (pbcCurrentStep === 1) goToStep(2);' in app_js
    assert "function hasPbcActiveMappings()" in app_js
    assert "(pbcCurrentStep === 2 && !hasPbcActiveMappings())" in app_js
    assert "pbcNextBtn?.addEventListener(\"click\", async () => {" in app_js
    assert 'const confirmed = await showConfirm("确认导入", "即将开始数据导入，是否确认？", { tone: "primary" });' in app_js
    assert "if (!confirmed) return;" in app_js
    assert re.search(r"else if \(pbcCurrentStep === 2\) \{(?P<body>.*?)goToStep\(3\);", app_js, re.S)
    assert "updatePbcStepUI();" in re.search(r"async function handlePbcFileUpload\(file\) \{(?P<body>.*?)function renderPbcFileList", app_js, re.S).group("body")
    assert "updatePbcStepUI();" in re.search(r"pbcFileListBody\?\.addEventListener\(\"click\", \(e\) => \{(?P<body>.*?)\}\);", app_js, re.S).group("body")
    assert "updatePbcStepUI();" in re.search(r"function renderPbcMappings\(\) \{(?P<body>.*?)\n\}", app_js, re.S).group("body")

    assert "#confirmModal" in css
    assert "z-index: 3000" in css
    assert 'id="confirmOk" class="btn-confirm-primary"' in html
    assert ".modal-confirm .modal-footer" in css
    assert ".btn-confirm-primary" in css
    assert '[data-color-mode="dark"] .btn-confirm-primary' in css


def test_pbc_import_job_completion_routes_success_and_failure_explicitly():
    app_js = _read(APP_JS)

    assert "function finishPbcImportSuccess(job, targetTable)" in app_js
    assert "function finishPbcImportFailure(job)" in app_js

    success_body = re.search(
        r"function finishPbcImportSuccess\(job, targetTable\) \{(?P<body>.*?)\n\}",
        app_js,
        re.S,
    )
    assert success_body is not None
    assert "pbcCurrentStep = 4;" in success_body.group("body")
    assert "updatePbcStepUI();" in success_body.group("body")
    assert "loadPbcImportSettings();" in success_body.group("body")

    failure_body = re.search(r"function finishPbcImportFailure\(job\) \{(?P<body>.*?)\n\}", app_js, re.S)
    assert failure_body is not None
    assert 'appendPbcLog(`导入失败: ${message}`, "error");' in failure_body.group("body")
    assert 'showToast(`导入失败: ${message}`, "error");' in failure_body.group("body")
    assert "pbcImportFailed = true;" in failure_body.group("body")
    assert 'setPbcImportProgressState("导入失败"' in failure_body.group("body")
    assert "pbcCurrentStep = 3;" in failure_body.group("body")
    assert "updatePbcStepUI();" in failure_body.group("body")

    poll_body = re.search(r"async function pollPbcImportJob\(jobId, targetTable\) \{(?P<body>.*?)\n\}", app_js, re.S)
    assert poll_body is not None
    assert "finishPbcImportFailure(job);" in poll_body.group("body")
    assert "finishPbcImportSuccess(job, targetTable);" in poll_body.group("body")


def test_pbc_file_list_counts_and_actions_are_centered():
    css = _read(STYLES_CSS)

    file_list = re.search(r"(?m)^\.pbc-file-list\s*\{(?P<body>.*?)\}", css, re.S)
    assert file_list is not None
    assert "--pbc-file-cols-width" in file_list.group("body")
    assert "--pbc-file-action-width" in file_list.group("body")
    assert "--pbc-file-scrollbar-width" in file_list.group("body")

    header_layout = re.search(r"(?m)^\.pbc-file-list-header\s*\{(?P<body>.*?)\}", css, re.S)
    assert header_layout is not None
    assert "minmax(0, 1fr)" in header_layout.group("body")
    assert "var(--pbc-file-cols-width)" in header_layout.group("body")
    assert "var(--pbc-file-action-width)" in header_layout.group("body")
    assert "var(--pbc-file-scrollbar-width)" in header_layout.group("body")

    body_layout = re.search(r"(?m)^#pbcFileListBody\s*\{(?P<body>.*?)\}", css, re.S)
    assert body_layout is not None
    assert "scrollbar-gutter: stable" in body_layout.group("body")

    row_layout = re.search(r"(?m)^\.pbc-file-list-row\s*\{(?P<body>.*?)\}", css, re.S)
    assert row_layout is not None
    assert "minmax(0, 1fr)" in row_layout.group("body")
    assert "var(--pbc-file-cols-width)" in row_layout.group("body")
    assert "var(--pbc-file-action-width)" in row_layout.group("body")

    header_centering = re.search(r"(?m)^\.pbc-file-list-header span:nth-child\(2\),\s*\n\.pbc-file-list-header span:nth-child\(3\)\s*\{(?P<body>.*?)\}", css, re.S)
    assert header_centering is not None
    assert "text-align: center" in header_centering.group("body")

    row_centering = re.search(r"(?m)^\.pbc-file-list-row > span:nth-child\(2\),\s*\n\.pbc-file-list-row > span:nth-child\(3\)\s*\{(?P<body>.*?)\}", css, re.S)
    assert row_centering is not None
    row_body = row_centering.group("body")
    assert "display: flex" in row_body
    assert "justify-content: center" in row_body
    assert "align-items: center" in row_body


def test_toast_deduplicates_same_message_and_type():
    app_js = _read(APP_JS)

    start = app_js.index('function showToast(message, type = "info")')
    end = app_js.index("// Theme Settings", start)
    body = app_js[start:end]
    assert "toast.dataset.message = message;" in app_js
    assert "toast.dataset.type = type;" in app_js
    assert 'toastContainer.querySelector(`[data-message="${cssEscape(message)}"][data-type="${cssEscape(type)}"]`)' in body
    assert "${message}</span>" not in body
    assert "messageEl.textContent = message;" in body


def test_login_page_uses_fixed_light_centered_layout_with_login_only_background():
    login_html = _read(ROOT / "src" / "auto_check" / "web" / "login.html")

    assert "<title>监管智核</title>" in login_html
    assert 'class="login-container"' in login_html
    assert 'class="left-panel"' not in login_html
    assert 'class="right-panel"' in login_html
    assert 'class="light-brand"' in login_html
    assert '<img class="login-brand-logo" src="/assets/logo-login.svg?v=1.2.12-regulatory-intelligence-core-horizontal" alt="监管智核" />' in login_html
    assert "开启您的智能工作台" not in login_html
    assert '<h1>欢迎登录</h1>' not in login_html
    assert "<p>开启您的智能工作台</p>" not in login_html
    assert '<h2 class="welcome-title" id="loginTitle">欢迎登录</h2>' in login_html
    assert '<p class="welcome-subtitle" id="loginSubtitle">请输入管理员密码继续访问系统。</p>' in login_html
    assert 'document.querySelector(".light-brand h1")' not in login_html
    assert 'document.querySelector(".light-brand p")' not in login_html
    assert 'document.getElementById("loginTitle").textContent = titleText;' in login_html
    assert 'document.getElementById("loginSubtitle").textContent = subtitleText;' in login_html
    assert 'class="deco deco-1"' not in login_html
    assert "--theme-page-background:" in login_html
    assert "radial-gradient" in login_html
    assert "linear-gradient(135deg, #EEF4FF 0%, #F8FAFC 58%, #FFFAF4 100%)" in login_html
    assert "background: var(--theme-page-background);" in login_html
    assert "max-width: 440px;" in login_html
    assert "padding: 34px 32px 34px;" in login_html
    assert "text-align: left;" in login_html
    assert "margin-bottom: 14px;" in login_html
    assert "width: min(360px, 100%);" in login_html
    assert "margin: 0;" in login_html
    assert "还没有账户？" not in login_html
    assert "去联系管理员" not in login_html
    assert 'class="forgot-password"' in login_html
    assert 'id="loginThemeToggle"' not in login_html
    assert "max-width: 860px;" not in login_html
    assert "min-height: 500px;" not in login_html
    assert "grid-template-columns" not in login_html
    assert 'class="title brand-wordmark' not in login_html
    assert '<html lang="zh-CN" data-login-theme="light">' in login_html
    assert ':root[data-login-theme="dark"] .form-input:-webkit-autofill' in login_html
    assert "-webkit-text-fill-color: var(--text-primary)" in login_html
    assert "var(--field-surface) inset" in login_html
    assert '"/api/auth/login"' in login_html
    assert '"/api/auth/setup"' in login_html
    assert "暂不支持" in login_html
    assert 'const normalized = "light";' in login_html
    assert 'root.setAttribute("data-login-theme", normalized);' in login_html
    assert 'localStorage.removeItem("autoCheckLoginTheme");' in login_html


def test_login_page_brand_restores_animated_blue_and_orange_bubbles():
    login_html = _read(ROOT / "src" / "auto_check" / "web" / "login.html")
    light_brand = re.search(r"\.light-brand\s*\{(?P<body>.*?)\n      \}", login_html, re.S)
    light_brand_circle = re.search(r"\.light-brand::before\s*\{(?P<body>.*?)\n      \}", login_html, re.S)
    light_brand_warm_circle = re.search(
        r"\.light-brand::after\s*\{(?P<body>\s*background:.*?)\n      \}",
        login_html,
        re.S,
    )
    login_logo = re.search(r"\.login-brand-logo\s*\{(?P<body>.*?)\n      \}", login_html, re.S)

    assert light_brand is not None
    light_brand_body = light_brand.group("body")
    assert "min-height: 128px;" in light_brand_body
    assert "display: flex;" in light_brand_body
    assert "align-items: center;" in light_brand_body
    assert light_brand_circle is not None
    circle_body = light_brand_circle.group("body")
    assert "linear-gradient" in circle_body
    assert "animation: lightBrandBubbleFloat" in login_html
    assert light_brand_warm_circle is not None
    assert "linear-gradient" in light_brand_warm_circle.group("body")
    assert "lightBrandBubbleWarmth" in light_brand_warm_circle.group("body")
    assert login_logo is not None
    assert "z-index: 1;" in login_logo.group("body")
    assert 'loginBrandLogo.src = "/assets/logo-login.svg?v=1.2.12-regulatory-intelligence-core-horizontal";' in login_html


def test_login_page_light_default_password_copy_and_eye_toggle_are_stable():
    login_html = _read(ROOT / "src" / "auto_check" / "web" / "login.html")

    assert '<html lang="zh-CN" data-login-theme="light">' in login_html
    assert 'applyTheme(localStorage.getItem("autoCheckLoginTheme") || "light")' in login_html
    assert 'placeholder="请输入密码"' in login_html
    assert 'id="loginPasswordToggle"' in login_html
    assert 'class="password-toggle"' in login_html
    assert "function togglePasswordVisibility" in login_html
    assert 'passwordInput.type = passwordInput.type === "password" ? "text" : "password";' in login_html
    assert "::-ms-reveal" in login_html


def test_auth_password_rule_copy_requires_six_chars_and_letter():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    login_html = _read(ROOT / "src" / "auto_check" / "web" / "login.html")
    readme = _read(README_MD)

    for text in [html, app_js, login_html, readme]:
        assert "至少 6 位且包含字母" in text
        assert "至少 8 位" not in text
        assert "至少 8 位密码" not in text

    assert "password must be at least 6 characters and include a letter" in app_js
    assert "密码长度至少 6 位，且需包含至少 1 个字母。" in app_js
    assert "password must be at least 6 characters and include a letter" in login_html
    assert "密码长度至少 6 位，且需包含至少 1 个字母。" in login_html


def test_login_page_uses_same_favicon_as_main_app():
    html = _read(INDEX_HTML)
    login_html = _read(ROOT / "src" / "auto_check" / "web" / "login.html")
    logo = ROOT / "src" / "auto_check" / "web" / "assets" / "logo-full.svg"
    favicon_asset = ROOT / "src" / "auto_check" / "web" / "assets" / "favicon-64x64.svg"

    favicon = re.search(r'<link rel="icon" href="(?P<href>[^"]+)" />', html)
    login_favicon = re.search(r'<link rel="icon" href="(?P<href>[^"]+)" />', login_html)
    assert favicon is not None
    assert login_favicon is not None
    assert login_favicon.group("href") == favicon.group("href")
    assert favicon.group("href") == "/assets/favicon-64x64.svg?v=1.2.12-regulatory-intelligence-core"
    assert logo.exists()
    assert favicon_asset.exists()


def test_login_remember_me_stores_username_without_defaulting_to_admin():
    login_html = _read(ROOT / "src" / "auto_check" / "web" / "login.html")

    assert 'const rememberLogin = document.getElementById("rememberLogin");' in login_html
    assert 'const REMEMBERED_USERNAME_KEY = "autoCheckRememberedUsername";' in login_html
    assert "function loadRememberedUsername()" in login_html
    assert 'localStorage.getItem(REMEMBERED_USERNAME_KEY)' in login_html
    assert 'localStorage.setItem(REMEMBERED_USERNAME_KEY, username);' in login_html
    assert 'localStorage.removeItem(REMEMBERED_USERNAME_KEY);' in login_html
    assert 'usernameInput.value = setupRequired ? "admin" : (usernameInput.value || loadRememberedUsername());' in login_html
    assert 'const username = setupRequired ? "admin" : usernameInput.value.trim();' in login_html
    assert 'usernameInput.value || "admin"' not in login_html
    assert 'usernameInput.value.trim() || "admin"' not in login_html


def test_login_submit_button_is_guarded_while_request_is_pending():
    login_html = _read(ROOT / "src" / "auto_check" / "web" / "login.html")

    assert "let loginSubmitting = false;" in login_html
    assert 'const MAIN_ENTRY_ANIMATION_KEY = "autoCheckMainEntryAnimation";' in login_html
    submit_handler = re.search(
        r'form\.addEventListener\("submit", async \(event\) => \{(?P<body>.*?)\n      \}\);',
        login_html,
        re.S,
    )
    assert submit_handler is not None
    body = submit_handler.group("body")
    assert "if (loginSubmitting) return;" in body
    assert "loginSubmitting = true;" in body
    assert "submitBtn.disabled = true;" in body
    assert body.index("loginSubmitting = true;") < body.index("submitBtn.disabled = true;")
    assert "let loginSucceeded = false;" in body
    assert "loginSucceeded = true;" in body
    assert 'sessionStorage.setItem(MAIN_ENTRY_ANIMATION_KEY, "login");' in body
    assert body.index("loginSucceeded = true;") < body.index('sessionStorage.setItem(MAIN_ENTRY_ANIMATION_KEY, "login");')
    assert "if (!loginSucceeded) {" in body
    assert "loginSubmitting = false;" in body
    assert "submitBtn.disabled = false;" in body


def test_index_hides_home_until_auth_check_finishes():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert '<html lang="zh-CN" class="auth-pending">' in html
    assert ".auth-pending body" in css
    assert 'const MAIN_ENTRY_ANIMATION_KEY = "autoCheckMainEntryAnimation";' in app_js
    assert "function consumeMainEntryAnimationFlag()" in app_js
    assert "function revealAuthenticatedApp()" in app_js
    assert 'sessionStorage.removeItem(MAIN_ENTRY_ANIMATION_KEY);' in app_js
    assert 'document.documentElement.classList.add("main-entry-animate");' in app_js
    assert 'document.documentElement.classList.remove("auth-pending");' in app_js
    assert 'document.documentElement.classList.remove("main-entry-animate");' in app_js
    assert "refreshReportNavigationScheduleLayout();" in app_js
    assert ".main-entry-animate .sidebar" in css
    assert ".main-entry-animate .top-nav" in css
    assert ".main-entry-animate .main-content" in css
    assert "@keyframes mainEntryContent" in css


def test_api_helper_sends_csrf_token_for_mutating_requests():
    app_js = _read(APP_JS)

    start = app_js.index("async function api(path, options = {})")
    end = app_js.index("function setStatus", start)
    body = app_js[start:end]
    assert '"X-CSRF-Token"' in body
    assert "authState.csrfToken" in body
    assert 'window.location.href = "/login.html";' in body


def test_logout_controls_exist_for_space_and_light_themes():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert 'class="sidebar-footer-main"' in html
    assert 'id="sidebarUserMenu"' in html
    assert 'id="topUserMenu"' in html
    assert 'class="user-menu-trigger"' in html
    assert 'class="user-menu-panel"' in html
    assert 'data-current-username' in html
    assert 'data-logout-btn' in html
    assert "退出登录" in html
    assert "/api/auth/logout" in app_js
    assert "async function logout()" in app_js
    assert 'await showConfirm(' in app_js
    logout_body = app_js[app_js.index("async function logout()"):app_js.index("function userDisplayRole")]
    assert "window.confirm" not in logout_body
    assert "resetInterfaceRadiusForAuthChange();" in logout_body
    assert logout_body.index("resetInterfaceRadiusForAuthChange();") < logout_body.index('api("/api/auth/logout"')
    assert 'window.location.href = "/login.html";' in app_js
    assert 'document.querySelectorAll("[data-logout-btn]")' in app_js
    assert ".sidebar-footer-main" in css
    assert ".user-menu-trigger" in css
    assert ".user-menu:hover .user-menu-panel" in css
    assert ".user-menu-panel" in css


def test_browser_native_dialogs_are_replaced_by_app_modals():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert 'id="confirmModal"' in html
    assert 'id="promptModal"' in html
    assert 'id="promptInput"' in html
    assert "function showConfirm(title, message, options = {})" in app_js
    assert "function showPrompt(title, message, options = {})" in app_js
    assert 'await showPrompt("重置密码"' in app_js
    assert 'await showConfirm("删除历史记录"' in app_js
    assert 'await showConfirm("删除数据源"' in app_js
    assert "#promptModal" in css
    assert ".prompt-input" in css
    assert not re.search(r"\b(?:alert|confirm|prompt)\s*\(", app_js)
    assert "window.alert" not in app_js
    assert "window.confirm" not in app_js
    assert "window.prompt" not in app_js


def test_system_prompt_supports_the_custom_date_component():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert 'id="promptDateControl" hidden' in html
    assert 'id="promptDateInput" class="prompt-input" type="date"' in html
    prompt_renderer = re.search(
        r"function showPrompt\(title, message, options = \{\}\) \{(?P<body>.*?)\n\}",
        app_js,
        re.S,
    )
    assert prompt_renderer is not None
    prompt_body = prompt_renderer.group("body")
    assert 'const dateInputEl = document.getElementById("promptDateInput");' in prompt_body
    assert 'const dateControlEl = document.getElementById("promptDateControl");' in prompt_body
    assert 'const isDate = options.type === "date";' in prompt_body
    assert "dateControlEl.hidden = !isDate;" in prompt_body
    assert "const activeInputEl = isDate ? dateInputEl : inputEl;" in prompt_body
    assert 'const dialogEl = modal.querySelector(".modal-prompt");' in prompt_body
    assert "const focusTargetEl = isDate ? dialogEl : activeInputEl;" in prompt_body
    assert "setTimeout(() => focusTargetEl?.focus(), 0);" in prompt_body
    assert "options.maxlength" in prompt_body
    assert "inputEl.maxLength = Number(options.maxlength);" in prompt_body
    assert 'inputEl.removeAttribute("maxLength");' in prompt_body
    assert "const submit = () => {" in prompt_body
    assert "options.required" in prompt_body
    assert "options.onInvalid" in prompt_body
    assert "prompt-required-hint" in prompt_body
    assert 'activeInputEl.setAttribute("aria-invalid", "true")' in prompt_body
    assert "okBtn.onclick = () => submit();" in prompt_body
    assert "cleanup(activeInputEl.value);" not in prompt_body
    assert "setTimeout(() => activeInputEl.focus(), 0);" not in prompt_body
    assert "closeCustomDatePicker(dateInputEl);" in prompt_body
    assert 'class="app-modal-shell modal modal-confirm modal-prompt" tabindex="-1"' in html

    assert ".prompt-date-control" in css
    assert ".prompt-date-control .custom-date-shell" in css


def test_user_management_page_and_role_based_navigation_are_present():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert 'data-page="users"' in html
    assert 'id="page-users"' in html
    assert 'id="userManagementPage"' in html
    assert 'id="userTableBody"' in html
    assert 'id="userModal"' in html
    assert 'id="userPassword"' in html
    assert 'id="userRoleCards"' in html
    assert '<select id="userRoleCards"' in html
    assert '<option value="admin">管理员</option>' in html
    assert '<option value="user">普通用户</option>' in html
    assert 'value="governance"' not in html
    assert 'value="regulatory_report"' not in html
    assert 'value="data_middle"' not in html
    assert 'value="fund_custody"' not in html
    assert 'class="nav-item admin-only" data-page="users"' in html
    assert 'class="top-nav-item top-nav-subitem" data-page="users"' in html
    assert "function applyRoleAccess" in app_js
    assert '.admin-only, [data-capability]' in app_js
    assert 'authState.user?.role === "admin"' in app_js
    assert 'api("/api/users"' in app_js
    assert 'api(`/api/users/${encodeURIComponent(userId)}`' in app_js
    assert 'api(`/api/users/${encodeURIComponent(userId)}/reset-password`' in app_js
    assert ".user-management" in css
    assert ".user-stats" in css
    assert ".user-table-card" in css
    assert ".role-badge" in css


def test_role_permissions_page_and_capability_access_are_present():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    # 系统管理分组与角色权限页
    assert 'data-nav-group="system-management"' in html
    assert 'data-page="role-permissions"' in html
    assert 'id="page-role-permissions"' in html
    assert 'data-capability="sys.settings"' in html
    assert 'data-capability="sys.users"' in html
    assert 'data-capability="sys.role_permissions"' in html
    # 能力判断与显隐
    assert "function hasCapability" in app_js
    assert "function applyCapabilityAccess" in app_js
    # canManageHistory 改为能力判断
    assert 'hasCapability("history.delete")' in app_js
    # 系统管理分组 toggle 分流
    assert '"system-management"' in app_js
    # 角色权限页加载
    assert "function loadRolePermissions" in app_js
    # 展示用月度版本号；更新日志条目使用 v1.2.x 小版本编号
    assert 'const DEFAULT_VERSION = "V1.2"' in app_js
    assert '<span class="changelog-version">v1.2.15</span>' in app_js
    # 用户角色仅保留管理员/普通用户；其余走自定义角色
    assert 'governance: "数据治理"' not in app_js
    assert 'regulatory_report: "监管报表"' not in app_js
    assert 'admin: "管理员"' in app_js
    assert 'user: "普通用户"' in app_js
    # 角色权限页样式与原生能力树
    assert ".role-permissions" in css
    assert ".capability-tree" in css
    assert "capability-tree-checkbox" in css
    assert "function createCapabilityTree" in app_js
    assert 'label: "页面查看"' in app_js
    assert 'code: "menu.history", label: "页面查看"' in app_js or (
        'code: "menu.history"' in app_js and 'label: "页面查看"' in app_js
    )
    assert "role-name-custom-hint" not in app_js
    assert 'code === "admin" ? 0 : code === "user" ? 1 : 2' in app_js
    assert "半选/全选按全部子能力勾选态计算" in app_js
    assert 'maxlength="20"' in html
    assert 'maxlength="10"' in html
    assert "角色备注不得超过 20 字" in app_js
    assert "角色名称不得超过 10 个字" in app_js
    assert "jstree" not in app_js.lower()
    assert "jquery.min.js" not in html
    assert "jstree.min.js" not in html
    assert "jstree-style.min.css" not in html
    assert 'input.type = "checkbox"' in app_js
    assert "getMatrix()" in app_js
    assert 'id="roleDefinitionModal"' in html
    assert "function openRoleDefinitionModal" in app_js
    assert "function saveRoleDefinition" in app_js
    assert 'showPrompt("新增角色"' not in app_js
    assert ".user-modal .custom-select-shell.user-role-cards" in css
    assert 'id="rolePermissionsAddBtn"' in html
    assert 'id="rolePermissionsEditBtn"' not in html
    assert 'id="rolePermissionsDeleteBtn"' not in html
    assert 'id="roleDefinitionCode"' in html
    assert 'roleTone = role === "admin" ? "admin" : "user"' in app_js
    assert "#roleDefinitionCode" in css


def test_user_management_table_keeps_action_column_compact():
    html = _read(INDEX_HTML)
    css = _read(STYLES_CSS)

    assert '<th class="user-actions-heading">' in html
    assert '<td class="user-actions-cell">' in _read(APP_JS)
    assert ".user-table col.user-actions-col" in css


def test_admin_local_storage_browser_page_and_api_hooks_are_removed():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)

    for markup in [
        'data-page="local-storage"',
        'id="page-local-storage"',
        "本地数据查询",
        "localStorageTableList",
        "localStorageDataHead",
        "localStorageSchemaBody",
        "localStorageInfoPanel",
        "localStorageJsonDrawer",
        "localStorageExportSchemaBtn",
        "localStorageMigrateHistoryBtn",
        "localStorageExportTableBtn",
        "localStorageBackupBtn",
        "localStorageRefreshBtn",
    ]:
        assert markup not in html

    for script_marker in [
        "/api/admin/storage",
        "function loadLocalStorageBrowser",
        "function localStorageColumnLabel",
        "function isLocalStorageBooleanField",
        "formatLocalStorageValue(field, value)",
        "localStorageBrowserState",
        "localStorageExportTableBtn",
        "localStorageMigrateHistoryBtn",
        "renderLocalStorageMigrationStatus",
        'name === "local-storage"',
    ]:
        assert script_marker not in app_js

    assert '.admin-only, [data-capability]' in app_js


def test_user_management_cards_and_rows_have_theme_glow_hover_motion():
    css = _read(STYLES_CSS)
    readme = _read(README_MD)

    for selector in [
        r"\.user-stat-card:hover",
        r"\.user-filter-bar:hover",
        r"\.user-table-card:hover",
        r"\.user-table tbody tr:not\(\.user-loading-row\):hover",
    ]:
        rule = re.search(rf"(?m)^{selector}\s*\{{(?P<body>.*?)\}}", css, re.S)
        assert rule is not None
        body = rule.group("body")
        assert "var(--card-hover-glow)" in body
        assert "var(--card-hover-shadow" in body
        assert "transform:" in body
        assert "scale(" not in body
        assert "rgba(0, 0, 0" not in body

    for selector in [
        r"\.user-stat-card",
        r"\.user-filter-bar",
        r"\.user-table-card",
        r"\.user-table tbody tr:not\(\.user-loading-row\)",
    ]:
        rule = re.search(rf"(?m)^{selector}\s*\{{(?P<body>.*?)\}}", css, re.S)
        assert rule is not None
        assert "transition:" in rule.group("body")

    assert "用户管理统计卡、筛选区和用户行加入主题化光晕及轻弹动效" in readme


def test_user_management_retains_reference_stats_filters_export_and_icon_actions():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    for text in ['id="exportUsersBtn"', 'id="userStatDisabled"', 'class="user-filter-pills"', 'data-user-filter="all"']:
        assert text in html
    for text in ["导出列表", "已停用", "全部"]:
        assert text in html
    assert "function exportUsers" in app_js
    assert "function setUserQuickFilter" in app_js
    assert 'document.querySelectorAll("[data-user-filter]")' in app_js
    assert 'class="user-icon-action edit-user"' in app_js
    assert 'class="user-icon-action toggle-user"' in app_js
    assert 'class="user-icon-action delete-user"' in app_js
    assert ".user-stat-card--disabled" in css
    assert ".user-filter-pills" in css
    assert ".user-icon-action" in css


def test_user_csv_export_escapes_formula_values():
    app_js = _read(APP_JS)

    assert "function escapeCsvValue" in app_js
    assert "formulaPrefixPattern.test(text.trimStart())" in app_js
    export_users = re.search(r"function exportUsers\(\) \{(?P<body>.*?)\n\}", app_js, re.S)
    assert export_users is not None
    assert "row.map(escapeCsvValue)" in export_users.group("body")
    assert 'String(value).replace(/"/g, \'""\')' not in export_users.group("body")


def test_user_menu_uses_random_initial_avatar_when_name_updates():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)
    login_html = _read(ROOT / "src" / "auto_check" / "web" / "login.html")

    assert 'class="user-menu-icon user-initial-avatar" data-current-user-avatar' in html
    assert '<span class="user-menu-icon" aria-hidden="true"><svg' not in html
    assert 'data-current-username-text' in html
    assert "const USER_AVATAR_SESSION_KEY = \"autoCheckUserAvatarVariant\";" in app_js
    assert "const USER_AVATAR_GRADIENTS = [" in app_js
    assert "function userAvatarInitial(value)" in app_js
    assert "function currentUserAvatarGradient()" in app_js
    assert "sessionStorage.getItem(USER_AVATAR_SESSION_KEY)" in app_js
    assert "sessionStorage.setItem(USER_AVATAR_SESSION_KEY, String(index));" in app_js
    assert "sessionStorage.removeItem(USER_AVATAR_SESSION_KEY)" in app_js
    assert 'querySelector("[data-current-username-text]")' in app_js
    assert 'querySelector("[data-current-user-avatar]")' in app_js
    username_body = re.search(r"function updateCurrentUsername\(\) \{(?P<body>.*?)\n\}", app_js, re.S)
    assert username_body is not None
    assert "item.textContent = username" not in username_body.group("body")
    assert "avatar.textContent = initial;" in username_body.group("body")
    assert 'avatar.style.setProperty("--avatar-from", from);' in username_body.group("body")
    assert 'avatar.style.setProperty("--avatar-to", to);' in username_body.group("body")
    assert ".user-menu-icon" in css
    assert "linear-gradient(135deg, var(--avatar-from, #6366f1), var(--avatar-to, #4338ca))" in css
    assert ".user-menu-icon svg" not in css
    dark_user_icon = re.search(r'\[data-color-mode="dark"\] \.user-menu-icon\s*\{(?P<body>.*?)\}', css, re.S)
    assert dark_user_icon is not None
    assert "var(--avatar-from, #6366f1)" in dark_user_icon.group("body")
    assert "box-shadow" in dark_user_icon.group("body")
    assert "const USER_AVATAR_SESSION_KEY = \"autoCheckUserAvatarVariant\";" in login_html
    assert "function refreshUserAvatarVariant()" in login_html
    assert "refreshUserAvatarVariant();" in login_html


def test_user_management_nav_icon_is_subtle_css_icon_in_light_and_dark_modes():
    html = _read(INDEX_HTML)
    css = _read(STYLES_CSS)

    assert '<span class="nav-icon nav-icon-users" aria-hidden="true"></span>' in html
    assert '&#128101;' not in html
    assert ".nav-icon-users::before" in css
    assert ".nav-icon-users::after" in css
    assert "[data-color-mode=\"dark\"] .nav-icon-users" in css


def test_regular_user_settings_are_compact_without_changing_admin_about_details():
    html = _read(INDEX_HTML)
    css = _read(STYLES_CSS)

    assert 'class="about-features about-admin-detail"' in html
    assert 'class="about-tech about-admin-detail"' in html
    assert "[data-role=\"user\"] #page-settings .about-admin-detail" in css
    assert "[data-role=\"user\"] #page-settings .card-about" in css
    assert "[data-role=\"user\"] #page-settings .card-system-info" in css
    assert "[data-role=\"user\"] #page-settings .card-interface" in css
    user_grid = re.search(r'\[data-role="user"\] #page-settings \.settings-dashboard-grid\s*\{(?P<body>.*?)\}', css, re.S)
    assert user_grid is not None
    assert "align-items: stretch" in user_grid.group("body")
    user_cards = re.search(
        r'\[data-role="user"\] #page-settings \.card-system-info,\s*'
        r'\[data-role="user"\] #page-settings \.card-interface,\s*'
        r'\[data-role="user"\] #page-settings \.card-about\s*\{(?P<body>.*?)\}',
        css,
        re.S,
    )
    assert user_cards is not None
    assert "height: 100%" in user_cards.group("body")
    assert "#page-settings .card-business,\n#page-settings .card-about {\n  height: 800px;" in css


def test_pbc_import_footer_shows_uploaded_file_total_near_next_button():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    footer = re.search(
        r'<div class="app-modal-footer pbc-modal-footer" id="pbcModalFooter">(?P<body>.*?)</div>',
        html,
        re.S,
    )
    assert footer is not None
    assert 'id="pbcUploadSummary"' in footer.group("body")
    assert 'id="pbcClearFilesBtn"' in footer.group("body")
    assert footer.group("body").index('id="pbcUploadSummary"') < footer.group("body").index('id="pbcNextBtn"')
    assert footer.group("body").index('id="pbcClearFilesBtn"') < footer.group("body").index('id="pbcNextBtn"')
    assert "function updatePbcUploadSummary()" in app_js
    assert 'pbcUploadSummary.textContent = `共 ${fileCount} 个文件`;' in app_js
    assert "function clearPbcUploadedFiles()" in app_js
    assert 'pbcClearFilesBtn.disabled = fileCount === 0;' in app_js
    assert "updatePbcUploadSummary();" in app_js
    assert ".pbc-upload-summary" in css
    assert ".pbc-btn--ghost" in css


def test_user_avatars_show_online_current_badge_and_reference_stat_icon():
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert "const isCurrentUser = user.id && user.id === authState.user?.id;" in app_js
    assert 'user-avatar-wrap' in app_js
    assert 'user-avatar-status' in app_js
    assert 'current-user-badge' in app_js
    assert '.user-avatar-wrap' in css
    assert '.user-avatar.is-online' in css
    assert '.user-avatar-status' in css
    assert '.current-user-badge' in css
    total_icon = re.search(r"(?m)^\.user-stat-icon--blue\s*\{(?P<body>.*?)\}", css, re.S)
    assert total_icon is not None
    assert "#eef2ff" in total_icon.group("body")
    assert "[data-theme=\"space-tech\"] .user-avatar" in css


def test_user_edit_modal_matches_reference_layout_and_does_not_close_on_blank_overlay():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    modal = re.search(
        r'<div class="app-modal-overlay modal-overlay" id="userModal" hidden>(?P<body>.*?)</div>\s*</div>\s*<!--',
        html,
        re.S,
    )
    assert modal is not None
    for token in [
        "user-modal-header",
        "user-modal-title",
        "user-modal-form",
        "user-role-cards",
        "user-enable-row",
        "user-modal-footer",
    ]:
        assert token in modal.group("body")
        assert f".{token}" in css
    assert "app-modal-shell" in modal.group("body")
    assert "user-modal-icon" not in modal.group("body")
    assert '<input id="userRole" type="hidden" value="user" />' in modal.group("body")
    assert '<input id="userEnabled" type="hidden" value="true" />' in modal.group("body")
    assert '<select id="userRoleCards"' in modal.group("body")
    assert 'class="filter-select user-role-cards"' in modal.group("body")
    assert "function renderUserRoleCards()" in app_js
    assert "function syncUserRoleCards()" in app_js
    assert "function syncUserEnabledSwitch()" in app_js
    assert "function isDelegatedAdminSession()" in app_js
    assert "初始管理员角色不可修改" in app_js
    assert "委派管理员不可创建或设置管理员" in app_js
    assert 'api(`/api/users/${encodeURIComponent(targetUserId)}/reset-password`' in app_js
    assert 'autocomplete="new-name"' in modal.group("body")
    assert 'const displayNameValue = isEdit ? userDisplayName(user) : "";' in app_js
    assert 'if (!userId.value && userDisplayNameInput) userDisplayNameInput.value = "";' in app_js
    user_events = re.search(
        r"userModalClose\?\.addEventListener\(\"click\", closeUserModal\);(?P<body>.*?)userModalSave\?\.addEventListener",
        app_js,
        re.S,
    )
    assert user_events is not None
    assert 'userModal?.addEventListener("click"' not in user_events.group("body")


def test_user_management_space_theme_toolbar_and_actions_match_reference():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert 'class="user-filter-actions-space"' in html
    assert 'data-new-user-btn' in html
    assert 'data-export-users-btn' in html
    assert '[data-theme="space-tech"] #page-users .user-toolbar-text' in css
    assert "display: none" in re.search(
        r'\[data-theme="space-tech"\] #page-users \.user-toolbar-text\s*\{(?P<body>.*?)\}',
        css,
        re.S,
    ).group("body")
    assert '[data-theme="space-tech"] #page-users .user-toolbar-actions' in css
    assert '[data-theme="space-tech"] #page-users .user-filter-actions-space' in css
    assert "border-radius: 999px" in re.search(
        r'\[data-theme="space-tech"\] #page-users \.user-filter-actions-space \.btn-outline,\s*\n\[data-theme="space-tech"\] #page-users \.user-filter-actions-space \.btn-primary\s*\{(?P<body>.*?)\}',
        css,
        re.S,
    ).group("body")
    assert 'document.querySelectorAll("[data-new-user-btn]")' in app_js
    assert 'document.querySelectorAll("[data-export-users-btn]")' in app_js


def test_user_management_list_uses_display_fields_pagination_and_admin_guards():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert "最近登录时间" in html
    assert "更新时间" not in html[html.index('id="page-users"'):html.index('id="page-settings"')]
    assert 'id="userPagination"' in html
    assert 'id="userPrevPage"' in html
    assert 'id="userNextPage"' in html
    assert "user-loading-row" in html
    assert "function renderUsersLoading()" in app_js
    assert "loadUsers({ force = false } = {})" in app_js
    assert "renderUsersLoading();" in app_js
    assert ".user-loading-row" in css
    assert ".user-skeleton" in css
    assert "function paginatedUsers" in app_js
    assert "let userCurrentPage = 1" in app_js
    assert "last_login_at" in app_js
    assert "role-badge-icon" in app_js
    assert 'id="userDisplayName"' in html
    assert "function userDisplayName" in app_js
    assert "const isInitialAdmin = user.username === \"admin\";" in app_js
    assert "const isAdminUser = role === \"admin\";" not in app_js
    assert 'button.classList.contains("reset-user")' not in app_js
    render_body = re.search(r"function renderUsers\(\) \{(?P<body>.*?)\n\}\n\nfunction renderUsersLoading", app_js, re.S).group("body")
    assert "const displayName = userDisplayName(user);" in render_body
    assert 'class="user-name-line"' in render_body
    assert '<span class="user-name-line">' in render_body
    assert "<strong>${escapeHtml(displayName)}</strong>" in render_body
    assert "<small>${escapeHtml(user.username || \"\")}</small>" in render_body
    assert "reset-user" not in render_body
    assert "disabled" in render_body
    assert 'title="${isAdminUser ? "管理员不可停用"' not in render_body
    assert 'title="${isAdminUser ? "管理员不可删除"' not in render_body
    assert "初始管理员不可停用" in render_body
    assert "初始管理员不可删除" in render_body
    assert "#page-users .user-management" in css
    page_users = re.search(r"(?m)^#page-users\s*\{(?P<body>.*?)\}", css, re.S)
    assert page_users is not None
    assert "height: 100%" in page_users.group("body")
    user_management = re.search(r"(?m)^\.user-management\s*\{(?P<body>.*?)\}", css, re.S)
    assert user_management is not None
    assert "flex: 1" in user_management.group("body")
    assert "min-height: 0" in user_management.group("body")
    assert ".user-pagination" in css


def test_user_display_name_drives_navigation_and_user_export():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)

    assert "data-current-username-text" in html
    username_body = re.search(r"function updateCurrentUsername\(\) \{(?P<body>.*?)\n\}", app_js, re.S)
    assert username_body is not None
    assert "const displayName = userDisplayName(authState.user);" in username_body.group("body")
    assert "nameText.textContent = displayName" in username_body.group("body")
    assert "item.title = `${displayName} (${username})`" in username_body.group("body")
    assert 'const headers = ["用户姓名", "用户账号", "角色", "状态", "创建时间", "最近登录时间"];' in app_js


def test_run_history_displays_executor_and_recent_run_summary():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)

    history_head = html[html.index('id="historyBody"') - 600:html.index('id="historyBody"')]
    assert history_head.index("报告期") < history_head.index("执行时间")
    assert "执行人" in history_head
    assert '<th class="history-source-only">数据源</th>' in history_head
    assert history_head.index("<th>总差异</th>") < history_head.index("<th>已解释</th>")
    assert history_head.index("<th>已解释</th>") < history_head.index("<th>新增差异</th>")
    assert history_head.index("<th>新增差异</th>") < history_head.index("<th>减少差异</th>")
    assert "<th>未解释</th>" not in history_head
    assert '<tr><td colspan="9" class="empty">' in html
    assert 'historySummaryItem("报告期", run.run_date)' in app_js
    assert 'historySummaryItem("执行时间", run.run_at)' in app_js
    assert '<td>${escapeHtml(run.run_date)}</td>' in app_js
    assert "<td>${escapeHtml(historyExecutorName(run))}</td>" in app_js
    assert 'function historyColumnCount() {\n  return canSeeHistorySource() ? 9 : 8;\n}' in app_js
    assert 'formatHistoryDiffCount(run, "added_count", { unit: false })' in app_js
    assert 'formatHistoryDiffCount(run, "removed_count", { unit: false })' in app_js
    assert '<td class="money-cell">${formatMoney(unresolved)}</td>' not in app_js
    row_start = app_js.index('return `<tr class="history-main-row"')
    row_end = app_js.index("</tr>`;", row_start)
    row_body = app_js[row_start:row_end]
    assert row_body.index("<td class=\"money-cell\">${formatMoney(run.total_count)}</td>") < row_body.index("<td class=\"money-cell\">${formatMoney(explained)}</td>")
    assert row_body.index("<td class=\"money-cell\">${formatMoney(explained)}</td>") < row_body.index("history-added")
    assert row_body.index("history-added") < row_body.index("history-removed")
    assert 'setLastRunTime(latestHistory.run_at, historyExecutorName(latestHistory))' in app_js
    assert 'lastRunTime.textContent = `最近执行：${executor}  ${latestRunAt}`;' in app_js

    css = _read(STYLES_CSS)
    added = re.search(r"(?m)^\.history-added\s*\{(?P<body>.*?)\}", css, re.S)
    removed = re.search(r"(?m)^\.history-removed\s*\{(?P<body>.*?)\}", css, re.S)
    assert added is not None
    assert removed is not None
    assert "color: var(--error)" in added.group("body")
    assert "color: var(--success-text)" in removed.group("body")


def test_history_restore_marks_result_list_and_keeps_latest_snapshot():
    app_js = _read(APP_JS)

    restore_history = re.search(r"function restoreHistoryRun\(run\) \{(?P<body>.*?)\n\}", app_js, re.S)
    assert restore_history is not None
    body = restore_history.group("body")
    assert "setResultHistoryRestoreState(run, results.length);" in body
    assert "setLastRunTime(run.run_at" not in body
    assert "saveLatestResultsSnapshot" not in body
    assert "lastRunTime.hidden = true;" in app_js
    assert 'showToast("结果列表已恢复到历史数据", "info")' in body
    assert "historyRestoreHintText(resultRestoreHistoryMeta, results.length)" in body

    restore_latest = re.search(r"async function restoreLatestResultsToResultList\(\) \{(?P<body>.*?)\n\}", app_js, re.S)
    assert restore_latest is not None
    latest_body = restore_latest.group("body")
    assert "const restored = await loadLatestHistoryResults();" in latest_body
    assert 'setStatus("结果列表已还原到最新结果")' in latest_body
    assert 'showToast("结果列表已还原到最新结果", "success")' in latest_body

    load_latest = re.search(r"async function loadLatestHistoryResults\(\) \{(?P<body>.*?)\n\}", app_js, re.S)
    assert load_latest is not None
    latest_loader = load_latest.group("body")
    assert "clearResultHistoryRestoreState();" in latest_loader
    assert 'homeResultListFilterLabel = "";' in latest_loader
    assert "setLastRunTime(latestHistory.run_at, historyExecutorName(latestHistory))" in latest_loader


def test_history_detail_opens_in_modal_and_respects_permissions():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert 'id="historyDetailCard"' not in html
    assert "let selectedHistoryId" in app_js
    assert "let historyDetailLoadingId" not in app_js
    assert "function historyDetailRow(innerHtml)" not in app_js
    assert "function renderHistoryDetailLoading(id)" in app_js
    assert "function renderHistoryDetailContent(run)" in app_js
    assert "function renderHistoryDetailFooter(run)" not in app_js
    assert "function showHistoryDetailModal(id)" in app_js
    assert 'showInfo("历史详情", renderHistoryDetailLoading(id), { modalClass: "modal-info--history-detail", closeOnBackdrop: false });' in app_js
    assert 'detailActionLabel: "恢复到结果页"' in app_js
    assert "onDetailAction: async () =>" in app_js
    assert "await restoreHistoryRun(history);" in app_js
    assert "rowHtml += historyDetailRow" not in app_js
    assert 'class="history-main-row"' in app_js
    assert 'class="history-detail-row"' not in app_js
    assert 'class="history-detail-title"' not in app_js
    assert 'class="btn-close close-history-detail"' not in app_js
    assert 'class="btn-outline btn-xs restore-history"' not in app_js
    assert 'id="infoFooter"' in html
    assert 'class="app-modal-footer history-detail-footer"' in html
    assert 'const footerEl = document.getElementById("infoFooter");' in app_js
    assert 'footerEl.innerHTML = options.footerContent || "";' in app_js
    assert 'footerEl.hidden = !options.footerContent;' in app_js
    assert "历史详情 -" not in app_js
    assert "function historyBaselineText(run = {})" in app_js
    assert "`${baselineRunAt}执行的同报告期记录`" in app_js
    assert "function historyHasBaseline(run = {})" in app_js
    assert "function formatHistoryDiffCount(run = {}, field = \"\", options = {})" in app_js
    assert "function historyDiffItems(run = {}, field = \"\")" in app_js
    detail_start = app_js.index("function renderHistoryDetailContent(run)")
    detail_end = app_js.index("function renderHistoryDetailLoading", detail_start)
    detail_body = app_js[detail_start:detail_end]
    assert "history-detail-footer" not in detail_body
    assert detail_body.index('historySummaryItem("报告期", run.run_date)') < detail_body.index('historySummaryItem("执行人", historyExecutorName(run))')
    assert detail_body.index('historySummaryItem("执行人", historyExecutorName(run))') < detail_body.index('historySummaryItem("执行时间", run.run_at)')
    assert detail_body.index('historySummaryItem("执行时间", run.run_at)') < detail_body.index('historySummaryItem("基准记录", historyBaselineText(run))')
    assert 'historySummaryItem("基准记录", historyBaselineText(run))' in app_js
    assert 'historySummaryItem("规则版本"' not in app_js
    assert 'historySummaryItem("执行人", historyExecutorName(run))' in app_js
    assert 'historySummaryItem("总差异", formatMoney(run.total_count))' not in app_js
    assert 'historySummaryItem("数据源", formatHistorySourceName(run))' not in app_js
    assert 'historySummaryItem("新增差异", formatMoney(run.added_count))' not in app_js
    assert 'historySummaryItem("减少差异", formatMoney(run.removed_count))' not in app_js
    assert "const sourceSummary = canSeeHistorySource()" not in app_js
    assert "function canManageHistory()" in app_js
    assert "function canSeeHistorySource()" in app_js
    assert "function historyColumnCount()" in app_js
    assert "const deleteAction = canManageHistory()" in app_js
    assert 'if (!canManageHistory()) {' in app_js
    assert '<td>${escapeHtml(formatHistorySourceName(run))}</td>' in app_js
    history_result_table = re.search(r"function historyResultTable\(items\) \{(?P<body>.*?)\n\}", app_js, re.S)
    assert history_result_table is not None
    assert "<th>差异类型</th><th>状态</th>" in history_result_table.group("body")
    assert "具体原因" not in history_result_table.group("body")
    assert "specificReasonText(item)" not in history_result_table.group("body")

    load_detail = re.search(r"async function loadHistoryDetail\(id\) \{(?P<body>.*?)\n\}", app_js, re.S)
    assert load_detail is not None
    assert "historyDetailLoadingId" not in load_detail.group("body")
    assert "renderHistoryList();" not in load_detail.group("body")
    assert "await api(" in load_detail.group("body")
    assert "options.closeOnBackdrop === false" in app_js

    assert ".history-detail-card" in css
    assert ".history-detail-card .history-detail" in css
    assert ".modal-info.modal-info--history-detail" in css
    assert "[data-color-mode=\"dark\"] .history-detail-card" in css
    assert "var(--surface-container-lowest)" in css
    assert "var(--on-surface)" in css
    assert 'history-section--${tone}' in app_js
    assert 'items.length > 10 ? " history-section--scroll" : ""' in app_js
    assert 'class="history-status ${statusClass}"' in app_js


def test_history_detail_modal_layout_keeps_tables_readable():
    html = _read(INDEX_HTML)
    css = _read(STYLES_CSS)

    card = re.search(r"(?m)^\.history-detail-card\s*\{(?P<body>.*?)\}", css, re.S)
    assert card is not None
    assert "display: flex" in card.group("body")
    assert "max-height" not in card.group("body")
    assert "overflow: hidden" in card.group("body")
    assert "width: 100%" in card.group("body")
    assert "height: 100%" in card.group("body")
    assert "border: 1px solid" not in card.group("body")
    assert "box-shadow" not in card.group("body")

    modal = re.search(r"(?m)^\.modal-info\.modal-info--history-detail\s*\{(?P<body>.*?)\}", css, re.S)
    assert modal is not None
    assert "width: min(960px, 88vw)" in modal.group("body")
    assert "max-height: 92vh" in modal.group("body")
    assert "display: flex" in modal.group("body")
    assert "flex-direction: column" in modal.group("body")

    modal_header = re.search(
        r"(?m)^\.modal-info\.modal-info--history-detail > \.app-modal-header\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert modal_header is not None
    assert "height: 58px" in modal_header.group("body")
    assert "min-height: 58px" in modal_header.group("body")

    modal_body = re.search(r"(?m)^\.modal-info\.modal-info--history-detail \.modal-body\s*\{(?P<body>.*?)\}", css, re.S)
    assert modal_body is not None
    assert "flex: 1 1 auto" in modal_body.group("body")
    assert "min-height: 0" in modal_body.group("body")
    assert "overflow: hidden" in modal_body.group("body")

    summary_grid = re.search(r"(?m)^\.history-summary-grid\s*\{(?P<body>.*?)\}", css, re.S)
    assert summary_grid is not None
    assert "display: flex" in summary_grid.group("body")
    assert "flex-wrap: wrap" in summary_grid.group("body")

    detail = re.search(r"(?m)^\.history-detail-card \.history-detail\s*\{(?P<body>.*?)\}", css, re.S)
    assert detail is not None
    assert "flex: 1 1 auto" in detail.group("body")
    assert "overflow: auto" in detail.group("body")
    assert "padding: 12px 8px" in detail.group("body")

    assert html.index('id="infoBody"') < html.index('id="infoFooter"')
    section = re.search(r"(?m)^\.history-section\s*\{(?P<body>.*?)\}", css, re.S)
    assert section is not None
    assert "flex: 0 0 auto" in section.group("body")

    section_table = re.search(r"(?m)^\.history-section-table\s*\{(?P<body>.*?)\}", css, re.S)
    assert section_table is not None
    assert "height:" not in section_table.group("body")
    assert "max-height:" not in section_table.group("body")
    assert "overflow-x: auto" in section_table.group("body")
    assert "overflow-y: visible" in section_table.group("body")

    scroll_table = re.search(
        r"(?m)^\.history-section--scroll \.history-section-table\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert scroll_table is not None
    assert "max-height: 360px" in scroll_table.group("body")
    assert "overflow: auto" in scroll_table.group("body")

    result_header = re.search(r"(?m)^\.history-result-table th\s*\{(?P<body>.*?)\}", css, re.S)
    assert result_header is not None
    assert "position: sticky" in result_header.group("body")
    assert "top: 0" in result_header.group("body")
    assert "z-index: 2" in result_header.group("body")
    assert "color: var(--on-surface-variant)" in result_header.group("body")
    assert "font-weight: 600" in result_header.group("body")
    assert "background: var(--surface-container-low)" in result_header.group("body")
    assert "border-bottom: 1px solid color-mix(in srgb, var(--outline-variant) 32%, var(--surface-container-lowest))" in result_header.group("body")

    status_col = re.search(
        r"(?m)^\.history-result-table th:nth-child\(5\),\s*\n\.history-result-table td:nth-child\(5\)\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert status_col is not None
    assert "width: 96px" in status_col.group("body")

    code_col = re.search(
        r"(?m)^\.history-result-table th:nth-child\(1\),\s*\n\.history-result-table td:nth-child\(1\)\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert code_col is not None
    assert "width: calc(11ch + 16px)" in code_col.group("body")
    assert "white-space: nowrap" in code_col.group("body")

    name_col = re.search(
        r"(?m)^\.history-result-table th:nth-child\(2\),\s*\n\.history-result-table td:nth-child\(2\)\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert name_col is not None
    assert "white-space: normal" in name_col.group("body")
    assert "overflow-wrap: anywhere" in name_col.group("body")
    assert "word-break: break-word" in name_col.group("body")

    money_col = re.search(
        r"(?m)^\.history-result-table th:nth-child\(3\),\s*\n\.history-result-table td:nth-child\(3\)\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert money_col is not None
    assert "white-space: nowrap" in money_col.group("body")
    assert "width: calc(18ch + 28px)" in money_col.group("body")
    assert "padding-right: 14px" in money_col.group("body")

    reason_col = re.search(
        r"(?m)^\.history-result-table th:nth-child\(4\),\s*\n\.history-result-table td:nth-child\(4\)\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert reason_col is not None
    assert "padding-left: 12px" in reason_col.group("body")
    assert "white-space: normal" in reason_col.group("body")

    history_status = re.search(r"(?m)^\.history-status\s*\{(?P<body>.*?)\}", css, re.S)
    assert history_status is not None
    assert "padding: 2px 8px" in history_status.group("body")
    assert "max-width: 100%" in history_status.group("body")

    cells = re.search(
        r"(?m)^\.history-result-table th,\s*\n\.history-result-table td\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert cells is not None
    assert "text-align: center" in cells.group("body")
    assert ".history-result-table td.money-cell" in css

    summary_value = re.search(r"(?m)^\.history-summary-item strong\s*\{\s*min-width: 0;(?P<body>.*?)\}", css, re.S)
    assert summary_value is not None
    assert "color: color-mix(in srgb, var(--on-surface) 90%, var(--surface-container-lowest))" in summary_value.group("body")
    assert "font-weight: 500" in summary_value.group("body")
    assert "overflow-wrap: break-word" in summary_value.group("body")
    assert "word-break: normal" in summary_value.group("body")

    summary_label = re.search(r"(?m)^\.history-summary-item span\s*\{(?P<body>.*?)\}", css, re.S)
    assert summary_label is not None
    assert "color: color-mix(in srgb, var(--on-surface-variant) 58%, var(--surface-container-lowest))" in summary_label.group("body")

    section_title = re.search(r"(?m)^\.history-section-title\s*\{(?P<body>.*?)\}", css, re.S)
    assert section_title is not None
    assert "color: color-mix(in srgb, var(--on-surface) 90%, var(--surface-container-lowest))" in section_title.group("body")
    assert "font-weight: 500" in section_title.group("body")

    section_count = re.search(r"(?m)^\.history-section-title > span:last-child\s*\{(?P<body>.*?)\}", css, re.S)
    assert section_count is not None
    assert "color: color-mix(in srgb, var(--on-surface-variant) 40%, var(--surface-container-lowest))" in section_count.group("body")
    assert "font-weight: 400" in section_count.group("body")
    assert '[data-color-mode="dark"] .history-summary-item' in css
    assert '[data-color-mode="dark"] .history-result-table td' in css
    dark_result_header = re.search(
        r'(?m)^\[data-color-mode="dark"\] \.history-result-table th\s*\{(?P<body>.*?)\}',
        css,
        re.S,
    )
    assert dark_result_header is not None
    assert "color: #cbd5e1" in dark_result_header.group("body")
    assert "background: rgba(30, 41, 59, 0.94)" in dark_result_header.group("body")
    for status_tone in ("done", "pending"):
        status = re.search(rf"(?m)^\.history-status--{status_tone}\s*\{{(?P<body>.*?)\}}", css, re.S)
        assert status is not None
        assert "background: transparent" in status.group("body")
    assert ".history-detail-counts" not in css


def test_history_detail_uses_inline_metadata_and_colored_sections():
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    start = app_js.index("function renderHistoryDetailContent(run)")
    end = app_js.index("function renderHistoryDetailLoading", start)
    detail = app_js[start:end]

    complete = 'historySection("本次完整核对结果", run.results || [], "complete")'
    added = 'historySection("本次新增差异", historyDiffItems(run, "added_results"), "added")'
    removed = 'historySection("本次减少差异", historyDiffItems(run, "removed_results"), "removed")'
    assert detail.index(complete) < detail.index(added) < detail.index(removed)
    assert "${historyDetailCounts(run)}" not in detail
    assert "function historyDetailCounts" not in app_js
    assert "function historyCountItem" not in app_js

    summary = re.search(r"(?m)^\.history-summary-grid\s*\{(?P<body>.*?)\}", css, re.S)
    assert summary is not None
    assert "display: flex" in summary.group("body")
    assert "flex-wrap: wrap" in summary.group("body")

    for tone in ("complete", "added", "removed"):
        assert f".history-section--{tone} .history-section-bar" in css
    assert ".history-status--done" in css
    assert ".history-status--pending" in css

    cells = re.search(
        r"(?m)^\.history-result-table th,\s*\n\.history-result-table td\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert cells is not None
    assert "text-align: center" in cells.group("body")
    assert ".history-result-table td.money-cell" in css


def test_shared_modal_shell_hides_scrollbars_without_changing_scroll_behavior():
    css = _read(STYLES_CSS)

    scrollbar_scope = re.search(
        r"(?m)^\.app-modal-shell,\s*\n\.app-modal-shell \*\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert scrollbar_scope is not None
    assert "scrollbar-width: none" in scrollbar_scope.group("body")
    assert "-ms-overflow-style: none" in scrollbar_scope.group("body")

    webkit_scrollbar = re.search(
        r"(?m)^\.app-modal-shell::\-webkit-scrollbar,\s*\n\.app-modal-shell \*::\-webkit-scrollbar\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert webkit_scrollbar is not None
    assert "display: none" in webkit_scrollbar.group("body")
    assert "width: 0" in webkit_scrollbar.group("body")
    assert "height: 0" in webkit_scrollbar.group("body")

    history_scroll = re.search(
        r"(?m)^\.app-modal-shell\.modal-info--history-detail \.history-detail\.is-scrolling,\s*\n"
        r"\.app-modal-shell\.modal-info--history-detail \.history-section--scroll \.history-section-table\.is-scrolling,\s*\n"
        r"\.app-modal-shell\.modal-info--history-detail \.history-section-table\.is-scrolling\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert history_scroll is not None
    assert "scrollbar-width: thin" in history_scroll.group("body")
    assert "scrollbar-color: var(--ui-thin-scrollbar-thumb, #c5d0e0) transparent" in history_scroll.group("body")
    assert "--ui-thin-scrollbar-size: 6px;" in css
    assert "--ui-thin-scrollbar-thumb: #c5d0e0;" in css
    assert (
        ".app-modal-shell.modal-info--history-detail .history-detail.is-scrolling::-webkit-scrollbar,\n"
        ".app-modal-shell.modal-info--history-detail .history-section--scroll .history-section-table.is-scrolling::-webkit-scrollbar,\n"
        ".app-modal-shell.modal-info--history-detail .history-section-table.is-scrolling::-webkit-scrollbar"
        in css
    )
    assert "width: var(--ui-thin-scrollbar-size, 6px);" in css
    assert ":hover, .is-scrolling" not in css or ".history-detail:is(:hover, .is-scrolling)" not in css
    assert "function bindHistoryDetailAutoHideScrollbars()" in _read(APP_JS)
    assert "bindHistoryDetailAutoHideScrollbars();" in _read(APP_JS)
    assert 'el.classList.add("is-scrolling");' in _read(APP_JS)
    assert "}, 2000);" in _read(APP_JS)

    modal_body = re.search(
        r"(?m)^\.app-modal-shell:not\(#pbcModal\):not\(#dbValidationModal\):not\(#flowModal\)"
        r" > \.app-modal-body\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert modal_body is not None
    assert "overflow: auto" in modal_body.group("body")


def test_report_navigation_and_settings_show_thin_main_content_scrollbar():
    css = _read(STYLES_CSS)
    html = _read(INDEX_HTML)

    page_scroll = re.search(
        r"(?m)^:root\[data-page=\"report-navigation\"\]\[data-theme=\"space-tech\"\] \.main-content,\s*\n"
        r":root\[data-page=\"settings\"\]\[data-theme=\"space-tech\"\] \.main-content\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert page_scroll is not None
    body = page_scroll.group("body")
    assert "scrollbar-width: thin" in body
    assert "scrollbar-color: var(--ui-thin-scrollbar-thumb, #c5d0e0) transparent" in body
    assert "margin-top: 76px" not in body
    assert "height: calc(100vh - 76px)" not in body
    assert (
        ':root[data-page="report-navigation"][data-theme="space-tech"] .main-content::-webkit-scrollbar'
        in css
    )
    assert (
        ':root[data-page="settings"][data-theme="space-tech"] .main-content::-webkit-scrollbar'
        in css
    )
    assert "width: var(--ui-thin-scrollbar-size, 6px);" in css
    assert "background: var(--ui-thin-scrollbar-thumb, #c5d0e0);" in css
    assert 'href="/styles.css?v=20260813l"' in html
    assert 'src="/app.js?v=20260813o"' in html


def test_space_tech_top_nav_is_edge_stuck_not_floating():
    css = _read(STYLES_CSS)

    space_top_nav = re.search(
        r'\[data-theme="space-tech"\] \.top-nav\s*\{(?P<body>.*?)\}',
        css,
        re.S,
    )
    assert space_top_nav is not None
    body = space_top_nav.group("body")
    assert "position: relative" in body
    assert "top: 0" in body
    assert "left: 0" in body
    assert "right: 0" in body
    assert "width: 100%" in body
    assert "border-radius: 0 !important" in body
    assert "border-bottom: 1px solid var(--outline-variant)" in body
    assert "box-shadow: none" in body
    assert "background: var(--surface-container-lowest)" in body
    assert "backdrop-filter: none" in body
    assert "overflow: visible" in body
    assert "z-index: 40" in body
    assert "padding: 8px 32px;" in body
    assert "flex: 0 0 auto" in body
    assert '[data-theme="space-tech"] .top-nav,' not in css.split("User interface radius preference: start", 1)[1].split("User interface radius preference: end", 1)[0]
    # Active tab keeps Logo gradient but never casts a blue glow under the nav.
    assert "0 8px 18px rgba(59, 130, 246, 0.24)" not in css
    assert "Keep Logo gradient fill; do not cast a blue glow under the active tab." in css
    assert (
        ".top-nav-item.active,\n.top-nav-group.active > .top-nav-group-toggle,\n.nav-item.active,"
        not in css
    )
    assert (
        ".nav-item.active,\n.nav-group.active > .nav-group-toggle,\n.btn-primary,"
        in css
    )
    assert "box-shadow: 0 6px 16px var(--theme-focus-ring);" in css
    glow_cleanup = re.search(
        r"/\* Top nav glow cleanup:[\s\S]*?\*/\n(?P<sel>[^{]+)\{(?P<body>[^}]+)\}",
        css,
    )
    assert glow_cleanup is not None
    assert ".top-nav .top-nav-item.active" in glow_cleanup.group("sel")
    assert "box-shadow: none !important;" in glow_cleanup.group("body")
    assert "filter: none !important;" in glow_cleanup.group("body")
    assert "top-nav-glow-off-20260810t" in css
    assert "-webkit-box-shadow: none !important;" in css
    assert "Clip Edge antialias fringe around the rounded gradient pill." in css
    assert "background-image: var(--theme-accent-gradient) !important;" in css
    assert "transition: color var(--transition-fast), background-color var(--transition-fast);" in css
    assert "transition: all var(--transition-fast);" not in re.search(
        r"(?m)^\.top-nav-item\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    ).group("body")

    space_body = re.search(r'\[data-theme="space-tech"\] body\s*\{(?P<body>.*?)\}', css, re.S)
    assert space_body is not None
    assert "flex-direction: column" in space_body.group("body")
    assert "height: 100vh" in space_body.group("body")
    assert "overflow: hidden" in space_body.group("body")

    main_content = re.search(
        r"\[data-theme=\"space-tech\"\] \.main-content\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert main_content is not None
    main_body = main_content.group("body")
    assert "flex: 1 1 auto" in main_body
    assert "margin: 0" in main_body
    assert "padding: 12px 32px 32px" in main_body
    assert "border-radius: 0" in main_body
    assert "background: transparent" in main_body
    assert "background: var(--surface-container-lowest)" not in main_body
    assert "height: calc(100vh - 12px)" not in main_body
    assert "margin: 12px 32px 0" not in main_body
    assert "padding: 64px 0 32px" not in main_body


def test_history_list_shows_loading_animation_while_fetching():
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert "function renderHistoryLoading()" in app_js
    load_history = re.search(r"async function loadHistoryList\(resetPage = false\) \{(?P<body>.*?)\n\}", app_js, re.S)
    assert load_history is not None
    assert "renderHistoryLoading();" in load_history.group("body")
    assert load_history.group("body").index("renderHistoryLoading();") < load_history.group("body").index('api("/api/history")')
    assert "getReconcileBusinessSourceName()" not in load_history.group("body")
    assert "filterRunsByReconcileBusinessSource" not in load_history.group("body")
    assert "historyRuns = payload.history || [];" in load_history.group("body")
    assert 'class="history-loading-row"' in app_js
    assert 'colspan="${historyColumnCount()}"' in app_js
    assert 'class="loading-spinner history-loading-spinner"' in app_js
    assert "加载核对历史..." in app_js

    assert ".history-loading-row td" in css
    assert ".history-loading-spinner" in css


def test_run_and_pbc_import_conflict_feedback_is_visible():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)

    assert "对数任务正在执行，请等待当前任务完成后再开始。" in app_js
    assert "showToast(message, \"warning\")" in app_js
    assert "您的执行失败，原因：有正在执行的任务" in app_js
    assert "用户正在执行中" in app_js
    assert "用户执行完成，您可再次执行。" in app_js
    assert "pollActiveRunConflict" in app_js
    assert "error.payload = p" in app_js
    assert "handlePbcImportStartError" in app_js
    assert "待插入表正在导入，请等待上一个任务完成后再导入。" in app_js
    assert 'id="pbcProgressTitle"' in html
    assert 'id="pbcRetryBtn"' in html
    assert "稍后再试" in app_js
    assert "pbcRetryBtn.hidden = false" in app_js
    assert "pbcRetryBtn?.addEventListener" in app_js


def test_db_validation_frontend_tool_settings_and_api_are_wired():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    for item_id in [
        'id="toolCardDbValidation"',
        'id="dbValidationModalOverlay"',
        'id="dbValidationReportDate"',
        'id="dbValidationPublicInfoCheck"',
        'id="dbValidationTableList"',
        'id="dbValidationRulesDocBtn"',
        'id="dbValidationDetailSource"',
        'id="dbValidationDetailSysManageId"',
        'id="dbValidationDetailClassificationId"',
        'id="dbValidationPublicInfoSource"',
        'id="dbValidationPublicInfoSysManageId"',
        'id="dbValidationPublicInfoClassificationId"',
        'id="dbValidationTemplateSource"',
        'id="dbValidationTemplateSysManageId"',
        'id="dbValidationTemplateClassificationId"',
        'id="dbValidationMetadataSource"',
        'id="dbValidationBaseinfoTable"',
        'id="dbValidationFieldInfoTable"',
        'id="dbValidationRefreshFieldMappingBtn"',
    ]:
        assert item_id in html

    assert 'id="dbValidationDataSource"' not in html
    assert 'id="dbValidationPublicInfoTable"' not in html
    assert 'class="db-validation-date-field"' in html
    assert 'id="dbValidationTemplateCheck"' in html
    assert 'id="dbValidationTemplateCheck" disabled' not in html

    for text in [
        "人行逐笔校验引擎",
        "/api/tools/db-validation/settings",
        "/api/tools/db-validation/start",
        "/api/tools/db-validation/status/",
        "/api/tools/db-validation/download/",
        "/api/tools/db-validation/history",
        "/api/tools/db-validation/history/download/",
        "/api/tools/db-validation/rules-document",
        "/api/tools/db-validation/field-mapping/refresh",
        "function loadDbValidationSettings",
        "function startDbValidation",
        "function pollDbValidationJob",
        "function openDbValidationHistory",
        "function renderDbValidationHistory",
        "function saveDbValidationSettings",
        "function refreshDbValidationFieldMapping",
        "function renderDbValidationFieldMappingStatus",
        "function readDbValidationDatasetSettings",
        "enable_public_info_check",
        "enable_template_check",
        "field_mapping_source_id",
        "unmapped_field_count",
    ]:
        assert text in app_js

    save_start = app_js.index("async function saveDbValidationSettings")
    save_end = app_js.index("async function refreshDbValidationFieldMapping", save_start)
    save_body = app_js[save_start:save_end]
    assert "const refreshMapping = options.refreshMapping !== false;" in save_body
    assert 'api("/api/tools/db-validation/field-mapping/refresh", { method: "POST" })' in save_body
    assert "已保存，正在刷新字段映射" in save_body
    assert "数据库校验配置已保存，字段映射已刷新" in save_body

    refresh_start = app_js.index("async function refreshDbValidationFieldMapping")
    refresh_end = app_js.index("function appendDbValidationLog", refresh_start)
    refresh_body = app_js[refresh_start:refresh_end]
    assert "saveDbValidationSettings(" not in refresh_body
    assert 'api("/api/tools/db-validation/field-mapping/refresh", { method: "POST" })' in refresh_body

    assert ".tool-card-db-validation" in css
    assert ".db-validation-grid" in css
    assert ".db-validation-table-list" in css
    assert 'id="dbValidationHistoryBtn"' in html
    assert 'id="dbValidationHistoryBody"' in html
    assert "<th>执行人</th>" in html
    assert "db-validation-history-count-link" in app_js
    assert "function dbValidationHistoryExecutorName" in app_js
    assert "dbValidationHistoryExecutorName(run)" in app_js
    assert "function formatDbValidationHistoryTime" in app_js
    assert '.replace("T", " ")' in app_js
    assert ".db-validation-history-modal" in css
    history_wrap = re.search(r"(?m)^\.db-validation-history-table-wrap\s*\{(?P<body>.*?)\}", css, re.S)
    assert history_wrap is not None
    assert "overflow-x: auto;" in history_wrap.group("body")
    assert ".db-validation-history-table {\n  table-layout: fixed;" in css
    assert ".db-validation-history-count-link" in css


def test_packaged_exe_includes_db_validation_resource_package():
    spec = _read(PYINSTALLER_SPEC)

    assert "src/auto_check/resources" in spec
    assert "auto_check/resources" in spec
    assert "'auto_check.resources'" in spec
    assert "'auto_check.resources.data'" in spec


def test_db_validation_history_sorts_by_execution_time_desc():
    app_js = _read(APP_JS)

    load_history = re.search(r"async function loadDbValidationHistory\(\) \{(?P<body>.*?)\n\}", app_js, re.S)
    assert load_history is not None
    assert "const sortedHistory = [...(payload.history || [])].sort(compareDbValidationHistoryRunsDesc);" in load_history.group("body")
    assert "renderDbValidationHistory(sortedHistory);" in load_history.group("body")
    assert "function compareDbValidationHistoryRunsDesc" in app_js
    assert "function dbValidationHistoryExecutionTimeValue" in app_js
    assert "const raw = dbValidationHistoryExecutionTime(run);" in app_js
    assert "Date.UTC(" in app_js
    assert "dbValidationHistoryExecutionTimeValue(right) - dbValidationHistoryExecutionTimeValue(left)" in app_js


def test_selects_use_scheme_5_glass_style_without_particles():
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert "custom-input-particle" not in app_js
    assert "custom-select-particle" not in app_js
    assert "custom-input-particle" not in css
    assert "custom-select-particle" not in css
    assert "float-particle" not in css

    for text in [
        "function initializeCustomSelects()",
        "function enhanceCustomSelect(select)",
        "function enhanceCustomInput(input)",
        "function shouldEnhanceCustomInput(input)",
        "function renderCustomDatePicker(input)",
        "function openCustomDatePicker(input)",
        "function toggleCustomDatePicker(input)",
        "function enhanceCustomDateInput(input, shell)",
        "function cleanupDetachedCustomDatePickers()",
        "function cleanupDetachedCustomSelects()",
        "cleanupDetachedCustomDatePickers();",
        "cleanupDetachedCustomSelects();",
        "dropdown._customControlOwner = input;",
        "dropdown._customControlOwner = select;",
        'dropdown.classList.add("rsp-compact-select-dropdown");',
        "const CUSTOM_INPUT_TYPES = new Set",
        "const customDateStates = new WeakMap();",
        "const CUSTOM_DATE_WEEKDAYS",
        "const customSelectStates = new WeakMap();",
        "custom-select-shell",
        "custom-select-trigger",
        "custom-select-dropdown",
        "custom-select-option",
        "custom-input-shell",
        "custom-input-native",
        "custom-date-shell",
        "custom-date-dropdown",
        "custom-date-day",
        "customSelectMeasure(select, shell)",
        "customInputMeasure(input, shell)",
        "shell.style.setProperty(\"--select-width\"",
        "shell.style.setProperty(\"--select-height\"",
        "shell.style.setProperty(\"--input-width\"",
        "shell.style.setProperty(\"--input-height\"",
        "const CUSTOM_INPUT_TYPES = new Set([\"text\", \"search\", \"number\", \"date\"",
        "if (type === \"date\") shell.classList.add(\"custom-date-shell\");",
        "input.type = \"text\";",
        "input.readOnly = true;",
        "input.classList.add(\"custom-date-input\");",
        "event.stopPropagation();",
        "positionCustomDateDropdown(input);",
        'input.addEventListener("click", () => toggleCustomDatePicker(input));',
        "setCustomDateValue(input, day.dataset.date || \"\")",
        "CUSTOM_INPUT_TYPES.has(type) && !input.hidden",
        "select.dispatchEvent(new Event(\"change\", { bubbles: true }))",
            "target.closest(\".custom-select-dropdown\")",
            "const dropdownWidth = compactRoleSelect",
            "enhanceCustomControls();",
            "initializeCustomSelects();",
        ]:
        assert text in app_js
    click_close = re.search(
        r'document\.addEventListener\("click", \(event\) => \{(?P<body>.*?)\n  \}, true\);',
        app_js,
        re.S,
    )
    assert click_close is not None
    assert "closeOtherCustomSelects()" in click_close.group("body")
    assert "closeOtherCustomDatePickers()" in click_close.group("body")
    assert 'input.addEventListener("focus", () => openCustomDatePicker(input));' not in app_js

    select_rule = re.search(r"(?m)^select\s*\{(?P<body>.*?)\}", css, re.S)
    assert select_rule is not None
    select_body = select_rule.group("body")
    for text in [
        "padding-right: 40px",
        "border-radius: 8px",
        "background-color: var(--surface-container-lowest)",
        "background-image: none",
        "caret-color: var(--theme-accent-readable)",
        "backdrop",
        "-webkit-appearance: none",
        "appearance: none",
    ]:
        assert text in select_body

    assert "select:hover" in css
    assert "select:focus" in css
    assert "select option" in css
    assert "select option:checked" in css
    assert '[data-color-mode="dark"] select' in css
    assert '[data-color-mode="dark"] select option:checked' in css
    assert "color-mix(in srgb, var(--theme-accent) 30%, var(--outline-variant))" in css

    for selector in [
        ".custom-select-shell",
        ".custom-select-native",
        ".custom-select-trigger",
        ".custom-select-trigger::after",
        ".custom-input-shell",
        "input.custom-input-native",
        ".custom-date-shell",
        ".date-picker .custom-date-shell",
        ".db-validation-date-field",
        ".db-validation-date-field .custom-date-shell",
        ".user-form-group .custom-input-shell.user-form-control",
        "input.custom-date-input",
        ".custom-date-shell::after",
        ".custom-date-shell::before",
        ".custom-date-dropdown",
        ".custom-date-head",
        ".custom-date-weekdays",
        ".custom-date-days",
        ".custom-date-day.active",
        ".custom-date-actions",
        ".custom-input-shell:focus-within input.custom-input-native",
        ".custom-select-dropdown",
        ".custom-select-option",
        ".custom-select-option::before",
        ".custom-select-option.active::after",
        '[data-color-mode="dark"] input.custom-input-native',
        '[data-color-mode="dark"] input.custom-date-input',
        '[data-color-mode="dark"] .custom-date-dropdown',
        '[data-color-mode="dark"] .custom-date-shell::after',
        ':root:not([data-theme="space-tech"]) input.custom-input-native',
        ':root:not([data-theme="space-tech"]) .custom-select-trigger',
        ':root:not([data-theme="space-tech"]) input.custom-date-input',
        ':root:not([data-theme="space-tech"]) .custom-date-day.active',
        '[data-color-mode="dark"] .custom-select-trigger',
        '[data-color-mode="dark"] .custom-select-dropdown',
        '[data-color-mode="dark"] .custom-select-option.active',
    ]:
        assert selector in css

    for text in [
        "width: var(--select-width)",
        "height: var(--select-height)",
        "width: var(--input-width)",
        "height: var(--input-height)",
        "flex: 0 0 180px",
        "width: 180px",
        "height: 38px",
        "overflow: hidden",
        "overscroll-behavior: contain",
        "scrollbar-gutter: stable",
        "background: var(--surface-container-lowest)",
        "color: var(--on-surface)",
        "caret-color: var(--theme-accent-readable)",
        "backdrop-filter: blur(10px)",
        "border-color: var(--theme-accent-readable)",
        "box-shadow: 0 0 0 3px var(--theme-focus-ring)",
        "grid-template-columns: repeat(7, 1fr)",
        "background: var(--theme-accent)",
        "color: var(--theme-on-accent)",
        "animation: dropdown-slide 0.3s ease-out",
        "padding-left: 24px",
        "content: \"✓\"",
        "@keyframes check-bounce",
    ]:
        assert text in css


def test_settings_uses_single_data_source_model():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)

    for item_id in [
        'id="mdSourceType"',
        'id="mdSourceHost"',
        'id="mdSourceDb"',
        'id="mdSourceSchemaField"',
        'id="mdSourceSchema"',
        'id="mdSourceUser"',
        'id="mdSourcePwd"',
    ]:
        assert item_id in html

    for removed_id in [
        'id="mdDwsType"',
        'id="mdBizType"',
        'id="dwsToggle"',
        'id="bizToggle"',
        'id="modalSetDefault"',
        'id="reconcileDwsSource"',
        'id="reconcileBusinessSource"',
        'id="saveReconcileSourcesBtn"',
        'id="reconcileSourcesStatus"',
    ]:
        assert removed_id not in html

    assert "function loadReconcileDataSourceSettings" not in app_js
    assert "function renderReconcileDataSourceSettings" not in app_js
    assert '"/api/settings/reconcile-data-sources", {' not in app_js
    assert "function syncDataSourceSchemaVisibility(prefix)" in app_js
    assert "function defaultPortForDataSourceType(type)" in app_js
    assert 'return String(type || "").toLowerCase() === "mysql" ? 3306 : 5432;' in app_js
    assert "function syncDataSourcePortForType(prefix, options = {})" in app_js
    assert 'document.getElementById("mdSourceType")?.addEventListener("change"' in app_js
    assert 'syncDataSourcePortForType("mdSource", { force: true });' in app_js
    assert 'schema: document.getElementById(prefix + "Type").value === "postgresql"' in app_js
    assert "source_id" in app_js
    assert "field_mapping_source_id" in app_js
    assert "set-def" not in app_js
    assert "设为默认" not in app_js
    assert "function parseDbValidationSource" not in app_js
    assert "field_mapping_config_name: selected.configName" not in app_js
    assert 'return `${item.config_name || ""}::${item.source || "dws"}`;' not in app_js


def test_user_name_stack_places_account_under_display_name():
    css = _read(STYLES_CSS)

    user_name_stack = re.search(r"(?m)^\.user-name-stack\s*\{(?P<body>.*?)\}", css, re.S)
    assert user_name_stack is not None
    assert "flex-direction: column" in user_name_stack.group("body")
    assert "align-items: flex-start" in user_name_stack.group("body")
    user_name_line = re.search(r"(?m)^\.user-name-line\s*\{(?P<body>.*?)\}", css, re.S)
    assert user_name_line is not None
    assert "display: inline-flex" in user_name_line.group("body")
    assert "align-items: center" in user_name_line.group("body")


def test_user_modal_explains_username_supported_characters():
    app_js = _read(APP_JS)

    assert "function userFriendlyError(message = \"\")" in app_js
    assert "username contains unsupported characters" in app_js
    assert "用户名仅支持英文字母、数字、下划线(_)、中横线(-)和点(.)" in app_js
    assert "不支持中文、空格及其他特殊字符" in app_js
    assert "userModalStatus.textContent = userFriendlyError(error.message);" in app_js


def test_history_action_column_keeps_table_cell_alignment():
    css = _read(STYLES_CSS)

    actions = re.search(r"(?m)^\.history-actions\s*\{(?P<body>.*?)\}", css, re.S)
    assert actions is not None
    assert "display: flex" not in actions.group("body")
    assert "text-align: center" in actions.group("body")
    assert "vertical-align: middle" in actions.group("body")
    assert ".history-actions .btn-xs" in css


def test_system_info_shows_runtime_status_and_history_count():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)

    start = app_js.index("async function loadSystemInfo()")
    end = app_js.index("function setSystemInfoFeedback", start)
    body = app_js[start:end]
    assert 'id="historyRunCount"' in html
    assert 'id="sessionStatusInfo"' not in html
    assert 'id="loginUserInfo"' in html
    assert 'id="autoRefreshInfo"' in html
    assert 'id="testAllConnBtn"' not in html
    assert 'id="dwsStatus"' not in html
    assert 'id="bizStatus"' not in html
    assert 'id="historyCount"' not in html
    assert 'api("/api/system-info")' in body
    assert 'api("/api/history")' not in body
    assert "historyRunCount" in body
    assert "authState.authenticated" not in body
    assert "userDisplayName(authState.user || {})" in body
    assert "settings.autoRefreshHome" in body
    assert 'api("/api/connection-status")' not in body
    assert "async function testConnectionStatusForFeedback()" not in app_js
    assert 'api("/api/connection-status")' not in app_js
    assert "仅管理员可测试" not in body
    assert 'if (authState.user?.role !== "admin")' not in body


def test_auth_passwords_are_encrypted_before_transport():
    login_html = _read(ROOT / "src" / "auto_check" / "web" / "login.html")
    index_html = _read(INDEX_HTML)
    app_js = _read(APP_JS)

    assert '"/api/auth/key"' in login_html
    assert "async function encryptPasswordForTransport" in login_html
    assert '<script src="/crypto_fallback.js"></script>' in login_html
    assert '<script src="/crypto_fallback.js"></script>' in index_html
    assert "window.autoCheckCrypto.encryptPasswordForTransport" in login_html
    assert "window.autoCheckCrypto.encryptPasswordForTransport" in app_js
    assert "password_encrypted" in login_html
    assert "body: JSON.stringify({ username, password_encrypted: encryptedPassword })" in login_html
    assert "body: JSON.stringify({ password_encrypted: encryptedPassword })" in login_html
    assert "body: JSON.stringify({ password })" not in login_html
    assert '"/api/auth/key"' in app_js
    assert "async function encryptPasswordForTransport" in app_js
    assert "password_encrypted" in app_js
    assert "password: userPassword.value" not in app_js


def test_data_source_passwords_are_encrypted_before_transport():
    app_js = _read(APP_JS)

    assert "async function encryptDataSourcePasswordsForTransport" in app_js
    assert "password_encrypted" in app_js
    assert "delete payload.password;" in app_js
    assert 'password: document.getElementById(prefix + "Pwd").value' not in app_js
    assert "body: JSON.stringify(await encryptDataSourcePasswordsForTransport(cfg))" in app_js
    assert "body: JSON.stringify(await encryptDataSourcePasswordsForTransport(body))" in app_js


def test_data_source_test_connection_modal_resets_pending_state():
    app_js = _read(APP_JS)

    assert "let modalTestRequestToken = 0;" in app_js
    assert "function resetModalTestConnectionState()" in app_js
    assert "if (modalTestBtn) modalTestBtn.disabled = false;" in app_js
    assert "function closeConfigModal()" in app_js
    assert "modalClose.addEventListener(\"click\", closeConfigModal);" in app_js

    open_start = app_js.index("function openModal(config)")
    open_end = app_js.index("function fillDs", open_start)
    open_body = app_js[open_start:open_end]
    assert "resetModalTestConnectionState();" in open_body

    test_start = app_js.index("modalTestBtn.addEventListener")
    test_end = app_js.index("modalSaveBtn.addEventListener", test_start)
    test_body = app_js[test_start:test_end]
    assert "const requestToken = ++modalTestRequestToken;" in test_body
    assert "if (requestToken !== modalTestRequestToken || configModal.hidden) return;" in test_body
    assert "if (requestToken === modalTestRequestToken && !configModal.hidden) modalStatus.textContent = e.message;" in test_body
    assert "if (requestToken === modalTestRequestToken && !configModal.hidden) modalTestBtn.disabled = false;" in test_body

    save_start = app_js.index("modalSaveBtn.addEventListener")
    save_end = app_js.index("/* ===== Tools: PBC full product import", save_start)
    save_body = app_js[save_start:save_end]
    assert "closeConfigModal();" in save_body


def test_login_error_messages_are_mapped_for_common_failures():
    login_html = _read(ROOT / "src" / "auto_check" / "web" / "login.html")

    assert "function getLoginErrorMessage" in login_html
    assert "账号或密码不正确" in login_html
    assert "账号已停用" in login_html
    assert "密码传输加密失败" in login_html
    assert "网络连接异常" in login_html


def test_regular_user_settings_are_limited_and_readonly_for_system_actions():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    for card in ["card-default", "card-db-validation", "card-datasource", "card-business"]:
        assert f'class="card settings-dashboard-card {card} admin-only"' in html
    assert 'class="card settings-dashboard-card card-data admin-only"' not in html
    assert 'class="card settings-dashboard-card card-interface"' in html
    assert 'class="card settings-dashboard-card card-interface admin-only"' not in html
    assert 'id="testAllConnBtn"' not in html
    assert 'id="refreshInfoBtn" type="button" class="btn-outline btn-sm admin-action"' in html
    assert "function applySettingsRoleAccess" in app_js
    assert 'document.querySelectorAll(".admin-action")' in app_js
    assert "[data-role=\"user\"] .admin-only" in css


def test_pbc_completed_step_and_importing_text_are_centered_and_green():
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert "const isDone = s < pbcCurrentStep || (pbcCurrentStep === 4 && s === 4);" in app_js
    assert 'el.classList.toggle("pbc-step--done", isDone);' in app_js

    progress_header = re.search(r"(?m)^\.pbc-progress-header\s*\{(?P<body>.*?)\}", css, re.S)
    assert progress_header is not None
    progress_header_body = progress_header.group("body")
    assert "justify-content: center" in progress_header_body
    assert "text-align: center" in progress_header_body

    assert "[data-color-mode=\"dark\"] .pbc-step--done .pbc-step-num" in css
    assert "[data-theme=\"space-tech\"][data-color-mode=\"dark\"] .pbc-step--done .pbc-step-num" in css


def test_pbc_import_dark_mode_and_mapping_layout_are_readable():
    app_js = _read(APP_JS)
    html = _read(INDEX_HTML)
    css = _read(STYLES_CSS)

    assert 'id="pbcColumnNotice"' in html
    assert "let pbcTableColumns = [];" in app_js
    assert "pbcTableColumns = payload.table_columns || [];" in app_js
    assert "const targetOptions = pbcTableColumns.map((target) =>" in app_js
    assert "pbcTableColumns.find((column) => column.name === target)" in app_js
    assert "function renderPbcColumnNotice()" in app_js
    assert "function hidePbcColumnNotice()" in app_js
    assert "function syncPbcUploadAggregate()" in app_js
    assert "function getPbcUploadIds()" in app_js
    assert "upload_ids: getPbcUploadIds()" in app_js
    assert "payload.upload_inspections" in app_js
    assert "pbcUploadedFiles = pbcUploadedFiles.map" in app_js
    assert "missingByFile" in app_js
    assert "renderPbcColumnNotice();" in app_js
    assert "pbcColumnNotice.title = fullDetails;" in app_js
    assert "hidePbcColumnNotice();" in re.search(r"pbcLoadMappingsBtn\?\.addEventListener\(\"click\", async \(\) => \{(?P<body>.*?)\n\}\);", app_js, re.S).group("body")
    assert ".pbc-column-notice" in css

    mapping_config = re.search(r"(?m)^\.pbc-mapping-config\s*\{(?P<body>.*?)\}", css, re.S)
    assert mapping_config is not None
    assert "minmax(300px, 0.42fr) minmax(430px, 0.58fr)" in mapping_config.group("body")
    assert "gap: 16px" in mapping_config.group("body")
    assert ".pbc-mapping-config-left,\n.pbc-mapping-config-right" in css

    mapping_item = re.search(r"(?m)^\.pbc-mapping-item\s*\{(?P<body>.*?)\}", css, re.S)
    assert mapping_item is not None
    assert "minmax(108px" in mapping_item.group("body")
    assert "minmax(0, 1.42fr)" in mapping_item.group("body")
    assert "22px" in mapping_item.group("body")
    assert "gap: 4px" in mapping_item.group("body")

    mapping_action = re.search(r"(?m)^\.pbc-mapping-action\s*\{(?P<body>.*?)\}", css, re.S)
    assert mapping_action is not None
    assert "width: 22px" in mapping_action.group("body")
    assert "height: 22px" in mapping_action.group("body")

    assert '[data-color-mode="dark"] .pbc-step-connector' in css
    assert '[data-theme="space-tech"][data-color-mode="dark"] .pbc-step-connector' in css
    assert '[data-color-mode="dark"] .pbc-import-log' in css
    assert '[data-color-mode="dark"] .pbc-log-entry' in css
    for variant in ["primary", "secondary", "success"]:
        assert f'[data-color-mode="dark"] .pbc-btn--{variant} {{' in css
        variant_rule = re.search(rf'\[data-color-mode="dark"\] \.pbc-btn--{variant}\s*\{{(?P<body>.*?)\}}', css, re.S)
        assert variant_rule is not None
        assert "background:" in variant_rule.group("body")
        assert "border:" in variant_rule.group("body")
        assert f'[data-color-mode="dark"] .pbc-btn--{variant}:hover' in css
    assert '[data-color-mode="dark"] .pbc-btn--primary:disabled' in css
    disabled_primary = re.search(r'\[data-color-mode="dark"\] \.pbc-btn--primary:disabled,\s*\n\[data-color-mode="dark"\] \.pbc-btn--primary:disabled:hover\s*\{(?P<body>.*?)\}', css, re.S)
    assert disabled_primary is not None
    assert "background:" in disabled_primary.group("body")
    assert "opacity: 1" in disabled_primary.group("body")
    assert '[data-color-mode="dark"] .pbc-btn:disabled' in css


def test_pbc_auto_mapping_remove_can_be_restored_without_affecting_manual_rows():
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert "function normalizePbcAutoMappings" in app_js
    assert "auto_target_column" in app_js
    assert "manual_unmapped_from_auto" in app_js
    assert "const canRestore = !mapping.target_column && mapping.auto_target_column && mapping.manual_unmapped_from_auto;" in app_js
    assert 'data-action="${canRestore ? "restore" : "remove"}"' in app_js
    assert 'title="${canRestore ? "还原自动映射" : "移除列"}"' in app_js
    assert "restorePbcAutoMapping(index);" in app_js
    assert "removePbcMapping(index);" in app_js
    assert "pbcMappings = normalizePbcAutoMappings(payload.mappings || []);" in app_js
    assert ".pbc-mapping-restore" in css
    assert ".pbc-mapping-action" in css


def test_pbc_mapping_modal_avoids_horizontal_clipping():
    css = _read(STYLES_CSS)

    modal = re.search(r"(?m)^\.pbc-modal\s*\{(?P<body>.*?)\}", css, re.S)
    assert modal is not None
    modal_body = modal.group("body")
    assert "width: 860px" in modal_body
    assert "max-width: 96vw" in modal_body

    mapping_config = re.search(r"(?m)^\.pbc-mapping-config\s*\{(?P<body>.*?)\}", css, re.S)
    assert mapping_config is not None
    mapping_config_body = mapping_config.group("body")
    assert "minmax(300px, 0.42fr) minmax(430px, 0.58fr)" in mapping_config_body

    mapping_list = re.search(r"(?m)^\.pbc-mapping-list\s*\{(?P<body>.*?)\}", css, re.S)
    assert mapping_list is not None
    assert "overflow-x: hidden" in mapping_list.group("body")

    mapping_item = re.search(r"(?m)^\.pbc-mapping-item\s*\{(?P<body>.*?)\}", css, re.S)
    assert mapping_item is not None
    mapping_item_body = mapping_item.group("body")
    assert "minmax(108px, 0.58fr) 10px minmax(0, 1.42fr) 22px" in mapping_item_body
    assert "min-width: 0" in mapping_item_body


def test_space_tech_top_nav_aligns_with_content_padding():
    css = _read(STYLES_CSS)
    app_js = _read(APP_JS)

    top_nav = re.search(r"(?m)^\.top-nav\s*\{(?P<body>.*?)\}", css, re.S)
    assert top_nav is not None
    # Base rule may still describe the legacy floating shell; space-tech overrides to edge-stuck.
    assert "position: fixed" in top_nav.group("body")

    space_top_nav = re.search(
        r'\[data-theme="space-tech"\] \.top-nav\s*\{(?P<body>.*?)\}',
        css,
        re.S,
    )
    assert space_top_nav is not None
    assert "background: var(--surface-container-lowest);" in space_top_nav.group("body")
    assert "backdrop-filter: none;" in space_top_nav.group("body")
    assert "-webkit-backdrop-filter: none;" in space_top_nav.group("body")
    assert "position: relative" in space_top_nav.group("body")
    assert "left: 0" in space_top_nav.group("body")
    assert "right: 0" in space_top_nav.group("body")
    assert "border-radius: 0 !important" in space_top_nav.group("body")

    assert "const topNav = document.querySelector(\".top-nav\");" in app_js
    assert "const mainContent = document.querySelector(\".main-content\");" in app_js
    assert "function updateSpaceTopNavFrost()" in app_js
    # Edge-stuck opaque nav clears frost instead of toggling it on scroll.
    assert 'document.documentElement.classList.remove("space-nav-over-content");' in app_js
    assert 'document.documentElement.classList.toggle("space-nav-over-content", shouldFrost);' not in app_js
    assert "scrollOffset > 1" not in app_js
    assert 'window.addEventListener("scroll", updateSpaceTopNavFrost, { passive: true });' in app_js
    assert 'mainContent?.addEventListener("scroll", updateSpaceTopNavFrost, { passive: true });' in app_js
    assert "function handleMainContentScroll()" not in app_js
    assert "function revealMainContentScrollbar()" not in app_js
    assert "mainContentScrollbarHideTimer" not in app_js

    frosted_nav = re.search(
        r"\[data-theme=\"space-tech\"\]\.space-nav-over-content \.top-nav\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert frosted_nav is not None
    frosted_nav_body = frosted_nav.group("body")
    assert "background: var(--surface-container-lowest);" in frosted_nav_body
    assert "border-bottom-color: var(--outline-variant);" in frosted_nav_body
    assert "linear-gradient" not in frosted_nav_body
    assert "backdrop-filter: none;" in frosted_nav_body
    assert "-webkit-backdrop-filter: none;" in frosted_nav_body
    assert '[data-theme="space-tech"].space-nav-over-content .top-nav::before' not in css
    assert '[data-theme="space-tech"].space-nav-over-content .top-nav::after' not in css
    assert '[data-theme="space-tech"].space-nav-over-content body::before' not in css

    assert '[data-theme="space-tech"]::before' not in css
    assert '[data-theme="space-tech"].space-nav-over-content body::after' not in css

    main_content = re.search(
        r"\[data-theme=\"space-tech\"\] \.main-content\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert main_content is not None
    main_content_body = main_content.group("body")
    assert "flex: 1 1 auto" in main_content_body
    assert "margin: 0" in main_content_body
    assert "padding: 12px 32px 32px" in main_content_body
    assert "border-radius: 0" in main_content_body
    assert "background: transparent" in main_content_body
    assert "background: var(--surface-container-lowest)" not in main_content_body
    assert "overflow-x: hidden" in main_content_body
    assert "overflow-y: auto" in main_content_body
    assert "--space-scrollbar-top-offset" not in main_content_body
    assert "scrollbar-width: none" in main_content_body
    assert "height: calc(100vh - 12px)" not in main_content_body
    assert "margin: 12px 32px 0" not in main_content_body
    assert "padding: 64px 0 32px" not in main_content_body

    hidden_scrollbar = re.search(
        r"\[data-theme=\"space-tech\"\] \.main-content::-webkit-scrollbar\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert hidden_scrollbar is not None
    assert "display: none" in hidden_scrollbar.group("body")
    assert "width: 0" in hidden_scrollbar.group("body")
    assert ".main-content.is-scrolling" not in css

    space_theme = re.search(
        r"\[data-theme=\"space-tech\"\]\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert space_theme is not None
    assert "--space-page-background:" not in space_theme.group("body")
    assert "--space-page-gutter-background" not in space_theme.group("body")
    assert "background: var(--theme-page-background)" in space_theme.group("body")
    assert "background-size: 100vw 100vh" in space_theme.group("body")
    assert "background-position: 0 0" in space_theme.group("body")
    assert "background-attachment: fixed" in space_theme.group("body")

    dark_space_theme = re.search(
        r"\[data-theme=\"space-tech\"\]\[data-color-mode=\"dark\"\]\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert dark_space_theme is not None
    assert "--space-page-gutter-background" not in dark_space_theme.group("body")

    space_body = re.search(
        r"\[data-theme=\"space-tech\"\] body\s*\{(?P<body>.*?)\}",
        css,
        re.S,
    )
    assert space_body is not None
    assert "background: var(--theme-page-background)" in space_body.group("body")

    mobile_top_nav = re.search(
        r"\[data-theme=\"space-tech\"\] \.top-nav\s*\{(?P<body>.*?)\}",
        css[css.index("@media (max-width: 900px)") :],
        re.S,
    )
    assert mobile_top_nav is not None
    assert "left: 0" in mobile_top_nav.group("body")
    assert "right: 0" in mobile_top_nav.group("body")
    assert "width: 100%" in mobile_top_nav.group("body")
    assert "border-radius: 0" in mobile_top_nav.group("body")

    mobile_main_content = re.search(
        r"\[data-theme=\"space-tech\"\] \.main-content\s*\{(?P<body>.*?)\}",
        css[css.index("@media (max-width: 900px)") :],
        re.S,
    )
    assert mobile_main_content is not None
    mobile_main_content_body = mobile_main_content.group("body")
    assert "height: auto" in mobile_main_content_body
    assert "margin: 0" in mobile_main_content_body
    assert "padding: 12px 14px 18px" in mobile_main_content_body
    assert "--space-scrollbar-top-offset" not in mobile_main_content_body
    assert "height: calc(100vh - 8px)" not in mobile_main_content_body
    assert "margin: 8px 14px 0" not in mobile_main_content_body
    assert "padding: 78px 0 18px" not in mobile_main_content_body

    compact_css = css[css.index("@media (max-width: 640px)") :]
    compact_top_nav = re.search(
        r"\[data-theme=\"space-tech\"\] \.top-nav\s*\{(?P<body>.*?)\}",
        compact_css,
        re.S,
    )
    assert compact_top_nav is not None
    assert "left: 0" in compact_top_nav.group("body")
    assert "right: 0" in compact_top_nav.group("body")
    assert "width: 100%" in compact_top_nav.group("body")
    assert "max-width: none" in compact_top_nav.group("body")
    assert "max-width: calc(100vw - 20px)" not in compact_top_nav.group("body")
    compact_tabs = re.search(
        r"\[data-theme=\"space-tech\"\] \.top-nav-tabs\s*\{(?P<body>.*?)\}",
        compact_css,
        re.S,
    )
    assert compact_tabs is not None
    assert "overflow-x: auto" in compact_tabs.group("body")
    compact_subtitle = re.search(
        r"\[data-theme=\"space-tech\"\] \.top-nav-wordmark \.brand-wordmark-sub\s*\{(?P<body>.*?)\}",
        compact_css,
        re.S,
    )
    assert compact_subtitle is not None
    assert "display: none" in compact_subtitle.group("body")


def test_flow_chain_ui_is_manual_only_and_uses_editor_modal_with_scrollable_list():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    server_py = _read(SERVER_PY)
    css = _read(STYLES_CSS)

    for element_id in [
        "flowChainEditorOverlay",
        "flowChainEditorName",
        "flowChainEditorEnabled",
        "flowDefinitionSearch",
        "flowDefinitionRefreshBtn",
        "flowDefinitionTable",
        "flowSelectedStepList",
        "flowChainEditorSave",
    ]:
        assert f'id="{element_id}"' in html

    for removed_element_id in [
        "flowChainEditorCron",
        "flowCronConfigBtn",
        "flowCronOverlay",
        "flowCronTabs",
        "flowCronPreview",
        "flowCronConfirm",
        "flowChainEditorScheduleEnabled",
        "flowChainEditorSteps",
    ]:
        assert f'id="{removed_element_id}"' not in html

    assert "flowChainEditorOverlay" in app_js
    assert "openFlowChainEditor" in app_js
    assert "/api/tools/flow/definitions" in app_js
    assert "loadFlowDefinitionsForEditor" in app_js
    assert "renderFlowDefinitionTable" in app_js
    assert "renderFlowSelectedSteps" in app_js
    assert "flowChainEditorSelectedSteps" in app_js
    assert "add-flow-definition" in app_js
    assert "move-step-up" in app_js
    assert "remove-selected-step" in app_js
    assert "/api/tools/flow/start" in app_js
    assert "/api/tools/flow/history" in app_js
    assert "schedule_cron" not in app_js
    assert "schedule_enabled" not in app_js
    assert "openFlowCronBuilder" not in app_js
    assert "renderFlowCronBuilder" not in app_js
    assert "parseFlowCronExpression" not in app_js
    assert "start_flow_scheduler" not in server_py
    assert "due_scheduled_chains" not in server_py
    assert 'trigger_type="scheduled"' not in server_py

    settings_list = re.search(r"(?m)^\.flow-chain-settings-list\s*\{(?P<body>.*?)\}", css, re.S)
    assert settings_list is not None
    assert "max-height:" in settings_list.group("body")
    assert "overflow-y: auto" in settings_list.group("body")

    editor_modal = re.search(r"(?m)^\.flow-chain-editor-modal\s*\{(?P<body>.*?)\}", css, re.S)
    assert editor_modal is not None
    assert "max-width:" in editor_modal.group("body")
    for selector in [
        ".flow-step-builder",
        ".flow-definition-table",
        ".flow-selected-step-list",
    ]:
        assert selector in css
    assert ".flow-cron-modal" not in css
    assert ".flow-cron-specific-grid" not in css


def test_flow_chain_editor_shows_only_flow_name_in_available_table():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)

    assert 'id="flowDefinitionTable"' in html
    assert 'id="flowManualFlowId"' in html
    assert 'id="addManualFlowBtn"' in html
    assert "搜索流程名称或 flow_id" in html
    assert "renderFlowDefinitionTable" in app_js
    assert "_renderFlowDefinitionTable" in app_js
    assert "renderFlowDefinitionLimitHint" in app_js
    assert "payload.truncated" in app_js
    assert "仅展示前 500 条" in app_js
    table_function = re.search(
        r"function _renderFlowDefinitionTable\(flows\) \{(?P<body>.*?)\n\}",
        app_js,
        re.S,
    )
    assert table_function is not None
    table_body = table_function.group("body")
    assert "<th>流程名称</th>" in table_body
    assert "<th>flow_id</th>" not in table_body
    assert "flow.name" in table_body
    assert "data-flow-id" in table_body
    assert "data-flow-name" in table_body


def test_flow_chain_editor_add_flow_uses_button_payload_fallback():
    app_js = _read(APP_JS)

    click_handler = re.search(
        r"flowDefinitionTable\?\.addEventListener\(\"click\", \(e\) => \{(?P<body>.*?)\n\}\);",
        app_js,
        re.S,
    )
    assert click_handler is not None
    click_body = click_handler.group("body")
    assert "addFlowDefinitionToSelected({" in click_body
    assert "flow_id: button.dataset.flowId" in click_body
    assert "name: button.dataset.flowName" in click_body
    assert "addFlowDefinitionToSelected(button.dataset.flowId" not in click_body

    add_function = re.search(
        r"function addFlowDefinitionToSelected\(flowInput = \{\}\) \{(?P<body>.*?)\n\}",
        app_js,
        re.S,
    )
    assert add_function is not None
    add_body = add_function.group("body")
    assert "const requestedFlow = normalizeFlowStep(flowInput);" in add_body
    assert "|| requestedFlow" in add_body
    assert "未找到该流程，请刷新流程列表后重试" in add_body
    assert "addManualFlowBtn?.addEventListener" in app_js
    assert "flowManualFlowId.value" in app_js


def test_flow_chain_editor_modal_fields_are_not_squeezed_by_global_modal_field_layout():
    css = _read(STYLES_CSS)

    editor_field = re.search(r"(?m)^\.flow-chain-editor-modal \.modal-field\s*\{(?P<body>.*?)\}", css, re.S)
    assert editor_field is not None
    editor_field_body = editor_field.group("body")
    assert "display: grid" in editor_field_body
    assert "align-items: stretch" in editor_field_body

    editor_field_label = re.search(r"(?m)^\.flow-chain-editor-modal \.modal-field span\s*\{(?P<body>.*?)\}", css, re.S)
    assert editor_field_label is not None
    assert "width: auto" in editor_field_label.group("body")

    editor_input = re.search(r"(?m)^\.flow-chain-editor-modal \.setting-input\s*\{(?P<body>.*?)\}", css, re.S)
    assert editor_input is not None
    editor_input_body = editor_input.group("body")
    assert "width: 100%" in editor_input_body
    assert "min-height: 38px" in editor_input_body

    editor_textarea = re.search(r"(?m)^\.flow-chain-editor-modal textarea\.setting-input\s*\{(?P<body>.*?)\}", css, re.S)
    assert editor_textarea is not None
    assert "min-height: 128px" in editor_textarea.group("body")


def test_flow_chain_editor_save_uses_single_function_set_and_visible_feedback():
    app_js = _read(APP_JS)

    for function_name in [
        "renderFlowChainSettings",
        "readFlowSettingsFromForm",
        "addFlowChainConfig",
    ]:
        assert app_js.count(f"function {function_name}") == 1

    save_function = re.search(
        r"function saveFlowChainFromEditor\(\) \{(?P<body>.*?)\n\}\n\ntoolCardFlow",
        app_js,
        re.S,
    )
    assert save_function is not None
    body = save_function.group("body")
    assert "setFlowChainEditorStatus(" in body
    assert "showToast(" in body
    assert "renderFlowChainSettings(chains)" in body
    assert "closeFlowChainEditor()" in body


def test_flow_chain_editor_blank_overlay_click_does_not_close_modal():
    app_js = _read(APP_JS)

    assert 'flowChainEditorOverlay?.addEventListener("click"' not in app_js


def test_flow_settings_source_select_uses_name_only_and_shows_execute_url_rule():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)

    source_select = re.search(r"function fillFlowSourceSelect\(select, dataSources, selected = \"\"\) \{(?P<body>.*?)\n\}\n\nasync function loadFlowSettings", app_js, re.S)
    assert source_select is not None
    source_select_body = source_select.group("body")
    assert "const label = item.name || value;" in source_select_body
    assert "item.db_type" not in source_select_body
    assert "item.database" not in source_select_body

    assert "系统会自动追加 ?id=flow_id" in html
    assert "validateFlowExecuteUrl" in app_js


def test_dark_mode_entry_points_are_removed_and_light_mode_is_forced():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert 'name="theme"' not in html
    assert html.count('data-theme-toggle-logo') == 0
    assert 'id="topDarkModeToggle"' not in html
    assert 'id="sidebarDarkModeToggle"' not in html
    assert "function applyDarkMode" in app_js
    apply_dark_mode = re.search(r"function applyDarkMode\([^)]*\) \{(?P<body>.*?)\n\}", app_js, re.S)
    assert apply_dark_mode is not None
    assert "const enabled = false;" in apply_dark_mode.group("body")
    assert 'document.documentElement.setAttribute("data-color-mode", "light")' in app_js
    assert ".top-nav-actions" in css


def test_flow_chain_background_toast_has_container_and_theme_styles():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    assert 'id="flowToastContainer"' in html
    assert "loadFlowToastStatus" in app_js
    assert "/api/flow-chain/status" in app_js
    assert "data-action=\"toggle-flow-toast\"" in app_js
    assert "data-action=\"close-flow-toast\"" in app_js
    assert "if (flowBgRunBtn) flowBgRunBtn.hidden = !flowCurrentJobId;" in app_js
    assert "流程任务正在提交" in app_js
    assert ".flow-toast-container" in css
    assert ".flow-toast.flow-toast--vitality.running .flow-toast-header" in css
    assert ".flow-toast.flow-toast--calm.running .flow-toast-header" in css
    assert "@keyframes flow-pulse-blue" in css
    assert "@keyframes flow-pulse-teal" in css
    assert "[data-color-mode=\"dark\"] .flow-toast" in css


def test_flow_modal_supports_background_progress_mode():
    app_js = _read(APP_JS)

    assert "showFlowModalProgressMode" in app_js
    assert "startFlowModalBackgroundPoll" in app_js
    assert "已提交，流程在后台运行中" in app_js
    assert "flowBgRunBtn" in app_js
    assert "后台运行" in app_js


def test_flow_cancel_uses_job_id_and_disables_button_while_stopping():
    app_js = _read(APP_JS)

    cancel_flow = re.search(r"async function cancelFlowChain\(\) \{(?P<body>.*?)\n\}", app_js, re.S)
    assert cancel_flow is not None
    body = cancel_flow.group("body")
    assert 'body: JSON.stringify({ job_id: flowCurrentJobId })' in body
    assert "if (flowCancelBtn?.disabled) return;" in body
    assert "flowCancelBtn.disabled = true" in body
    assert "停止中" in body


def test_provider_managed_report_card_retires_manual_entry_and_never_fakes_zero():
    app_js = _read(APP_JS)

    assert "function reportNavigationCardMaintainable(cardCode)" in app_js
    assert "maintenance.editable !== false" in app_js
    assert "const unavailable = card.source === \"provider\" && card.available === false;" in app_js
    assert "统计暂不可用" in app_js
    assert "result.provider_issues || []" in app_js
    assert "renderReportNavigationProviderIssues" in app_js


def test_db_validation_mapping_modal_has_three_filtered_fixed_views():
    html = _read(INDEX_HTML)
    app_js = _read(APP_JS)
    css = _read(STYLES_CSS)

    for button_id, label in (
        ("dbValidationViewTableMappingBtn", "查看表映射"),
        ("dbValidationViewFieldMappingBtn", "查看字段映射"),
        ("dbValidationViewCrossTableMappingBtn", "查看跨表映射"),
    ):
        assert f'id="{button_id}"' in html
        assert label in html
    assert '>刷新映射</button>' in html
    assert '刷新逐笔字段映射' not in html
    assert 'id="dbValidationMappingFilter"' in html
    assert 'id="dbValidationMappingFilterCount"' in html
    assert 'id="dbValidationMappingQuickFilters"' in html
    assert 'id="promptTemplateTableInput"' in html
    assert 'id="promptTemplateFieldInput"' in html
    assert "#crossTableMappingPromptModal" in css
    assert re.search(r"#crossTableMappingPromptModal\s*\{[^}]*z-index:\s*3000", css, re.S)
    assert "openDbValidationMappingView" in app_js
    assert "toggleDbValidationMappingView" not in app_js
    assert '["表类型", "逻辑表", "物理表", "修改"]' in app_js
    assert '["逻辑表", "中文名", "映射字段", "修改"]' in app_js
    assert '["逐笔表", "逐笔字段", "模板表", "模板字段", "修改"]' in app_js
    assert 'data-kind="cross_table"' in app_js
    assert '"请输入逐笔字段"' not in app_js
    assert "showCrossTableMappingPrompt" in app_js
    assert "dbValidationMappingSummaryText" in app_js
    assert "dbValidationMappingQuickFilterOptions" in app_js
    assert "dbValidationMappingMatchesQuickFilter" in app_js
    assert 'function dbValidationMappingIsUnmapped(item)' in app_js
    assert 'String(item.mapping_status || "mapped") !== "mapped"' in app_js
    assert 'const unmapped = rows.filter(dbValidationMappingIsUnmapped).length;' in app_js
    assert 'if (filter === "unmapped") return dbValidationMappingIsUnmapped(item);' in app_js
    assert "const totalFieldCount = fieldCount + requiredMissingCount + missingPhysicalCount;" in app_js
    assert "const totalUnmappedCount = unmappedCount + requiredMissingCount + missingPhysicalCount;" in app_js
    assert "async function loadDbValidationMappingPayload()" in app_js
    load_settings = re.search(
        r"async function loadDbValidationSettings\(\)\s*\{(?P<body>.*?)\n\}",
        app_js,
        re.S,
    )
    assert load_settings is not None
    assert "await loadDbValidationMappingPayload()" in load_settings.group("body")
    for filter_label in ("全部", "已映射", "未映射", "人工修改", "与自动映射不同"):
        assert filter_label in app_js
    assert '["required_missing", "必需缺失"]' in app_js
    assert '["semantic", "语义匹配"]' in app_js
    assert 'function dbValidationMappingIsSemantic(item)' in app_js
    assert 'if (filter === "semantic") return dbValidationMappingIsSemantic(item);' in app_js
    assert "mapping-semantic-badge" in app_js
    assert ".mapping-semantic-badge" in css
    assert 'if (view === "field")' in app_js
    assert 'await showPrompt("修改跨表映射", "请输入模板表"' not in app_js
    assert 'await showPrompt("修改跨表映射", "请输入模板字段"' not in app_js
    assert 'JSON.stringify({ template_table: templateTable, template_field: templateField })' in app_js
    assert '请输入人工修改原因' not in app_js
    assert '请输入恢复自动映射的原因' not in app_js
    assert '>恢复自动</button>' not in app_js
    assert '>恢复</button>' in app_js
    assert "mapping-has-difference" in app_js
    assert "当前条数" not in app_js  # 使用紧凑的 当前/总数 数字形式
    assert "overflow-x: hidden" in css
    assert "table-layout: fixed" in css
    assert "white-space: nowrap" in css
    assert "text-overflow: ellipsis" in css
    assert ".db-validation-mapping-entry.mapping-has-difference::after" in css
    filterbar = re.search(r"(?m)^\.db-validation-mapping-filterbar\s*\{(?P<body>.*?)\}", css, re.S)
    assert filterbar is not None
    assert "flex: 0 0 auto" in filterbar.group("body")
    assert "min-height: 40px" in filterbar.group("body")
    filter_input = re.search(r"(?m)^\.db-validation-mapping-filterbar input\s*\{(?P<body>.*?)\}", css, re.S)
    assert filter_input is not None
    assert "height: 40px" in filter_input.group("body")
    assert "min-height: 40px" in filter_input.group("body")
    assert "padding: 0 14px" in filter_input.group("body")
    filter_shell = re.search(r"(?m)^\.db-validation-mapping-filterbar \.custom-input-shell\s*\{(?P<body>.*?)\}", css, re.S)
    assert filter_shell is not None
    assert "height: 40px" in filter_shell.group("body")
    assert "min-height: 40px" in filter_shell.group("body")
    assert ".db-validation-mapping-table.mapping-view-cross_table th:nth-child(1) { width: 8%; }" in css
    assert ".db-validation-mapping-table.mapping-view-cross_table th:nth-child(2) { width: 18%; }" in css
    assert ".db-validation-mapping-table.mapping-view-cross_table th:nth-child(3) { width: 34%; }" in css
    assert ".db-validation-mapping-table.mapping-view-cross_table th:nth-child(4) { width: 22%; }" in css
    assert ".db-validation-mapping-table.mapping-view-cross_table th:nth-child(5) { width: 18%; }" in css
    assert ".db-validation-mapping-table.mapping-view-table th:nth-child(3) { width: 52%; }" in css
    assert ".db-validation-mapping-table tr.mapping-row-different td" in css
    assert "background: #fff8e8" in css
    assert "background: #f97316" in css
    assert "function dbValidationMappingDifferenceBadge(item, view, differenceField)" in app_js
    assert 'class="mapping-actions-cell"' not in app_js
    assert 'dbValidationMappingValueCell(item.effective_template_table_name, item, "cross_table", "template_table")' in app_js
    assert 'dbValidationMappingValueCell(item.effective_template_field_name, item, "cross_table", "template_field")' in app_js
    value_cell = re.search(
        r"function dbValidationMappingValueCell\(value, item, view, differenceField\)\s*\{(?P<body>.*?)\n\}",
        app_js,
        re.S,
    )
    assert value_cell is not None
    assert "dbValidationMappingDifferenceBadge" in value_cell.group("body")
    assert "mapping-value-wrap" in value_cell.group("body")
    assert ".mapping-value-wrap .mapping-difference-badge" in css
    assert "position: absolute" in css
    value_wrap = re.search(r"(?m)^\.mapping-value-wrap\s*\{(?P<body>.*?)\}", css, re.S)
    assert value_wrap is not None
    assert "display: inline-block" in value_wrap.group("body")
    assert "max-width: 100%" in value_wrap.group("body")
    assert not re.search(r"(?m)^\s*width:\s*100%;", value_wrap.group("body"))
