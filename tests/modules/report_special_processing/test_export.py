from datetime import date, datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from auto_check.app.module_system.contracts import ModuleManifest, ModuleRequest
from auto_check.app.module_system.permissions import default_permission_evaluator
from auto_check.app.module_system.routing import ModuleRouter
from auto_check.modules.report_special_processing.contracts import ValidationError
from auto_check.modules.report_special_processing.export_workbook import (
    EXPORT_HEADERS,
    build_export_xlsx,
    export_rows,
)


NOW = datetime(2026, 8, 6, 11, 20, tzinfo=ZoneInfo("Asia/Shanghai"))


def test_export_rows_follow_visible_ledger_columns():
    rows = export_rows(
        [
            {
                "report_period": date(2026, 7, 31),
                "report_process_name_snapshot": "人行大集中报送、1104报送",
                "dimension": "project",
                "summary": "摘要",
                "table_name": "t_demo",
                "field_name": "amt",
                "value_before": "1",
                "value_after": "2",
                "special_handling_at": NOW,
                "handler_display_name_snapshot": "管理员",
                "governance_owner_display_name_snapshot": "治理负责人",
                "status": "pending",
                "record_no": "RSP-should-not-export",
                "reports": ["报表A"],
                "processing_content": "说明",
                "processing_script": "select 1;",
            }
        ]
    )
    assert EXPORT_HEADERS == (
        "报送期", "处理表名", "修改字段", "修改前", "修改后", "所属业务系统", "关联报送", "处理摘要", "状态", "处理人", "处理时间",
    )
    assert rows == [[
        "2026-07-31", "t_demo", "amt", "1", "2", "—", "人行大集中报送、1104报送", "摘要", "待确认", "管理员", "2026-08-06 11:20:00",
    ]]
    assert "RSP-should-not-export" not in str(rows)
    assert "报表A" not in str(rows)
    assert "说明" not in str(rows)
    assert "select 1;" not in str(rows)


def test_build_export_xlsx_writes_ledger_headers_and_styles():
    payload = build_export_xlsx(
        [
            {
                "report_period": "2026-07-31",
                "report_process_name_snapshot": "1104报送",
                "dimension": "fund",
                "summary": "s",
                "table_name": "t_fund",
                "field_name": "bal",
                "value_before": "a",
                "value_after": "b",
                "special_handling_at": "2026-08-06T11:20:00+08:00",
                "handler_username_snapshot": "admin",
                "governance_owner_display_name_snapshot": "治理资金",
                "status": "completed",
            }
        ]
    )
    workbook = load_workbook(BytesIO(payload))
    sheet = workbook.active
    assert [cell.value for cell in sheet[1]] == list(EXPORT_HEADERS)
    assert sheet["A2"].value == "2026-07-31"
    assert sheet["B2"].value == "t_fund"
    assert sheet["C2"].value == "bal"
    assert sheet["F2"].value == "—"
    assert sheet["G2"].value == "1104报送"
    assert sheet["H2"].value == "s"
    assert sheet["I2"].value == "已完成"
    assert sheet["J2"].value == "admin"
    assert sheet["K2"].value == "2026-08-06 11:20:00"
    assert sheet.max_column == 11
    assert sheet.freeze_panes == "A2"
    assert all(cell.alignment.wrap_text for cell in sheet[2])
    assert sheet.column_dimensions["C"].width >= 40
    for row in sheet:
        for cell in row:
            assert cell.alignment.horizontal == cell.alignment.vertical == "center"
            assert all(getattr(cell.border, edge).style == "thin" for edge in ("left", "right", "top", "bottom"))



class _ExportService:
    def catalog(self):
        return {}

    def list_records(self, query, user):
        return {"items": []}

    def export_records(self, query):
        if query.get("empty") == "1":
            raise ValidationError(message="无数据可导出")
        return "报表特殊处理_2026-07-31_20260806_112000.xlsx", build_export_xlsx(
            [
                {
                    "report_period": query.get("report_period"),
                    "report_process_name_snapshot": "1104报送",
                    "dimension": "asset",
                    "summary": "摘要",
                    "table_name": "t_asset",
                    "field_name": "qty",
                    "value_before": "3",
                    "value_after": "4",
                    "special_handling_at": NOW,
                    "handler_display_name_snapshot": "管理员",
                    "governance_owner_display_name_snapshot": "治理资产",
                    "status": "pending",
                }
            ]
        )

    def create(self, body, user, request_id):
        return {"id": 1}

    def get(self, record_id, user):
        return {"id": record_id}

    def update(self, record_id, body, user, request_id):
        return {"id": record_id}

    def change_status(self, record_id, body, user, request_id):
        return {"id": record_id}

    def void(self, record_id, body, user, request_id):
        return {"id": record_id}

    def reopen(self, record_id, body, user, request_id):
        return {"id": record_id}

    def audit(self, record_id, query):
        return {"items": []}

    def summary(self, query):
        return {"total": 0}


def _manifest():
    import json
    from importlib import resources

    return ModuleManifest.from_mapping(
        json.loads(
            resources.files("auto_check.modules.report_special_processing")
            .joinpath("manifest.json")
            .read_text(encoding="utf-8")
        )
    )


def _dispatch(suffix, *, query=None, user=None):
    from auto_check.modules.report_special_processing.api import register_routes

    router = ModuleRouter(_manifest(), default_permission_evaluator)
    register_routes(router, lambda: _ExportService())
    user = dict(user or {"role": "user"})
    if str(user.get("role")) != "admin" and "capabilities" not in user:
        user["capabilities"] = ["rsp.view", "rsp.detail"]
    return router.dispatch(
        request=ModuleRequest(
            "GET",
            _manifest().api_prefix + suffix,
            {},
            query or {},
            None,
            user,
        ),
        body_size=0,
    )


def test_api_export_returns_xlsx_attachment_before_record_id_route():
    response = _dispatch(
        "/records/export",
        query={"report_period": "2026-07-31", "status": "pending"},
    )
    assert response.status == 200
    assert response.content_type.startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert any(name.lower() == "content-disposition" for name, _ in response.headers)
    assert response.body[:2] == b"PK"


def test_api_export_empty_result_returns_domain_message():
    response = _dispatch("/records/export", query={"empty": "1"})
    assert response.status == 400
    assert response.body["error"]["message"] == "无数据可导出"


def test_export_rows_render_legacy_chinese_fields_like_ledger():
    rows = export_rows(
        [
            {
                "report_period": date(2026, 7, 31),
                "report_process_name_snapshot": "人行大集中报送",
                "dimension": "project",
                "summary": "摘要",
                "table_name": "资产表｜fa；估值表｜va",
                "field_name": "金额｜amt；数量｜qty；；汇率｜rate",
                "value_before": "1",
                "value_after": "2",
                "special_handling_at": NOW,
                "handler_display_name_snapshot": "管理员",
                "governance_owner_display_name_snapshot": "治理负责人",
                "status": "pending",
                "record_no": "RSP-1",
                "reports": ["报表A"],
                "processing_content": "说明",
                "processing_script": "select 1;",
            }
        ]
    )
    assert [row[1:5] for row in rows] == [["资产表", "金额\n数量", "1", "2"], ["估值表", "汇率", "1", "2"]]


def test_structured_export_groups_values_and_merges_record_metadata():
    record = {
        "report_period": date(2026, 8, 31),
        "summary": "第一行摘要\n第二行摘要",
        "business_system_name_snapshot": "衡泰",
        "status": "completed",
        "report_processes": [{"name": "人行大集中报送"}, {"name": "1104报送"}],
        "structured_content": {"tables": [
            {"chinese_table_name": "资产表", "table_name": "fa", "fields": [
                {"chinese_column_name": "金额", "column_name": "amt", "value_before": "001", "value_after": "=1+1"},
                {"chinese_column_name": "数量", "column_name": "qty", "value_before": "001", "value_after": "=1+1"},
            ]},
            {"chinese_table_name": "估值表", "table_name": "va", "fields": [
                {"chinese_column_name": "金额", "column_name": "amt2", "value_before": "001", "value_after": "=1+1"},
                {"column_name": "x", "value_before": 0, "value_after": ""},
            ]},
        ]},
    }
    rows = export_rows([record])
    assert [row[1:5] for row in rows] == [["资产表", "金额\n数量", "001", "=1+1"], ["估值表", "x", "0", "—"], ["估值表", "金额", "001", "=1+1"]]
    assert rows[0][5:9] == ["衡泰", "1104报送\n人行大集中报送", "第一行摘要\n第二行摘要", "已完成"]
    sheet = load_workbook(BytesIO(build_export_xlsx([record]))).active
    assert {str(item) for item in sheet.merged_cells.ranges} == {"A2:A4", "F2:F4", "G2:G4", "H2:H4", "I2:I4", "J2:J4", "K2:K4"}
    assert sheet["A2"].value == "2026-08-31"
    assert sheet["H2"].value == "第一行摘要\n第二行摘要"
    assert sheet["H2"].alignment.wrap_text
    assert sheet["D2"].value == "001"
    assert sheet["E2"].value == "=1+1"
    assert sheet["E2"].data_type == "s"


def test_legacy_export_keeps_unmatched_multiline_values_intact():
    rows = export_rows([{"field_name": "金额｜amt；数量｜qty", "value_before": "a\nb\nc", "value_after": "z"}])
    assert rows[0][2:5] == ["金额\n数量", "a\nb\nc", "z"]


def test_legacy_export_pairs_fields_and_groups_equal_values():
    rows = export_rows([{"field_name": "金额｜amt；数量｜qty；汇率｜rate", "value_before": "1\n1\n2", "value_after": "3\n3\n4"}])
    assert [row[2:5] for row in rows] == [["汇率", "2", "4"], ["金额\n数量", "1", "3"]]


def test_legacy_tables_remain_multiline_when_field_mapping_is_unknown():
    rows = export_rows([{"table_name": "资产表｜fa；估值表｜va", "field_name": "旧字段", "value_before": "1\n2", "value_after": "3"}])
    assert rows[0][1:5] == ["资产表\n估值表", "旧字段", "1\n2", "3"]
