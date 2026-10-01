"""20260930 逐笔业务同步：原始值边界、跨期保护与真实映射执行链。"""
from datetime import date
from collections import Counter
from decimal import Decimal
import re

from openpyxl import load_workbook
import pytest

from auto_check.db_validation.engine import DbValidationEngine
from auto_check.db_validation.excel import HEADERS
from auto_check.db_validation.mapping_models import TableMapping
from auto_check.db_validation.mapping_service import DbValidationMappingService
from auto_check.db_validation.metadata import TableFieldCatalog
from auto_check.db_validation.rules.basic import (
    REQUIRED_CHINESE_FIELDS_BY_SCOPE,
    OPTIONAL_CHINESE_FIELDS_BY_SCOPE,
    run_basic_rules,
)
from tests.test_db_validation_mapping_service import (
    FakeDataClient,
    FakeMetadataClient,
    FakeStorage,
)


REPORT_DATE = date(2026, 9, 30)
YIELD_FIELDS = {
    "产品代码": "physical_product",
    "地区": "physical_region",
    "地区代码": "physical_region",
    "客户类型": "physical_client",
    "币种": "physical_currency",
    "当月年化收益率": "physical_yield",
    "期末产品金额折人民币": "physical_amount",
}
CODE_RULES = (
    ("ZG06", "资产收益权内部编码", "Zg06_Rule17", "Zg06_Rule18"),
    ("ZG07", "贷款借据编码", "Zg07_Rule19", "Zg07_Rule20"),
    ("ZG12", "除资产收益权外其他债权内部编码", "Zg12_Rule19", "Zg12_Rule20"),
)
CODE_OUTPUT_CONTRACTS = {
    "ZG06": {
        "form": "资产收益权明细信息",
        "empty_error": "资产收益权内部编码一般不应为空",
        "duplicate_error": "资产收益权内部编码应该唯一",
        "empty_detail": "产品代码_资产收益权内部编码_基础资产出让机构名称:P_OUT_{code}_ISSUER_OUT",
        "duplicate_detail": "产品代码_资产收益权内部编码:P_OUT_{code}",
    },
    "ZG07": {
        "form": "除回购和拆借外贷款明细信息",
        "empty_error": "贷款借据编码不应为空",
        "duplicate_error": "贷款借据编码应该唯一",
        "empty_detail": "产品代码_借款人代码_贷款借据编码:P_OUT_BORROWER_OUT_{code}",
        "duplicate_detail": "产品代码_贷款借据编码:P_OUT_{code}",
    },
    "ZG12": {
        "form": "除资产收益权外其他债权明细信息",
        "empty_error": "除资产收益权外其他债权内部编码不应为空",
        "duplicate_error": "除资产收益权外其他债权内部编码应该唯一",
        "empty_detail": "产品代码_借款人代码_除资产收益权外其他债权内部编码:P_OUT_BORROWER_OUT_{code}",
        "duplicate_detail": "产品代码_除资产收益权外其他债权内部编码:P_OUT_{code}",
    },
}
FLAGS = (
    "科技相关产业标识", "绿色领域标识", "普惠领域标识",
    "养老产业标识", "数字经济核心产业标识",
)


def _rule_id(row):
    return re.search(r"Zg\d{2}_Rule\d+", row.rule).group()


def _run(code, current, previous=(), *, fields, related=None):
    table = f"physical_{code.lower()}"
    return run_basic_rules(
        code, REPORT_DATE, current, list(previous), related,
        field_catalog=TableFieldCatalog({table: fields}), table_name=table,
    )


def _only(rows, rule):
    return [row for row in rows if _rule_id(row) == rule]


def _yield_row(value, *, product="P1", region="", client="", currency=""):
    return {
        "physical_product": product, "physical_region": region,
        "physical_client": client, "physical_currency": currency,
        "physical_yield": value,
    }


@pytest.mark.parametrize("current,previous,expected", [
    ("12", "2", False), ("12.00001", "2", True),
    ("2", "12", False), ("1.99999", "12", True),
    ("10", "0", False), ("10.00001", "0", True),
    ("-10", "0", False), ("-10.00001", "0", True),
    ("4.5", "1", False), ("100", "89", True),
])
def test_zg04_rule15_uses_strict_absolute_difference(current, previous, expected):
    rows = _run("ZG04", [_yield_row(current)], [_yield_row(previous)], fields=YIELD_FIELDS)
    hits = _only(rows, "Zg04_Rule15")
    assert bool(hits) is expected
    if expected:
        assert hits[0].value1 == f"当月年化收益率:{float(current)}"
        assert hits[0].value2 == f"当月年化收益率上期数:{float(previous)}"
        assert "200%" not in hits[0].rule + hits[0].error
        assert "波动幅度" not in hits[0].detail


@pytest.mark.parametrize("value_type", [str, Decimal, float], ids=["string", "decimal-db", "float-db"])
@pytest.mark.parametrize("current,previous,expected", [
    ("16.01", "6.01", True), ("6.01", "16.01", True),
    ("16.01001", "6.01", True), ("6.01", "16.01001", True),
])
def test_zg04_rule15_keeps_the_exe_float_boundary(value_type, current, previous, expected):
    rows = _run(
        "ZG04", [_yield_row(value_type(current))], [_yield_row(value_type(previous))],
        fields=YIELD_FIELDS,
    )
    assert bool(_only(rows, "Zg04_Rule15")) is expected


@pytest.mark.parametrize("current,expected", [("10", False), ("10.01", True), ("-10.01", True)])
def test_zg04_rule15_missing_previous_compares_to_zero_without_extra_note(current, expected):
    rows = _run("ZG04", [_yield_row(current)], fields=YIELD_FIELDS)
    hits = _only(rows, "Zg04_Rule15")
    assert bool(hits) is expected
    if expected:
        assert hits[0].value2.endswith(":0.0")
        assert hits[0].note == ""


@pytest.mark.parametrize("changed_field,changed_value", [
    ("physical_product", "P2"), ("physical_client", "2"),
    ("physical_currency", "USD"), ("physical_region", "320101"),
])
def test_zg04_rule15_keeps_the_four_business_matching_dimensions(changed_field, changed_value):
    current = _yield_row("11", client="1", currency="CNY")
    previous = _yield_row("11", client="1", currency="CNY")
    previous[changed_field] = changed_value
    hits = _only(_run("ZG04", [current], [previous], fields=YIELD_FIELDS), "Zg04_Rule15")
    assert len(hits) == 1
    assert hits[0].value2 == "当月年化收益率上期数:0.0"
    assert hits[0].note == ""


@pytest.mark.parametrize("field,current_key,previous_key,expected", [
    ("physical_region", None, "", True),
    ("physical_region", None, float("nan"), False),
    ("physical_client", None, "", True),
    ("physical_client", None, None, False),
    ("physical_client", "1", " 1", True),
    ("physical_currency", None, "", True),
    ("physical_currency", None, float("nan"), False),
    ("physical_product", "P1 ", "P1", True),
])
def test_zg04_rule15_raw_four_keys_distinguish_empty_null_and_spaces(field, current_key, previous_key, expected):
    current = _yield_row("11", currency="CNY")
    previous = _yield_row("11", currency="CNY")
    current[field] = current_key
    previous[field] = previous_key
    hits = _only(_run("ZG04", [current], [previous], fields=YIELD_FIELDS), "Zg04_Rule15")
    assert bool(hits) is expected
    if expected:
        assert hits[0].value2 == "当月年化收益率上期数:0.0"


@pytest.mark.parametrize("raw_key,expected_detail", [
    ("", "P1____11.0"),
    (None, "P1_0_0_0_11.0"),
])
def test_zg04_rule15_detail_retains_raw_empty_keys_and_only_fills_null_with_zero(raw_key, expected_detail):
    current = _yield_row("11", region=raw_key, client=raw_key, currency=raw_key)
    previous = _yield_row("0", region=raw_key, client=raw_key, currency=raw_key)
    hits = _only(_run("ZG04", [current], [previous], fields=YIELD_FIELDS), "Zg04_Rule15")
    assert len(hits) == 1
    assert hits[0].detail == "产品代码_地区_客户类型_币种_当月年化收益率跨期差值:" + expected_detail


def test_zg04_rule15_keeps_each_identical_previous_candidate_in_the_left_join():
    current = [_yield_row("20", currency="CNY")]
    previous = [_yield_row("5", currency="CNY"), _yield_row("5", currency="CNY")]
    hits = _only(_run("ZG04", current, previous, fields=YIELD_FIELDS), "Zg04_Rule15")
    assert len(hits) == 2
    assert all(hit.value2 == "当月年化收益率上期数:5.0" and hit.note == "" for hit in hits)


@pytest.mark.parametrize("current,previous,expected", [
    ("11", None, True), ("11", float("nan"), True),
    ("11", "nan", False), (None, "11", True),
    ("", "11", True), ("nan", "11", False),
    (float("inf"), "0", True),
])
def test_zg04_rule15_fillna_precedes_float_but_does_not_fill_literal_nan(current, previous, expected):
    hits = _only(_run("ZG04", [_yield_row(current)], [_yield_row(previous)], fields=YIELD_FIELDS), "Zg04_Rule15")
    assert bool(hits) is expected


@pytest.mark.parametrize("current,previous,zero_amount,expected", [
    ("1.2", "1", True, False), ("1.20001", "1", True, True),
    ("0.8", "1", True, False), ("0.79999", "1", True, True),
    ("-1.20001", "-1", True, True), ("12", "0", True, False),
    ("12", "1", False, False),
])
def test_zg04_rule19_retains_relative_twenty_percent_and_nonzero_previous(current, previous, zero_amount, expected):
    amount_row = _yield_row("", region="000000", client="1", currency="BWB")
    amount_row["physical_amount"] = "0" if zero_amount else "100"
    rows = _run("ZG04", [_yield_row(current), amount_row], [_yield_row(previous)], fields=YIELD_FIELDS)
    hits = _only(rows, "Zg04_Rule19")
    assert bool(hits) is expected
    if expected:
        assert "20%" in hits[0].rule
        assert "波动幅度" in hits[0].detail


def test_zg04_rule19_keeps_the_same_duplicate_previous_candidates_as_rule15():
    amount_row = _yield_row("", region="000000", currency="BWB")
    amount_row["physical_amount"] = "0"
    current = [_yield_row("0"), amount_row]
    previous = [_yield_row("5"), _yield_row("5")]
    hits = _only(_run("ZG04", current, previous, fields=YIELD_FIELDS), "Zg04_Rule19")
    assert len(hits) == 2
    assert all(hit.value2 == "当月年化收益率上期数:5.0" and hit.note == "" for hit in hits)


@pytest.mark.parametrize("current,previous,expected", [
    ("0", "inf", False), ("0", "-inf", False),
    ("inf", "5", True), ("nan", "5", False),
    ("0", "nan", False), ("0", None, False),
])
def test_zg04_rule19_keeps_shared_exe_float_nan_and_infinity_conditions(current, previous, expected):
    amount_row = _yield_row("", region="000000", currency="BWB")
    amount_row["physical_amount"] = "0"
    hits = _only(_run("ZG04", [_yield_row(current), amount_row], [_yield_row(previous)], fields=YIELD_FIELDS), "Zg04_Rule19")
    assert bool(hits) is expected
    if expected:
        assert hits[0].value1 == f"当月年化收益率:{float(current)}"


@pytest.mark.parametrize("region,currency", [(" 000000", "BWB"), ("000000 ", "BWB"), ("000000", " BWB"), ("000000", "BWB ")])
def test_zg04_rule19_zero_amount_source_uses_raw_area_and_currency(region, currency):
    amount_row = _yield_row("", region=region, currency=currency)
    amount_row["physical_amount"] = "0"
    rows = _run("ZG04", [_yield_row("0"), amount_row], [_yield_row("5")], fields=YIELD_FIELDS)
    assert not _only(rows, "Zg04_Rule19")


@pytest.mark.parametrize("amount", ["bad", " ", "None", "1,000"])
def test_zg04_rule19_zero_amount_source_exposes_exe_float_conversion_errors(amount):
    amount_row = _yield_row("", region="000000", currency="BWB")
    amount_row["physical_amount"] = amount
    with pytest.raises(ValueError):
        _run("ZG04", [_yield_row("0"), amount_row], [_yield_row("5")], fields=YIELD_FIELDS)


@pytest.mark.parametrize("amount_values,expected", [
    (("nan",), True), ((None,), True), ((float("nan"),), True),
    (("nan", "100"), False), (("50", "-50"), True),
    (("0.1", "0.2", "-0.3"), True),
    (("10000000000000000", "1", "1", "-10000000000000000"), False),
])
def test_zg04_rule19_zero_amount_sum_skips_null_and_nan_like_exe_groupby(amount_values, expected):
    amount_rows = [_yield_row("", region="000000", currency="BWB") for _ in amount_values]
    for amount_row, amount in zip(amount_rows, amount_values):
        amount_row["physical_amount"] = amount
    hits = _only(_run("ZG04", [_yield_row("0"), *amount_rows], [_yield_row("5")], fields=YIELD_FIELDS), "Zg04_Rule19")
    assert bool(hits) is expected


@pytest.mark.parametrize("rule,current_yield,previous_yield", [
    ("Zg04_Rule15", "11", "0"),
    ("Zg04_Rule17", "0", "5"),
    ("Zg04_Rule19", "0", "5"),
])
def test_zg04_yield_rules_mark_contains_the_complete_exe_rule_title(rule, current_yield, previous_yield):
    amount_row = _yield_row("", region="000000", currency="BWB")
    amount_row["physical_amount"] = "0"
    rows = _run("ZG04", [_yield_row(current_yield), amount_row], [_yield_row(previous_yield)], fields=YIELD_FIELDS)
    hits = _only(rows, rule)
    assert len(hits) == 1
    assert hits[0].mark == f"20260930-D1003632000013-ZG04-{hits[0].rule}"


@pytest.mark.parametrize("previous,expected_value", [
    ("0", None), ("0.0000", None), ("", None),
    (None, "nan"), (float("nan"), "nan"), (Decimal("NaN"), "nan"),
    ("nan", "nan"), (float("inf"), "inf"), (float("-inf"), "-inf"),
    (Decimal("Infinity"), "inf"), (Decimal("-Infinity"), "-inf"),
    ("5", "5.0"), ("-3", "-3.0"),
])
def test_zg04_rule17_keeps_the_exe_previous_float_and_missing_values(previous, expected_value):
    rows = _run("ZG04", [_yield_row("0")], [_yield_row(previous)], fields=YIELD_FIELDS)
    hits = _only(rows, "Zg04_Rule17")
    assert bool(hits) is (expected_value is not None)
    if expected_value is not None:
        assert len(hits) == 1
        assert hits[0].value1 == "当月年化收益率:0.0"
        assert hits[0].value2 == f"当月年化收益率_上期数:{expected_value}"
        assert hits[0].note == ""


def test_zg04_rule17_keeps_the_exe_left_merge_missing_previous_nan_result():
    hits = _only(_run("ZG04", [_yield_row("0")], fields=YIELD_FIELDS), "Zg04_Rule17")
    assert len(hits) == 1
    assert hits[0].value2 == "当月年化收益率_上期数:nan"
    assert hits[0].note == ""


def test_zg04_rule17_synthetic_68_results_split_into_16_matched_and_52_missing():
    # Preserve the observed result pattern without copying business identifiers or records.
    current = [_yield_row("0", product=f"MISSING_{index}") for index in range(52)]
    current += [_yield_row("0", product=f"MATCHED_{index}") for index in range(16)]
    previous = [_yield_row(str(index + 1), product=f"MATCHED_{index}") for index in range(16)]
    hits = _only(_run("ZG04", current, previous, fields=YIELD_FIELDS), "Zg04_Rule17")
    missing = [hit for hit in hits if hit.value2 == "当月年化收益率_上期数:nan"]
    assert len(hits) == 68
    assert len(missing) == 52
    assert len(hits) - len(missing) == 16
    assert {hit.detail for hit in missing} == {f"产品代码:MISSING_{index}" for index in range(52)}


def test_zg04_rule17_matches_product_total_without_requiring_equal_client_type():
    rows = _run(
        "ZG04", [_yield_row("0", client="1")],
        [_yield_row("4.5", client="2")], fields=YIELD_FIELDS,
    )
    hits = _only(rows, "Zg04_Rule17")
    assert len(hits) == 1
    assert hits[0].value2 == "当月年化收益率_上期数:4.5"


@pytest.mark.parametrize("previous", [
    _yield_row("5", product="P2"),
    _yield_row("5", region="320101"),
    _yield_row("5", currency="CNY"),
    _yield_row("5", region=" "),
    _yield_row("5", currency=" "),
    _yield_row("5", product=" P1"),
])
def test_zg04_rule17_excludes_other_raw_products_and_non_totals_then_reports_missing(previous):
    rows = _run("ZG04", [_yield_row("0")], [previous], fields=YIELD_FIELDS)
    hits = _only(rows, "Zg04_Rule17")
    assert len(hits) == 1
    assert hits[0].value2 == "当月年化收益率_上期数:nan"


@pytest.mark.parametrize("region,currency,current", [("320101", "", "0"), ("", "CNY", "0"), (" ", "", "0"), ("", " ", "0"), ("", "", "1")])
def test_zg04_rule17_keeps_current_total_and_zero_conditions(region, currency, current):
    rows = _run("ZG04", [_yield_row(current, region=region, currency=currency)], [_yield_row("5")], fields=YIELD_FIELDS)
    assert not _only(rows, "Zg04_Rule17")


@pytest.mark.parametrize("current", [None, float("nan"), Decimal("NaN"), "nan", float("inf"), Decimal("Infinity")])
def test_zg04_rule17_keeps_the_exe_missing_and_nonzero_current_conditions(current):
    rows = _run("ZG04", [_yield_row(current)], [_yield_row("5")], fields=YIELD_FIELDS)
    assert not _only(rows, "Zg04_Rule17")


def test_zg04_rule17_converts_raw_empty_current_string_to_zero():
    hits = _only(_run("ZG04", [_yield_row("")], [_yield_row("5")], fields=YIELD_FIELDS), "Zg04_Rule17")
    assert len(hits) == 1
    assert hits[0].value1 == "当月年化收益率:0.0"
    assert hits[0].value2 == "当月年化收益率_上期数:5.0"


@pytest.mark.parametrize("region,currency", [(None, None), (float("nan"), float("nan")), ("", None)])
def test_zg04_rule17_accepts_null_and_raw_empty_total_keys(region, currency):
    current = _yield_row("0", region=region, currency=currency)
    previous = _yield_row("5", region=region, currency=currency)
    hits = _only(_run("ZG04", [current], [previous], fields=YIELD_FIELDS), "Zg04_Rule17")
    assert len(hits) == 1
    assert hits[0].value2 == "当月年化收益率_上期数:5.0"


@pytest.mark.parametrize("invalid", ["invalid", "None", " ", "0,0"])
@pytest.mark.parametrize("invalid_side,current", [("current", "0"), ("previous", "0"), ("previous", "3")])
def test_zg04_rule17_exposes_exe_float_conversion_errors_before_filtering(invalid, invalid_side, current):
    current_value = invalid if invalid_side == "current" else current
    previous_value = invalid if invalid_side == "previous" else "5"
    with pytest.raises(ValueError):
        _run("ZG04", [_yield_row(current_value)], [_yield_row(previous_value)], fields=YIELD_FIELDS)


def test_zg04_rule17_does_not_convert_an_unmatched_other_product_invalid_yield():
    rows = _run("ZG04", [_yield_row("0")], [_yield_row("invalid", product="P2")], fields=YIELD_FIELDS)
    hits = _only(rows, "Zg04_Rule17")
    assert len(hits) == 1
    assert hits[0].value2 == "当月年化收益率_上期数:nan"


@pytest.mark.parametrize("previous_values,expected_values", [
    (("5", "0", "nan"), ("5.0", "nan")),
    (("nan", "0", "5"), ("nan", "5.0")),
    (("5", "7", None, "0"), ("5.0", "7.0")),
    (("0", None, "7", "5"), ("7.0", "5.0")),
    (("5", "5"), ("5.0", "5.0")),
    ((None, float("nan")), ("nan",)),
    (("", float("nan")), ()),
    (("nan", "nan"), ("nan", "nan")),
    ((None, float("nan"), "5"), ("5.0",)),
])
def test_zg04_rule17_keeps_each_exe_candidate_after_filtering_without_overwriting(previous_values, expected_values):
    rows = _run("ZG04", [_yield_row("0")], [_yield_row(value) for value in previous_values], fields=YIELD_FIELDS)
    hits = _only(rows, "Zg04_Rule17")
    assert Counter(row.value2 for row in hits) == Counter(f"当月年化收益率_上期数:{value}" for value in expected_values)
    assert len(hits) == len(expected_values)
    assert all(row.note == "" for row in hits)


@pytest.mark.parametrize("issuer_type", ["4", "5"])
def test_zg06_rule14_does_not_report_when_all_five_flags_are_null_or_empty(issuer_type):
    fields = {"产品代码": "physical_product", "基础资产出让机构类型": "physical_type"}
    row = {"physical_product": "P1", "physical_type": issuer_type}
    for index, flag in enumerate(FLAGS):
        fields[flag] = f"physical_flag_{index}"
        row[f"physical_flag_{index}"] = None if index % 2 else ""
    assert not _only(_run("ZG06", [row], fields=fields), "Zg06_Rule14")


@pytest.mark.parametrize("flag", FLAGS)
@pytest.mark.parametrize("value", [0, "0", " ", "1"])
def test_zg06_rule14_treats_each_filled_raw_flag_including_zero_and_space_as_filled(flag, value):
    fields = {
        "产品代码": "physical_product", "基础资产出让机构类型": "physical_type",
        flag: "physical_flag",
    }
    row = {"physical_product": "P1", "physical_type": "5", "physical_flag": value}
    assert len(_only(_run("ZG06", [row], fields=fields), "Zg06_Rule14")) == 1


@pytest.mark.parametrize("issuer_type", ["1", "2", "3"])
def test_zg06_rule14_does_not_apply_to_other_issuer_types(issuer_type):
    fields = {"产品代码": "physical_product", "基础资产出让机构类型": "physical_type", FLAGS[0]: "physical_flag"}
    row = {"physical_product": "P1", "physical_type": issuer_type, "physical_flag": "1"}
    assert not _only(_run("ZG06", [row], fields=fields), "Zg06_Rule14")


@pytest.mark.parametrize("field", ["转让预计终止日期", "转让展期到期日期"])
@pytest.mark.parametrize("value,expected", [("2089-12-31", False), ("2090-01-01", True)])
def test_zg06_rule9_retains_2090_boundary_and_removes_only_error_prefix(field, value, expected):
    fields = {"产品代码": "physical_product", "资产收益权内部编码": "physical_code", field: "physical_end"}
    rows = _run("ZG06", [{"physical_product": "P1", "physical_code": "B1", "physical_end": value}], fields=fields)
    hits = _only(rows, "Zg06_Rule9")
    assert bool(hits) is expected
    if expected:
        assert hits[0].mark.endswith("Zg06_Rule9")
        assert hits[0].rule.startswith("Zg06_Rule9:")
        assert hits[0].error == "转让预计终止日期，转让展期到期日期大于、等于2090，需核实"


@pytest.mark.parametrize("issuer_type,expected", [("4", True), ("5", True), (" 4", False), ("4 ", False), (" 5 ", False)])
def test_zg06_rule14_does_not_strip_raw_issuer_type_and_keeps_repeated_hits(issuer_type, expected):
    fields = {"产品代码": "physical_product", "资产收益权内部编码": "physical_code",
              "基础资产出让机构类型": "physical_type", "科技相关产业标识": "physical_flag"}
    row = {"physical_product": "P_OUT", "physical_code": "CODE_OUT", "physical_type": issuer_type, "physical_flag": "1"}
    hits = _only(_run("ZG06", [dict(row), dict(row)], fields=fields), "Zg06_Rule14")
    assert len(hits) == (2 if expected else 0)


@pytest.mark.parametrize("field", ["转让预计终止日期", "转让展期到期日期"])
@pytest.mark.parametrize("value,expected", [("2090-01-01", True), (" 2090-01-01", False), (None, False), (float("nan"), False)])
def test_zg06_rule9_uses_raw_first_four_date_characters_and_keeps_repeated_hits(field, value, expected):
    fields = {"产品代码": "physical_product", "资产收益权内部编码": "physical_code", field: "physical_end"}
    row = {"physical_product": "P_OUT", "physical_code": "CODE_OUT", "physical_end": value}
    hits = _only(_run("ZG06", [dict(row), dict(row)], fields=fields), "Zg06_Rule9")
    assert len(hits) == (2 if expected else 0)


@pytest.mark.parametrize("scope,field,empty_rule,duplicate_rule", CODE_RULES)
@pytest.mark.parametrize("code,expected", [
    (None, True), (float("nan"), True), (Decimal("NaN"), True), ("", True),
    (" ", False), ("None", False), ("nan", False), ("VALID1", False),
])
def test_new_code_empty_rules_preserve_raw_null_empty_space_and_literal_semantics(scope, field, empty_rule, duplicate_rule, code, expected):
    fields = {"产品代码": "physical_product", field: "physical_code"}
    rows = _run(scope, [{"physical_product": "P1", "physical_code": code}], fields=fields)
    hits = _only(rows, empty_rule)
    assert bool(hits) is expected
    assert not _only(rows, duplicate_rule)
    if expected:
        assert hits[0].mark.endswith(empty_rule)
        assert field in hits[0].detail
        assert hits[0].value1.startswith(field + ":")
        assert field in hits[0].error


@pytest.mark.parametrize("scope,field,empty_rule,duplicate_rule", CODE_RULES)
@pytest.mark.parametrize("code", ["CODE1", "", " "])
def test_new_code_duplicates_execute_without_institution_or_manager_fields(scope, field, empty_rule, duplicate_rule, code):
    fields = {"产品代码": "physical_product", field: "physical_code"}
    records = [{"physical_product": "P1", "physical_code": code} for _ in range(3)]
    rows = _run(scope, records, fields=fields)
    hits = _only(rows, duplicate_rule)
    assert hits, "三条同产品同原始编码的业务明细应真正执行查重并输出"
    assert len(hits) == len(records), "相同记录也须逐条输出，不能按结果去重折叠重复明细"
    assert all(row.mark.endswith(duplicate_rule) for row in hits)
    assert all(row.value2 == "" for row in hits)
    assert all(field in row.detail and field in row.error for row in hits)
    assert bool(_only(rows, empty_rule)) is (code == "")
    if code == "":
        assert len(_only(rows, empty_rule)) == len(records)


@pytest.mark.parametrize("scope,field,empty_rule,duplicate_rule", CODE_RULES)
@pytest.mark.parametrize("code", [None, float("nan"), Decimal("NaN")])
def test_new_code_null_values_do_not_form_duplicate_groups(scope, field, empty_rule, duplicate_rule, code):
    fields = {"产品代码": "physical_product", field: "physical_code"}
    records = [{"physical_product": "P1", "physical_code": code} for _ in range(2)]
    rows = _run(scope, records, fields=fields)
    assert _only(rows, empty_rule)
    assert not _only(rows, duplicate_rule)


@pytest.mark.parametrize("scope,field,empty_rule,duplicate_rule", CODE_RULES)
def test_new_code_duplicate_groups_do_not_strip_original_codes(scope, field, empty_rule, duplicate_rule):
    fields = {"产品代码": "physical_product", field: "physical_code"}
    records = [{"physical_product": "P1", "physical_code": code} for code in ("CODE1", " CODE1", "CODE1 ", "CODE2")]
    rows = _run(scope, records, fields=fields)
    assert not _only(rows, empty_rule)
    assert not _only(rows, duplicate_rule)


@pytest.mark.parametrize("scope,field,empty_rule,duplicate_rule", CODE_RULES)
@pytest.mark.parametrize("product", [None, float("nan"), Decimal("NaN")])
def test_new_code_null_products_do_not_form_duplicate_groups(scope, field, empty_rule, duplicate_rule, product):
    fields = {"产品代码": "physical_product", field: "physical_code"}
    records = [{"physical_product": product, "physical_code": "CODE1"} for _ in range(2)]
    rows = _run(scope, records, fields=fields)
    assert not _only(rows, duplicate_rule)


@pytest.mark.parametrize("scope,field,empty_rule,duplicate_rule", CODE_RULES)
@pytest.mark.parametrize("raw_code,display_code", [("", ""), (None, "nan"), (float("nan"), "nan"), (Decimal("NaN"), "nan")])
def test_new_code_empty_output_matches_exact_exe_contract(scope, field, empty_rule, duplicate_rule, raw_code, display_code):
    contract = CODE_OUTPUT_CONTRACTS[scope]
    fields = {
        "产品代码": "physical_product", field: "physical_code",
        "借款人代码": "physical_borrower", "基础资产出让机构名称": "physical_issuer",
    }
    current = [{"physical_product": "P_OUT", "physical_code": raw_code,
                "physical_borrower": "BORROWER_OUT", "physical_issuer": "ISSUER_OUT"}]
    hits = _only(_run(scope, current, fields=fields), empty_rule)
    assert len(hits) == 1
    assert hits[0].form == contract["form"]
    assert hits[0].detail == contract["empty_detail"].format(code=display_code)
    assert hits[0].value1 == f"{field}:{display_code}"
    assert hits[0].value2 == ""
    assert hits[0].rule == f"{empty_rule}:{field}为空，需核实"
    assert hits[0].error == contract["empty_error"]


@pytest.mark.parametrize("scope,field,empty_rule,duplicate_rule", CODE_RULES)
@pytest.mark.parametrize("original_excel_cell", ["", None])
def test_new_code_imported_excel_empty_cell_becomes_db_null_and_exe_nan_output(scope, field, empty_rule, duplicate_rule, original_excel_cell):
    from auto_check.app.pbc_import import _normalize_cell
    fields = {"产品代码": "physical_product", field: "physical_code",
              "借款人代码": "physical_borrower", "基础资产出让机构名称": "physical_issuer"}
    imported_code = _normalize_cell(original_excel_cell)
    assert imported_code is None
    row = {"physical_product": "P_OUT", "physical_code": imported_code,
           "physical_borrower": "BORROWER_OUT", "physical_issuer": "ISSUER_OUT"}
    hits = _only(_run(scope, [row], fields=fields), empty_rule)
    assert len(hits) == 1
    assert hits[0].value1 == f"{field}:nan"
    assert hits[0].detail == CODE_OUTPUT_CONTRACTS[scope]["empty_detail"].format(code="nan")
    assert _normalize_cell("None") == "None"


@pytest.mark.parametrize("scope,field,empty_rule,duplicate_rule", CODE_RULES)
@pytest.mark.parametrize("raw_code", ["CODE_OUT", " ", "None", "nan"])
def test_new_code_duplicate_output_matches_exact_exe_contract(scope, field, empty_rule, duplicate_rule, raw_code):
    contract = CODE_OUTPUT_CONTRACTS[scope]
    fields = {
        "产品代码": "physical_product", field: "physical_code",
        "借款人代码": "physical_borrower", "基础资产出让机构名称": "physical_issuer",
    }
    original = {"physical_product": "P_OUT", "physical_code": raw_code,
                "physical_borrower": "BORROWER_OUT", "physical_issuer": "ISSUER_OUT"}
    hits = _only(_run(scope, [dict(original), dict(original)], fields=fields), duplicate_rule)
    assert len(hits) == 2
    for hit in hits:
        assert hit.form == contract["form"]
        assert hit.detail == contract["duplicate_detail"].format(code=raw_code)
        assert hit.value1 == f"{field}:{raw_code}"
        assert hit.value2 == ""
        assert hit.rule == f"{duplicate_rule}:{field}不唯一，需核实"
        assert hit.error == contract["duplicate_error"]


@pytest.mark.parametrize("scope,field,empty_rule,duplicate_rule", CODE_RULES)
def test_new_code_outputs_preserve_source_order_with_empty_before_duplicate(scope, field, empty_rule, duplicate_rule):
    fields = {"产品代码": "physical_product", field: "physical_code"}
    products = ["PZ", "PA", "PZ", "PA"]
    current = [{"physical_product": product, "physical_code": ""} for product in products]
    rows = _run(scope, current, fields=fields)
    hits = [row for row in rows if _rule_id(row) in {empty_rule, duplicate_rule}]
    assert [_rule_id(hit) for hit in hits] == [empty_rule] * 4 + [duplicate_rule] * 4
    assert [hit.detail.split(":", 1)[1].split("_", 1)[0] for hit in hits] == products * 2


@pytest.mark.parametrize("scope,field,empty_rule,duplicate_rule,old_rule,old_fields", [
    ("ZG06", "资产收益权内部编码", "Zg06_Rule17", "Zg06_Rule18", "Zg06_Rule14", {"基础资产出让机构类型": "physical_type", "科技相关产业标识": "physical_flag"}),
    ("ZG07", "贷款借据编码", "Zg07_Rule19", "Zg07_Rule20", "Zg07_Rule18", {"贷款种类": "physical_type", "贷款合同原始发放机构代码": "physical_original_issuer"}),
    ("ZG12", "除资产收益权外其他债权内部编码", "Zg12_Rule19", "Zg12_Rule20", "Zg12_Rule18", {"担保方式": "physical_guarantee"}),
])
def test_new_code_rules_follow_existing_table_rules(scope, field, empty_rule, duplicate_rule, old_rule, old_fields):
    fields = {"产品代码": "physical_product", field: "physical_code", **old_fields}
    original = {"physical_product": "P_OUT", "physical_code": "", "physical_type": "4",
                "physical_flag": "1", "physical_original_issuer": "", "physical_guarantee": "Z"}
    rows = _run(scope, [dict(original), dict(original)], fields=fields)
    ids = [_rule_id(row) for row in rows]
    assert old_rule in ids
    assert max(i for i,rule in enumerate(ids) if rule == old_rule) < min(i for i,rule in enumerate(ids) if rule in {empty_rule, duplicate_rule})


@pytest.mark.parametrize("loan_type", ["3", "4"])
def test_zg07_rule19_does_not_reuse_loan_type_or_original_issuer_condition(loan_type):
    fields = {
        "产品代码": "physical_product", "贷款借据编码": "physical_code",
        "贷款种类": "physical_type", "贷款合同原始发放机构代码": "physical_original_issuer",
    }
    rows = _run("ZG07", [{
        "physical_product": "P1", "physical_code": "LOAN1", "physical_type": loan_type,
        "physical_original_issuer": "",
    }], fields=fields)
    if loan_type == "4":
        assert _only(rows, "Zg07_Rule18"), "应证明原机构字段规则条件实际已触发"
    assert not _only(rows, "Zg07_Rule19")


def _refreshed_catalog(scope, *, extra_fields=None):
    table = f"business_{scope.lower()}"
    fields = {name: f"column_{index:02d}" for index, name in enumerate(sorted(REQUIRED_CHINESE_FIELDS_BY_SCOPE[scope]))}
    fields.update(extra_fields or {})
    mappings = [TableMapping("detail", scope, "", table)]
    if scope == "ZG12":
        mappings.extend([TableMapping("detail", "ZG01", "", "business_zg01"), TableMapping("detail", "ZG05", "", "business_zg05")])
    storage = FakeStorage(mappings)
    service = DbValidationMappingService(database=object())
    service.storage = storage
    metadata = FakeMetadataClient(
        tables=[{"table_name_en": table}],
        fields=[{"table": table, "chinese": name, "english": column} for name, column in fields.items()],
    )
    catalog = service.refresh(
        metadata_client=metadata, data_clients={"detail": FakeDataClient({table: list(fields.values())})},
        baseinfo_table="xt_reg_table_baseinfo", field_info_table="xt_reg_table_field_info",
        sys_manage_id="", classification_id="", signature=("20260930-offline",), source="manual",
        required_chinese_fields_by_scope={scope: REQUIRED_CHINESE_FIELDS_BY_SCOPE[scope]},
        optional_chinese_fields_by_scope={scope: OPTIONAL_CHINESE_FIELDS_BY_SCOPE.get(scope, frozenset())},
    )
    return table, fields, catalog, storage


class _EngineClient:
    """只返回显式的物理表行；上期表不会因为表名前缀意外取得当期行。"""
    class config:
        db_type = "postgresql"
        schema = "dws"

    def __init__(self, rows_by_table):
        self.rows_by_table = rows_by_table
        self.calls = []

    def fetch_all(self, sql, params=()):
        self.calls.append((sql, params))
        for table, rows in self.rows_by_table.items():
            if f'"{table}"' in sql:
                return rows
        return []


@pytest.mark.parametrize("scope,field,empty_rule,duplicate_rule", CODE_RULES)
def test_code_rules_run_through_refresh_catalog_engine_and_excel_without_institution_mapping(tmp_path, scope, field, empty_rule, duplicate_rule):
    table, fields, catalog, storage = _refreshed_catalog(scope)
    assert not [item for item in storage.saved["fields"] if item.is_required and item.mapping_status == "required_missing"]
    assert all(name not in REQUIRED_CHINESE_FIELDS_BY_SCOPE[scope] for name in ("金融机构编码", "数据管理机构", "法人金融机构名称"))
    assert "金融机构编码" not in catalog.fields_for_table(table)
    assert "数据管理机构" not in catalog.fields_for_table(table)
    raw = {column: "" for column in fields.values()}
    raw[fields["产品代码"]] = "P1"
    client = _EngineClient({table: [raw, dict(raw)]})
    result = DbValidationEngine(data_client=client, field_catalog=catalog, output_dir=tmp_path).run(
        report_date=REPORT_DATE, selected_tables=[scope],
    )
    hits = _only(result.rows, empty_rule) + _only(result.rows, duplicate_rule)
    assert {_rule_id(row) for row in hits} == {empty_rule, duplicate_rule}
    assert len(_only(result.rows, empty_rule)) == 2
    assert len(_only(result.rows, duplicate_rule)) == 2
    assert all(row.value2 == "" for row in _only(result.rows, duplicate_rule))
    assert "Ver.20260930" in result.excel_path.name
    assert all(field in row.detail and row.value1.startswith(field + ":") for row in hits)
    assert all(row.form == {"ZG06": "资产收益权明细信息", "ZG07": "除回购和拆借外贷款明细信息", "ZG12": "除资产收益权外其他债权明细信息"}[scope] for row in hits)
    assert all(field in row.error and "规则执行" not in row.error for row in hits)
    wb = load_workbook(result.excel_path, read_only=True, data_only=True)
    try:
        exported = list(wb.active.iter_rows(values_only=True))
        assert list(exported[0]) == HEADERS
        assert all(len(row) == 12 for row in exported)
        excel_hits = [dict(zip(HEADERS, row)) for row in exported[1:] if any(rule in str(row[9]) for rule in (empty_rule, duplicate_rule))]
        assert len(excel_hits) == len(hits)
        assert {re.search(r"Zg\d{2}_Rule\d+", row["校验规则"]).group() for row in excel_hits} == {empty_rule, duplicate_rule}
        assert all(field in row["明细数据相关信息"] and row["数据值1"].startswith(field + ":") for row in excel_hits)
        assert all(row["校验标识"].endswith(re.search(r"Zg\d{2}_Rule\d+", row["校验规则"]).group()) for row in excel_hits)
        assert all(row["数据值2"] is None for row in excel_hits if duplicate_rule in row["校验规则"])
    finally:
        wb.close()


def test_zg13_removed_financial_institution_code_rules_do_not_execute_and_other_rules_still_do():
    fields = {
        "产品代码": "physical_product", "标的企业代码": "physical_target",
        "标的企业名称": "physical_name", "股权出让方代码": "physical_out",
        "股权出让方名称": "physical_out_name", "地区代码": "physical_region",
        "其他股权投资内部编码": "physical_code",
    }
    records = [{
        "physical_product": "P1", "physical_target": "91310000100019382F",
        "physical_name": "光大证券股份有限公司", "physical_out": "91310000100019382F",
        "physical_out_name": "光大证券股份有限公司", "physical_region": "320100",
        "physical_code": "EQ1",
    }]
    rows = _run("ZG13", records, fields=fields)
    assert not _only(rows, "Zg13_Rule15")
    assert not _only(rows, "Zg13_Rule16")
    assert _only(rows, "Zg13_Rule1")


@pytest.mark.parametrize("zg05_balance,zg12_balance,expected", [
    ("0.1", "0", False), ("0.10001", "0", True),
    ("0", "0.1", False), ("0", "0.10001", True),
    ("100", "100.05", False), ("100", "101", True), ("101", "100", True),
])
def test_zg12_rule16_retains_both_difference_directions_and_strict_tenth_boundary(zg05_balance, zg12_balance, expected):
    rows = _run("ZG12", [{"physical_product": "P1", "physical_balance": zg12_balance}], fields={
        "产品代码": "physical_product", "除资产收益权外其他债权余额折人民币": "physical_balance",
    }, related={"ZG05": [{"产品代码": "P1", "币种": "BWB", "AD200_除资产收益权外其他债权": zg05_balance}]})
    assert bool(_only(rows, "Zg12_Rule16")) is expected


def test_empty_zg12_still_checks_zg05_ad200_in_the_formal_engine(tmp_path):
    catalog = TableFieldCatalog({
        "business_zg12": {"产品代码": "product_id", "除资产收益权外其他债权余额折人民币": "debt_balance"},
        "business_zg05": {"产品代码": "product_id", "币种": "currency_id", "AD200_除资产收益权外其他债权": "amount_ad200"},
    }, table_mappings={
        ("detail", "ZG12", ""): "business_zg12", ("detail", "ZG01", ""): "business_zg01",
        ("detail", "ZG05", ""): "business_zg05",
    })
    client = _EngineClient({"business_zg12": [], "business_zg05": [{"product_id": "P1", "currency_id": "BWB", "amount_ad200": "100"}]})
    result = DbValidationEngine(data_client=client, field_catalog=catalog, output_dir=tmp_path).run(report_date=REPORT_DATE, selected_tables=["ZG12"])
    hits = _only(result.rows, "Zg12_Rule16")
    assert len(hits) == 1
    assert hits[0].value1.endswith(":0.0")
    assert hits[0].value2.endswith(":100.0")
    assert any("ZG12 当期表无数据" in message for message in result.warnings)


def test_formal_engine_result_filename_identifies_the_20260930_rule_version(tmp_path):
    catalog = TableFieldCatalog({}, table_mappings={("detail", "ZG06", ""): "business_zg06"})
    result = DbValidationEngine(data_client=_EngineClient({}), field_catalog=catalog, output_dir=tmp_path).run(
        report_date=REPORT_DATE, selected_tables=["ZG06"],
    )
    assert "Ver.20260930" in result.excel_path.name
